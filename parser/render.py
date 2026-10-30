# -*- coding: utf-8 -*-
"""칸 2.8 — 쪽 전체 렌더 — ④ 이미지 요약에 **화면 전체**를 함께 주기 위한 것 (B53).

**왜 필요한가.** 슬라이드에서 잘라낸 그림 하나에는 축 이름·범례·단위가 없는 일이
흔하다. 그것만 보내면 모델이 「무엇의 그림인가」를 지어내고, 그 문장이 청크 텍스트가
되어 답변의 「문서 근거」로 되돌아온다. 쪽 전체를 함께 주면 지어낼 자리가 줄어든다.

**LibreOffice는 파이프라인 의존이 아니라 캡처 길 중 하나다**(문서 7 §7.6-B-6) — 길이 있는지는 `capture_path` 한 함수가
답하고(B128 ①), 실호출 PPT · PDF 인입은 사전 점검에서 그 답으로 멈춘다. 문서 중간에 렌더가 없으면 `None`을 돌려주고 호출부가
`slide_render="none"`을 데이터에 남긴다 — **조용히 다르게 돌지 않는다.**

**이 모듈은 렌더 바이트를 저장하지 않는다**(원본 바이너리 비저장 — 문서 1 B16) — 쪽 캡처를 ④작업 단에 두는 일은
`core/state/captures`(손잡이 `keep_page_captures` — 원본에서 다시 만들 수 있는 쪽 JPEG · B156 ⑤)의 몫이다.
"""
from __future__ import annotations

import logging
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

_LOG = logging.getLogger("onto.parser.render")

MAX_EDGE = 1600         # 긴 변 상한(px) — 넘으면 축소해 보낸다
SOFFICE_TIMEOUT = 180   # 변환이 걸리면 문서를 붙잡지 않고 포기한다


def _pdf_beside(path):
    """같은 폴더·같은 이름의 `.pdf`가 **원본보다 새것**이면 그것을 쓴다.

    사내에서 PPTX와 함께 배포된 PDF가 있으면 변환이 필요 없다. 오래된 PDF는
    쓰지 않는다 — 옛 판의 그림을 새 문서의 근거로 삼게 된다.
    """
    p = Path(path)
    cand = p.with_suffix(".pdf")
    try:
        if cand.exists() and cand.stat().st_mtime >= p.stat().st_mtime:
            return str(cand)
    except OSError:
        pass
    return None


#: **사내가 지정하는 한 자리**(B129 ①) — soffice 실행 파일(또는 그 폴더)의 경로 · 셸 설정 파일에 둔다(`ONTO_*`가 아니라 점검 ·
#: 시험의 격리가 지우지 않는다 — 상태가 아니라 도구 자리다). 파서는 환경을 읽지 않는다 — 값은 `core/paths.bind_parser()`가 주입한다
SOFFICE_ENV = "SOFFICE_PATH"
_GIVEN = None


def use_soffice_given(fn):
    """지정 자리 주입 — 부르는 쪽은 `core/paths.bind_parser()` 하나다(주입이 없으면 지정 없음 — 킷 · 단독 실행)."""
    global _GIVEN
    _GIVEN = fn


def _given():
    try:
        return str((_GIVEN() if _GIVEN else "") or "").strip()
    except Exception:                                       # noqa: BLE001 — 읽지 못하면 지정 없음
        return ""
#: **표준 설치 자리**(B129 ①) — 후보 목록 하나(OS 갈래 0 · 있는 것만 쓴다): macOS · Linux 고정 자리 · `/opt/libreoffice*` ·
#: Windows `%ProgramFiles%` · `%ProgramFiles(x86)%` 아래
SOFFICE_PLACES = ("/Applications/LibreOffice.app/Contents/MacOS/soffice", "/usr/lib/libreoffice/program/soffice",
                  "/usr/lib64/libreoffice/program/soffice", "/opt/libreoffice*/program/soffice",
                  "%ProgramFiles%/LibreOffice/program/soffice.exe", "%ProgramFiles(x86)%/LibreOffice/program/soffice.exe")
PATH_, PLACE, GIVEN = "PATH", "표준 자리", "지정"


def _places():
    import glob
    out = []
    for c in SOFFICE_PLACES:
        c = os.path.expandvars(c)
        if "%" in c:                                     # 그 환경변수가 없는 기계 — 그 후보는 없다
            continue
        out += sorted(glob.glob(c)) if "*" in c else [c]
    return out


def find_soffice():
    """**soffice 찾기 한 함수**(B129 ①) — `(경로 또는 None, 어디서 — 지정 | PATH | 표준 자리, 찾아본 자리 글)`.

    순서는 **지정(`SOFFICE_PATH`) → PATH(`soffice` · `libreoffice`) → 표준 설치 자리**다 — 사람이 지정했으면 그것이 이긴다
    (여러 판이 깔린 기계). 캡처 판정 · 변환 · 그림 형식 변환 · 환경 점검이 이 함수 하나를 부른다."""
    tried = []
    given = _given()
    if given:
        g = Path(given)
        cand = [g] if not g.is_dir() else [g / n for n in ("soffice", "soffice.exe", "soffice.bin")]
        hit = next((str(c) for c in cand if c.is_file() and os.access(c, os.X_OK)), None)
        if hit:
            return hit, GIVEN, None
        tried.append(f"지정 {SOFFICE_ENV}={given}(없음)")
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe, PATH_, None
    tried.append("PATH(soffice · libreoffice)")
    places = _places()
    for c in places:
        if os.path.isfile(c) and os.access(c, os.X_OK):
            return c, PLACE, None
    tried.append("표준 자리 " + (" · ".join(places) if places else "0곳"))
    if not given:
        tried.append(f"지정 {SOFFICE_ENV} 없음")
    return None, None, " · ".join(tried)


def _soffice():
    return find_soffice()[0]


def pymupdf():
    """PyMuPDF 모듈 — `import pymupdf`(지금 이름 · B128 ①) · 없으면 `fitz`(옛 판) · 둘 다 없으면 None. 렌더 · 리더가 이것 하나로 부른다."""
    try:
        import pymupdf as mod                               # noqa: F401
        return mod
    except ImportError:
        pass
    try:
        import fitz as mod                                  # 옛 판(1.24 전)
        return mod
    except ImportError:
        return None


#: soffice가 **PPTX를 변환할 수 있나** — Impress 구성 요소(프로그램 폴더의 라이브러리)가 있어야 한다(배포판은 나눠 깐다)
IMPRESS_LIBS = ("libsdlo.so", "libsdlo.dylib", "sdlo.dll")
#: 캡처 길 — 판정 함수(`capture_path`)의 답
PDF_DOC, BESIDE, SOFFICE, NONE = "PDF 문서", "옆 PDF", "soffice", "없음"
NEED_PYMUPDF = "PyMuPDF(pip install PyMuPDF)"
NEED_ROUTE = "LibreOffice(Impress 포함 — soffice) 또는 같은 폴더 · 같은 이름의 더 최신 PDF(PowerPoint에서 PDF로 저장)"


def _impress(exe):
    real = Path(os.path.realpath(exe)).parent
    return any((d / n).exists() for d in (real, real.parent / "Frameworks", real.parent / "program") for n in IMPRESS_LIBS)


def capture_path(path=None, probe=False):
    """**캡처 가능 판정 한 자리**(B128 ①) — `(가능?, 길, 빠진 것)` · 길 = PDF 문서 | 옆 PDF | soffice | 없음.

    PyMuPDF가 있고 「PDF 문서 · 같은 이름의 더 최신 PDF가 옆에 · soffice(Impress 포함)」 중 하나면 가능이다. 그림 요약의 쪽 렌더 ·
    추출 캡처 · 사전 점검 · 환경 점검이 이 함수를 부른다(렌더는 길을 그대로 따른다). `path`가 없으면 환경만(soffice 길) ·
    `probe`면 soffice로 작은 PPTX를 실제로 변환해 본다(환경 점검 — 몇 초 · 그 밖은 변환 0)."""
    if pymupdf() is None:
        return False, NONE, NEED_PYMUPDF
    if path is not None:
        p = str(path)
        if p.lower().endswith(".pdf"):
            return True, PDF_DOC, None
        if _pdf_beside(p):
            return True, BESIDE, None
    exe, _where, tried = find_soffice()
    if not exe:
        return False, NONE, f"{NEED_ROUTE} — 찾아본 자리: {tried}"
    if not _impress(exe):
        return False, NONE, f"LibreOffice Impress — soffice({exe})는 있는데 PPTX 변환 구성 요소가 없다 · 또는 옆 PDF"
    if probe:
        why = _probe(exe)
        if why:
            return False, NONE, f"soffice 변환 실패 — {why}"
    return True, SOFFICE, None


#: 캡처 한글 글꼴(B168 ④) — 상태 셋
FONT_OK, FONT_NONE, FONT_UNKNOWN = "있음", "없음", "확인 불가"
FONT_NEXT = "캡처의 한글이 깨져 판독이 배치만 본다 — 한글 글꼴(예: 나눔 · Noto CJK)을 깔고 fontconfig 대체 지정"


def korean_font():
    """**PPT 캡처의 한글 글꼴**(B168 ④) — `fc-match :lang=ko`가 고른 글꼴이 한국어를 덮는가 · `(상태, 글꼴 이름)`.
    `fc-match`가 없으면 확인 불가(soffice는 fontconfig로 글꼴을 고른다 — 없는 기계는 다른 길)."""
    import shutil
    import subprocess
    exe = shutil.which("fc-match")
    if not exe:
        return FONT_UNKNOWN, None
    try:
        r = subprocess.run([exe, ":lang=ko", "--format=%{family}\n%{lang}\n"], capture_output=True, text=True,
                           timeout=10, stdin=subprocess.DEVNULL)
    except (OSError, subprocess.SubprocessError):
        return FONT_UNKNOWN, None
    lines = (r.stdout or "").splitlines()
    fam = (lines[0].split(",")[0].strip() if lines else "") or None
    langs = set((lines[1] if len(lines) > 1 else "").split("|"))
    return (FONT_OK, fam) if fam and "ko" in langs else (FONT_NONE, fam)


def _probe(exe):
    """작은 PPTX 한 장을 실제로 변환해 본다 — 실패 사유 또는 None(python-pptx가 없으면 그 사유)."""
    try:
        from pptx import Presentation
    except ImportError:
        return "python-pptx 없음 — 변환을 재지 못했다(pip install python-pptx)"
    with tempfile.TemporaryDirectory(prefix="onto_probe_") as td:
        src = Path(td) / "probe.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "probe"
        prs.save(src)
        pdf = _soffice_pdf(src, td, exe=exe)
        if not pdf:
            return "PDF가 나오지 않았다(로그 참조)"
        doc = pymupdf().open(pdf)
        try:
            return None if doc.page_count == 1 else f"쪽 {doc.page_count} ≠ 1"
        finally:
            doc.close()


def _soffice_pdf(path, outdir, exe=None):
    exe = exe or _soffice()
    if not exe:
        return None
    try:
        subprocess.run([exe, "--headless", "--convert-to", "pdf",
                        "--outdir", str(outdir), str(path)],
                       check=True, capture_output=True, timeout=SOFFICE_TIMEOUT, stdin=subprocess.DEVNULL)
    except Exception as e:                                  # noqa: BLE001
        # **변환 실패는 문서 실패가 아니다** — 결함 한 줄로 남기고 계속한다.
        _LOG.warning("soffice 변환 실패 — %s: %s", type(e).__name__, e)
        return None
    out = Path(outdir) / (Path(path).stem + ".pdf")
    return str(out) if out.exists() else None


def _pdf_to_png(pdf):
    """PDF → {쪽번호: PNG 바이트}. PyMuPDF는 **지연 import**다(선택 의존)."""
    mod = pymupdf()
    if mod is None:
        _LOG.info("PyMuPDF 없음 — 쪽 렌더를 건너뛴다")
        return None
    try:
        doc = mod.open(pdf)
    except Exception as e:                                  # noqa: BLE001
        _LOG.warning("PDF 열기 실패 — %s: %s", type(e).__name__, e)
        return None
    out = {}
    try:
        for i, page in enumerate(doc, start=1):
            r = page.rect
            long_edge = max(r.width, r.height) or 1
            zoom = min(1.0, MAX_EDGE / long_edge) if long_edge > MAX_EDGE else 1.0
            pix = page.get_pixmap(matrix=mod.Matrix(zoom, zoom))
            out[i] = pix.tobytes("png")
    finally:
        doc.close()
    return out or None


def render_pages(path):
    """`{쪽번호: png 바이트}` 또는 `None` — 길은 `capture_path`가 정한다(싼 것부터: PDF 문서 · 옆 최신 PDF · soffice 변환)."""
    ok, route, _miss = capture_path(path)
    if not ok:
        return None
    p = str(path)
    if route == PDF_DOC:
        return _pdf_to_png(p)
    if route == BESIDE:
        return _pdf_to_png(_pdf_beside(p))
    with tempfile.TemporaryDirectory() as td:
        pdf = _soffice_pdf(p, td)
        return _pdf_to_png(pdf) if pdf else None


# ---------------------------------------------------------------- 그림 형식 변환 (B88 ①)
#: 모델이 받는 형식 — 나머지(WMF·EMF·BMP·TIFF …)는 PNG로 바꿔 보낸다.
MODEL_MIMES = ("image/png", "image/jpeg")
_EXT_OF = {"image/x-wmf": ".wmf", "image/x-emf": ".emf", "image/bmp": ".bmp",
           "image/tiff": ".tif", "image/gif": ".gif", "image/svg+xml": ".svg"}
NO_CONVERTER, CONVERT_FAILED = "변환기 없음", "변환 실패"


def to_png(data, mime):
    """그림 바이트 → `(PNG 바이트, None)` 또는 `(None, 사유)` — `soffice`가 있으면 쓴다.

    쪽 렌더와 같은 규칙이다: `soffice`는 **있으면 쓰는 외부 프로그램**이고, 없으면 조용히
    다르게 돌지 않고 **사유를 돌려준다**(호출부가 「건너뜀 — 변환기 없음」으로 화면·보고에 싣는다).
    바이트는 저장하지 않는다 — ④에 보내고 버린다(문서 5 §5.2-2).
    """
    exe = _soffice()                                        # 찾기 한 함수 (B129 ①)
    if not exe:
        return None, NO_CONVERTER
    ext = _EXT_OF.get(mime, ".bin")
    with tempfile.TemporaryDirectory(prefix="img_conv_") as td:
        src = Path(td) / f"image{ext}"
        src.write_bytes(data)
        out = Path(td) / "out"
        try:
            subprocess.run([exe, "--headless", "--convert-to", "png",
                            "--outdir", str(out), str(src)],
                           check=True, capture_output=True, timeout=SOFFICE_TIMEOUT)
        except Exception as e:                              # noqa: BLE001
            _LOG.warning("그림 변환 실패 — %s: %s", type(e).__name__, e)
            return None, CONVERT_FAILED
        png = out / "image.png"
        if not png.exists() or not png.stat().st_size:
            return None, CONVERT_FAILED
        return png.read_bytes(), None
