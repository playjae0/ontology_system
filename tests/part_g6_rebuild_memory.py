# -*- coding: utf-8 -*-
"""G6 ㉔ 재구축 기억 — 옛 사람 판단 기록에 이름 채우기 · 좌표 태깅 null 답 기억(닫힌 목록 판) · 재시도가 판정 대장을 갱신 (B107).

창작 표본(mock · 시험 주입)으로 메커니즘을 잰다 — 수치는 근거가 아니다([정정] 50). 표 하나 · 산문 하나가 층 둘에 걸친다:
표 B107CP(CP01 사본 · cp · process — 공정명 둘을 목록 밖으로: 극성 cathode 행 하나 · 극성 없는 행 하나) · 산문 B107Q(RFQ01
통합문서의 사본 · quality · 문서 좌표 노칭 — 집이 process인 개체와 그 끝점으로 가는 품질 그래프 엣지: 노드는 집 층 · 엣지는 문서 층).
층 · 형태 곱은 층 중립 자리라 대표 한 벌이다(B121 ②ⓒ — 층 호환은 층 어휘 정적 점검 · 네 벌이 잰다). 좌표 태깅 `pick`은 **가짜
실호출**이다(게이트웨이 사용량에 센다 · 답은 늘 null) — mock 태깅은 정확 일치만이라 이 축이 비어 있다. 산문 조각은 실제 파싱에서
좌표가 정확 일치(제목)·문서 좌표로만 서서 목록 밖이 생기지 않는다 — 산문 몫은 같은 함수(`tagger.tag` · `run_document`)에 조각을 넣는다.

잠그는 성질:
  ⓐ 이름 채움: B106 전 모양의 기록(이름 없음 · id만 — merge(canonical 지정 포함) · confirm(뒤에 개명) · alias · obsolete(대체 포함) ·
     edge 삭제 · 표 · 산문 · 층 둘)을 옛 자리에 → 첫 실행(옮김)이 채운 이름 = B106이 연산 때 적은 이름(출처 `backfill`) / ②에 옮겨진
     이름 없는 기록 → `ops replay --dry-run`(그래프 쓰기 0)과 `rebuild` 계획(쓰기 0)이 채운다 · 그래프에 없는 id는 채우지 않고
     「이름 없음」 목록 / 그 뒤 재구축 ⇒ 연산 전부 재생(「기록에 이름 없음」 0) · canonical 기준 상태 = 재구축 전
  ⓑ null 기억: 목록 밖 표기 null → 같은 조각 다시 ⇒ 묻지 않는다(「목록 밖(기억)」 · 표 · 산문 한 벌씩) / 지우는 손잡이(`ops learn-reject`
     · --actor · 기록) ⇒ 그 표기만 / 골격 별칭 하나(판 문자열 그대로) ⇒ 다시 묻는다 / 실호출 모양의 재구축 ⇒ 끝 보고 「태깅 0」
  ⓒ 재시도 대장: 좌표가 목록 밖이라 orphan이던 값 → 별칭 + bootstrap → 재시도가 붙임(판정 호출 > 0 · 인입과 같은 이름 — 극성 하강) →
     대장 행이 그 판정으로(`by: retry` · 표 · 산문 × 층 둘) → 재구축 ⇒ 그 값 재생(「키 없음」 0 · 「대상 없음」 0 · 판정 호출 0)
"""
from __future__ import annotations

import contextlib
import io
import re
import tempfile
from pathlib import Path

sys_path = __import__("sys").path
sys_path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done   # noqa: F401
from g6_common import RFQ, _register, _run, _unregister   # noqa: E402 — 산문 엑셀 표본 · 등록 · CLI

import run as RUN                                                  # noqa: E402
from cli import ingest as IG, parse as PA, rebuild as RB           # noqa: E402
from core.build import ledger as LG, replay as RP                  # noqa: E402
from core.build.entry import finalize                              # noqa: E402
from core.llm import gateway, points                               # noqa: E402
from core.state import oplog                                       # noqa: E402
from core.state.bootstrap import coord_layer                       # noqa: E402
from core.state.status import is_live                              # noqa: E402
from core.state.world import World                                 # noqa: E402
from openpyxl import load_workbook                                 # noqa: E402
from parser import tagger as TG                                    # noqa: E402

FX = ROOT / "tests" / "fixtures"
RAW = FX / "raw"
from core.state import fixtures as _FXR                   # noqa: E402 — 픽스처 뿌리(회귀는 스위트마다 사본 — B133)
HINTS = _FXR.EXTRACT_HINTS
A = "시험자"
X = "노칭"                                         # 문서 좌표 — 산문 둘
_hinted = []

# ── 가짜 게이트웨이 — 판정 · 추출은 mock 호출 수(`gateway.mock`) · 좌표 태깅은 실호출 모양(사용량에 센다 · 답 null)
CALLS, PICKED = {}, []
_mock0 = gateway.mock


def _counted(point, detail=""):
    CALLS[point] = CALLS.get(point, 0) + 1
    return _mock0(point, detail)


def _fake_pick(surface, choices):
    gateway._account("coord_tag", {"usage": {"prompt_tokens": 8, "completion_tokens": 2, "total_tokens": 10}})
    PICKED.append(surface)
    return None


gateway.mock = _counted
points.coord_picker = lambda: _fake_pick           # 인입 · 재구축 · 리허설이 같은 주입(`cli/parse.injections`)


def _quiet(f, *a, **k):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        r = f(*a, **k)
    return r, buf.getvalue()


def _hint(doc_id, by_loc):
    (HINTS / f"{doc_id}.json").write_text(json.dumps(by_loc, ensure_ascii=False), encoding="utf-8")
    _hinted.append(doc_id)


def _snap():
    """canonical 기준 상태 — 노드(카테고리 · status · 값) · 별칭(출처) · 엣지(status) · 큐(id는 이름으로 · 시각 뺌) + 결정."""
    w = World()
    nm = {n["id"]: f"{lay}:{n['canonical']}" for lay, n in w.nodes()}

    def canon(x):
        if isinstance(x, dict):
            return {k: canon(v) for k, v in x.items() if k not in ("created", "at", "first_seen", "last_seen")}
        if isinstance(x, list):
            return [canon(v) for v in x]
        return nm.get(x, x) if isinstance(x, str) else x
    return {"노드": sorted(f"{lay}:{n['canonical']}|{n['category']}|{n['status']}|"
                          + json.dumps(n.get("attrs") or {}, ensure_ascii=False, sort_keys=True)
                          for lay, n in w.nodes()),
            "별칭": sorted(f"{lay}:{n['canonical']}|{a['surface']}|{a.get('by')}" for lay, n in w.nodes()
                          for a in n.get("aliases") or []),
            "엣지": sorted(f"{e['rel']}|{nm.get(e['src'], e['src'])}|{nm.get(e['dst'], e['dst'])}|{e.get('status')}"
                          for g in w.graphs.values() for e in g.edges),
            "큐": sorted(json.dumps([x["kind"], x.get("doc_id"), canon(x.get("payload")),
                                    (x.get("resolution") or {}).get("decision")], ensure_ascii=False, sort_keys=True)
                        for x in store.read(store.QUEUE, []))}


def _diff(a, b):
    return {k: (sorted(set(a[k]) - set(b[k]))[:3], sorted(set(b[k]) - set(a[k]))[:3]) for k in a if a[k] != b[k]}


def _data_hash():
    """③진실의 바이트 지문(락 제외) — 계획 · dry-run이 그래프를 건드리지 않았나."""
    import hashlib
    d = _P.home() / "data"
    return hashlib.sha256(b"".join(p.read_bytes() for p in sorted(d.rglob("*"))
                                   if p.is_file() and not p.name.endswith(".lock"))).hexdigest()


def _end(out):
    """재구축 끝 보고의 수 — 판정(재생 · 재판정 · 사유) · LLM 호출(추출 · 판정 · 태깅)."""
    m = re.search(r"판정 — 재생 ([\d,]+) · 재판정 ([\d,]+)\(키 없음 ([\d,]+) · 대상 없음 ([\d,]+) · 끔 ([\d,]+)\)", out)
    c = re.search(r"LLM 호출 — 추출 ([\d,]+) · 판정 ([\d,]+) · 태깅 ([\d,]+)", out)
    keys = ("재생", "재판정", "키 없음", "대상 없음", "끔", "추출", "판정", "태깅")
    vals = (m.groups() if m else ()) + (c.groups() if c else ())
    return dict(zip(keys, (int(v.replace(",", "")) for v in vals))) if m and c else {}


def _rebuild():
    CALLS.clear()
    PICKED.clear()
    rc, out = _quiet(RB.main, ["--allow-mock", "--yes"], bootstrap=RUN.cmd_bootstrap, screen=RUN.fresh_screen)
    return rc, out


def _end_lines(out):
    ls = out.splitlines()
    i = next((k for k, ln in enumerate(ls) if ln.startswith("■ 재구축 끝")), len(ls))
    return ls[i:]


# ────────────────────────────────────────────────────────────── 바닥 — 표 · 산문 · 층 둘
print("\n■ B107 바닥 — 표 · 산문 · 층 둘 (B107CP · B107Q) · 좌표 태깅은 가짜 실호출(답 null)")
with contextlib.redirect_stdout(io.StringIO()):
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
QDT = "b107pq"
_register("quality", QDT)
TMP = Path(tempfile.mkdtemp(prefix="b107_"))          # 클린(init --fresh)이 지우지 않는 자리
QX = TMP / "B107Q.xlsx"
shutil.copy(RFQ, QX)
CPX = TMP / "B107CP.xlsx"                             # CP01 사본 — 공정명 둘을 목록 밖으로(표 × process의 목록 밖 좌표)
MISS_CP = ("B107 미지 공정 가", "B107 미지 공정 나")
_wb = load_workbook(RAW / "CP01.xlsx")
_ws = _wb.worksheets[0]
#: cathode 탭용접 첫 행(극성 cathode — 재시도의 극성 하강 · 병합 C23:C24를 풀어 그 한 행만) · 사이드 실링 행. 원본 CP01 ·
#: PFMEA01은 넣지 않는다 — 사람 연산의 표 노드는 이 사본이 세운다(B121 ②ⓒ)
_ws.unmerge_cells("C23:C24")
_ws["C24"] = _ws["C23"].value
_ws["C23"], _ws["C9"] = MISS_CP
_wb.save(CPX)
#: 품질 산문 — 집이 process인 개체(「가압력」 · 「금형 간극」 · 「버 높이」 — 새 노드)와 그 끝점으로 가는 품질 그래프 엣지(`controlled_by`)
_hint("B107Q", {
    "사양_기계!R8-R17": {"entities": [{"surface": "가압력 부족", "category": "Failure"},
                                      {"surface": "금형 클리어런스 과다", "category": "Failure"},
                                      {"surface": "가압력", "category": "Property"},
                                      {"surface": "금형 간극", "category": "Property"}],
                        "relations": [{"src": "금형 클리어런스 과다", "rel": "causes", "dst": "가압력 부족"},
                                      {"src": "가압력 부족", "rel": "controlled_by", "dst": "가압력"}]},
    "사양_기계!R19-R23": {"entities": [{"surface": "버 발생량 과다", "category": "Failure"},
                                       {"surface": "버 높이", "category": "Property"}], "relations": []},
    "사양_전장!R5-R6": {"entities": [{"surface": "전원 차단 실패", "category": "Failure"}], "relations": []}})
SHEETS = "2-4:prose 5:ref *:skip"
DOCS = [("B107CP", CPX, "cp", None, None), ("B107Q", QX, QDT, SHEETS, X)]
_first = {}
for did, f, dt, sh, co in DOCS:
    _r, _o = _quiet(IG.ingest_file, str(f), dt, ask=False, sheets=sh, coord=co)
    _first[did] = _r.get("status")
P_FIRST = sorted(set(PICKED))
print(f"  첫 인입 — {_first} · 좌표 태깅 가짜 실호출 {len(PICKED)}회 {P_FIRST} · mock 판정·추출 {dict(CALLS)}")

# ── 사람 연산 — 종류마다 하나 이상(표 · 산문 · 층 둘) · merge는 canonical 지정 하나 포함 · confirm 뒤 개명(그때의 이름)
_by = {(lay, n["canonical"]): n["id"] for lay, n in World().nodes() if is_live(n)}


def _id(lay, canon):
    return _by[(lay, canon)]


def _live_edge(src, rel=None, dst=None):
    return next((e for g in World().graphs.values() for e in g.edges if e["src"] == src
                 and e.get("status") != "deleted_by_user" and rel in (None, e["rel"]) and dst in (None, e["dst"])), None)


with contextlib.redirect_stdout(io.StringIO()):
    ops.confirm("process", _id("process", "노칭::노칭 프레스"), A, "B107 확인(표)")
    _k, _g = _id("process", "노칭::노칭 프레스::금형 클리어런스"), _id("process", "노칭::금형 간극")   # 기대 변경(B144 ② — 표의 관리항목은 같은 행의 설비 아래 이름)
    ops.merge("process", _g, _k, A, canonical="노칭::금형 간극", override=_k,
              reason="B107 합침(품질 문서의 노드 → 표 · 이름은 없어지는 쪽)")
    ops.merge("process", _id("process", "노칭::버 높이"), _id("process", "노칭::anode::노칭 프레스::버 높이"), A,
              reason="B107 합침(산문 → 표)")
    ops.alias("process", _id("process", "노칭::노칭 프레스::노칭 정밀도"), "노칭 정확도", A, "B107 표기(표)")
    ops.alias("quality", _id("quality", "전원 차단 실패"), "전원 차단 불량", A, "B107 표기(산문)")
    ops.obsolete("process", _id("process", "스태킹::스태커"), A, reason="B107 폐기(표)")
    _e1 = _live_edge(_id("process", "패키징::사이드 실링::실러"))
    ops.delete_edge("process", _e1["src"], _e1["rel"], _e1["dst"], A, "B107 엣지 삭제(표)")
    _e2 = _live_edge(_id("quality", "가압력 부족"), "controlled_by", _id("process", "노칭::가압력"))   # 품질 그래프 · 끝점은 process
    ops.delete_edge("quality", _e2["src"], _e2["rel"], _e2["dst"], A, "B107 엣지 삭제(산문 · 층 건너는 엣지)")
    ops.confirm("quality", _id("quality", "가압력 부족"), A, "B107 확인(산문)")
    ops.obsolete("quality", _id("quality", "금형 클리어런스 과다"), A, _id("quality", "가압력 부족"), "B107 폐기(산문 · 대체)")
    ops.rename("process", _id("process", "노칭::노칭 프레스"), "노칭::노칭 프레스 장비", A, "B107 개명(확인 뒤)")
S1 = _snap()
RECS = oplog.read()
ORIG = [r.get("names") for r in RECS]
OLD = [{k: v for k, v in r.items() if k != "names"} for r in RECS]         # B106 전 모양 — 이름 없음 · id만
FAKE = {"op": "I5:confirm", "actor": "옛사람", "at": "2026-01-01T00:00:00+00:00",
        "targets": ["01ZZZZZZZZZZZZZZZZZZZZZZZZ"], "reason": "그래프에 없는 id", "detail": {"from": "auto"}}
print(f"  사람 연산 {len(RECS)}건 — " + " · ".join(r["op"] for r in RECS))


def _filled():
    """채운 이름이 B106이 연산 때 적은 이름과 같은가 · 출처 표시 · 그래프에 없는 id는 이름 없음 그대로."""
    cur = oplog.read()
    same = [c.get("names") == o for c, o in zip(cur, ORIG)]
    marks = [c.get("names_from") == oplog.NAMES_FROM for c in cur[:len(ORIG)]]
    fake = [c for c in cur[len(ORIG):]]
    return all(same) and all(marks) and all(not c.get("names") for c in fake), sum(same), len(ORIG), cur


# ────────────────────────────────────────────────────────────── ⓐ 이름 채움
print("\n■ B107 ① 옛 사람 판단 기록에 이름 채우기 — 옮길 때 · ops replay(계획) · 재구축 계획 (표 · 산문 · 층 둘)")
_old = _P.data(store.OPS_LOG)
_old.write_text(json.dumps(OLD + [FAKE], ensure_ascii=False, indent=1), encoding="utf-8")
store.path(store.OPS_LOG).unlink()
_o1 = _run("show", "learned")                          # 아무 명령 — 공통 진입이 옛 자리를 옮긴다
_ok1, _n1, _all1, _ = _filled()
show("ⓐ 옛 자리(③ data/ops_log.json)의 이름 없는 기록 → 첫 실행이 ②등록으로 옮기며 이름을 채운다 · 채운 이름 = B106이 연산 때 "
     "적은 이름(merge canonical 지정 · confirm 뒤 개명 포함) · 출처 표시 · 그래프에 없는 id는 이름 없음 그대로 · 옛 자리 없음",
     _ok1 and not _old.exists() and "이름 채움" in _o1,
     f"같은 이름 {_n1}/{_all1} · 옛 자리 {'남음' if _old.exists() else '없음'}")
for ln in _o1.splitlines():
    if ln.startswith("[사람 판단 기록]"):
        print(f"      {ln}")
store.write(store.OPS_LOG, OLD + [FAKE])               # 이미 ②에 옮겨진 이름 없는 기록(B106 반입 뒤의 사내)
_h0 = _data_hash()
_o2 = _run("ops", "replay", "all", "--actor", A, "--dry-run")
_ok2, _n2, _all2, _ = _filled()
show("ⓐ ②에 옮겨진 이름 없는 기록 → `ops replay --dry-run`이 시작에서 채운다(그래프 쓰기 0 · 채운 이름은 기록에)",
     _ok2 and _h0 == _data_hash() and "이름 채움" in _o2,
     f"같은 이름 {_n2}/{_all2} · ③ 쓰기 {'0' if _h0 == _data_hash() else '있음'}")
for ln in _o2.splitlines():
    if ln.strip() and not ln.startswith("[실행"):
        print(f"      {ln}")
store.write(store.OPS_LOG, OLD + [FAKE])
_o3 = _run("rebuild", "--allow-mock")                  # 비대화형 — 계획만
_ok3, _n3, _all3, _cur3 = _filled()
_plan = [ln for ln in _o3.splitlines() if "사람 판단 기록" in ln or "이름 없음 —" in ln]
show("ⓐ 재구축 계획(fresh 앞)이 채운다 · 그래프에 없는 id는 채우지 않고 계획에 「이름 없음」 목록(연산 · id · 행위자 · 시각) · "
     "비대화형은 계획만(③ 쓰기 0)",
     _ok3 and _h0 == _data_hash() and any(FAKE["targets"][0] in ln and "옛사람" in ln for ln in _plan)
     and "(비대화형 — 계획만 · 쓰기 0)" in _o3,
     f"같은 이름 {_n3}/{_all3} · 이름 없음 {sum(1 for c in _cur3 if not c.get('names'))}")
for ln in _plan:
    print(f"      {ln.strip()}")
store.write(store.OPS_LOG, OLD)                        # 그래프에 없는 id 기록은 빼고 — 재구축이 전부 재생하는가
_rc, OUT_A = _rebuild()
S2, _ops = _snap(), [x for x in _end_lines(OUT_A) if "사람 판단" in x]
_endA = _end(OUT_A)
_rep = re.search(r"사람 판단 재생 — 한 일 ([\d,]+) · 이미 그 상태 ([\d,]+) · 대상 없음 ([\d,]+) · 기록에 이름 없음 ([\d,]+)",
                 OUT_A)
show("ⓐ 그 뒤 재구축 ⇒ 연산 전부 재생(「기록에 이름 없음」 0 · 대상 없음 0) · canonical 기준 상태(노드 · 별칭 · 엣지 · 큐) = "
     "재구축 전",
     _rc == 0 and _rep and int(_rep[1]) == len(OLD) and int(_rep[3]) == 0 and int(_rep[4]) == 0 and not _diff(S1, S2),
     " · ".join(f"{k} {len(v)}" for k, v in S1.items()) + (f" · 다른 것 {_diff(S1, S2)}" if _diff(S1, S2) else " · 같음"))
print("  ── 재구축 계획의 사람 판단 줄 · 끝 보고(잘리지 않은 줄)")
for ln in [x for x in OUT_A.splitlines() if x.startswith("  사람 판단 기록")] + _end_lines(OUT_A):
    print(f"      {ln}")
show("ⓑ 실호출 모양의 재구축 ⇒ 끝 보고 「태깅 0」(첫 인입은 목록 밖 표기마다 한 번 물었다 — 같은 판이라 기억) · 판정·추출 0",
     _endA.get("태깅") == 0 and not PICKED and set(MISS_CP) <= set(P_FIRST) and _endA.get("판정") == 0
     and CALLS.get("judge", 0) == CALLS.get("extract", 0) == 0,
     f"첫 인입 태깅 {len(P_FIRST)}종 {P_FIRST} → 재구축 태깅 {_endA.get('태깅')}(가짜 실호출 {len(PICKED)}) · 끝 줄 {_endA}")
for ln in OUT_A.splitlines():
    if "목록 밖(기억)" in ln:
        print(f"      {ln.strip()}")

# ────────────────────────────────────────────────────────────── ⓑ null 기억 — 같은 함수(tag) · 표·산문 × 층 둘
print("\n■ B107 ② 좌표 태깅 null 답 기억 — 닫힌 목록 판(내용 해시) · 표 × process · 산문 × quality")
with contextlib.redirect_stdout(io.StringIO()):
    RUN.cmd_bootstrap([])                              # 지금 그래프로 스냅샷을 다시 쓴다(판의 출발점)
LAY = coord_layer()
COMBOS = [("표", "process"), ("산문", "quality")]     # 형태 × 층 한 벌씩 — 태깅 기억은 층 중립(B121 ②ⓒ)


def _pieces(form, lay):
    """조각 셋 — 목록 밖 표기 둘(조합마다 다른 표기) + 정확 일치 하나. 표는 레코드 · 산문은 청크(`text`)."""
    tag = f"B107 {form}{lay[0]}"
    refs = [f"{tag} 미지 가", f"{tag} 미지 나", "노칭"]
    if form == "표":
        return [dict(CPREC, source_locator=f"{tag}-{i}", process_ref=r, process_group=None, meta={})
                for i, r in enumerate(refs)], f"b107t{lay[0]}", refs[:2]
    return [dict(C1, source_locator=f"{tag}-{i}", process_ref=r, process_group=None, meta={},
                 text=f"{r} — B107 조각") for i, r in enumerate(refs)], f"ppt_{lay}", refs[:2]


def _tag(form, lay):
    """같은 조각을 같은 함수로 — `(호출 수, 예고 그릇, 화면 머리 줄)`."""
    pcs, dt, _m = _pieces(form, lay)
    plan, n0 = {}, len(PICKED)
    scr, _prog = PA.coord_screen()

    def note(info):
        if info.get("단계") == "예고":
            plan.update(info)
        scr(info)
    _r, out = _quiet(TG.tag, pcs, layer=LAY, nodes=TG.closed_list(LAY), pick=_fake_pick, doc_type=dt,
                     doc_id=f"B107{form}{lay[0]}", notice=note)
    return len(PICKED) - n0, plan, next((ln.strip() for ln in out.splitlines() if "좌표 태깅 —" in ln), "")


_first_b = {c: _tag(*c) for c in COMBOS}
_ver_a = TG.list_version(TG.closed_list(LAY))
_book = TG.learned()["목록밖"]
_again = {c: _tag(*c) for c in COMBOS}
show("ⓑ 목록 밖 표기 2종이 null → 같은 조각 다시(같은 함수 `tag`) ⇒ 태깅 호출 0 · 「목록 밖(기억) 2종」 · 묻는 종수 0 "
     "(표 · 산문 · 기록의 판 = 닫힌 목록 내용 해시)",
     all(_first_b[c][0] == 2 and _again[c][0] == 0 and _again[c][1].get("기억_종수") == 2
         and _again[c][1].get("묻는_종수") == 0 for c in COMBOS)
     and all(_book.get(s, {}).get("list_version") == _ver_a for c in COMBOS for s in _pieces(*c)[2]),
     " ‖ ".join(f"{f}×{l} 처음 {_first_b[(f, l)][0]} → 다시 {_again[(f, l)][0]}" for f, l in COMBOS) + f" · 판 {_ver_a}")
for c in COMBOS:
    print(f"      {c[0]}×{c[1]}  {_again[c][2]}")
_forget = _pieces("표", "process")[2][0]
_o4 = _run("ops", "learn-reject", LAY, _forget, "--actor", A)
_o4b = _run("ops", "learn-reject", LAY, _forget)                 # --actor 없이는 거부(인자 오류)
_rej = [r for r in oplog.read() if r.get("op") == "learn:reject" and (r.get("detail") or {}).get("surface") == _forget]
_after = {c: _tag(*c) for c in COMBOS}
show("ⓑ 지우는 손잡이 `ops learn-reject <층> <표기> --actor`(기록에 남는다 · --actor 없으면 거부) ⇒ 그 표기 하나만 다시 묻는다 · "
     "새 공정 후보에는 그대로",
     _after[("표", "process")][0] == 1 and all(_after[c][0] == 0 for c in COMBOS if c != ("표", "process"))
     and "목록 밖 기억" in _o4 and "--actor" in _o4b and len(_rej) == 1
     and _forget in TG.learned()["목록밖"],
     f"「{next((ln.strip() for ln in _o4.splitlines() if ln.startswith('[거부]')), _o4.strip()[:120])}」 · 다시 묻기 "
     + " · ".join(f"{f}×{l} {_after[(f, l)][0]}" for f, l in COMBOS))
SK = _P.layers(LAY, "skeleton.json")
_SK0 = SK.read_bytes()
_sv0 = (store.read(store.SKELETON_LIST, {}).get(LAY) or {}).get("skeleton_version")
_sk = json.loads(_SK0.decode("utf-8"))
_sk.setdefault("ALIASES", {}).setdefault("노칭", []).append("B107 판 별칭")
SK.write_text(json.dumps(_sk, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
with contextlib.redirect_stdout(io.StringIO()):
    RUN.cmd_bootstrap([])
_sv1 = (store.read(store.SKELETON_LIST, {}).get(LAY) or {}).get("skeleton_version")
_ver_b = TG.list_version(TG.closed_list(LAY))
_alias = {c: _tag(*c) for c in COMBOS}
show("ⓑ 골격에 별칭 하나 더함(판 문자열 `skeleton_version`은 그대로) ⇒ 판(내용 해시)이 바뀌어 다시 묻는다(호출 > 0 · 표 · 산문)",
     _sv0 == _sv1 and _ver_a != _ver_b and all(_alias[c][0] == 2 for c in COMBOS),
     f"판 문자열 {_sv0} → {_sv1} · 판 {_ver_a} → {_ver_b} · 다시 묻기 "
     + " · ".join(f"{f}×{l} {_alias[(f, l)][0]}" for f, l in COMBOS))
SK.write_bytes(_SK0)

# ────────────────────────────────────────────────────────────── ⓒ 재시도 대장
print("\n■ B107 ③ 재시도가 판정 대장을 갱신 — 같은 키 · by retry · 인입과 같은 이름(표 · 산문 × 층 둘)")
PMISS = "B107 산문 미지 공정"


def _prose_env(lay):
    """산문 조각 — 좌표가 목록 밖(`PMISS`) · 개체는 이름 규칙 카테고리(좌표가 없으면 만들지 않는다 — 재료는 큐가 든다)."""
    did = f"B107P{lay[0].upper()}"
    loc = f"{did}-C1"
    _hint(did, {loc: {"entities": [{"surface": f"B107 산문 속성 {lay[0]}", "category": "Property"}],
                      "relations": [], "attach": []}})
    ch = dict(C1, source_locator=loc, doc_type=f"ppt_{lay}", process_ref=PMISS, process_group=None, meta={},
              text=f"{PMISS}의 B107 산문 속성 {lay[0]}을 관리한다.")
    return {**PROSE, "doc_id": did, "doc_type": f"ppt_{lay}", "source_path": f"{did}.pptx", "chunks": [ch]}


PENVS = {lay: _prose_env(lay) for lay in ("process", "quality")}
for _env in PENVS.values():
    with contextlib.redirect_stdout(io.StringIO()):
        run_document(_env)


def _rows(did, pred):
    return [r for r in (LG.read(did) or {}).get("rows") or [] if r.get("role") == "entity" and pred(r)]


ALIASED = (MISS_CP[0], PMISS)                        # 아래에서 별칭을 줄 좌표 표기 — 「B107 미지 공정 나」는 그대로 목록 밖


def _held(did):
    """별칭을 줄 좌표가 선 행(대장 anchor 행)의 보류 개체 행 — 재시도가 갈아 끼울 대상."""
    rows = (LG.read(did) or {}).get("rows") or []
    locs = {r.get("locator") for r in rows if r.get("role") == "anchor" and r.get("surface") in ALIASED}
    return [r for r in rows if r.get("role") == "entity" and r.get("locator") in locs
            and r.get("verdict") in ("orphan", "pending")]


_before = {d: _held(d) for d in ("B107CP", *[e["doc_id"] for e in PENVS.values()])}
_sk = json.loads(_SK0.decode("utf-8"))
_al = _sk.setdefault("ALIASES", {})
_al.setdefault("노칭", []).append(ALIASED[0])                  # 표 × process(B107CP) — 극성 cathode 행(그 이름 노드는 아직 없다)
_al.setdefault("탭용접", []).append(ALIASED[1])                # 산문 × 층 둘
SK.write_text(json.dumps(_sk, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
with contextlib.redirect_stdout(io.StringIO()):
    RUN.cmd_bootstrap([])
CALLS.clear()
_fin, _ = _quiet(finalize)
_judge_retry = CALLS.get("judge", 0)
_after_c = {d: _rows(d, lambda r: r.get("by") == "retry") for d in _before}
_keyed = all({(r["locator"], r["field"], r["surface"]) for r in _before[d]}
             <= {(r["locator"], r["field"], r["surface"]) for r in _after_c[d]} for d in _before)
_bf = _after_c["B107CP"]
show("ⓒ 좌표가 목록 밖이라 orphan이던 값 → 별칭 + bootstrap → 재시도가 붙인다(판정 호출 > 0) · 그 문서 대장의 같은 키 행이 그 판정으로 "
     "갈린다(`by: retry` · 경로 · 결과 · 대상 canonical — 표 · 산문 × 층 둘) · 인입과 같은 이름(극성 하강)",
     _judge_retry > 0 and _keyed and all(_before[d] for d in _before)
     and all(r["verdict"] in ("match", "new", "uncertain") and r.get("target") for d in _after_c for r in _after_c[d])
     and any(str(r.get("target")).startswith("노칭::cathode::") for r in _bf),
     f"판정 호출 {_judge_retry} · {_fin.get('retry_line')} ‖ "
     + " · ".join(f"{d} {len(_before[d])}행 → " + ", ".join(f"{r['surface']}({r['field']}) {r['verdict']}·{r['path']}→{r.get('target')}"
                                                         for r in _after_c[d]) for d in _after_c))
_rc, OUT_C = _rebuild()
_endC = _end(OUT_C)
_rp = {d: _rows(d, lambda r: (r["locator"], r["field"], r["surface"]) in
                {(x["locator"], x["field"], x["surface"]) for x in _after_c[d]}) for d in ("B107CP",)}
show("ⓒ 재구축 ⇒ 재시도가 붙인 값이 재생된다(그 행 경로 replay · 끝 보고 「키 없음」 0 · 「대상 없음」 0 · 판정 호출 0) — 표",
     _rc == 0 and _endC.get("키 없음") == 0 and _endC.get("대상 없음") == 0 and _endC.get("판정") == 0
     # 기대 변경(B144 ② — 새 설비 아래 첫 값은 후보 0으로 끝난 새 노드라 재생하지 않고 다시 잰다(B142 · LLM 0) · 나머지는 재생)
     and CALLS.get("judge", 0) == 0 and all(_rp[d] and all(r.get("path") in ("replay", "none") for r in _rp[d])
                                            and any(r.get("path") == "replay" for r in _rp[d]) for d in _rp),
     f"끝 줄 {_endC} · mock 판정 {CALLS.get('judge', 0)} · " + " · ".join(
         f"{d} " + ", ".join(f"{r['surface']} {r.get('path')}" for r in _rp[d]) for d in _rp))
for ln in _end_lines(OUT_C)[:3]:
    print(f"      {ln}")
RP.reset()                                           # 산문 조각은 원본 파일이 없다 — 재구축과 같은 함수(`run_document`)로 다시
CALLS.clear()
for _env in PENVS.values():
    with contextlib.redirect_stdout(io.StringIO()):
        run_document(_env)
_rpp = {e["doc_id"]: _rows(e["doc_id"], lambda r: r["surface"].startswith("B107 산문 속성")) for e in PENVS.values()}
show("ⓒ 산문(층 둘) — 같은 함수로 다시 넣으면 재시도가 적은 행이 재생된다(키 없음 0 · 판정 호출 0)",
     RP.STATS["키 없음"] == 0 and RP.STATS["대상 없음"] == 0 and CALLS.get("judge", 0) == 0
     and all(rs and all(r.get("path") == "replay" for r in rs) for rs in _rpp.values()),
     f"재생 {RP.STATS['재생']} · 키 없음 {RP.STATS['키 없음']} · 대상 없음 {RP.STATS['대상 없음']} · mock 판정 {CALLS.get('judge', 0)} · "
     + " · ".join(f"{d} " + ", ".join(f"{r['surface']}→{r.get('target')} {r.get('path')}" for r in rs) for d, rs in _rpp.items()))

# ── 뒷정리 — 시험이 세운 것은 시험이 치운다
for _d in _hinted:
    (HINTS / f"{_d}.json").unlink(missing_ok=True)
_unregister(QDT)
SK.write_bytes(_SK0)
shutil.rmtree(TMP, ignore_errors=True)
with contextlib.redirect_stdout(io.StringIO()):
    init.init(fresh_=True, all_=True)
done()
