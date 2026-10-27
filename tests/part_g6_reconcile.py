# -*- coding: utf-8 -*-
"""G6 ㉓ 맞춤표 · 시트 숨김 · 청크 목차 · 소속 진단 · 문서 좌표 경고 · 체크포인트 자리 (B112).

창작 표본(mock · 시험 주입)으로 메커니즘을 잰다 — 수치는 근거가 아니다([정정] 50). 화면은 장부의 투영이라(새 판정 0 ·
저장 0) 시험은 **화면의 식을 다시 더하고**, 같은 수를 장부(대장 · 결과 묶음 · 체크포인트 · 그래프)에서 따로 센다.
맞춤표는 층 중립 자리라 형태마다 대표 하나(표 CP01 · 산문 RFQ01)이고 층 둘은 소속 진단(ⓓ)이 잰다(B121 ②ⓒ — 층 호환은 층
어휘 정적 점검 · 네 벌이 잰다).

잠그는 성질:
  ⓐ 맞춤: 식마다 칸 합 = 전체 — CP01(표) · RFQ01(산문) 인입 · `show report` · 재구축(재생) · 값 전체 = 대장 개체 행 ·
     보류 전체 = 결과 묶음의 보류(보류가 있는 묶음 — ⓔ의 인입 화면) · 명령 로그에도 · 실행 사실과 체크포인트가
     어긋난 묶음은 `[결함]`
  ⓑ 시트: 숨김 · 아주 숨김이 표에 · 기본 제안 skip · 숨김 시트는 LLM을 부르지 않는다 · `--sheets`로 살림 · 파싱 끝
     시트별 조각 · 기록이 있어도 표(묻지 않는다 · 다시 고르는 명령줄) · 기록 prose인데 숨김이면 그 줄에 표시
  ⓒ 청크: 목차 행 수 = 청크 수 · 상태 합 = 청크 맞춤 줄 · 라벨 — 엑셀 `시트 › 경로` · PPT `슬라이드 n` · 개체 0 목록 =
     체크포인트의 개체 0 청크
  ⓓ 소속: 소속 대상 합 = 소속을 낸 개체 수 · 지어낸 이름은 버림 · 엣지 없는 노드 목록 = 그래프에서 센 것 ·
     대장 `attached.dir` = 실제 엣지 방향 · `--trace`의 소속 줄 (층 둘)
  ⓔ 문서 좌표: 띄어쓰기로 어긋나는 파일명 — 경고 줄 · 후보(보여만 · 기록 0) · 경고의 행 = 결과 묶음의 「좌표 미해소」
  ⓕ 체크포인트: 추출 끝에 자리 · 있음 · 쓰기가 빠진 상태에서 `[결함]`(화면 · 결함 장부)
  B149 진단 — ⓐ 문서 좌표 없음(원인 하나 · 합 = 보류 · 고아 행) ⓑ 구간 좌표 없음(다른 시트의 가장 위 제목) ⓒ 소속 이름 없음(= 고아 큐
     항목 수) ⓓ 목차 나무(문제 가지만 · 보류 수 = 대장 · -v 전부) ⓔ 표(골격 밖 좌표 표기 한 줄 · 좌표 표기 목록) ⓕ 인입 끝 = 같은 함수
     ⓖ 읽기 전용(상태 루트 바이트 같음 · LLM 0)
"""
from __future__ import annotations

import contextlib
import io
import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, _rw, done, overlay   # noqa: F401
from g6_common import _register, _unregister, _run, DT, RFQ   # noqa: E402

from openpyxl import Workbook                           # noqa: E402
from cli import ingest as IG, ingest_screen as IS, result_screen as RS, sheet_gate as SG   # noqa: E402
from core.build import extract as EX, ledger as LG, reconcile as RC   # noqa: E402
from core.state.world import World                      # noqa: E402
from parser import pipeline, struct_map, tagger as TG   # noqa: E402
from parser.adapters import basic_prose_xlsx as BX      # noqa: E402
from router import discover                             # noqa: E402

from cli import doc_coord as DC                         # noqa: E402
from core.state import sheets as SH                     # noqa: E402

from core.state import fixtures as _FXR                   # noqa: E402 — 픽스처 뿌리(회귀는 스위트마다 사본 — B133)
HINTS = _FXR.EXTRACT_HINTS
RAWS = ROOT / "tests" / "fixtures" / "raw"
TMP = Path(__import__("tempfile").mkdtemp(prefix="b112_"))   # 클린(init --fresh)이 지우지 않는 자리
_made = []
# 시트 역할 · 문서 좌표 기록은 ②등록이라 클린이 지우지 않는다 — 시험이 쓰기 전 판을 잡아 두고 끝에 되돌린다
_KEEP = {p: (p.read_bytes() if p.exists() else None) for p in [DC._file()]}

_FORM = re.compile(r"^\s+(\S.*?) ([\d,]+) = (.+)$")


def _n(x):
    return int(x.replace(",", ""))


def formulas(text):
    """화면의 맞춤 묶음(`   맞춤 —` 아래)에서 식마다 `(줄, 전체, 칸 합)` — **화면 문면을 다시 더한다**."""
    out, on = [], False
    for ln in text.splitlines():
        if ln.startswith("   맞춤 —"):
            on = True
            continue
        if on and not ln.startswith("     "):
            on = False
        m = _FORM.match(ln) if on else None
        if not m:
            continue
        cells = re.sub(r"  \(.*\)$", "", m.group(3)).split(" + ")
        got = [re.match(r"^(.*\S)\s+([\d,]+)$", re.sub(r"\(.*\)$", "", c.strip())) for c in cells]
        out.append((m.group(1).strip(), _n(m.group(2)), sum(_n(g.group(2)) for g in got if g),
                    all(got)))
    return out


def _sum_ok(text, doc, need_chunk=False):
    fs = formulas(text)
    data = LG.read(doc) or {}
    ent = sum(1 for r in data.get("rows") or [] if r.get("role") == "entity")
    held = sum(((data.get("result") or {}).get("보류") or {}).values())
    val = next((t for lb, t, _s, _p in fs if lb == "값(개체 행)"), None)
    hold = next((t for lb, t, _s, _p in fs if lb == "보류(행)"), None)
    bad = [(lb, t, s) for lb, t, s, p in fs if t != s or not p]
    ok = (len(fs) >= (8 if need_chunk else 6) and not bad and "[결함]" not in text and val == ent and hold == held
          and all(f"그래프 {lay} —" in text for lay in ("process", "quality"))
          and (not need_chunk or any(lb.startswith("청크(산문)") for lb, *_x in fs)))
    return ok, f"{doc} 식 {len(fs)} · 어긋남 {bad} · 값 {val}/대장 {ent} · 보류 {hold}/결과 {held}"


def _edgeless_in_graph(doc):
    gs = {lay: open_graph(lay) for lay in discover()}
    ends = {x for g in gs.values() for e in g.edges for x in (e["src"], e["dst"])}
    return sum(1 for g in gs.values() for nid, n in g.nodes.items() if nid not in ends
               and any(str(p) == doc or str(p).startswith(f"{doc}#") for p in n.get("provenance") or []))


def _book(path, sheets):
    """창작 엑셀 — `{시트: (상태, [줄])}`."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, (st, rows) in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.sheet_state = st
        for r, v in enumerate(rows, 1):
            ws.cell(row=r, column=1, value=v)
    _P.ensure(path)
    wb.save(path)
    return path


def _hint(doc_id, by_loc):
    _made.append(doc_id)
    (HINTS / f"{doc_id}.json").write_text(json.dumps(by_loc, ensure_ascii=False), encoding="utf-8")


# ────────────────────────────────────────────────────────────── ⓐ
print("\n■ B112 ① 맞춤표 — 식마다 칸 합 = 전체 (표 · 산문 · 인입 · show report · 재구축)")
init.init(fresh_=True, all_=True)
for _lay in ("process", "quality"):
    bootstrap(_lay, echo=False)
_register("process", DT)
_KEEP_RFQ = SH.path("RFQ01").read_bytes() if SH.path("RFQ01").exists() else None
#: 형태마다 대표 하나 — 표 × quality(PFMEA01)는 같은 함수의 곱이라 뺀다(B121 ②ⓒ)
RUNS = {"CP01": ("--allow-mock",),
        "RFQ01": ("--doc-type", DT, "--allow-mock", "--coord", "노칭", "--sheets", "1:skip 2-4:prose 5:ref *:skip")}
for _d in RUNS:
    shutil.copy(RAWS / f"{_d}.xlsx", _P.ensure(_P.raw(f"{_d}.xlsx")))
OUTS = {d: _run("ingest-file", str(_P.raw(f"{d}.xlsx")), *a) for d, a in RUNS.items()}
_ck = [(d, *_sum_ok(OUTS[d], d, need_chunk=(d == "RFQ01"))) for d in RUNS]
show("ⓐ 인입 — 식마다 칸 합 = 전체 · 값 = 대장 개체 행 · 보류 = 결과 묶음 · 그래프 층마다 (CP01 표 · RFQ01 산문)",
     all(o for _d, o, _x in _ck), " ‖ ".join(x for _d, _o, x in _ck))
_succ = [ln.strip() for d in RUNS for ln in OUTS[d].splitlines() if ln.strip().startswith("[성공]")]
show("ⓐ 성공 줄 — 「그래프 노드 이 문서 +n(그 층 … 전체 m)」 (총수만 두지 않는다)",
     len(_succ) == len(RUNS) and all(re.search(r"그래프 노드 이 문서 [+-]\d[\d,]*\(그 층 \S+ 전체 [\d,]+\)", s) for s in _succ),
     " / ".join(s[s.find("그래프"):][:60] for s in _succ))
_rep = {d: _run("show", "report", d) for d in RUNS}
_ckr = [(d, *_sum_ok(_rep[d], d, need_chunk=(d == "RFQ01"))) for d in RUNS]
show("ⓐ show report — 같은 함수 · 식마다 합 = 전체 (장부만 — 청크 줄은 「뽑음」)",
     all(o for _d, o, _x in _ckr) and "뽑음" in _rep["RFQ01"], " ‖ ".join(x for _d, _o, x in _ckr))
_logs = "".join(p.read_text(encoding="utf-8", errors="replace") for p in sorted(_P.work("logs").glob("*.log")))
show("ⓐ 맞춤 묶음은 명령 로그에도 있다 (화면에 나온 것은 로그에도 — B109)",
     "맞춤 — 줄마다 단위" in _logs and "값(개체 행)" in _logs)
_rb = _run("rebuild", "--yes", "--allow-mock")
_fr = formulas(_rb)
_replay = [int(m.group(1).replace(",", "")) for m in re.finditer(r"값\(개체 행\) [\d,]+ = .*?재생 ([\d,]+)", _rb)]
show("ⓐ 재구축(재생) — 문서마다 맞춤 묶음 · 식마다 합 = 전체 · 재생 칸에 수가 든다",
     _rb.count("   맞춤 —") >= len(RUNS) and _fr and all(t == s and p for _l, t, s, p in _fr) and "[결함]" not in _rb
     and any(n > 0 for n in _replay),
     f"묶음 {_rb.count('   맞춤 —')} · 식 {len(_fr)} · 어긋남 {[(l, t, s) for l, t, s, _p in _fr if t != s]} · 재생 칸 {_replay}")
_mt = RC.for_report("RFQ01", LG.read("RFQ01"))
_clean = RS.reconcile_lines(_mt)
_cks = RC._checkpoints("RFQ01")
_mt["줄"][0] = RC.chunk_line(_cks[0][1], {"부름": 2, "재사용": 0})[0]   # 실행 사실이 체크포인트와 어긋난 묶음
_broken = RS.reconcile_lines(_mt)
show("ⓐ 일부러 깨뜨린 묶음 — 그 줄 아래 「[결함] <줄> 합 x ≠ 전체 y」 (깨지 않은 묶음은 0)",
     not any("[결함]" in ln for ln in _clean)
     and any(re.match(r"^\s+\[결함\] 청크\(산문\) 합 \d+ ≠ 전체 \d+$", ln) for ln in _broken),
     [ln.strip() for ln in _broken if "[결함]" in ln][:1])

# ────────────────────────────────────────────────────────────── ⓒ
print("\n■ B112 ③ 청크 목차 — 행 수 = 청크 수 · 상태 합 = 맞춤 줄 · 라벨(엑셀 · PPT) · 개체 0 목록")
_o = OUTS["RFQ01"].splitlines()
_h = next(i for i, ln in enumerate(_o) if re.match(r"^   #\s+라벨\s+좌표\s+상태\s+글자$", ln))
_e = next(i for i in range(_h + 1, len(_o)) if _o[i].startswith("   청크(산문) "))
_toc = [ln.split() for ln in _o[_h + 1:_e]]
_tot = _FORM.match(_o[_e])
_st = {k: sum(1 for r in _toc if r[-2] == k) for k in ("부름", "ref", "거름", "재사용")}
_want = " + ".join(f"{k} {v}" for k, v in _st.items())
_chunks = chunks_of("RFQ01")
_sheets = [r["name"] for r in SG.rows_of(RFQ)]
_labels = [" ".join(r[1:-3]) for r in _toc]
_end = next(ln for ln in _o if ln.strip().startswith("청크(산문) ") and ln.startswith("     "))
show("ⓒ 엑셀 — 목차 행 수 = 청크 수 · 상태 합 = 표 아래 청크 맞춤 줄 = 끝의 맞춤 줄 · 라벨 「시트 › 경로」",
     len(_toc) == len(_chunks) == _n(_tot.group(2)) and _tot.group(3) == _want and _end.strip() == _o[_e].strip()
     and all(any(lb.startswith(s) for s in _sheets) for lb in _labels) and any(" › " in lb for lb in _labels),
     f"행 {len(_toc)} · 청크 {len(_chunks)} · {_o[_e].strip()} · 라벨 {_labels[:2]}")
_ckp = json.loads(EX.checkpoint_path("RFQ01").read_text(encoding="utf-8"))
_zero = sum(1 for c in _ckp["candidates"] if not c.get("failed") and not c.get("lens_skipped") and not c.get("entities"))
_zl = next((ln for ln in _o if ln.startswith("   개체 0 청크 ")), "")
show("ⓒ 추출 끝 — 개체 0 청크 목록 = 체크포인트의 개체 0 청크 (라벨 + 앞부분)",
     _zero > 0 and _zl.startswith(f"   개체 0 청크 {_zero} — ")
     and sum(1 for ln in _o if re.match(r"^     \S.*「.*」$", ln)) >= _zero, _zl.strip())
_notes = []
run_document(load("PPT01"), notice=_notes.append)
_pv = next((n for n in _notes if n.get("단계") == "추출예고"), {})
_pl = [t[1] for t in _pv.get("목차") or []]
show("ⓒ PPT — 라벨 「슬라이드 n」 · 목차 행 수 = 청크 수 (포맷 전부 한 함수)",
     _pl and len(_pl) == len(load("PPT01")["chunks"]) and all(re.match(r"^슬라이드 \d+$", x) for x in _pl), _pl[:3])

# ────────────────────────────────────────────────────────────── ⓑ
print("\n■ B112 ② 시트 표 — 숨김 · 아주 숨김 · 기록이 있어도 표")
HID = _book(TMP / "B112_시트숨김.xlsx", {
    "사양": ("visible", ["1. 개요", "노칭 설비는 전극을 타발한다.", "2. 구동", "모터는 서보 모터를 쓴다."]),
    "다른공정": ("hidden", ["1. 개요", "스태킹 설비는 적층한다.", "2. 구동", "모터는 리니어 모터다."]),
    "메모": ("veryHidden", ["내부 메모", "가격은 별도 협의."])})
_did_h = IG.doc_id_of(HID)
SH.path(_did_h).unlink(missing_ok=True)   # 지난 실행의 기록이 남아 있으면 「기록 없음」 표가 아니다
_rows = SG.rows_of(HID)
from core.llm import points as _PT        # noqa: E402
_called, _orig = [], _PT.sheet_judge
_PT.sheet_judge = lambda: (lambda name, sample: _called.append(name) or {"role": "prose", "reason": "시험"})   # 기대 변경(B148 ④ — 표본 줄만)
try:
    SG.judge(_rows, True)
finally:
    _PT.sheet_judge = _orig
_byn = {r["name"]: r for r in _rows}
show("ⓑ 리더가 시트 상태를 싣는다 · 숨김 시트는 기본 skip · LLM을 부르지 않는다(보이는 시트만 부름)",
     {n: r.get("state", "visible") for n, r in _byn.items()} == {"사양": "visible", "다른공정": "hidden", "메모": "veryHidden"}
     and _byn["다른공정"]["suggest"] == _byn["메모"]["suggest"] == "skip" and _called == ["사양"]
     and _byn["다른공정"]["llm_reason"].startswith("숨김 시트"),
     f"부름 {_called} · 제안 {[(n, r['suggest']) for n, r in _byn.items()]}")
_live = _run("ingest-file", str(HID), "--doc-type", DT, "--allow-mock", "--coord", "노칭",
             "--sheets", "1:prose 2:prose 3:skip")
_pc = next((ln for ln in _live.splitlines() if ln.strip().startswith("시트별 조각 — ")), "")
show("ⓑ `--sheets`로 숨김 시트를 살린다 — 파싱 끝 시트별 조각(역할과 함께) · 그 시트의 청크가 선다",
     re.search(r"다른공정\(prose\) [1-9]", _pc) and "메모(skip) 0" in _pc
     and any((c.get("meta") or {}).get("sheet") == "다른공정" for c in chunks_of(_did_h).values()), _pc.strip())
_again = _run("ingest-file", str(HID), "--doc-type", DT, "--allow-mock")
#: 숨김 칸은 기록이 있는 실행의 표에서 잰다 — 같은 표 함수 · dry-run을 한 번 더 돌리지 않는다(B121 ②ⓓ) · 제안 skip은 위 판정 줄
_ln = {n: next((ln for ln in _again.splitlines() if re.match(rf"^   \d+\s+{n}\s", ln)), "") for n in _byn}
show("ⓑ 시트 표에 숨김 칸 — 「숨김」 · 「아주 숨김」 (보이는 시트는 「-」)",
     re.search(r"\s-\s", _ln["사양"]) and " 숨김 " in _ln["다른공정"] and " 아주 숨김 " in _ln["메모"],
     " / ".join(re.sub(r"\s+", " ", x.strip()) for x in _ln.values()))
_hl = _ln["다른공정"]
show("ⓑ 기록이 있어도 표를 낸다 — 묻지 않고 기록대로 · 다시 고르는 명령줄 · 기록 prose인데 숨김이면 그 줄에 표시",
     "■ 시트 역할" in _again and "Enter=" not in _again and "기록대로 진행" in _again and "다시 고르려면" in _again
     and "⚠ 기록 prose · 숨김 시트" in _hl and "[성공]" in _again, re.sub(r"\s+", " ", _hl.strip()))


# ────────────────────────────────────────────────────────────── ⓓ
def _setup():
    """덧칠(Component · 겸 렌즈) + `nest_categories` + 소속 엣지의 카테고리쌍 — 상태 루트에만(B102 시험과 같은 바닥)."""
    overlay(component=True)
    _rw(_P.common(), lambda c: c.setdefault("canonical_scope", {}).__setitem__("nest_categories", ["Component"]))
    pats = [{"src": s_, "rel": "part_of", "dst": d_, "symmetric": False, "정의문": f"{s_}가 {d_}에 속한다"}
            for s_, d_ in (("Component", "Unit"), ("Component", "Process"), ("Unit", "Process"))]
    pairs = {f"{p_['src']},{p_['dst']}": "part_of" for p_ in pats}

    def _q(c):
        c["categories"]["Component"] = "(품질층의 렌즈) 불량이 일어나는 구성 기계"
        c["relation_patterns"] += pats
        c.setdefault("category_pair_map", {}).update(pairs)

    def _p(c):
        have = {(p_["src"], p_["rel"], p_["dst"]) for p_ in c["relation_patterns"]}
        c["relation_patterns"] += [p_ for p_ in pats if (p_["src"], p_["rel"], p_["dst"]) not in have]
        for k, v in pairs.items():
            c.setdefault("category_pair_map", {}).setdefault(k, v)
    _rw(_P.layers("quality", "config.json"), _q)
    _rw(_P.layers("process", "config.json"), _p)
    for lay in ("process", "quality"):
        bootstrap(lay, echo=False)


UNITS = _book(TMP / "B112_유닛.xlsx", {
    "이송 유닛": ("visible", ["1. 구동부", "모터는 서보 모터를 쓴다.", "2. 안전", "커버를 단다."]),
    "가압 유닛": ("visible", ["모터는 유압 모터를 쓴다.", "실린더는 공압 실린더다."])})


def _unit_env(doc_id, lay, doc_coord="노칭"):
    struct_map.invalidate(doc_id)
    env = pipeline.parse(BX, doc_id, str(UNITS), closed_list=TG.closed_list("process"), doc_coord=doc_coord,
                         sheet_roles={"이송 유닛": "prose", "가압 유닛": "prose"}).envelope
    env["doc_type"] = "ppt_" + lay
    for c in env["chunks"]:
        c["doc_type"] = env["doc_type"]
    hints = {}
    for c in env["chunks"]:
        ents = []
        unit = (c.get("meta") or {}).get("sheet")
        if "모터" in c.get("text", ""):
            ents.append({"surface": "모터", "category": "Component",
                         "belongs_to": {"name": unit, "category": "Unit", "from": "시트"}})
        if "실린더" in c.get("text", ""):        # 근거 없는 소속(지어낸 이름) — 버림(출처 없음)
            ents.append({"surface": "실린더", "category": "Component",
                         "belongs_to": {"name": "없는 유닛 B112", "category": "Unit", "from": "본문"}})
        if ents:
            hints[c["source_locator"]] = {"entities": ents, "relations": [], "attach": []}
    _hint(doc_id, hints)
    return env, sum(1 for h in hints.values() for e in h["entities"] if e.get("belongs_to"))


print("\n■ B112 ④⑥ 소속 진단 · 엣지 없는 노드 · 화살표 = 엣지 방향 (층 둘)")
_od = []
for _lay in ("process", "quality"):
    _setup()
    _did = f"B112U{_lay[0].upper()}"
    _env, _decl = _unit_env(_did, _lay)
    _nt = []
    run_document(_env, notice=_nt.append)
    _mtd = next(n for n in _nt if n.get("단계") == "끝")["맞춤"]
    _bl = next(x for x in _mtd["줄"] if x["줄"].startswith("소속 대상"))
    _cl = {k: v for k, v, _d in _bl["칸"]}
    _rows_d = [r for r in (LG.read(_did) or {}).get("rows") or [] if r.get("role") == "entity"]
    _ids = {n["canonical"]: n["id"] for _l, n in World().nodes()}
    _edges = {(e["rel"], e["src"], e["dst"]) for lay in discover() for e in open_graph(lay).edges}
    _dirs = []
    for r in _rows_d:
        for a in r.get("attached") or []:
            if str(a.get("path") or "").startswith("소속") and r.get("node_id"):
                oid = _ids.get(a.get("other"))
                real = ("←" if (a["rel"], oid, r["node_id"]) in _edges else
                        "→" if (a["rel"], r["node_id"], oid) in _edges else None)
                _dirs.append((r["surface"], a.get("dir"), real))
    _g = _edgeless_in_graph(_did)
    _od.append((_lay, _bl["전체"] == _decl and _bl["맞음"] and _cl["버림(출처 없음)"] == 1
                and _cl["유닛·부품(노드)"] == 2 and len(_mtd["엣지 없는 노드"]) == _g
                and _dirs and all(d == real for _s, d, real in _dirs),
                f"{_lay}: 소속 대상 {_bl['전체']}/낸 개체 {_decl} · {_cl} · 엣지 없는 노드 {len(_mtd['엣지 없는 노드'])}"
                f"/그래프 {_g} · 화살표(대장 → 실제) {_dirs}"))
show("ⓓ 소속 대상 합 = 소속을 낸 개체 수 · 유닛·부품(노드) · 지어낸 이름은 버림 · 엣지 없는 노드 = 그래프에서 센 수 · "
     "대장 dir = 실제 엣지 방향 (층 둘)", all(o for _l, o, _d in _od), " ‖ ".join(d for _l, _o, d in _od))
_register("process", DT)
_tr = _run("show", "report", _did, "--trace")
_tl = [ln.strip() for ln in _tr.splitlines() if ln.strip().startswith("소속 ") and "'" in ln]
_dir_q = {s: d for s, d, _r in _dirs}
show("ⓓ --trace — 행마다 「소속 이름 → 대상(종류) → 결과」 · 화살표는 대장 dir 그대로",
     any("유닛·부품 · 노드" in ln and f"엣지 part_of {_dir_q.get('모터')}" in ln for ln in _tl)
     and any("버림(출처 없음" in ln for ln in _tl), _tl[:2])

# ────────────────────────────────────────────────────────────── ⓔ
print("\n■ B112 ⑤ 문서 좌표 경고 — 파일명 조각이 띄어쓰기로 어긋남")
XE = _book(TMP / "B112_전극리와인딩_사양.xlsx", {"사양": ("visible", ["장력은 매 시간 기록한다.", "커버를 단다."])})
_did_e = IG.doc_id_of(XE)
_env_e = pipeline.parse(BX, _did_e, str(XE), closed_list=TG.closed_list("process"), doc_coord=None,
                        sheet_roles={"사양": "prose"}).envelope
struct_map.invalidate(_did_e)
_hint(_did_e, {c["source_locator"]: {"entities": [{"surface": "장력", "category": "Property"}], "relations": [], "attach": []}
               for c in _env_e["chunks"] if "장력" in c.get("text", "")})
_oe = _run("ingest-file", str(XE), "--doc-type", DT, "--allow-mock", "--sheets", "1:prose")
_res_e = (LG.read(_did_e) or {}).get("result") or {}
_unres = (_res_e.get("보류") or {}).get("좌표 미해소로 스코프 노드 미생성", 0)
# 기대 변경(B152 ⑤ — 맞춤 아래 옛 경고 두 줄 0 · 같은 사실은 원인 묶음의 「문서 좌표 없음」 줄 — 가까운 골격 이름 · 실경로 처방)
_wl = next((ln for ln in _oe.splitlines() if ln.strip().startswith("문서 좌표 없음 — 문서 좌표 · 못 붙은 행")), "")
_cand = _wl
#: 보류가 0이 아닌 맞춤 묶음 — ⓐ의 대표 둘은 보류 0이라 이 인입 화면(이미 돈 실행)에서 잰다(B121 ②ⓒ — PFMEA01 몫)
_oke, _whye = _sum_ok(_oe, _did_e)
show("ⓔ 문서 좌표 없음 — 원인 묶음의 「문서 좌표 없음」 줄에 가까운 골격 이름(띄어쓰기 무시) · 옛 경고 줄 0 · 처방 --coord · "
     "자동 적용 0(문서 좌표 기록은 없음) · 그 화면의 맞춤 묶음도 식마다 합 = 전체 · 보류(행) = 결과 묶음(ⓐ)",
     _unres > 0 and _wl and "좌표 미해소 보류" not in _oe and "--coord <골격 이름>" in _oe
     and "노칭::전극 리와인딩" in _cand and "문서 좌표 — 없음 (출처: 없음)" in _oe
     and (DC.recorded(_did_e) or {}).get("coord") is None and _oke,
     f"{_wl.strip()} / {_cand.strip()} / {_whye}")

# ────────────────────────────────────────────────────────────── ⓕ
print("\n■ B112 ⑦ 추출 체크포인트 — 끝 줄의 자리 · 있음 · 없으면 [결함]")
_setup()
_env_f, _x = _unit_env("B112CKP", "process")
_buf = io.StringIO()
with contextlib.redirect_stdout(_buf):
    run_document(_env_f, notice=IS.build_screen(prose=True))
_ok_line = next((ln for ln in _buf.getvalue().splitlines() if ln.strip().startswith("추출 체크포인트 — ")), "")
_env_f2, _x = _unit_env("B112CKX", "process")
_wc = EX._write_checkpoint


def _vanish(env, *a, **k):
    """쓴 뒤 바로 없앤다 — 「추출은 끝났는데 체크포인트가 없다」 상태(구축은 돌려받은 묶음으로 간다)."""
    out = _wc(env, *a, **k)
    EX.checkpoint_path(env["doc_id"], a[1] if len(a) > 1 else None).unlink(missing_ok=True)
    return out


EX._write_checkpoint = _vanish
_buf2, _nt2 = io.StringIO(), []
try:
    with contextlib.redirect_stdout(_buf2):
        _scr = IS.build_screen(prose=True)
        run_document(_env_f2, notice=lambda i: (_nt2.append(i), _scr(i))[1], allow_duplicate=True)   # 같은 표본
finally:
    EX._write_checkpoint = _wc
_sv = next((n for n in _nt2 if n.get("단계") == "추출저장"), {})
_dfx = store.path(store.DEFECTS).read_text(encoding="utf-8") if store.path(store.DEFECTS).exists() else ""
show("ⓕ 추출 끝에 체크포인트 자리와 있음 · 쓰기가 빠지면 화면 [결함] · 결함 장부 (같은 자리 `checkpoint_path`)",
     _ok_line.endswith("(있음)") and "B112CKP.json" in _ok_line and _sv.get("있음") is False
     and _sv.get("자리") == EX.checkpoint_path("B112CKX")
     and "[결함] 추출이 끝났는데 체크포인트가 없다" in _buf2.getvalue()
     and "B112CKX: 추출이 끝났는데 체크포인트가 없다" in _dfx,
     f"{_ok_line.strip()} / {[ln.strip() for ln in _buf2.getvalue().splitlines() if '[결함]' in ln][:1]}")

# ────────────────────────────────────────────────────────────── B149 진단 출력
print("\n■ B149 진단 출력 — 원인 묶음(못 찾은 것 · 값은 가장 위 원인에서 한 번) · 목차 나무 · 표의 좌표 표기 목록 · 읽기 전용")
from cli import report_cause as RCZ                     # noqa: E402
from core.build import outline as OL149                 # noqa: E402
from core.build.result import hold_reason as _hr149     # noqa: E402
from openpyxl import load_workbook as _lw149            # noqa: E402


def _ingest149(path, sheets, hint_of, *extra):
    """창작 엑셀 → 추출 힌트(위치마다 · `hint_of(조각)`) → 인입(화면) — `(doc_id, 화면)`."""
    did = IG.doc_id_of(path)
    from core.build import structure as _ST149
    env = pipeline.parse(BX, did, str(path), closed_list=TG.closed_list("process"), doc_coord=None,
                         sheet_roles={s: "prose" for s in sheets}, structure=_ST149.structurer()).envelope   # 인입과 같은 조각 위치
    struct_map.invalidate(did)
    _hint(did, {c["source_locator"]: {"entities": hint_of(c), "relations": [], "attach": []} for c in env["chunks"]
                if hint_of(c)})
    return did, _run("ingest-file", str(path), "--doc-type", DT, "--allow-mock",
                     "--sheets", " ".join(f"{i + 1}:prose" for i in range(len(sheets))), *extra)


def _lost149(did):
    """못 붙은 행 — 보류 행 ∪ 고아 행(대장 · 큐에서 따로 센다 · 좌표 행은 세지 않는다 — 기대 변경 B152 ①)."""
    rows = [r for r in (LG.read(did) or {}).get("rows") or [] if r.get("role") != "anchor"]
    items = {((it.get("provenance") or "").split("#")[-1], it.get("surface"))
             for x in store.read(store.QUEUE, []) if x.get("kind") == "orphan_attach" and x.get("doc_id") == did
             for it in (x.get("payload") or {}).get("items") or []}
    return sum(1 for r in rows if _hr149(r) or ((r.get("locator"), r.get("surface")) in items))


_ents149 = lambda c: [{"surface": w, "category": "Property"} for w in ("장력", "속도") if w in c.get("text", "")]
init.init(fresh_=True, all_=True)
for _lay in ("process", "quality"):
    bootstrap(_lay, echo=False)
_register("process", DT)
# ⓐ 문서 좌표 없음 — 파일명 · --coord · 제목 어디에도 골격 이름이 없다
XA = _book(TMP / "B149_설비_사양서.xlsx", {"일반": ("visible", ["1. 범위", "장력은 매 시간 기록한다.", "2. 운전",
                                                               "속도는 분당 기록한다."])})
_da, _oa = _ingest149(XA, ["일반"], _ents149)
_ca = RCZ.causes(RCZ.material(_da))
_ra = _run("show", "report", _da)
# 기대 변경(B152 ① — 수의 단위는 못 붙은 행(좌표 행 제외) · 원인 줄 「못 붙은 행 m(종류별)」)
show("ⓐ 문서 좌표 없음 — 원인 묶음 첫 줄 「문서 좌표 없음 — 못 붙은 행 n」 + 가까운 골격 이름 + 처방(--coord · 실경로) · 원인 하나 · "
     "합 = 보류 · 고아 행 수(좌표 행 제외 · 아래 원인에서 다시 세지 않는다)",
     _ca and _ca[0]["원인"] == RCZ.DOC and len(_ca) == 1 and sum(g["행"] for g in _ca) == _lost149(_da) > 0
     and "    문서 좌표 없음 — 문서 좌표 · 못 붙은 행" in _ra and "가까운 골격 이름(보여만 준다)" in _ra
     and f"ingest-file {XA} --doc-type {DT} --coord <골격 이름>" in _ra,
     f"원인 {[(g['원인'], g['행']) for g in _ca]} · 못 붙은 행 {_lost149(_da)}")
# ⓕ 인입 끝 화면 — 원인 묶음 위 n줄(같은 함수 · 전체는 show report)
_lines_a = RCZ.cause_block(_da, limit=IS.LIST_LIMIT)
show("ⓕ 인입 끝 화면이 원인 묶음 위 n줄을 낸다(`show report`와 같은 함수 · 값 중심 보류 목록 0)",
     _lines_a and all(ln in _oa for ln in _lines_a if ln.strip()) and f"전체는 python run.py show report {_da}" in _oa
     and "보류 목록 — 사유" not in _oa, [ln.strip()[:60] for ln in _lines_a if ln.strip()][:2])
# ⓖ 읽기 전용 — show report 전후 상태 루트 바이트 같음(명령 로그 제외) · LLM 0
_h0 = {p: p.read_bytes() for p in _P.home().rglob("*") if p.is_file() and "logs" not in p.parts}
_rg = _run("show", "report", _da, "-v")
_h1 = {p: p.read_bytes() for p in _P.home().rglob("*") if p.is_file() and "logs" not in p.parts}
show("ⓖ 읽기 전용 — show report(-v 포함) 전후 상태 루트 바이트 같음(명령 로그 제외) · LLM 0",
     _h0 == _h1 and "LLM 사용량" not in _rg, f"파일 {len(_h0)} · 바뀐 것 {[p.name for p in set(_h0) | set(_h1) if _h0.get(p) != _h1.get(p)][:3]}")
# ⓑ 구간 좌표 없음 · ⓓ 목차 나무 — 시트 하나만 골격 이름(시트 「노칭」의 첫 줄 = 시트 제목 「노칭」 — 그 조각의 제목 사슬이 골격 이름)
XB = _book(TMP / "B149_구간.xlsx", {"노칭": ("visible", ["노칭", "장력은 매 시간 기록한다."]),
                                   "기타 사양": ("visible", ["1. 운전", "속도는 분당 기록한다.", "2. 점검", "속도를 본다."])})
_db, _ob = _ingest149(XB, ["노칭", "기타 사양"], _ents149)
_mb = RCZ.material(_db)
_cb = RCZ.causes(_mb)
_rows_b = (LG.read(_db) or {}).get("rows") or []
_heldb = {r.get("locator") for r in _rows_b if _hr149(r)}
_other_locs = {c["source_locator"] for c in _mb["env"]["chunks"] if (c.get("meta") or {}).get("sheet") == "기타 사양"}
# 기대 변경(B152 ① — 못 붙은 행 · 좌표 행 제외)
show("ⓑ 구간 좌표 없음 — 골격 이름 시트의 값은 붙고 · 다른 시트의 가장 위 제목이 「구간 좌표 없음 — <경로> 못 붙은 행 n」 한 줄",
     [(g["원인"], g["못 찾은 것"]) for g in _cb] == [(RCZ.SECTION, "기타 사양")]
     and _cb[0]["행"] == sum(1 for r in _rows_b if _hr149(r) and r.get("role") != "anchor" and r.get("locator") in _other_locs) > 0
     and not {c["source_locator"] for c in _mb["env"]["chunks"] if (c.get("meta") or {}).get("sheet") == "노칭"} & _heldb
     and any(r.get("surface") == "장력" and r.get("node_id") for r in _rows_b),
     f"원인 {[(g['원인'], g['못 찾은 것'], g['행']) for g in _cb]}")
_tb, _tv = RCZ.tree_lines(_mb), RCZ.tree_lines(_mb, verbose=True)
_line_o = next((ln for ln in _tb if ln.strip().startswith("1. 운전 — ")), "")
_tr149 = OL149.tree(_mb["env"])
_hid_o = next(h["id"] for h in _tr149["headings"] if h["path"].endswith("1. 운전"))
# 기대 변경(B152 ④ — 「붙음 · 못 붙음」(못 붙은 행 — 좌표 행 제외))
_n_o = sum(1 for r in _rows_b if _hr149(r) and r.get("role") != "anchor"
           and _tr149["of"].get(r.get("locator")) == _hid_o)   # 그 제목 바로 아래 조각의 못 붙은 행
show("ⓓ 목차 나무 — 문제 가지만 펼침(골격 시트는 접힘) · 줄의 좌표 출처 · 노드 상태 · 못 붙음 수 = 봉투 · 대장 · -v면 전부",
     any(ln.strip().startswith("기타 사양 — 좌표 없음") for ln in _tb) and f"못 붙음 {_n_o:,}" in _line_o and "노드: " in _line_o
     and not any(ln.strip().startswith("노칭 — ") for ln in _tb) and any("다 붙은 가지 1" in ln for ln in _tb)
     and any(ln.strip().startswith("노칭 — 좌표 노칭(자기)") for ln in _tv),
     f"{_line_o.strip()} · 접음 {[ln.strip() for ln in _tb if '다 붙은' in ln]}")
# ⓒ 소속 이름 없음 — 그 이름이 청크에 있지만 노드가 없다(카테고리 없는 소속 — 고아 큐)
XC = _book(TMP / "B149_소속.xlsx", {"사양": ("visible", ["1. 반송", "장력은 반송 유닛에서 잰다.", "속도는 반송 유닛의 값이다."])})
_dc, _oc = _ingest149(XC, ["사양"], lambda c: [{"surface": w, "category": "Property",
                                               "belongs_to": {"name": "반송 유닛", "category": None, "from": "본문"}}
                                              for w in ("장력", "속도") if w in c.get("text", "")], "--coord", "노칭")
_cc = RCZ.causes(RCZ.material(_dc))
_q_c = sum(int((x.get("payload") or {}).get("rows") or 0) for x in store.read(store.QUEUE, [])
           if x.get("kind") == "orphan_attach" and x.get("doc_id") == _dc and (x.get("payload") or {}).get("key") == "반송 유닛")
# 기대 변경(B152 ① — 못 붙은 행)
show("ⓒ 소속 이름 없음 — 기다리는 이름마다 한 줄 · 못 붙은 행 = 고아 큐의 그 이름 항목 수 · 처방은 사람 후보(--step)",
     [(g["원인"], g["못 찾은 것"]) for g in _cc if g["원인"] == RCZ.BELONGS] == [(RCZ.BELONGS, "반송 유닛")]
     and next(g["행"] for g in _cc if g["원인"] == RCZ.BELONGS) == _q_c > 0
     and "--step" in next(g["처방"] for g in _cc if g["원인"] == RCZ.BELONGS),
     f"원인 {[(g['원인'], g['못 찾은 것'], g['행']) for g in _cc]} · 큐 {_q_c}")
# ⓔ 표 — 좌표 표기 하나가 골격 밖(CP01 사본의 「프리 실링」 한 행)
_wb = _lw149(RAWS / "CP01.xlsx")
_ws = _wb.worksheets[0]
for _row in _ws.iter_rows(min_row=4):
    if _row[2].value == "프리 실링":
        _row[2].value = "없는공정B149"
XT = TMP / "B149_CP.xlsx"
_wb.save(XT)
_ot = _run("ingest-file", str(XT), "--doc-type", "cp", "--allow-mock")
_dt = IG.doc_id_of(XT)
_mt149 = RCZ.material(_dt)
_ct = RCZ.causes(_mt149)
_rows_t = (LG.read(_dt) or {}).get("rows") or []
_lt = [ln for ln in RCZ.tree_lines(_mt149) if "없는공정B149" in ln]
# 기대 변경(B152 ① ④ — 못 붙은 행(좌표 행 제외) · 목록 줄 「레코드 r · 붙음 · 못 붙음」)
_nt149 = sum(1 for r in _rows_t if _hr149(r) and r.get("role") != "anchor"
             and RC.row_coord(r, _mt149["by_loc"]) == "없는공정B149")
show("ⓔ 표 — 골격 밖 좌표 표기 하나가 한 줄(못 붙은 행 = 그 표기의 보류 행) · 목차 나무 자리는 좌표 표기 목록(그 표기 → 없음 · 레코드 1)",
     (RCZ.SECTION, "좌표 표기 '없는공정B149'") in [(g["원인"], g["못 찾은 것"]) for g in _ct]
     and next(g["행"] for g in _ct if g["못 찾은 것"] == "좌표 표기 '없는공정B149'") == _nt149 > 0
     and len(_lt) == 1 and "· 레코드 1 ·" in _lt[0] and "좌표 표기 목록" in RCZ.tree_lines(_mt149)[1],
     f"원인 {[(g['원인'], g['못 찾은 것'], g['행']) for g in _ct]} · {(_lt or [''])[0].strip()}")

# ────────────────────────────────────────────────────────────── B152 진단 화면의 수 단위
print("\n■ B152 — 수의 단위는 못 붙은 행(좌표 행 제외 · 맞춤표 행 종류) · 머리 · 원인 줄 · 목차 나무가 같은 행 집합")
_ma152 = RCZ.material(_da)
_hd152 = RCZ.head_line(_ma152) or ""
_ra_lines = _ra.splitlines()
_doc152 = next((ln for ln in _ra_lines if ln.strip().startswith("문서 좌표 없음 — 문서 좌표 · 못 붙은 행")), "")
_ent152 = sum(1 for r in (LG.read(_da) or {}).get("rows") or [] if r.get("role") == "entity")
show("B152 ⓐ 단위 — 문서 좌표 없는 산문(값 2 · 조각 2) → 원인 줄 「못 붙은 행 2(개체 2)」 · 머리 「개체 2/2」 · 같은 화면 맞춤 「값(개체 행) 2」와 "
     "같은 수 · 좌표 행 0",
     _ent152 == 2 and "못 붙은 행 2(개체 2)" in _doc152 and "못 붙은 행 2 / 전체 2 — 개체 2/2" in _hd152 and _hd152.strip() in _ra
     and re.search(r"값\(개체 행\) 2\b", _ra) and not any("좌표" in k for k in (RCZ.head_count(_ma152)[2] or [])),
     f"{_hd152.strip()} · {_doc152.strip()[:80]}")
from collections import Counter                          # noqa: E402
from tools import four_way as _FW152                      # noqa: E402 — B147 부착 대상 표본(창작 — 한 자리)
_f152 = _FW152.CPU_SCHEMA["fields"]
_x152 = _FW152.cpu_setup({**_FW152.CPU_SCHEMA, "fields": {**_f152, "규격": {**_f152["규격"], "attach_to_field": "process_ref"}}})
_run("ingest-file", str(_x152), "--doc-type", _FW152.CPU_DT, "--allow-mock")
_mx152 = RCZ.material("FWCPU")
_cx152 = RCZ.causes(_mx152)
_gx = [g for g in _cx152 if g["원인"] == RCZ.OTHER and "개체 열" in str(g["못 찾은 것"])]
_hx = RCZ.head_line(_mx152) or ""
show("B152 ⓑ 속성만 못 붙은 원인이 보인다 — 부착 대상 표본(개체는 붙고 속성이 보류) → 「그 밖 — 부착 대상」 줄에 속성·내용 수 · 개체 칸 없음 · "
     "머리 「속성·내용 b/B」",
     len(_gx) == 1 and set(_gx[0]["종류"]) == {"속성·내용"} and _gx[0]["행"] > 0
     and f"속성·내용 {_gx[0]['행']}/" in _hx and " 개체 " not in _hx,
     f"{_hx.strip()} · {[(g['원인'], g['못 찾은 것'], g['종류']) for g in _cx152]}")
_sums = []
for _d152 in (_db, _dc, _dt):
    _m152 = RCZ.material(_d152)
    _n152, _tot, _kinds, _st = RCZ.head_count(_m152)
    _c152 = RCZ.causes(_m152)
    _kc = Counter()
    for g in _c152:
        _kc.update(g["종류"])
    _tree = RCZ.tree_lines(_m152, verbose=True)
    _lost_tree = sum(int(m.group(1).replace(",", "")) for ln in _tree for m in [re.search(r"못 붙음 ([\d,]+)", ln)] if m)
    _sums.append((_d152, _n152, dict(_kc) == {k: a for k, a, _b in _kinds}, sum(_kc.values()), _lost_tree))
show("B152 ⓒ 합 — B149 ⓑ · ⓒ · ⓔ 표본에서 원인 줄의 종류별 합 = 머리의 분자 · 목차 나무(펼친 줄 + 접은 줄) · 좌표 표기 목록의 못 붙음 합 = 머리의 n",
     all(ok and n == s == lt and n > 0 for _d, n, ok, s, lt in _sums), str(_sums))
for _x in (XA, XB, XC, XT):
    _d149 = IG.doc_id_of(_x)
    SH.path(_d149).unlink(missing_ok=True)

for _d in set(_made):
    (HINTS / f"{_d}.json").unlink(missing_ok=True)
for _d in set(_made) | {_did_h}:
    SH.path(_d).unlink(missing_ok=True)
_KEEP[SH.path("RFQ01")] = _KEEP_RFQ
for _p, _b in _KEEP.items():
    if _b is None:
        _p.unlink(missing_ok=True)
    else:
        _p.write_bytes(_b)
for _d in RUNS:
    _P.raw(f"{_d}.xlsx").unlink(missing_ok=True)
_unregister(DT)
shutil.rmtree(TMP, ignore_errors=True)
done()
