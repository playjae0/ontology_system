# -*- coding: utf-8 -*-
"""G6 ㉑ 공유 별칭의 범위 안 해소 · 질의 하이브리드 — 임베딩 후보 + LLM 선별 링킹 · 문서 검색 채널 · 짧은 답과
링크 결과 · 채점 축 (B104).

창작 표본(mock 그래프 + 시험이 만든 엑셀 · 시험 주입)으로 메커니즘을 잰다 — 수치는 근거가 아니다([정정] 50).
실호출 갈래는 **가짜 게이트웨이**(`tests/fake_gateway.py`)로 잰다: 임베딩은 동의어 표를 아는 결정적 벡터(전송
`_post` 대역) · 선별·답변은 결정적 규칙(`chat` 대역). 형태(표·산문) × 층(process · quality)에서 같은 성질을 잰다(CLAUDE.md §4).

잠그는 성질:
  ⓐ 공유 별칭: 문서 좌표 범위 안에서 하나인 공유 별칭 시트·행의 조각 좌표 = 그 노드 canonical(원 표기 meta) ·
     구축 보류 0 · 그 시트의 개체가 그 노드 아래 · 범위 안에서도 여럿이면 지금처럼(산문은 건너뜀 · 표는
     보류) · 보류 문면 「표기 모호(골격 노드 여럿)」 + 다음 줄 `--coord`
  ⓑ 링킹: 사전에 없는 동의 표현이 임베딩 후보 + 선별로 노드에 닿는다 · 후보 밖 id 0 · LLM 후보 수 ≤ k(노드 전부
     아님) · 보충은 사전 노드 유지 + 추가 · 폴백은 사전이 잡으면 선별 호출 0 · 미스율은 사전 단 기준
  ⓒ 문서 채널: 링킹 0인 질문이 문서 검색 청크로 답(「근거 없음」 아님) · 두 채널이 다 비면 「근거 없음」 ·
     캐시 재사용(두 번째 질의 벡터 재계산 0) · 청크 추가 시 그 몫만
  ⓓ 답: 5줄 이내 표시 · used_facts·used_chunks 표시 · CLI가 ⓐ~ⓕ를 나란히 · `--json`과 CLI가 같은 수
  ⓔ 채점: 링킹 사전만 vs 보충 · 문서 검색 doc@k가 BM25 대조군 옆에
  ⓕ 쓰기 0(③진실·②등록 해시 불변 — 캐시는 `work/vectors/`)
  B161 질의 벡터 색인: 색인 행 = 노드 + 청크(글 = 재료 함수의 출력 · 글 해시가 v에) · 지문이 같으면 재료 함수 호출 0이고 결과가
     색인 전의 길과 같다 · 재료 판 · 한도가 바뀌면 다시 쓴다 · 참조 없는 벡터는 세기만 하고 --gc만 지운다(그 판만) · 옛 저장에서
     만든 벡터 0 · show near · dist embed는 색인에서(쓰기 0) · 두 프로세스가 함께 갈아 끼워도 깨짐 0 · 그래프 직렬화는 graph.py 안
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import re
import tempfile
from pathlib import Path

sys_path = __import__("sys").path
sys_path.insert(0, str(Path(__file__).resolve().parent))
sys_path.insert(0, str(Path(__file__).resolve().parent / "fixtures"))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done, overlay   # noqa: F401

from openpyxl import Workbook                           # noqa: E402
from cli import golden as GD, query as CQ               # noqa: E402
from cli.platform import orphan_next_lines              # noqa: E402
from core.build import ledger as LG                     # noqa: E402
from core.build.result import hold_reason               # noqa: E402
from core.query import vectors as V                     # noqa: E402
from core.state import knobs as KB                      # noqa: E402
from parser import pipeline, struct_map, tagger as TG   # noqa: E402
from parser.adapters import basic_prose_xlsx as BX      # noqa: E402
from fake_gateway import CALLS, Live                    # noqa: E402 — 실호출 갈래의 대역(전송·채팅·설정만)

from core.state import fixtures as _FXR                   # noqa: E402 — 픽스처 뿌리(회귀는 스위트마다 사본 — B133)
HINTS = _FXR.EXTRACT_HINTS
TMP = Path(tempfile.mkdtemp(prefix="b104_"))
_made = []


def _state_hash():
    h = hashlib.sha256()
    for d in (_P.data(), _P.registry()):
        for p in sorted(d.rglob("*")) if d.exists() else []:
            if p.is_file() and not p.name.endswith(".lock"):
                h.update(p.relative_to(d).as_posix().encode("utf-8"))
                h.update(p.read_bytes())
    return h.hexdigest()


def _knobs(d):
    if d:
        _P.knobs().write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    else:
        _P.knobs().unlink(missing_ok=True)
    KB.apply()


# ────────────────────────────────────────────────────────────── ⓐ
print("\n■ B104 ① 공유 별칭의 범위 안 해소 (엑셀 산문 · 표 × 층 둘)")


def _xlsx(path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in {"비전 검사": ["카메라는 B104 비전 카메라를 쓴다."],
                       "노칭 타발": ["타발 금형은 매 교대 점검한다."]}.items():
        ws = wb.create_sheet(title=name)
        for r, v in enumerate(rows, 1):
            ws.cell(row=r, column=1, value=v)
    wb.save(path)
    return path


XL = _xlsx(TMP / "B104_노칭_사양.xlsx")
ROLES = {"비전 검사": "prose", "노칭 타발": "prose"}


def _prose_env(doc_id, lay, doc_coord):
    struct_map.invalidate(doc_id)
    env = pipeline.parse(BX, doc_id, str(XL), closed_list=TG.closed_list("process"),
                         sheet_roles=ROLES, doc_coord=doc_coord).envelope
    dt = "ppt_" + lay
    env["doc_type"] = dt
    for c in env["chunks"]:
        c["doc_type"] = dt
    hints = {c["source_locator"]: {"entities": [{"surface": "B104 비전 카메라", "category": "Unit"}],
                                   "relations": [], "attach": []}
             for c in env["chunks"] if "카메라" in c.get("text", "")}
    (HINTS / f"{doc_id}.json").write_text(json.dumps(hints, ensure_ascii=False), encoding="utf-8")
    _made.append(doc_id)
    return env


def _table_env(doc_id, lay, ref, doc_coord):
    dt = f"b104t{lay[0]}"
    p = _P.schemas(f"{dt}.json")
    _P.ensure(p)
    p.write_text(json.dumps({"doc_type": dt, "schema_version": 1, "layer": lay,
                             "use_blocks": ["common_core", "process_coord"],
                             "fields": {"개체": {"role": "entity", "category": "Unit"}}, "edges": []},
                            ensure_ascii=False), encoding="utf-8")
    dts = store.read(store.DOC_TYPES, {})
    dts[dt] = {"doc_type": dt, "status": "registered", "layer": lay, "schema": f"schemas/{dt}.json",
               "adapter": "-", "schema_version": 1}
    store.write(store.DOC_TYPES, dts)
    base = {k: v for k, v in CPREC.items() if k not in ("설비", "관리항목")}
    recs = [dict(base, source_locator="T1", doc_type=dt, process_ref=ref, process_group=None,
                 개체="B104 비전 카메라")]
    # 파서와 같은 함수(좌표 태깅 — 문서 좌표 범위 안 해소)를 지난 행이 구축으로 간다
    recs = TG.tag(recs, nodes=TG.closed_list("process"), doc_coord=doc_coord)
    return {**TABLE, "doc_id": doc_id, "doc_type": dt, "records": recs}


_oa, _ob = [], []
for form, lay in (("산문", "process"), ("산문", "quality"), ("표", "process"), ("표", "quality")):
    overlay()                                         # 품질층이 Unit을 렌즈로 쓴다(층 둘)
    did = f"B104S{form[0] == '표' and 'T' or 'P'}{lay[0].upper()}"
    env = _prose_env(did, lay, "노칭") if form == "산문" else _table_env(did, lay, "비전 검사", "노칭")
    pcs = env.get("chunks") or env.get("records")
    one = [p for p in pcs if (p.get("meta") or {}).get("coord_tag_from") == "비전 검사"]
    run_document(env)
    rows = (LG.read(did) or {}).get("rows") or []
    holds = [x for x in store.read(store.QUEUE, []) if x["doc_id"] == did and x["kind"] == "orphan_anchor"]
    ents = [r for r in rows if r.get("role") == "entity" and r.get("surface") == "B104 비전 카메라"]
    _oa.append((form, lay, one and all(p["process_ref"] == "노칭::비전검사" for p in one) and not holds
                and ents and all((r.get("target") or "").startswith("노칭::비전검사::") for r in ents)
                and not any(r.get("hold") for r in rows),
                f"좌표 {sorted({p['process_ref'] for p in one})}(원 표기 비전 검사) · 보류 {len(holds)} · "
                f"개체 → {sorted({r.get('target') for r in ents})}"))
    # 범위 안에서도 여럿 — 문서 좌표 「조립」(노칭·스태킹이 다 그 아래)
    did2 = did + "M"
    env2 = _prose_env(did2, lay, "조립") if form == "산문" else _table_env(did2, lay, "비전 검사", "조립")
    pcs2 = env2.get("chunks") or env2.get("records")
    run_document(env2, allow_duplicate=True)
    rows2 = (LG.read(did2) or {}).get("rows") or []
    q2 = [x for x in store.read(store.QUEUE, []) if x["doc_id"] == did2 and x["kind"] == "orphan_anchor"]
    if form == "산문":
        sk = [p for p in pcs2 if "비전 검사" in ((p.get("meta") or {}).get("coord_shared_skip") or [])]
        ok2 = sk and all(p["process_ref"] == "조립" for p in sk) and not q2
        det = f"건너뜀 {len(sk)} · 문서 좌표 물려받음 {sorted({p['process_ref'] for p in sk})}"
    else:
        hr = [hold_reason(r) for r in rows2 if r.get("hold")]
        nxt = orphan_next_lines(q2[0]) if q2 else ""
        ok2 = (pcs2[0]["process_ref"] == "비전 검사" and q2 and "표기 모호" in q2[0]["reason"]
               and hr and all(h == "표기 모호(골격 노드 여럿)" for h in hr) and "--coord" in nxt)
        det = f"보류 {q2[0]['reason'] if q2 else '-'} · 대장 {sorted(set(hr))} · 다음 줄 --coord {'--coord' in nxt}"
    _ob.append((form, lay, ok2, det))
    if form == "표":
        dts = store.read(store.DOC_TYPES, {})
        dts.pop(f"b104t{lay[0]}", None)
        store.write(store.DOC_TYPES, dts)
        _P.schemas(f"b104t{lay[0]}.json").unlink(missing_ok=True)
show("ⓐ 범위 안에서 하나인 공유 별칭 — 조각 좌표가 그 노드 canonical(원 표기 meta) · 보류 0 · 개체가 그 노드 아래 "
     "(엑셀 산문 · 표 × 층 둘)", all(o for _f, _l, o, _d in _oa), " ‖ ".join(f"{f}×{l} {d}" for f, l, _o, d in _oa))
show("ⓐ 범위 안에서도 여럿이면 지금처럼 — 산문은 건너뛰고 문서 좌표 · 표는 보류이되 문면 「표기 모호(골격 노드 여럿)」 "
     "+ 다음 줄 --coord (층 둘)", all(o for _f, _l, o, _d in _ob), " ‖ ".join(f"{f}×{l} {d}" for f, l, _o, d in _ob))

# ────────────────────────────────────────────────────────────── ⓑ
print("\n■ B104 ② 링킹 — 임베딩 후보 + LLM 선별 (가짜 게이트웨이)")
with contextlib.redirect_stdout(io.StringIO()):
    fresh()                                           # mock 그래프(CP01·PFMEA01·PPT01~03·QPPT01)
_h0 = _state_hash()
_knobs({"query_link_top_k": 5})
_nall = sum(1 for lay in ("process", "quality") for n in open_graph(lay).nodes.values() if ops.is_live(n))
_lm = store.path(store.LINK_MISS)
_lm0 = _lm.read_text(encoding="utf-8").count("\n") if _lm.exists() else 0
with Live():
    r1 = CQ.answer("틈새는 어떻게 관리해?")
    c1 = [b for p, b in CALLS if p == "link"]
    r2 = CQ.answer("노칭 틈새는 어떻게 관리해?")
    _knobs({"query_link_top_k": 5, "query_link_mode": "폴백"})
    del CALLS[:]
    r3 = CQ.answer("노칭 틈새는 어떻게 관리해?")
    c3 = [p for p, _b in CALLS if p == "link"]
_knobs({})
_lm1 = _lm.read_text(encoding="utf-8").count("\n") if _lm.exists() else 0
L1 = r1["trace"]["linking"]
show("ⓑ 사전에 없는 동의 표현(「틈새」)이 임베딩 후보 + 선별로 노드에 닿는다 — method embed+llm · 점수 · 이유 · 어디의 무엇",
     [x["canonical"].split("::")[-1] for x in L1] == ["금형 클리어런스"] and L1[0]["method"] == "embed+llm"  # 기대 변경(B153 ② — 노드 글이 맨 이름 · 이름 사슬로 바뀌어 가짜 임베딩 후보 차례가 달라짐 · 닿는 것은 같은 이름 노드 하나 — 성질로 잰다)
     and L1[0].get("why") and L1[0].get("score") is not None and L1[0].get("where"),
     " ‖ ".join(f"{x['canonical']} [{x['method']} · 점수 {x.get('score')} · 이유 「{x.get('why')}」 · "
                f"위치 {x.get('where')}]" for x in L1))
show("ⓑ 후보 밖 id 0 · LLM이 받은 후보 수 ≤ k(5) < 노드 전부 — 노드 전부를 보내지 않는다",
     not any("NOPE" in x["node_id"] for x in L1) and c1 and len(c1[0]["candidates"]) == 5 < _nall
     and len(r1["trace"]["link_stage"]["candidates"]) == 5,
     f"후보 {len(c1[0]['candidates']) if c1 else '-'} / 노드 {_nall} · 고른 {r1['trace']['link_stage'].get('picked')}")
m2 = {x["canonical"]: x["method"] for x in r2["trace"]["linking"]}
m3 = {x["canonical"]: x["method"] for x in r3["trace"]["linking"]}
show("ⓑ 보충(기본) — 사전 노드는 그대로 + 선별이 더한다 · 폴백 — 사전이 잡으면 선별 호출 0(지금 규칙)",
     m2.get("노칭") == "dict" and any(k.endswith("::금형 클리어런스") and v == "embed+llm" for k, v in m2.items())  # 기대 변경(B153 ② — 같은 까닭 · 성질로 잰다)
     and m3 == {k: v for k, v in m2.items() if v == "dict"} and not c3,
     f"보충 {m2} · 폴백 {m3} · 폴백 선별 호출 {len(c3)}")
show("ⓑ 링킹 미스율은 사전 단 기준 — 선별이 찾았어도 사전이 못 찾은 질문은 link_miss에 적재(사전이 잡은 질문은 0)",
     _lm1 - _lm0 == 1 and r1["trace"]["miss"] == ["틈새는 어떻게 관리해?"] and not r2["trace"]["miss"],
     f"link_miss +{_lm1 - _lm0} · trace.miss {r1['trace']['miss']} / {r2['trace']['miss']}")

# ────────────────────────────────────────────────────────────── ⓒ
print("\n■ B104 ③ 문서 검색 채널 · 벡터 캐시")
_r = CQ.answer("육안 점검 교체 주기는?")
_g = CQ.answer("리튬이온 배터리 동작 원리 설명해줘")
show("ⓒ 링킹 0인 질문이 문서 검색 청크로 답한다(「근거 없음」 아님 · path chunk) · 두 채널이 다 비면 「근거 없음」",
     not _r["trace"]["linking"] and _r["path"] == "chunk" and _r["doc_search"] and not _r.get("note")
     and _g["path"] == "general_knowledge" and not _g["doc_search"] and "근거를 찾지 못했다" in (_g["note"] or ""),
     f"{_r['path']} · 문서 검색 {len(_r['doc_search'])} {[(c['doc_id'], c['by']) for c in _r['doc_search'][:3]]} · "
     f"리튬이온 → {_g['path']}")
with Live():
    _rl = CQ.answer("틈새는 어떻게 관리해?")
_emb = [c for c in _rl["doc_search"] if "embed" in c["by"]]
show("ⓒ 실호출 갈래 — 임베딩이 동의 표현으로 청크를 찾는다(문턱 이상 · BM25와 합쳐 중복 제거)",
     _emb and any("클리어런스" in c["text"] for c in _emb)
     and len({c["chunk_id"] for c in _rl["doc_search"]}) == len(_rl["doc_search"]),
     f"임베딩 {len(_emb)} · BM25 {sum(1 for c in _rl['doc_search'] if 'bm25' in c['by'])} · "
     f"{[(c['doc_id'], c['embed']) for c in _emb[:3]]}")
CQ.answer("육안 점검 교체 주기는?")              # 판이 바뀌었다(가짜 게이트웨이 → mock) — 통째로 다시
_s0 = {k: dict(v) for k, v in V.STATS.items()}
CQ.answer("육안 점검 교체 주기는?")
_s1 = {k: dict(v) for k, v in V.STATS.items()}
_extra = {**PROSE, "doc_id": "B104ADD", "source_path": "B104ADD.pptx", "doc_type": "ppt_quality",
          "chunks": [{**C1, "source_locator": "B104ADD-C001", "text": "육안 점검 기준은 사진 견본과 대조한다."},
                     {**C1, "source_locator": "B104ADD-C002", "text": "교체 주기는 타발 횟수로 정한다."}]}
(HINTS / "B104ADD.json").write_text("{}", encoding="utf-8")
_made.append("B104ADD")
run_document(_extra)
CQ.answer("육안 점검 교체 주기는?")
_s2 = {k: dict(v) for k, v in V.STATS.items()}
# 기대 변경(B136 ③ — 저장의 열쇠가 판 + 글 해시라 판마다 따로 담기고 섞이지 않는다 · 돌아온 판은 제 벡터를 다시 쓴다)
import sqlite3 as _sq0                                   # noqa: E402
_c0 = _sq0.connect(str(_P.vectors(V.DB)))
_bks = sorted(r[0] for r in _c0.execute("SELECT DISTINCT backend FROM v"))
_c0.close()
show("ⓒ 벡터 재사용 — 두 번째 질의는 벡터 재계산 0 · 청크를 더하면 그 몫만 다시(노드 그대로면 0) · "
     "판(갈래·모델)이 바뀌면 다른 공간(판마다 따로 담긴다 · 섞이지 않는다)",
     _s1["chunks"]["made"] == 0 and _s1["nodes"]["made"] == 0 and _s2["chunks"]["made"] == 2
     and _s2["chunks"]["reused"] == _s1["chunks"]["reused"] and _s2["nodes"]["made"] == 0
     and len(_bks) == 2 and V.cost_line().startswith("벡터 — "),
     f"판 {_bks} · 두 번째 {_s1['chunks']} · 청크 2 추가 뒤 {_s2['chunks']} · {V.cost_line()}")

# ────────────────────────────────────────────────────────────── ⓓ
print("\n■ B104 ④ 답 — 짧게 · 쓴 근거 · 링크 결과와 나란히")
with Live():
    _out = io.StringIO()
    with contextlib.redirect_stdout(_out):
        CQ.main(["노칭 틈새는 어떻게 관리해?"])
    _txt = _out.getvalue()
    _jo = io.StringIO()
    with contextlib.redirect_stdout(_jo), contextlib.redirect_stderr(io.StringIO()):
        CQ.main(["노칭 틈새는 어떻게 관리해?", "--json"])
_j = json.loads(_jo.getvalue())
_jt = _j["trace"]
_need = ["ⓐ 링킹", "[링킹]", "[링킹 추가]", "ⓑ 확장", "[그래프 사실]", "[문서 근거]", "ⓔ 문서 검색",
         "[문서 검색]", "ⓕ 답 — LLM(live)"]
_cnt = lambda t, k: sum(1 for l in t.splitlines() if l.strip().startswith(k))   # noqa: E731
show("ⓓ CLI가 링크 결과(ⓐ~ⓔ)와 LLM 답(ⓕ)을 나란히 낸다",
     all(k in _txt for k in _need), " / ".join(l.strip()[:50] for l in _txt.splitlines()
                                                if any(l.strip().startswith(k) for k in ("ⓐ", "ⓑ", "ⓔ", "ⓕ"))))
_ds = [c for c in _j["doc_search"] if not c.get("in_graph")]
_tick = lambda k: _cnt(_txt, k + " ✓")                                          # noqa: E731
show("ⓓ `--json`과 CLI가 같은 수 — 사실(trace 전부 · ✓ = 답이 쓴 묶음 facts) · 노드 근거 · 문서 검색 · 링킹 추가 · "
     "쓴 청크 ✓ = used",
     _cnt(_txt, "[그래프 사실]") == len(_jt["facts"]) and _tick("[그래프 사실]") == len(_j["facts"])
     and _cnt(_txt, "[문서 근거]") == len(_j["chunks"]) and _cnt(_txt, "[문서 검색]") == len(_ds)
     and _cnt(_txt, "[링킹 추가]") == sum(1 for x in _jt["linking"] if x["method"] == "embed+llm")
     and _tick("[문서 근거]") + _tick("[관련 원문]") + _tick("[문서 검색]") == len(_jt["answer"]["used_chunks"]),
     f"사실 {_cnt(_txt, '[그래프 사실]')}={len(_jt['facts'])}(✓ {_tick('[그래프 사실]')}={len(_j['facts'])}) · "
     f"노드 근거 {_cnt(_txt, '[문서 근거]')}={len(_j['chunks'])} · 문서 검색 {_cnt(_txt, '[문서 검색]')}={len(_ds)} · "
     f"쓴 청크 ✓ {_tick('[문서 근거]') + _tick('[관련 원문]') + _tick('[문서 검색]')}")
_a = _jt["answer"]
_used_c = [c for c in _jt["collection"] if c.get("used")] + [c for c in _jt["doc_search"] if c.get("used")]
show("ⓓ 답 — 5줄 이내(줄 수 표시) · used_facts·used_chunks가 사실·청크에 표시된다 · 지시문 a-1.2 · l-1.1",
     _a["mode"] == "live" and _a["lines"] == 3 <= CQ.ANSWER_LINES and _a["used_chunks"] == [0]
     and len(_used_c) == 1 and sum(1 for f in _jt["facts"] if f["used"]) >= 1
     and "5줄" in (ROOT / "prompts/4.4_answer.md").read_text(encoding="utf-8")
     and re.search(r"^version: a-1\.2", (ROOT / "prompts/4.4_answer.md").read_text(encoding="utf-8"), re.M)
     and re.search(r"^version: l-1\.1", (ROOT / "prompts/4.1_link.md").read_text(encoding="utf-8"), re.M),
     f"{_a['lines']}줄 · 쓴 청크 {_a['used_chunks']} → 표시 {[c['chunk_id'][-10:] for c in _used_c]}")

# ────────────────────────────────────────────────────────────── ⓔ
print("\n■ B104 ⑤ 채점 — 링킹 사전만 vs 보충 · 문서 검색 doc@k (BM25 대조군 옆)")
_qs = [{"id": "B104-1", "type": "Q1", "q": "틈새는 어떻게 관리해?", "expected_path": "chunk",
        "expected_linked": [f"{L1[0].get('layer') or 'process'}:{L1[0]['canonical']}" if L1 else "-"], "expected_docs": ["CP01"]},  # 기대 변경(B153 ② — ⓑ가 닿은 노드를 기대로 · 가짜 임베딩 차례에 묶지 않는다)
       {"id": "B104-2", "type": "Q1", "q": "육안 점검 교체 주기는?", "expected_path": "chunk",
        "expected_docs": ["PPT03"]}]
with Live():
    _rows = GD.score_set(_qs)
_agg = GD.aggregate(_rows)
_ren = GD.render(_agg, _rows, src="B104", is_mock=True, skipped=[], k=8)
show("ⓔ 링킹 두 칸(사전만 0 → 보충 1) · 문서 검색 doc@k가 근거@k · BM25@k 옆에 계산된다",
     _agg["linking_dict_recall"] == 0.0 and _agg["linking_recall"] == 1.0 and _agg["doc_n"] == 2
     and _agg["doc_at_k"] is not None and _agg["bm25_at_k"] is not None
     and "link사전" in _ren and "link보충" in _ren and "doc@k" in _ren and "bm25@k" in _ren,
     f"링킹 사전 {_agg['linking_dict_recall']} → 보충 {_agg['linking_recall']} · doc@k {_agg['doc_at_k']} · "
     f"evid@k {_agg['evidence_at_k']} · bm25@k {_agg['bm25_at_k']}")

# ────────────────────────────────────────────────────────────── ⓕ
print("\n■ B104 ⑥ 쓰기 0")
# 기대 변경(B136 ③ — 벡터 저장은 `work/vectors/vectors.sqlite` 하나)
show("ⓕ 질의는 ③진실·②등록에 쓰지 않는다(해시 불변 — 벡터 저장은 work/vectors) — 시험이 더한 문서(B104ADD)는 제외",
     _P.vectors(V.DB).exists() and V.count("nodes") > 0 and V.count("chunks") > 0
     and not (_P.data() / "vectors").exists(), str(sorted(p.name for p in _P.vectors().iterdir())))
_hq = _state_hash()
with Live():
    CQ.answer("노칭 틈새는 어떻게 관리해?")
CQ.answer("육안 점검 교체 주기는?")
show("ⓕ 질의 전후 ③진실·②등록 해시 같다 (mock · 가짜 게이트웨이 둘 다)", _state_hash() == _hq)

# ────────────────────────────────────────────────────────────── B136 임베딩 정리
print("\n■ B136 — 노드 벡터 글 · 판정 좁히기 글 · 벡터 저장 · run.py vectors · show near (해시 벡터는 메커니즘만 · 순위는 가짜 임베딩)")
import base64 as _b64                                    # noqa: E402
import os as _os                                         # noqa: E402
import shutil as _sh                                     # noqa: E402
import subprocess as _sp                                 # noqa: E402
import sys as _sys                                       # noqa: E402
from array import array as _arr                          # noqa: E402
from cli import show_near as _SN, vectors as _CV         # noqa: E402
from core import matcher as _MT                          # noqa: E402
from core.llm import embeddings as _EM, narrow as _NR    # noqa: E402
from core.query import hybrid as _HY                     # noqa: E402
from router import discover as _disc                     # noqa: E402
_gr = {l: open_graph(l) for l in _disc()}
_cf = {l: load_config(l) for l in _gr}
_tx, _inf = _HY.node_texts(_gr, _cf)
_defs = [d for c in _cf.values() for d in (c.get("categories") or {}).values() if isinstance(d, str) and len(d) > 8]
_sh.rmtree(_P.vectors(), ignore_errors=True)
V.begin()
_HY.warm(_gr, _cf)
_nch = V.STATS["chunks"]["made"]
_lay0, _nid0 = next((l, i) for l, g in sorted(_gr.items()) for i in sorted(g.nodes))
_gr[_lay0].nodes[_nid0].setdefault("aliases", []).append({"surface": "B136 시험 별칭"})   # 메모리만 — 저장하지 않는다
V.begin()
_HY.warm(_gr, _cf)
show("B136 ⓐ 노드 벡터 글에 정의문 글자 0 · 별칭 하나 바꾼 노드만 다시 만든다(1) · 청크 벡터는 다시 쓴다(다시 쓴 수 = 청크 전부)",
     _defs and not any(d in t for t in _tx.values() for d in _defs) and V.STATS["nodes"]["made"] == 1
     and V.STATS["chunks"]["made"] == 0 and V.STATS["chunks"]["reused"] == _nch > 0, f"{V.STATS}")
_gr[_lay0].nodes[_nid0]["aliases"].pop()


def _bag(text, dim=1024):
    """가짜 임베딩 — 낱말 주머니(같은 낱말이 든 글끼리 가깝다 · 순위 판정 전용)."""
    v = [0.0] * dim
    for w in re.split(r"[\s|,:]+", text or ""):
        if w:
            v[int(hashlib.sha256(w.encode("utf-8")).hexdigest(), 16) % dim] += 1.0
    return v


_pool = [{"id": f"P{i:02d}", "canonical": f"노칭 부품 {i}", "aliases": ["NP-100"] if i == 7 else [f"별칭{i}"],
          "category": "Unit"} for i in range(15)]
_emb0 = _EM.embed
_EM.embed = _bag
_NR.set_narrow("embed")
_MT._VEC.clear()
try:
    _kept, _how = _MT._narrow("NP-100", _pool, 5)
finally:
    _EM.embed = _emb0
    _NR.set_narrow(None)
    _MT._VEC.clear()
_sp4 = (ROOT / "docs" / "spec" / "4_쓰기절차.md").read_text(encoding="utf-8")
show("B136 ⓑ 판정 좁히기 후보 글 = 이름 + 별칭(정의문 0) · 별칭 표기와 같은 후보가 맨 앞(가짜 임베딩) · 주석 · 명세가 같은 규칙",
     _MT.cand_text(_pool[7]) == "노칭 부품 7 | 별칭: NP-100" and _how == "임베딩" and _kept[0]["id"] == "P07"
     and "이름 + 별칭" in _sp4 and "임베딩 대상은 노드의 canonical과 정의문" not in _sp4,
     f"{_how} · 앞 {[c['id'] for c in _kept[:3]]}")
_env = {**_os.environ, "USE_MOCK": "1"}


def _vec_cmd(*a):
    r = _sp.run([_sys.executable, str(ROOT / "run.py"), "vectors", *a], cwd=str(ROOT), env=_env, capture_output=True,
                text=True, stdin=_sp.DEVNULL)
    return r.returncode, r.stdout + r.stderr


_sh.rmtree(_P.vectors(), ignore_errors=True)
_r1 = _vec_cmd()
_r2 = _vec_cmd()
_sh.rmtree(_P.vectors(), ignore_errors=True)
_save0, _many0, _k = V.SAVE_EVERY, _EM.embed_many, {"n": 0}


def _cut(texts, batch=32):
    _k["n"] += 1
    if _k["n"] > 2:
        raise KeyboardInterrupt("끊음(시험)")
    return _many0(texts, batch)


V.SAVE_EVERY, _EM.embed_many = 16, _cut
try:
    with contextlib.redirect_stdout(io.StringIO()):
        _CV.main([])
except KeyboardInterrupt:
    pass
V.SAVE_EVERY, _EM.embed_many = _save0, _many0
_mid = V.count()
_r3 = _vec_cmd()
_m1 = [int(x.replace(",", "")) for x in re.findall(r"만든 ([\d,]+)", _r1[1])]
_m2 = [int(x.replace(",", "")) for x in re.findall(r"만든 ([\d,]+)", _r2[1])]
_m3 = [int(x.replace(",", "")) for x in re.findall(r"만든 ([\d,]+)", _r3[1])]
show("B136 ⓒ run.py vectors — 처음 n 만들고 · 다시 치면 0 · 중간에 끊고 다시 치면 남은 것만 · 끝 줄이 선다",
     _r1[0] == 0 and sum(_m1) > 0 and _r2[0] == 0 and sum(_m2) == 0 and 0 < _mid < sum(_m1)
     and sum(_m3) == sum(_m1) - _mid and all("벡터 끝 —" in r[1] for r in (_r1, _r2, _r3)),
     f"처음 {_m1} · 다시 {_m2} · 끊긴 뒤 저장 {_mid} · 남은 것 {_m3}")
fresh()                                                   # 클린 바닥 — 저장도 지운다
_clean_gone = not _P.vectors(V.DB).exists()
V.begin()
_HY.warm()
_before = V.count()
init.init(fresh_=True)                                    # 시험 대상: 기본 fresh — 벡터는 남는다(보존 일곱)
_kept_fresh = V.count()
for _ly in ("process", "quality"):
    bootstrap(_ly, echo=False)
for _dd in DOCS:
    run_document(load(_dd))
finalize()                                                # 같은 문서를 다시 넣는다(재구축)
V.begin()
_HY.warm()
_after = {k: dict(v) for k, v in V.STATS.items()}
_oldtxt = "B136 옛 캐시 글"
_old = _P.vectors("nodes.json")
_P.ensure(_old)
_old.write_text(json.dumps({"backend": _EM.backend_id(), "kind": "nodes", "items": {"x": {
    "ver": V._ver(_oldtxt), "v": _b64.b64encode(_arr("f", _EM.embed(_oldtxt)).tobytes()).decode("ascii")}}}), encoding="utf-8")
V.begin()
V.vectors("nodes", {"x": _oldtxt})
_mig = (dict(V.STATS["nodes"]), _old.exists())
show("B136 ⓓ 기본 fresh는 벡터를 남기고 재구축 뒤 다시 만든 벡터 0(글이 같다) · 클린은 지운다 · 옛 JSON을 한 번 옮겨 담는다",
     _clean_gone and _before > 0 and _kept_fresh == _before and _after and all(x["made"] == 0 for x in _after.values())
     and _mig[0]["made"] == 0 and _mig[0]["reused"] == 1 and not _mig[1],
     f"클린 뒤 저장 없음 {_clean_gone} · fresh 전 {_before} → 뒤 {_kept_fresh} · 재구축 뒤 {_after} · 옛 JSON {_mig}")
_ch0 = V.count("chunks")
_w = ("import sys; from core.query import vectors as V; V.SAVE_EVERY = 16; "
      "V.vectors('chunks', {f'{sys.argv[1]}{i}': f'동시 {sys.argv[1]} {i}' for i in range(200)})")
_ps = [_sp.Popen([_sys.executable, "-c", _w, t], cwd=str(ROOT), env=_env, stdin=_sp.DEVNULL, stdout=_sp.DEVNULL,
                 stderr=_sp.PIPE, text=True) for t in ("갑", "을")]
_errs = [p.communicate()[1] for p in _ps]
import sqlite3 as _sq                                    # noqa: E402
_con = _sq.connect(str(_P.vectors(V.DB)))
_ok = _con.execute("PRAGMA integrity_check").fetchone()[0]
_con.close()
show("B136 ⓔ 동시 저장 — 두 프로세스가 같은 저장에 동시에 쓰면 깨짐 0 · 둘의 몫이 다 남는다",
     all(p.returncode == 0 for p in _ps) and _ok == "ok" and V.count("chunks") == _ch0 + 400,
     f"{_ok} · {_ch0} → {V.count('chunks')} · {_errs}")
_gr3 = {l: open_graph(l) for l in _disc()}
_cf3 = {l: load_config(l) for l in _gr3}
_t3, _i3 = _HY.node_texts(_gr3, _cf3)
_raw3 = {i: n for g in _gr3.values() for i, n in g.nodes.items()}
from core.build.naming import scope_of as _SO3                    # noqa: E402 — 부모 경로 칸(B142 ⑤ · 기대 변경)
_key = lambda i: (_i3[i]["layer"], _i3[i]["category"], _SO3(_raw3[i]))   # noqa: E731
_grp = {}
for _i in sorted(_i3):
    _grp.setdefault(_key(_i), []).append(_i)
_big = next(k for k, v in sorted(_grp.items(), key=str) if len(v) >= 2)
_a, _b = _grp[_big][:2]
_d = next(i for i in sorted(_i3) if _i3[i]["layer"] == _big[0] and _i3[i]["category"] != _big[1])
_e = next((i for i in sorted(_i3) if _key(i)[:2] == _big[:2] and _key(i)[2] != _big[2]), None)
_same = {_t3[x] for x in (_a, _b, _d, _e) if x}
_order = {t: n for n, t in enumerate(sorted(_t3.values()))}
_EM.embed_many = lambda texts, batch=32: [[1.0 if (t in _same and j == 0) or (t not in _same and j == _order.get(t, 0) + 1)
                                           else 0.0 for j in range(len(_order) + 2)] for t in texts]
_sh.rmtree(_P.vectors(), ignore_errors=True)               # 심을 벡터만 담는다(앞 시험의 해시 벡터를 비운다)
V.vectors("nodes", _t3)
_EM.embed_many = _many0


def _fp():
    return {p.relative_to(_P.home()).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(_P.home().rglob("*")) if p.is_file()}


_f0 = _fp()
_bn = io.StringIO()
with contextlib.redirect_stdout(_bn):
    _SN.cmd_near(["--top", "50"])
_f1 = _fp()
_lines = [ln for ln in _bn.getvalue().splitlines() if "↔" in ln]
_nm = lambda i: _i3[i]["canonical"]                      # noqa: E731
show("B136 ⓕ 중복 후보 보고 — 심은 가까운 쌍이 나오고(가짜 임베딩) 다른 카테고리 · 다른 부모 쌍은 나오지 않는다 · 상태 지문 전후 같다(쓰기 0)",
     len(_lines) == 1 and _nm(_a) in _lines[0] and _nm(_b) in _lines[0] and "100.0%" in _lines[0] and _f0 == _f1
     and "ops merge" in _bn.getvalue(), f"쌍 {_lines} · 다른 카테고리 {_nm(_d)} · 다른 부모 {_nm(_e) if _e else '없음'}")

# ────────────────────────────────────────────────────────────── B153
print("\n■ B153 — 임베딩 재료: 표 청크 머리 · 노드 글 · 입력 길이 (창작 표본 · mock — 수치는 근거가 아니다)")
from core.state.ids import doc_hash as _DH153                     # noqa: E402
with contextlib.redirect_stdout(io.StringIO()):
    fresh()
_raw153 = TMP / "b153"
_raw153.mkdir(exist_ok=True)
_ing153 = []
for _src, _dst, _dt in (("CP01.xlsx", "B153CP.xlsx", "cp"), ("PFMEA01.xlsx", "B153FM.xlsx", "pfmea")):
    _sh.copy(_FXR.RAW / _src, _raw153 / _dst)                    # 원본 엑셀에서 파싱하는 길(같은 doc_id 충돌을 피해 이름만 바꾼다)
    _ing153.append(_sp.run([_sys.executable, str(ROOT / "run.py"), "ingest-file", str(_raw153 / _dst), "--doc-type", _dt,
                            "--allow-mock"], cwd=str(ROOT), env=_env, capture_output=True, text=True, stdin=_sp.DEVNULL,
                           timeout=600).returncode)
_ch153 = store.read(store.CHUNKS, {"chunks": {}})["chunks"]
_ev153 = {d: json.loads(_P.parsed(f"{d}.json").read_text(encoding="utf-8")) for d in ("B153CP", "B153FM")}
_cpk = next(k for k, c in _ch153.items() if c["doc_id"] == "B153CP" and c["text"] == "공정 정지 후 금형 점검")
_fm = next(c for c in _ch153.values() if c["doc_id"] == "B153FM" and c["meta"].get("field") == "prevention_control")
_bm153 = _HY.doc_search("노칭 정밀도 대응계획", 20, 101)          # 문턱 101% — 임베딩 0 · BM25만
show("B153 ⓐ 표 청크 — 원본 엑셀(ingest-file)의 봉투 field_labels = 실물 헤더 글 · meta.label · owner · 재료 「[열 머리 · 주인 · 좌표] 원문」"
     "(행 위치 0) · 청크 id · 문서 해시 그대로 · 열 이름 · 주인만 든 질문을 BM25가 같은 재료로 찾는다",
     _ing153 == [0, 0] and _ev153["B153FM"]["field_labels"].get("prevention_control") == "예방관리"
     and _fm["meta"].get("label") == "예방관리" and _fm["meta"].get("owner")
     and _HY._chunk_text(_ch153[_cpk]) == "[대응계획 · 노칭 정밀도 · 노칭] 공정 정지 후 금형 점검"
     and _ch153[_cpk]["section"] not in _HY._chunk_text(_ch153[_cpk]) and _cpk.endswith("-대응계획")
     and _DH153(_ev153["B153CP"]) == _DH153({k: v for k, v in _ev153["B153CP"].items() if k != "field_labels"})
     and any(x["chunk_id"] == _cpk and x["by"] == ["bm25"] for x in _bm153),
     f"인입 rc {_ing153} · {_fm['meta']} · {_HY._chunk_text(_ch153[_cpk])} · BM25 {[x['chunk_id'][-12:] for x in _bm153][:5]}")
_coord153 = __import__("parser.normalizer", fromlist=["coord_text"]).coord_text
_old153 = lambda c: (lambda h: f"[{h}] {c.get('text') or ''}" if h else (c.get("text") or ""))(   # noqa: E731 — B152까지의 재료
    " · ".join(x for x in (c.get("section"), _coord153(c.get("process_ref"))) if x))
_pr153 = {k: c for k, c in _ch153.items() if not c["meta"].get("field") and not c["meta"].get("subject")}
V.vectors("chunks", {k: _old153(c) for k, c in _pr153.items()})
V.begin()
V.vectors("chunks", {k: _HY._chunk_text(c) for k, c in _pr153.items()})
show("B153 ⓑ 산문 청크 — 머리 그대로(재료가 같아 만든 수 0 · 다시 씀 = 전부) · PPT 다루는 것은 머리 끝에",
     _pr153 and V.STATS["chunks"]["made"] == 0 and V.STATS["chunks"]["reused"] == len(_pr153)
     and _HY._chunk_text({"text": "t", "section": "S", "meta": {"subject": "노칭 커터"}}) == "[S · 노칭 커터] t",
     f"산문 {len(_pr153)} · {V.STATS['chunks']}")
_gr153 = {l: open_graph(l) for l in _disc()}
_t153, _i153 = _HY.node_texts(_gr153)
_by153 = {_i153[i]["canonical"]: (t, _i153[i]["where"]) for i, t in _t153.items()}
_np = _by153.get("노칭::노칭 프레스::노칭 정밀도", ("", ""))
show("B153 ⓒ 노드 글 — 「맨 이름 | Property | 별칭 | 위치: 노칭 > 노칭 프레스」(관계 이름 · 화살표 0 · 경로 한 번) · 극성 갈래 마디 · "
     "골격 노드는 골격 경로 · 화면의 where는 엣지 꼴 그대로",
     _np[0].startswith("노칭 정밀도 | Property | ") and _np[0].endswith(" | 위치: 노칭 > 노칭 프레스") and "→" not in _np[0]
     and "노칭::" not in _np[0] and "—has_property→" in _np[1]
     and any(c.startswith("탭용접::cathode::") and "위치: 탭용접 > cathode" in t for c, (t, _w) in _by153.items())
     and _by153.get("탭용접::cathode", ("",))[0].endswith("| 위치: 탭용접"),
     f"{_np[0]} · 화면 {_np[1][:60]}")
_knobs({"embed_max_chars": 40})
_r153 = [_vec_cmd(), _vec_cmd()]
_d153 = _sp.run([_sys.executable, str(ROOT / "run.py"), "show", "dist", "embed"], cwd=str(ROOT), env=_env, capture_output=True,
                text=True, stdin=_sp.DEVNULL)
_cut153 = [int(x.replace(",", "")) for x in re.findall(r"(?:노드|청크) ([\d,]+)", (re.findall(r"넘어 자른 글 — (.*?)\(", _r153[0][1]) or [""])[0])]
_over153 = re.findall(r"넘은 글 ([\d,]+)", _d153.stdout)
_fit153 = _EM.fit(_HY._chunk_text(_fm))[0]
show("B153 ⓔ 길이 — 한도를 작게(embed_max_chars 40) · 넘는 글은 머리를 지키고 꼬리가 잘려 한도 안 · 끝 줄 「넘어 자른 글 — 노드 a · 청크 b"
     "(가장 긴 m)」 · show dist embed가 같은 수 · 다시 치면 재사용(만든 수 0)",
     _r153[0][0] == 0 and _cut153 and sum(_cut153) > 0 and "가장 긴" in _r153[0][1] and "한도 40자(손잡이)" in _r153[0][1]
     and _over153 and int(_over153[0].replace(",", "")) == sum(_cut153)
     and sum(int(x.replace(",", "")) for x in re.findall(r"만든 ([\d,]+)", _r153[1][1])) == 0
     and _fit153.startswith("[예방관리 · ") and len(_fit153) == 40 < len(_HY._chunk_text(_fm)),
     f"끝 줄 자른 {_cut153} · dist 넘은 글 {_over153} · 맞춘 글 {_fit153!r} · rc {_d153.returncode}")
_knobs({})

# ────────────────────────────────────────────────────────────── B161 질의 벡터 색인
print("\n■ B161 — 질의 벡터 색인(id → 글 해시 · 지문): 같으면 조립 0 · 바뀐 것만 · 정리는 --gc만 (해시 벡터는 메커니즘만)")
import sqlite3 as _sq161                                  # noqa: E402
from cli import show_knobs as _SK                         # noqa: E402
from core.query import bm25 as _BM                        # noqa: E402


def _db(q, a=(), write=False):
    c = _sq161.connect(str(_P.vectors(V.DB)))
    try:
        got = c.execute(q, a).fetchall()
        if write:
            c.commit()
        return got
    finally:
        c.close()


def _made161(out):
    return sum(int(x.replace(",", "")) for x in re.findall(r"만든 ([\d,]+)", out))


_calls161 = {"n": 0}
_nt0, _ct0 = _HY.node_texts, _HY._chunk_text


def _counted(on):
    """재료 함수 둘을 세는 대역으로 바꾼다(on) · 되돌린다(off) — 셈만 하고 결과는 그대로."""
    if not on:
        _HY.node_texts, _HY._chunk_text = _nt0, _ct0
        return

    def wrap(f):
        def w(*a, **k):
            _calls161["n"] += 1
            return f(*a, **k)
        return w
    _HY.node_texts, _HY._chunk_text = wrap(_nt0), wrap(_ct0)


_sh.rmtree(_P.vectors(), ignore_errors=True)
_r161 = _vec_cmd()
_gr161 = {l: open_graph(l) for l in _disc()}
_t161 = _HY.node_texts(_gr161)[0]
_ch161 = store.read(store.CHUNKS, {"chunks": {}})["chunks"]
_bid161 = _EM.backend_id()
_ix161 = {(k, i): (ver, t) for k, i, ver, t in _db("SELECT kind, id, ver, text FROM idx WHERE backend = ?", (_bid161,))}
_vs161 = {r[0] for r in _db("SELECT ver FROM v WHERE backend = ?", (_bid161,))}
_want161 = ({("nodes", i): _EM.fit(t)[0] for i, t in _t161.items()}
            | {("chunks", c): _EM.fit(_HY._chunk_text(x))[0] for c, x in _ch161.items()})
_mt161 = _db("SELECT kind FROM meta WHERE backend = ? ORDER BY kind", (_bid161,))
show("B161 ⓐ 색인이 선다 — idx 행 = 노드 + 청크 · 행마다 글 해시가 v에 있다 · 글 = 재료 함수의 출력(자른 뒤) · meta 두 줄",
     _r161[0] == 0 and len(_ix161) == len(_t161) + len(_ch161) > 0 and all(v in _vs161 for v, _t in _ix161.values())
     and {k: t for k, (_v, t) in _ix161.items()} == _want161 and _mt161 == [("chunks",), ("nodes",)],
     f"idx {len(_ix161)} = 노드 {len(_t161)} + 청크 {len(_ch161)} · meta {_mt161}")
_q161 = "노칭 금형 교체 주기는?"
_qv161 = V.question(_q161)
_lv161 = V.lookup("nodes", _t161)                          # 색인 전의 길 — 글 → 해시 → 벡터(대조군)
_old_lk = [(i, s) for s, i in sorted(((round(_EM.cosine(_qv161, v), 4), i) for i, v in _lv161.items()),
                                     key=lambda t: (-t[0], t[1]))[:5]]
_ctx161 = {c: _HY._chunk_text(x) for c, x in _ch161.items()}
_cv161 = V.lookup("chunks", _ctx161)
_old_es = {c: round(_EM.cosine(_qv161, v), 4) for c, v in _cv161.items()}
_old_bs = dict(_BM.search(_q161, 8, index=_BM.Index(_ctx161)))
_counted(True)
try:
    V.begin()
    _HY.warm()
    _w161 = V.index_line()
    _lk161 = _HY.link_candidates(_q161, _gr161, {}, 5)
    _ds161 = _HY.doc_search(_q161, 8, 0)
    _an161 = CQ.answer(_q161)
finally:
    _counted(False)
_top_es = [(c, x) for c, x in sorted(_old_es.items(), key=lambda t: (-t[1], t[0])) if x * 100 >= 0][:8]
show("B161 ⓑ 지문이 같으면 조립 0 — 두 번째 벡터 준비 · 링킹 후보 · 문서 검색 · 질의 한 건에서 재료 함수 호출 0 · 「색인 — 그대로」 · "
     "링킹 후보 · 문서 검색이 색인 전의 길(글 → 해시 → 벡터 · 같은 재료의 BM25)과 같은 순위 · 같은 점수",
     _calls161["n"] == 0 and _w161 == "색인 — 노드 그대로 · 청크 그대로"
     and [(x["id"], x["score"]) for x in _lk161] == _old_lk
     and {x["chunk_id"] for x in _ds161} == {c for c, _s in _top_es} | set(_old_bs)
     and all(x["embed"] == _old_es[x["chunk_id"]] for x in _ds161 if "embed" in x["by"])
     and all(x["bm25"] == round(_old_bs[x["chunk_id"]], 4) for x in _ds161 if "bm25" in x["by"]) and _an161.get("path"),
     f"재료 호출 {_calls161['n']} · {_w161} · 링킹 {len(_lk161)} · 문서 검색 {len(_ds161)}")
_d161 = []
_mv161 = _HY.MATERIAL_VER
for _knob, _mat in (({"embed_max_chars": 3999}, _mv161), ({}, "m-시험")):   # 하나씩만 바꾼다 — 한도 · 재료 판
    _knobs(_knob)
    _HY.MATERIAL_VER = _mat
    V.begin()
    _HY.warm()
    _d161 += [V.index_line(), sum(s["made"] for s in V.STATS.values())]
    _knobs({})
    _HY.MATERIAL_VER = _mv161
    _HY.warm()                                             # 되돌린 입력으로 색인을 다시 맞춘다
show("B161 ⓓ 재료 판(MATERIAL_VER) · 입력 한도(embed_max_chars)가 바뀌면 지문이 같아도 색인을 다시 쓴다(글이 같으면 만든 벡터 0)",
     _d161 == ["색인 — 노드 다시 씀 · 청크 다시 씀", 0] * 2, f"{_d161}")
_lay161, _nid161 = next((l, i) for l, g in sorted(_gr161.items()) for i in sorted(g.nodes) if i in _t161)
_db("INSERT OR IGNORE INTO v VALUES ('다른 판(시험)', 'x161', 'nodes', ?)", (b"\0\0\0\0",), write=True)
_n161 = V.count()


def _alias161(add):
    g = open_graph(_lay161)
    al = g.nodes[_nid161].setdefault("aliases", [])
    al.append({"surface": "B161 시험 별칭"}) if add else al.pop()
    g.save()


_alias161(True)
_e161 = [_vec_cmd()]
_kept161 = V.count()
_e161 += [_vec_cmd("--gc"), _vec_cmd()]
_alias161(False)
_vec_cmd()                                                 # 별칭을 되돌린 그래프로 색인 · 벡터를 맞춘다(ⓕ의 바닥)
_orph = [re.findall(r"참조 없는 벡터 ([\d,]+)\(이 판\)", o) for _rc, o in _e161]
_gone = re.findall(r"지운 벡터 ([\d,]+)", _e161[1][1])
_other = _db("SELECT COUNT(*) FROM v WHERE backend = '다른 판(시험)'")[0][0]
show("B161 ⓔ 정리 — 별칭을 바꿔 저장하면 끝 줄 「참조 없는 벡터 1」(지우지 않는다) · --gc 「지운 벡터 1」 → 다시 0 · 다른 판의 행은 남는다",
     [r for r, _o in _e161] == [0, 0, 0] and _orph == [["1"], ["0"], ["0"]] and _gone == ["1"] and _kept161 == _n161 + 1
     and _other == 1, f"참조 없는 {_orph} · 지운 {_gone} · 저장 {_n161} → {_kept161} · 다른 판 {_other}")
_db("DROP TABLE idx", write=True)
_db("DROP TABLE meta", write=True)
_f161 = _vec_cmd()
show("B161 ⓕ 옛 저장(idx · meta 없음)에서 vectors → 색인이 서고 만든 벡터 0(벡터는 다 다시 쓴다)",
     _f161[0] == 0 and _made161(_f161[1]) == 0 and "색인 — 노드 다시 씀 · 청크 다시 씀" in _f161[1]
     and len(_db("SELECT * FROM meta")) >= 2, f"{_f161[1].splitlines()[-1][:160]}")
_knobs({"embed_max_chars": 40})
_g161 = _vec_cmd()
_fp161 = _fp()
_counted(True)
_out161 = io.StringIO()
try:
    with contextlib.redirect_stdout(_out161):
        _SN.cmd_near([])
        _SK.dist_embed()
finally:
    _counted(False)
_fp161b = _fp()                                            # 손잡이를 되돌리기 전에 잰다(knobs.json도 상태다)
_knobs({})
_cutg = sum(int(x.replace(",", "")) for x in re.findall(r"(?:노드|청크) ([\d,]+)",
                                                        (re.findall(r"넘어 자른 글 — (.*?)\(", _g161[1]) or [""])[0]))
_overg = re.findall(r"넘은 글 ([\d,]+)", _out161.getvalue())
show("B161 ⓖ 읽기 전용 — show near · show dist embed가 색인에서 읽는다(재료 함수 호출 0 · 다시 재지 않는다) · 상태 지문 전후 같다 · "
     "dist 넘은 수 = vectors 끝 줄 「넘어 자른 글」",
     _calls161["n"] == 0 and _fp161 == _fp161b and _overg and int(_overg[0].replace(",", "")) == _cutg > 0
     and "색인이 지금 입력과 달라" not in _out161.getvalue(), f"재료 호출 {_calls161['n']} · 넘은 {_overg} · 자른 {_cutg}")
_db("DELETE FROM meta", write=True)                        # 두 프로세스가 동시에 색인을 갈아 끼운다
_ph = [_sp.Popen([_sys.executable, str(ROOT / "run.py"), "vectors"], cwd=str(ROOT), env=_env, stdin=_sp.DEVNULL,
                 stdout=_sp.DEVNULL, stderr=_sp.PIPE, text=True) for _ in range(2)]
_eh = [p.communicate()[1] for p in _ph]
_ih = _db("PRAGMA integrity_check")[0][0]
_gr161 = {l: open_graph(l) for l in _disc()}
show("B161 ⓗ 동시 쓰기 — 두 프로세스가 같은 저장의 색인을 함께 갈아 끼워도 깨짐 0 · 색인 행 = 노드 + 청크 · meta 두 줄(마지막에 쓴다)",
     all(p.returncode == 0 for p in _ph) and _ih == "ok"
     and len(_db("SELECT id FROM idx WHERE backend = ?", (_bid161,))) == len(_HY.node_texts(_gr161)[0]) + len(_ch161)
     and len(_db("SELECT kind FROM meta WHERE backend = ?", (_bid161,))) == 2, f"{_ih} · {[e[-120:] for e in _eh if e]}")
_ser = [f"{p.relative_to(ROOT)}:{n}" for d in ("core", "cli", "parser") for p in sorted((ROOT / d).rglob("*.py"))
        if p.as_posix() != (ROOT / "core" / "graph.py").as_posix()
        for n, ln in enumerate(p.read_text(encoding="utf-8").splitlines(), 1)
        if re.search(r"(dumps|sha256)\([^\n]*\.(nodes|edges)\b", ln)]
show("B161 ⓘ 경계 — graph.py 밖에서 그래프 노드 · 엣지를 직렬화 · 해시하는 코드 0(정적 점검 · 지문은 GraphStore.content_fp 하나)",
     not _ser, f"{_ser}")
# 기대 변경(B161 PR 리뷰 — 동시 실행): 확인과 읽기가 한 판 · 갈아 끼우는 쪽 공유 잠금 · gc 배타 잠금
V.begin()
_HY.warm(_gr161)
_lv2 = V.lookup("nodes", _HY.node_texts(_gr161)[0])
_want2 = [(i, s) for s, i in sorted(((round(_EM.cosine(_qv161, v), 4), i) for i, v in _lv2.items()),
                                    key=lambda t: (-t[0], t[1]))[:5]]
_ens0 = _HY._ensure
_HY._ensure = lambda *a, **k: (_ens0(*a, **k), V.index_write("nodes", {"남의 판": "남의 판 글"}, {}, "남의 지문", _HY.MATERIAL_VER))[0]
try:
    _race = [(x["id"], x["score"]) for x in _HY.link_candidates(_q161, _gr161, {}, 5)]   # 맞춘 뒤 · 읽기 전에 남이 갈아 끼웠다
finally:
    _HY._ensure = _ens0
_HY.warm(_gr161)
_pend = ("import sys, time; from core.query import vectors as V; from core.llm import embeddings as E; b = E.backend_id()\n"
         "with V._gc_guard(exclusive=False):\n"
         "    import sqlite3; from core import paths; c = sqlite3.connect(str(paths.vectors(V.DB)))\n"
         "    c.execute(\"INSERT OR IGNORE INTO v VALUES (?, 'pend161', 'nodes', ?)\", (b, bytes(4))); c.commit()\n"
         "    open(sys.argv[1], 'w').close(); time.sleep(1.0)\n"
         "    c.execute(\"INSERT OR REPLACE INTO idx VALUES (?, 'nodes', 'pend161', 'pend161', NULL, NULL, 't', 1, 0)\", (b,)); c.commit()\n")
_mark = TMP / "pend161.held"
_pp = _sp.Popen([_sys.executable, "-c", _pend, str(_mark)], cwd=str(ROOT), env=_env, stdin=_sp.DEVNULL, stdout=_sp.DEVNULL,
                stderr=_sp.PIPE, text=True)
for _ in range(200):
    if _mark.exists() or _pp.poll() is not None:
        break
    __import__("time").sleep(0.05)
_gone2 = V.gc()                                            # 새 벡터가 색인에 오를 때까지 기다린다
_pe = _pp.communicate()[1]
_kept2 = _db("SELECT COUNT(*) FROM v WHERE ver = 'pend161'")[0][0]
_db("DELETE FROM idx WHERE id = 'pend161'", write=True)
_db("DELETE FROM v WHERE ver = 'pend161'", write=True)
show("B161 ⓗ′ 동시 실행 — 맞춘 뒤 다른 입력으로 색인이 갈렸어도 링킹 후보는 제 입력의 결과(남의 판을 읽지 않는다) · 다른 프로세스가 새 벡터를 "
     "색인에 올리는 동안 --gc는 기다리고 그 벡터를 지우지 않는다",
     _race == _want2 and _pp.returncode == 0 and _mark.exists() and _kept2 == 1,
     f"후보 {_race[:2]} · 기대 {_want2[:2]} · gc 지운 {_gone2} · 새 벡터 남음 {_kept2} · {_pe[-120:]}")

for d in set(_made):
    (HINTS / f"{d}.json").unlink(missing_ok=True)
import shutil                                       # noqa: E402
shutil.rmtree(TMP, ignore_errors=True)
done()
