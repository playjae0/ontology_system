# -*- coding: utf-8 -*-
"""G6 — **산문의 값** — 「항목 = 값」을 주인 아래 노드의 속성으로 · 값을 받는 카테고리 선언 · 추출 판은 config 칸 해시 (B120).

창작 표본(TW 꼴 — 유닛 둘 아래 같은 이름 사양 · 수치는 근거가 아니다 [정정] 50)을 mock 힌트(`<doc>@structure.json` · `<doc>.json`
— 값은 지시문 「그 항목에 건 값 · 원문 글자 그대로」의 대역)로 넣는다. 대표 사례 하나 · 항목당 어서션 하나(가벼운 회차):
  ⓐ 값 — `노칭::언와인더 unit::장력` · `노칭::리와인더 unit::장력`이 각자 원문 값 · 출처 청크(충돌 0) · 같은 청크의 날짜 · 쪽수는
     어디에도 값으로 앉지 않는다
  ⓑ 버림 · 충돌 — 원문에 없는 값은 버림(결함 로그 · 행의 `value.verdict` · 맞춤표 수) · 다른 문서가 같은 노드 같은 맥락에 다른 값을
     내면 `spec_conflict`(있던 값 그대로)
  ⓒ 선언 · 판 — `relation_patterns` · `config_version`만 바꾸면 추출 재사용(힌트를 지워도 값이 그대로) · 정의문을 바꾸면 다시 뽑는다
     (힌트가 없으니 값이 사라진다) · 선언을 지우면 힌트에 값이 있어도 0
  ⓓ 재구축 — 값이 재생된다(추출 · 판정 LLM 0 · ⓒ 첫 재구축) · 네 벌 A · B 같음 · C는 값 몫은 마감 한 줄의 네 벌 단계가 잰다
"""
from __future__ import annotations

import contextlib
import io
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, _rw, done   # noqa: F401

from g6_common import _register, _unregister             # noqa: E402
from openpyxl import Workbook                           # noqa: E402
import run as RUN                                       # noqa: E402
from cli import ingest as IG, rebuild as RB             # noqa: E402
from core.build import ledger as LG, reconcile as RC, structure as ST   # noqa: E402
from core.state.world import World                      # noqa: E402
from parser import pipeline, struct_map, tuning         # noqa: E402
from parser.adapters import basic_prose_xlsx as BX      # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="b120_"))
HINTS = ROOT / "tests" / "fixtures" / "extract_hints"
DT = "b120_process"
SHEET = "TW 사양"
UNITS = [{"text": "언와인더 unit", "depth": 1, "role": "유닛", "name": "언와인더 unit", "category": "Unit"},
         {"text": "1) 기본사양", "depth": 2, "role": "사양 묶음"}, {"text": "2) 지정사양", "depth": 2, "role": "사양 묶음"},
         {"text": "리와인더 unit", "depth": 1, "role": "유닛", "name": "리와인더 unit", "category": "Unit"},
         {"text": "1) 기본사양", "depth": 2, "role": "사양 묶음"}, {"text": "2) 지정사양", "depth": 2, "role": "사양 묶음"}]
#: 창작 표본 — 둘째 문서는 언와인더 장력만 다른 값(충돌)
DOCS = {"B120TW": ["언와인더 unit", "1) 기본사양", "속도는 100ppm이다.", "2) 지정사양", "장력은 30N이다. 작성일 2026-10-16 · 3쪽",
                   "리와인더 unit", "1) 기본사양", "속도는 80ppm이다.", "2) 지정사양", "장력은 20N이다."],
        "B120TW2": ["언와인더 unit", "2) 지정사양", "장력은 35N이다."]}
#: 지어낸 값(원문에 없다 — 버려져야 한다) — `(문서, 소속, 말) → 값`
INVENTED = {("B120TW", "리와인더 unit", "속도"): "85ppm"}
WORDS = ("속도", "장력")


def _value(word, text):
    m = re.search(re.escape(word) + r"\S*\s*(\d[\d.,]*\s*[A-Za-z%]+)", text or "")
    return m.group(1) if m else None


def _hints(did, p):
    """구조 힌트로 자른 조각마다 추출 힌트(소속 = 사슬의 가장 깊은 대상 · 값 = 말 뒤의 숫자+단위 · 지어낸 값은 INVENTED)."""
    env = pipeline.parse(BX, did, str(p), closed_list=[], sheet_roles={SHEET: "prose"},
                         structure=ST.structurer(["process"])).envelope
    out = {}
    for c in env["chunks"]:
        own = next((x for x in reversed(c["meta"].get("structure") or []) if x.get("name")), None)
        ents = []
        for w in WORDS:
            if w + "은" not in c["text"] and w + "는" not in c["text"]:
                continue
            v = INVENTED.get((did, own and own["name"], w)) or _value(w, c["text"])
            ents.append({"surface": w, "category": "Property", "value": v,
                         "belongs_to": {"name": own["name"], "category": own["category"], "from": "제목"} if own else None})
        if ents:
            out[c["source_locator"]] = {"entities": ents, "relations": [], "attach": []}
    struct_map.invalidate(did)
    return out


def _put_hints(did):
    (HINTS / f"{did}.json").write_text(json.dumps(HINTS_OF[did], ensure_ascii=False), encoding="utf-8")


def _drop_hints():
    for did in DOCS:
        (HINTS / f"{did}.json").unlink(missing_ok=True)


def _spec(canonical):
    n = next((n for _l, n in World().nodes() if n["canonical"] == canonical and n.get("status") != "obsolete"), None)
    return [(it["value"], tuple(it["provenance"])) for it in ((n or {}).get("attrs") or {}).get("spec") or []], n


def _prose_values():
    """산문 문서가 실은 값 전부 — `{(canonical, 값)}`(출처가 이 시험의 문서)."""
    return {(n["canonical"], it["value"]) for _l, n in World().nodes() for it in (n.get("attrs") or {}).get("spec") or []
            if any(str(p).split("#")[0] in DOCS for p in it["provenance"])}


def _rebuild():
    with contextlib.redirect_stdout(io.StringIO()):
        return RB.execute(RB.make_plan(), bootstrap=RUN.cmd_bootstrap, screen=RUN.fresh_screen)


def _cfg(fn):
    _rw(_P.layers("process", "config.json"), fn)


with contextlib.redirect_stdout(io.StringIO()):
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register("process", DT)
tuning.apply({"chunk_rows": [1, 2]})
HINTS_OF = {}
for _did, _rows in DOCS.items():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(title=SHEET)
    for i, v in enumerate(_rows, 1):
        ws.cell(row=i, column=1, value=v)
    _p = TMP / f"{_did}.xlsx"
    wb.save(_p)
    _st = [x for x in UNITS if x["text"] in _rows]
    if _did == "B120TW2":
        _st = [UNITS[0], UNITS[2]]
    (HINTS / f"{_did}@structure.json").write_text(json.dumps({SHEET: _st}, ensure_ascii=False), encoding="utf-8")
    HINTS_OF[_did] = _hints(_did, _p)
    _put_hints(_did)
    with contextlib.redirect_stdout(io.StringIO()):
        _r = IG.ingest_file(str(_p), DT, ask=False, sheets="1:prose", coord="노칭")
    print(f"  인입 {_did} — {(_r or {}).get('status')}")

print("\n■ ⓐ 값 · ⓑ 버림 · 충돌")
_un, _n1 = _spec("노칭::언와인더 unit::장력")
_re, _n2 = _spec("노칭::리와인더 unit::장력")
_all = [str(it["value"]) for _l, n in World().nodes() for v in (n.get("attrs") or {}).values() for it in v
        if isinstance(v, list) and isinstance(it, dict)]
show("ⓐ 유닛마다 같은 이름 사양이 각자 원문 값 · 출처 청크(B120TW#…) — 날짜 · 쪽수는 어디에도 값이 아니다",
     [v for v, _p in _un] == ["30N"] and [v for v, _p in _re] == ["20N"]
     and all(p[0].startswith("B120TW#") for _v, p in _un + _re)
     and not any("2026-10-16" in v or "3쪽" in v for v in _all),
     f"언와인더 {_un} · 리와인더 {_re}")

_defects = store.path(store.DEFECTS).read_text(encoding="utf-8") if store.path(store.DEFECTS).exists() else ""
_rows_v = [r for d in DOCS for r in (LG.read(d) or {}).get("rows") or [] if r.get("value")]
_drop = [r for r in _rows_v if r["value"]["verdict"] == "버림"]
_q = [x for x in q_of("spec_conflict", "B120TW2") if (x.get("payload") or {}).get("node_id") == (_n1 or {}).get("id")]
_sv = RC.summary(RC.for_report("B120TW", LG.read("B120TW") or {})["줄"], {}, {})["값"]
show("ⓑ 원문에 없는 값은 버림(노드 attrs 0 · 행 `value.verdict` 버림 · 맞춤표 수) · 다른 문서의 다른 값은 spec_conflict(있던 값 그대로)",
     len(_drop) == 1 and _drop[0]["value"]["value"] == "85ppm" and not _spec("노칭::리와인더 unit::속도")[0]
     and "값 버림 — 청크 원문에 없다 '85ppm'" in _defects
     and _sv[1].get("버림") == 1 and _sv[0] == sum(_sv[1].values())
     and len(_q) == 1 and _q[0]["payload"]["incoming"] == "35N" and _q[0]["payload"]["existing"] == "30N"
     and [v for v, _p in _spec("노칭::언와인더 unit::장력")[0]] == ["30N"],
     f"버림 {[r['value'] for r in _drop]} · 요약 {_sv} · 충돌 {[(x['payload']['existing'], x['payload']['incoming']) for x in _q]}")

print("\n■ ⓒ 선언 · 판 · ⓓ 재구축")
_before = _prose_values()
_drop_hints()                                             # 다시 뽑으면 값이 사라진다 — 재사용만이 값을 지킨다
_cfg(lambda c: (c["relation_patterns"].append({"src": "Unit", "rel": "has_property", "dst": "Unit"}),
                c.__setitem__("config_version", c["config_version"] + "-b120")))
_r1 = _rebuild()
_c1 = _r1.get("calls") or {}
show("ⓒⓓ relation_patterns · config_version만 바꾸고 재구축 — 추출 재사용(힌트 없이 값 그대로) · LLM 0(추출 · 판정)",
     _r1.get("done") and _prose_values() == _before and _before and _c1.get("추출") == 0 and _c1.get("판정") == 0,
     f"값 {len(_before)} → {len(_prose_values())} · 호출 {_c1}")
_cfg(lambda c: c["categories"].__setitem__("Property", c["categories"]["Property"] + " (B120 시험 — 정의문을 바꿨다)"))
_r2 = _rebuild()
_gone = _prose_values()
_cfg(lambda c: c.pop("value_categories", None))
for _did in DOCS:
    _put_hints(_did)
_r3 = _rebuild()
_nodecl = _prose_values()
_ent = sum(1 for d in DOCS for r in (LG.read(d) or {}).get("rows") or [] if r.get("role") == "entity" and r.get("surface") == "장력")
show("ⓒ 정의문을 바꾸면 다시 뽑는다(힌트가 없어 값 0) · 선언을 지우면 힌트에 값이 있어도 0(개체는 그대로 선다)",
     _r2.get("done") and not _gone and _r3.get("done") and not _nodecl and _ent >= 2,
     f"정의문 뒤 값 {len(_gone)} · 선언 없음 값 {len(_nodecl)} · 장력 개체 행 {_ent}")

for _did in DOCS:
    for _p in (HINTS / f"{_did}.json", HINTS / f"{_did}@structure.json", HINTS / f"{_did}@outline.json"):
        _p.unlink(missing_ok=True)
_unregister(DT)                                           # 등록 단은 클린이 남긴다 — 시험 몫을 치운다
done()
