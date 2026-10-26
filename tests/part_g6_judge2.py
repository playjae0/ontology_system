# -*- coding: utf-8 -*-
"""G6 — **다듬기** — 잘리지 않는 md 보고서 · 붙은 자리 칸에 관계 · 불확실만 청크 문장을 보고 다시 판정 (B123).

대표 사례 하나씩(가벼운 회차 · 창작 표본 — 수치는 근거가 아니다 [정정] 50):
  ⓐ md — 산문 표본(TW 꼴 · mock 힌트) 하나의 `export/report_<id>.md`: 값 표 행 수 = 대장 개체 행 · 잘린 칸 0 ·
     카테고리 쌍별 수의 합 = 그 문서 몫 엣지(`show summary`의 문서 줄과 같은 셈)
  ⓑ 칸 — 같은 표본: 붙은 자리 칸이 관계 이름으로 시작한다 · 연결이 하나인 행은 `└` 줄 0 · 둘 이상인 행만 연결마다 한 줄
  ⓒ 2단 — 가짜 게이트웨이(실호출 갈래 · 판정 지시문 · 후보 · 계약은 실물): 1차 불확실 → 2차(청크 문장)로 매칭 · 2차도
     불확실이면 지금 동작(신규 + `uncertain_match`) · 2차는 1차 불확실에만 · 재구축(기본 fresh → 다시 인입 · 재생) 판정 LLM 0
네 벌(ⓓ)은 마감 한 줄의 네 벌 단계가 잰다(mock 판정은 문장을 읽지 않아 2차 0).
"""
from __future__ import annotations

import contextlib
import io
import os
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done   # noqa: F401

from fake_gateway import CALLS, Live                     # noqa: E402
from g6_common import _register, _unregister             # noqa: E402
from openpyxl import Workbook                            # noqa: E402
import run as RUN                                        # noqa: E402
from cli import ingest as IG, ingest_screen as SCR, report_md as MD, show_summary as SS   # noqa: E402
from core.build import ledger as LG, replay as RP, structure as ST   # noqa: E402
from parser import pipeline, struct_map, tuning          # noqa: E402
from parser.adapters import basic_prose_xlsx as BX       # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="b123_"))
from core.state import fixtures as _FXR                   # noqa: E402 — 픽스처 뿌리(회귀는 스위트마다 사본 — B133)
HINTS = _FXR.EXTRACT_HINTS


def _quiet():
    return contextlib.redirect_stdout(io.StringIO())


# ────────────────────────────────────────────────────────────── ⓒ 2단 판정 (가짜 게이트웨이)
print("\n■ ⓒ 2단 판정 — 불확실만 청크 문장을 보고 다시")
_TEXT = {"S1": "노칭 정밀도는 금형 클리어런스로 정한다.", "S2": "앞서 말한 노칭 정밀도, 줄여서 정밀도는 매일 잰다.",
         "S3": "피치는 2mm로 둔다."}
_ENTS = {_TEXT["S1"]: "노칭 정밀도", _TEXT["S2"]: "정밀도", _TEXT["S3"]: "피치"}
ENV = {"source_path": "(합성) B123J.pptx", "revision": "R1", "parsed_at": "2026-10-19T00:00:00", "parser_version": "m",
       "adapter_version": "m", "context": {}, "payload_kind": "prose", "doc_type": "b123_live", "doc_id": "B123J",
       "chunks": [{"source_locator": k, "doc_type": "b123_live", "process_group": "조립", "process_ref": "노칭",
                   "electrode_type": "both", "text": t, "section": f"슬라이드 {i}", "meta": {}}
                  for i, (k, t) in enumerate(_TEXT.items(), 1)]}


def _extract(body):
    text = body["chunk"].split("\n", 1)[-1] if body["chunk"].startswith("[") else body["chunk"]
    return {"entities": [{"surface": _ENTS[text], "category": "Property", "parent": None, "belongs_to": None,
                          "value": None}], "relations": [], "about": []}


def _judge(body):
    """가짜 판정(창작): 1차는 늘 불확실 · 2차는 문장이 후보 이름(접두 뺀)을 글로 말하면 그 후보."""
    cands = body["candidates"]
    if "context" in body:
        hit = next((c for c in cands if c["canonical"].split("::")[-1] in body["context"]
                    and c["canonical"].split("::")[-1] != body["mention"].split("::")[-1]), None)
        if hit:
            return {"type": "match", "matched_id": hit["id"], "confidence": 0.95}
    return {"type": "uncertain", "matched_id": cands[0]["id"] if cands else None, "confidence": 0.5}


def _live():
    return Live(by={"extract": _extract, "judge": _judge, "outline": lambda b: {"headings": []},
                    "reinforce": lambda b: {"pick": None}})


def _row(surface):
    return next((r for r in (LG.read("B123J") or {}).get("rows") or []
                 if r.get("role") == "entity" and r.get("surface") == surface), {})


with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register("process", "b123_live")                         # 실호출 갈래에서는 내장(mock) 종류가 숨는다 — 등록 단에
with _live(), _quiet():
    run_document(dict(ENV))
_judged = [b for p, b in CALLS if p == "judge"]
_r2, _r3 = _row("정밀도"), _row("피치")
_sp = (_r2.get("judge2") or {}).get("span") or [0, 0]
show("ⓒ 1차 불확실 → 2차(그 표기가 나온 청크 문장)로 매칭 · 2차도 불확실이면 신규 + uncertain_match · 2차는 불확실에만",
     _r2.get("verdict") == "match" and _r2.get("target") == "노칭::노칭 정밀도"
     and _r2["judge2"]["was"] == "uncertain" and _r2["judge2"]["now"] == "match"
     and "정밀도" in _TEXT["S2"][_sp[0]:_sp[1]]
     and _r3.get("verdict") == "uncertain" and (_r3.get("judge2") or {}).get("now") == "uncertain"
     and q_of("uncertain_match", "B123J")
     and sum(1 for b in _judged if "context" in b) == 2 and not (_row("노칭 정밀도").get("judge2")),
     f"정밀도 {_r2.get('verdict')} → {_r2.get('target')} · 2차 {_r2.get('judge2')} · 피치 {_r3.get('verdict')} · "
     f"판정 호출 {len(_judged)}(문장 {sum(1 for b in _judged if 'context' in b)})")

with _quiet():
    init.init(fresh_=True)                                # 시험 대상: 기본 fresh — 재구축의 바닥(추출 · 판정 대장이 남는다)
    RUN.cmd_bootstrap([])
RP.reset()
with _live(), _quiet():
    run_document(dict(ENV))
_again = [p for p, _b in CALLS if p in ("judge", "extract")]
show("ⓒ 재구축(기본 fresh → 다시 인입) — 최종 판정을 재생(판정 · 추출 LLM 0 · 같은 붙음)",
     not _again and _row("정밀도").get("target") == "노칭::노칭 정밀도" and _row("피치").get("verdict") == "uncertain",
     f"호출 {_again} · 정밀도 → {_row('정밀도').get('target')} · 재생 {dict(RP.DOC)}")
RP.OFF, RP.COUNT_ALL = None, False
_unregister("b123_live")

# ────────────────────────────────────────────────────────────── ⓐ md · ⓑ 칸 (mock 힌트)
print("\n■ ⓐ md 보고서 · ⓑ 붙은 자리 칸")
DT, DID, SHEET = "b123_process", "B123TW", "TW 사양"
ROWS = ["언와인더 unit", "1) 기본사양", "속도는 100ppm이다.", "2) 지정사양", "장력은 30N이다.",
        "리와인더 unit", "1) 기본사양", "속도는 80ppm이다."]
STRUCT = [{"text": "언와인더 unit", "depth": 1, "role": "대상", "name": "언와인더 unit", "category": "Unit"},
          {"text": "1) 기본사양", "depth": 2, "role": "묶음"}, {"text": "2) 지정사양", "depth": 2, "role": "묶음"},
          {"text": "리와인더 unit", "depth": 1, "role": "대상", "name": "리와인더 unit", "category": "Unit"},
          {"text": "1) 기본사양", "depth": 2, "role": "묶음"}]
with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register("process", DT)
tuning.apply({"chunk_rows": [1, 2]})
_wb = Workbook()
_wb.remove(_wb.active)
_ws = _wb.create_sheet(title=SHEET)
for _i, _v in enumerate(ROWS, 1):
    _ws.cell(row=_i, column=1, value=_v)
_xp = TMP / f"{DID}.xlsx"
_wb.save(_xp)
(HINTS / f"{DID}@structure.json").write_text(json.dumps({SHEET: STRUCT}, ensure_ascii=False), encoding="utf-8")
_env = pipeline.parse(BX, DID, str(_xp), closed_list=[], sheet_roles={SHEET: "prose"},
                      structure=ST.structurer()).envelope
_hints = {}
for _c in _env["chunks"]:
    _own = next((x for x in reversed(_c["meta"].get("structure") or []) if x.get("name")), None)
    _ents = [{"surface": w, "category": "Property",
              "belongs_to": {"name": _own["name"], "category": "Unit", "from": "제목"} if _own else None}
             for w in ("속도", "장력") if w + "은" in _c["text"] or w + "는" in _c["text"]]
    # 둘째 연결(창작) — 언와인더의 장력 청크가 리와인더 unit과의 관계도 말한다(연결이 둘인 행)
    _rel = [{"src": "리와인더 unit", "rel": "has_property", "dst": "장력"}] if "장력은" in _c["text"] else []
    if _ents:
        _hints[_c["source_locator"]] = {"entities": _ents, "relations": _rel, "attach": []}
(HINTS / f"{DID}.json").write_text(json.dumps(_hints, ensure_ascii=False), encoding="utf-8")
struct_map.invalidate(DID)
with _quiet():
    IG.ingest_file(str(_xp), DT, ask=False, sheets="1:prose", coord="노칭")

_data = LG.read(DID) or {}
_ent = [r for r in _data.get("rows") or [] if r.get("role") == "entity"]
_mdp = MD.write(DID, _data)
_md = _mdp.read_text(encoding="utf-8")
_vals = _md.split("## 판정 값", 1)[-1]
_vrows = [ln for ln in _vals.splitlines() if ln.startswith("| ") and not ln.startswith("| # |")]
from core.state.bootstrap import open_graph                 # noqa: E402
from router import discover                                 # noqa: E402
_graphs = {lay: open_graph(lay) for lay in discover()}
_pairs = MD.pair_counts(_graphs, DID, _data.get("rows") or [])
show("ⓐ md — 값 표 행 = 대장 개체 행 · 잘린 칸 0 · 카테고리 쌍별 수의 합 = 그 문서 몫 엣지(show summary와 같은 셈)",
     _ent and len(_vrows) == len(_ent) and "…" not in _vals and sum(_pairs.values()) == SS._mine(_graphs, DID)[1] > 0
     and _mdp.name == f"report_{DID}.md" and _mdp.parent == _P.export(),
     f"값 행 {len(_vrows)}/{len(_ent)} · 쌍 {dict(_pairs)} · 엣지 {SS._mine(_graphs, DID)[1]} · {_mdp.name}")

_keep_loud, _cols = SCR._loud, os.environ.get("COLUMNS")
SCR._loud, os.environ["COLUMNS"] = (lambda r: True), "220"
_buf = io.StringIO()
with contextlib.redirect_stdout(_buf):
    _pr = SCR.row_printer(_env)
    for _r in _ent:
        _pr(_r)
    _pr.flush(None)
SCR._loud = _keep_loud
if _cols is None:
    os.environ.pop("COLUMNS", None)
else:
    os.environ["COLUMNS"] = _cols
LG.ON_CHUNK = None
_tree = [ln for ln in _buf.getvalue().splitlines() if re.match(r"\s+└ ", ln) and "└ 소속" not in ln]
_multi = [r for r in _ent if len(r.get("attached") or []) > 1]
_stuck = [SCR.value_cells(r, {}, {})[0][SCR.VALUE_HEADS.index("붙은 자리")] for r in _ent if r.get("attached")]
show("ⓑ 붙은 자리 칸 = 관계 상대(방식) · 연결이 하나인 행은 └ 줄 0 · 둘 이상인 행만 연결마다 한 줄",
     _multi and len(_tree) == sum(len(r["attached"]) for r in _multi)
     and all(s.startswith(r["attached"][0]["rel"] + " ") for s, r in zip(_stuck, [r for r in _ent if r.get("attached")])),
     f"└ 줄 {len(_tree)} · 둘 이상인 행 {len(_multi)} · 칸 {_stuck[:2]}")

for _p in (HINTS / f"{DID}.json", HINTS / f"{DID}@structure.json", HINTS / f"{DID}@outline.json"):
    _p.unlink(missing_ok=True)
_unregister(DT)                                            # 등록 단은 클린이 남긴다 — 시험 몫을 치운다
done()
