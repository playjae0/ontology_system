# -*- coding: utf-8 -*-
"""산문 엑셀의 계층 표본 둘 — **B87의 재료** (번호 군 · 글자 상한 · 규칙 선언).

사내 문서의 모양을 창작으로 옮긴 것이다 — **사내 정보는 한 글자도 들어가지 않는다**
(mock 자산의 규율 — 문서 7 §7.5).

    python tests/fixtures/make_b87.py
      → tests/fixtures/raw/HIER01.xlsx  ⓐ `Ⅰ./1./가./(1)` 혼용 · 60자 넘는 번호 행 ·
                                          글머리표 · 3,000자 넘는 절 · 혼자 상한을 넘는 행
      → tests/fixtures/raw/FONT01.xlsx  ⓑ 번호·병합·굵게 없이 **글자 크기만** 다른 제목
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

RAW = Path(__file__).resolve().parent / "raw"
HIER = RAW / "HIER01.xlsx"
FONT = RAW / "FONT01.xlsx"

_S = "설비는 규격 범위 안에서 연속 운전하며 이상이 생기면 경보를 내고 안전하게 정지한다"

#: ⓐ 시트 「사양」 — 군 넷이 처음 나온 순서가 곧 레벨이다(Ⅰ.=1 · 1.=2 · 가.=3 · (1)=4)
MIXED = [
    "Ⅰ. 총칙",
    "이 사양서는 노칭 설비의 요구 조건을 정한다.",
    "1. 목적",
    "설비의 성능과 인수 기준을 명확히 한다.",
    "가. 적용 범위",
    "노칭 프레스와 이송 장치에 적용한다.",
    "(1) 대상 설비",
    "프레스 본체와 금형 교환 장치를 포함한다.",
    # 60자를 넘는 번호 행 — **제목이 아니다**(목록 항목·본문)
    "1) " + _S + "고 그 이력은 상위 시스템에 모두 남겨야 한다.",
    "- 글머리 항목은 목록이지 제목이 아니다.",
    "• 두 번째 글머리 항목.",
    "(2) 부속 장치",
    "이송 장치와 스크랩 배출 장치를 포함한다.",
    "Ⅱ. 기술 사양",
    "1. 기계",
    "가. 프레스",
    "(1) 가압력",
    "가압력은 120 kN 이상이어야 한다.",
    "(2) 정밀도",
    "금형 클리어런스는 0.02 mm 이하로 관리한다.",
    "나. 이송",
    "이송 정밀도는 ±0.05 mm 이내로 한다.",
]

#: ⓐ 시트 「장문」 — 점 번호만 · 2절은 행 30개가 3,000자를 넘는다 · 3절 한 행은 혼자 넘는다
LONG = (["1. 개요"] + [f"개요 본문 {i}행 — {_S}." for i in range(1, 7)]
        + ["2. 상세 조건"] + [f"상세 {i:02d} — {_S}. 점검 주기와 기록 양식은 부록을 따르고 "
                          f"변경이 있으면 사전에 협의하며 기록은 삼 년간 보관한다." for i in range(1, 31)]
        + ["3. 부록"] + ["부록 원문 — " + (_S + ". ") * 70]
        + [f"부록 비고 {i}행." for i in range(1, 5)])

#: ⓑ 시트 「사양서」 — 제목은 **글자 크기 14**뿐이다(본문 10) · 번호·병합·굵게 없음
FONT_ROWS = [("개요", 14)] + [(f"개요 본문 {i} — {_S}.", 10) for i in range(1, 7)] \
    + [("기계 사양", 14)] + [(f"기계 본문 {i} — {_S}.", 10) for i in range(1, 9)] \
    + [("전장 사양", 14)] + [(f"전장 본문 {i} — {_S}.", 10) for i in range(1, 7)] \
    + [("인수 조건", 14)] + [(f"인수 본문 {i} — {_S}.", 10) for i in range(1, 6)]


def _sheet(wb, name, rows, heads=()):
    """본문 행은 **셀 들여쓰기 1** — 실물 사양서의 모양이고(RFQ01과 같다), 형태 판정의
    indent 신호가 보는 것이 이것이다. 제목 행(`heads`의 번호 · 글자 크기 14)은 들여쓰지 않는다.
    """
    ws = wb.create_sheet(title=name)
    for r, v in enumerate(rows, start=1):
        text, size = v if isinstance(v, tuple) else (v, 10)
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = Font(size=size)
        if size == 10 and not str(text).startswith(heads):
            cell.alignment = Alignment(indent=1)
    ws.column_dimensions["A"].width = 80


def build():
    RAW.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "사양", MIXED, ("Ⅰ.", "Ⅱ.", "1. ", "가.", "나.", "(1)", "(2)"))
    _sheet(wb, "장문", LONG, ("1. ", "2. ", "3. "))
    wb.save(HIER)
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "사양서", FONT_ROWS)
    wb.save(FONT)
    return HIER, FONT


if __name__ == "__main__":
    print("만들었다 —", *build())
