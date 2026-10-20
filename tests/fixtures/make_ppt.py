# -*- coding: utf-8 -*-
"""판독 픽스처 생성 — **shape 종류를 전부 담은 PPTX 한 벌** (B53 완료판정 a-ⓐ).

바이너리를 레포에 커밋하지 않고 **코드로 만든다**: ①무엇이 들었는지가 diff에
보인다 ②`python-pptx` 판이 올라가도 같은 것을 다시 만든다 ③외부 LLM 실산출
스냅샷(`tests/fixtures/fixtures/`, D-26)과 달리 이것은 **입력**이라 재생성이 정당하다.

만드는 것 — 10슬라이드:
  1 제목 슬라이드      2 구획 헤더(Section)   3 본문      4 표 5×2
  5 꺾은선 차트 2계열   6 그림 1               7 그룹 안 텍스트
  8 시각 순서 시험(아래 상자를 XML 앞에)       9 목차       10 숨김

사용: python tests/fixtures/make_ppt.py [출력경로]
"""
from __future__ import annotations

import struct
import sys
import zlib
from pathlib import Path

OUT = Path(__file__).resolve().parent / "raw" / "PPT_shapes.pptx"


def _png(w=64, h=48, rgb=(0x2E, 0x74, 0xB5)):
    """의존 없이 만드는 최소 PNG — Pillow를 끌어오지 않는다(선택 의존 0)."""
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(tag, data):
        c = tag + data
        return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c))

    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


def build(out=OUT):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    blank = prs.slide_layouts[6]
    title_only = prs.slide_layouts[5]

    # 1 — 제목 슬라이드
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "조립 공정 판독 시험"
    s.placeholders[1].text = "B53 픽스처 · 형태 판독 전 종류"

    # 2 — 구획 헤더. **레이아웃 이름으로 판정한다**(결정적).
    s = prs.slides.add_slide(prs.slide_layouts[2])
    s.shapes.title.text = "노칭 구획"
    try:                      # 레이아웃 이름에 Section이 들어가야 한다
        s.slide_layout.name = "Section Header"
    except Exception:         # noqa: BLE001
        pass

    # 3 — 본문
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "노칭 설비 구성"
    s.placeholders[1].text = ("노칭 프레스는 금형으로 전극 원단을 타발한다.\n"
                              "금형 클리어런스는 20±2㎛로 관리한다.")

    # 4 — 표 5×2 (헤더 + 4행)
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "노칭 관리 규격"
    tb = s.shapes.add_table(5, 2, Inches(1), Inches(2), Inches(6), Inches(2.5)).table
    rows = [("항목", "규격"), ("클리어런스", "20±2㎛"), ("장력", "25±3N"),
            ("사행량", "±1.0mm"), ("버 높이", "10㎛ 이하")]
    for r, (a, b) in enumerate(rows):
        tb.cell(r, 0).text = a
        tb.cell(r, 1).text = b

    # 5 — 꺾은선 차트 2계열. **값은 있는 그대로** 실려야 한다(a-ⓑ).
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "월별 불량률"
    cd = CategoryChartData()
    cd.categories = ["1월", "2월", "3월"]
    cd.add_series("A라인", (0.8, 1.1, 0.95))
    cd.add_series("B라인", (1.2, 0.7, 1.05))
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(2),
                            Inches(7), Inches(4), cd)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = "월별 불량률"

    # 6 — 그림
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "노칭 프레스 외관"
    import io as _io
    s.shapes.add_picture(_io.BytesIO(_png()), Inches(1), Inches(2),
                         Inches(3), Inches(2.25))

    # 7 — 그룹 안 텍스트
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "그룹 판독"
    g = s.shapes.add_group_shape()
    for n, (top, txt) in enumerate([(3.0, "그룹 안 두 번째"), (2.0, "그룹 안 첫 번째")]):
        b = g.shapes.add_textbox(Inches(1), Inches(top), Inches(4), Inches(0.6))
        b.text_frame.text = txt

    # 8 — **시각 순서 시험**: 아래쪽(top 큰) 상자를 XML에 먼저 넣는다.
    #     판독이 XML 순서를 그대로 쓰면 여기서 순서가 뒤집힌다(a-ⓒ).
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "순서 시험"
    for top, txt in ((4.5, "아래 상자 — XML에는 먼저 있다"),
                     (2.0, "위 상자 — XML에는 나중에 있다")):
        s.shapes.add_textbox(Inches(1), Inches(top), Inches(6), Inches(0.6)) \
            .text_frame.text = txt

    # 9 — 목차
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "목차"
    s.placeholders[1].text = "1. 노칭\n2. 스태킹\n3. 탭용접"

    # 10 — 숨김. 어댑터가 걸러야 한다(a-ⓓ).
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "숨긴 슬라이드 — 청크에 없어야 한다"
    s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)) \
        .text_frame.text = "이 문장은 어느 청크에도 없어야 한다."
    s._element.set("show", "0")

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def build_report(out):
    """**보고서 꼴 덱**(B125 — 창작 표본) — 표지 · 글 상자 제목 + 작은 상자 여섯 · (계속) · 본문 + 표 + 그림 + 노트 · 두 공정."""
    import io
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]

    def box(s, text, top, left=0.5, w=4.0):
        s.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(0.5)).text_frame.text = text

    s = prs.slides.add_slide(prs.slide_layouts[0])                       # 1 — 표지
    s.shapes.title.text, s.placeholders[1].text = "노칭 · 스태킹 점검 보고", "B125 창작"
    s = prs.slides.add_slide(blank)                                       # 2 — 글 상자 제목 + 작은 상자 여섯
    box(s, "노칭 금형 점검 결과", 0.3, w=8)
    for i, t in enumerate(("마모 0.05mm", "→ 버 발생", "교체 2주", "클리어런스 0.02mm", "윤활 1회", "점검 완료")):
        box(s, t, 1.5 + (i // 3) * 1.0, 0.5 + (i % 3) * 3.0, 2.5)
    s = prs.slides.add_slide(blank)                                       # 3 — (계속)
    box(s, "노칭 금형 점검 결과 (계속)", 0.3, w=8)
    box(s, "다이 균열 없음", 1.5)
    s = prs.slides.add_slide(prs.slide_layouts[5])                        # 4 — 본문 + 표 + 그림 + 노트
    s.shapes.title.text = "커터 사양"
    box(s, "아래 표는 커터 사양이다.", 1.4)
    tb = s.shapes.add_table(2, 2, Inches(0.5), Inches(2.2), Inches(4), Inches(1)).table
    for (r, c), v in {(0, 0): "항목", (0, 1): "값", (1, 0): "장력", (1, 1): "30N"}.items():
        tb.cell(r, c).text = v
    s.shapes.add_picture(io.BytesIO(_png()), Inches(5.5), Inches(2.2))
    s.notes_slide.notes_text_frame.text = "점검 메모 — 압력 5bar"
    s = prs.slides.add_slide(blank)                                       # 5 — 두 공정 · 언급 하나
    box(s, "노칭 · 스태킹 연동 점검", 0.3, w=8)
    box(s, "노칭 커터 마모와 스태킹 매거진 정렬을 함께 본다. 탭용접 설비와 연동한다.", 1.5, w=8)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return Path(out)


if __name__ == "__main__":
    p = build(sys.argv[1] if len(sys.argv) > 1 else OUT)
    print(f"[fixture] {p}  [{p.stat().st_size / 1024:.0f}KB]")
