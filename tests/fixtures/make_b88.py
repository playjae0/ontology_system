# -*- coding: utf-8 -*-
"""리더의 폭 표본 셋 — **B88의 재료** (엑셀 그림 · Word 두 벌).

창작 내용이다 — **사내 정보는 한 글자도 들어가지 않는다**(mock 자산의 규율 — 문서 7 §7.5).
그림(PNG·WMF)은 표준 라이브러리로 만든다(PIL 없이).

    python tests/fixtures/make_b88.py
      → tests/fixtures/raw/IMG01.xlsx  ⓐ 산문 시트 「사양」에 PNG 1 · WMF 1 / 참조 시트 「도면」에 PNG 1
      → tests/fixtures/raw/DOC01.docx  ⓑ 한글 스타일 제목(스타일 id `1`·`2` · 상속 한 벌) + 표 + 그림 +
                                          머리글 + 변경 추적 삭제
      → tests/fixtures/raw/DOC02.docx  ⓑ 스타일 없이 번호만(`1.` · `1.1` · `2.`) + 번호 목록 한 줄
"""
import io
import struct
import zipfile
import zlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

RAW = Path(__file__).resolve().parent / "raw"
IMG, DOC1, DOC2 = RAW / "IMG01.xlsx", RAW / "DOC01.docx", RAW / "DOC02.docx"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_S = "설비는 규격 범위 안에서 연속 운전하며 이상이 생기면 경보를 내고 안전하게 정지한다."


def png(w=8, h=8, rgb=(200, 30, 30)):
    """단색 PNG — 표준 라이브러리로(PIL 없이)."""
    def chunk(tag, data):
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def wmf():
    """placeable WMF — 머리 + 빈 메타파일(EOF 레코드). openpyxl은 이것을 읽기 전에 버린다."""
    head = struct.pack("<IHhhhhHI", 0x9AC6CDD7, 0, 0, 0, 100, 100, 1440, 0)
    chk = 0
    for i in range(0, 20, 2):
        chk ^= struct.unpack("<H", head[i:i + 2])[0]
    return head + struct.pack("<H", chk) + struct.pack("<HHHIHIH", 1, 9, 0x0300, 12, 0, 3, 0) \
        + struct.pack("<IH", 3, 0)


# ---------------------------------------------------------------- ⓐ IMG01.xlsx
SPEC = ["1. 개요", "   " + _S, "   노칭 설비의 기계 사양을 정한다.", "   그림 1은 설비 배치다.",
        "   배치는 공급 → 타발 → 배출 순이다.", "   각 구간의 간격은 협의로 정한다.",
        "2. 설비", "   프레스와 이송 장치로 이루어진다.", "   그림 2는 금형 단면이다.",
        "   금형 클리어런스는 0.02 mm 이하로 관리한다.", "   금형 교환은 10분 안에 끝낸다."]
DWG = [f"DWG-{i:03d}" for i in range(1, 7)]


def _drawing(anchors):
    """`anchors = [(열0, 행0, rId)]` → drawing XML."""
    pics = "".join(
        f'<xdr:oneCellAnchor><xdr:from><xdr:col>{c}</xdr:col><xdr:colOff>0</xdr:colOff>'
        f'<xdr:row>{r}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        f'<xdr:ext cx="95250" cy="95250"/><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{i + 1}" '
        f'name="p{i}"/><xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="{rid}"/>'
        f'<a:stretch><a:fillRect/></a:stretch></xdr:blipFill><xdr:spPr><a:prstGeom prst="rect">'
        f'<a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor>'
        for i, (c, r, rid) in enumerate(anchors))
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/'
            'spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            f'xmlns:r="{_REL}">{pics}</xdr:wsDr>')


def _rels(items):
    body = "".join(f'<Relationship Id="{i}" Type="{_REL}/{t}" Target="{g}"/>' for i, t, g in items)
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
            f'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>')


def build_img():
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in (("사양", SPEC), ("도면", DWG)):
        ws = wb.create_sheet(title=name)
        for r, v in enumerate(rows, start=1):
            deep = v.startswith("   ")
            cell = ws.cell(row=r, column=1, value=v.strip())
            if deep:
                cell.alignment = Alignment(indent=1)
        ws.column_dimensions["A"].width = 70
    buf = io.BytesIO()
    wb.save(buf)
    # 그림을 zip에 직접 넣는다 — 사양 A4(PNG) · A9(WMF) / 도면 A2(PNG)
    plan = {"xl/worksheets/sheet1.xml": ("drawing1.xml", [(0, 3, "rId1", "image1.png", png()),
                                                         (0, 8, "rId2", "image2.wmf", wmf())]),
            "xl/worksheets/sheet2.xml": ("drawing2.xml", [(0, 1, "rId1", "image3.png",
                                                          png(rgb=(30, 30, 200)))])}
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(IMG, "w", zipfile.ZIP_DEFLATED) as z:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.decode().replace(
                    "</Types>",
                    '<Default Extension="png" ContentType="image/png"/>'
                    '<Default Extension="wmf" ContentType="image/x-wmf"/>'
                    + "".join(f'<Override PartName="/xl/drawings/{d}" ContentType="application/'
                              f'vnd.openxmlformats-officedocument.drawing+xml"/>'
                              for d, _a in plan.values()) + "</Types>").encode()
            if it.filename in plan:
                data = data.decode().replace(
                    "</worksheet>", f'<drawing xmlns:r="{_REL}" r:id="rIdD"/></worksheet>').encode()
            z.writestr(it, data)
        for sheet, (dname, pics) in plan.items():
            sname = sheet.rsplit("/", 1)[1]
            z.writestr(f"xl/worksheets/_rels/{sname}.rels",
                       _rels([("rIdD", "drawing", f"../drawings/{dname}")]))
            z.writestr(f"xl/drawings/{dname}", _drawing([(c, r, rid) for c, r, rid, _m, _b in pics]))
            z.writestr(f"xl/drawings/_rels/{dname}.rels",
                       _rels([(rid, "image", f"../media/{m}") for _c, _r, rid, m, _b in pics]))
            for _c, _r, _rid, m, b in pics:
                z.writestr(f"xl/media/{m}", b)
    return IMG


# ---------------------------------------------------------------- ⓑ DOCX
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
       '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
       '<Default Extension="xml" ContentType="application/xml"/>'
       '<Default Extension="png" ContentType="image/png"/>'
       '<Override PartName="/word/document.xml" ContentType="application/'
       'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
       '<Override PartName="/word/styles.xml" ContentType="application/'
       'vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
       '<Override PartName="/word/header1.xml" ContentType="application/'
       'vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/></Types>')
_ROOT_RELS = _rels([("rId1", "officeDocument", "word/document.xml")])


def _t(text):
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _p(text, style=None, extra=""):
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{ppr}{_t(text) if text else ''}{extra}</w:p>"


def _doc(body, sect=""):
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:document xmlns:w="{_W}" xmlns:r="{_REL}" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<w:body>{body}{sect}</w:body></w:document>')


#: 한글 Word의 모양 — 스타일 **id**가 `1`·`2`이고 이름이 「제목 1」「제목 2」다(로캘마다 다르다).
#: `요약제목`은 개요 수준 없이 `1`을 상속한다(상속 체인 한 벌).
_STYLES = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<w:styles xmlns:w="{_W}">'
           '<w:style w:type="paragraph" w:default="1" w:styleId="a"><w:name w:val="바탕글"/></w:style>'
           '<w:style w:type="paragraph" w:styleId="1"><w:name w:val="제목 1"/><w:basedOn w:val="a"/>'
           '<w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>'
           '<w:style w:type="paragraph" w:styleId="2"><w:name w:val="제목 2"/><w:basedOn w:val="a"/>'
           '<w:pPr><w:outlineLvl w:val="1"/></w:pPr></w:style>'
           '<w:style w:type="paragraph" w:styleId="요약제목"><w:name w:val="요약 제목"/>'
           '<w:basedOn w:val="1"/></w:style></w:styles>')

_PIC = ('<w:r><w:drawing><wp:inline><wp:extent cx="95250" cy="95250"/><wp:docPr id="1" name="그림 1"/>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic><pic:nvPicPr><pic:cNvPr id="1" name="그림 1"/><pic:cNvPicPr/></pic:nvPicPr>'
        '<pic:blipFill><a:blip r:embed="rIdImg1"/></pic:blipFill><pic:spPr/></pic:pic>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')


def _table(rows):
    return "<w:tbl>" + "".join(
        "<w:tr>" + "".join(f"<w:tc><w:p>{_t(c)}</w:p></w:tc>" for c in r) + "</w:tr>"
        for r in rows) + "</w:tbl>"


def build_doc1():
    body = "".join([
        _p("개요", "1"),
        _p("본 문서는 노칭 설비의 사양과 인수 조건을 정한다."), _p(_S),
        _p("적용 범위", "2"),
        _p("노칭 프레스와 이송 장치에 적용한다."), _p("금형 교환 장치도 포함한다."),
        _table([["항목", "값", "비고"], ["가압력", "120 kN", "이상"], ["정밀도", "±0.05 mm", "이내"]]),
        _p("그림은 설비 배치다.", extra=_PIC),
        # 변경 추적 — 삭제된 글자는 본문에 넣지 않는다
        "<w:p>" + _t("금형은 분말 고속도강을 쓴다.")
        + '<w:del w:id="1" w:author="검토자"><w:r><w:delText>삭제된 옛 문장이다.</w:delText></w:r></w:del>'
        + "</w:p>",
        _p("기계 사양", "1"),
        _p("프레스", "2"),
        _p("가압력은 120 kN 이상이어야 한다."), _p("편차는 3% 이내로 관리한다."),
        _p("정리", "요약제목"),
        _p("인수 시험은 연속 8시간 운전으로 한다."),
    ])
    sect = f'<w:sectPr><w:headerReference w:type="default" r:id="rIdHdr"/></w:sectPr>'
    header = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:hdr xmlns:w="{_W}">'
              f'<w:p>{_t("머리글 — 양식 번호 F-001")}</w:p></w:hdr>')
    with zipfile.ZipFile(DOC1, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _CT)
        z.writestr("_rels/.rels", _ROOT_RELS)
        z.writestr("word/document.xml", _doc(body, sect))
        z.writestr("word/styles.xml", _STYLES)
        z.writestr("word/header1.xml", header)
        z.writestr("word/_rels/document.xml.rels", _rels([
            ("rIdImg1", "image", "media/image1.png"), ("rIdHdr", "header", "header1.xml"),
            ("rIdSty", "styles", "styles.xml")]))
        z.writestr("word/media/image1.png", png(rgb=(30, 160, 30)))
    return DOC1


def build_doc2():
    numpr = '<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr>'
    body = "".join([
        _p("1. 개요"), _p("본 문서는 이송 장치의 사양을 정한다."), _p(_S),
        _p("1.1 적용 범위"), _p("이송 장치와 그 제어반에 적용한다."), _p("보조 장치는 협의한다."),
        _p("2. 사양"), _p("이송 정밀도는 ±0.05 mm 이내로 한다."),
        # 번호 목록 — 목록이지 제목이 아니다(w:numPr)
        f"<w:p>{numpr}{_t('1. 속도는 분당 30 미터 이상')}</w:p>",
        _p("이송 속도는 온도 보정을 적용한다."),
    ])
    with zipfile.ZipFile(DOC2, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _CT)
        z.writestr("_rels/.rels", _ROOT_RELS)
        z.writestr("word/document.xml", _doc(body))
        z.writestr("word/_rels/document.xml.rels", _rels([]))
    return DOC2


if __name__ == "__main__":
    RAW.mkdir(parents=True, exist_ok=True)
    print("만들었다 —", build_img(), build_doc1(), build_doc2())
