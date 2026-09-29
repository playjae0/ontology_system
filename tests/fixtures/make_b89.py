# -*- coding: utf-8 -*-
"""산문 엑셀 표본 셋 — **B89의 재료** (셀 안 줄바꿈 · 관문의 시트 역할 · 형태 판정의 시트 역할).

사내 문서의 모양을 창작으로 옮긴 것이다 — **사내 정보는 한 글자도 들어가지 않는다**
(mock 자산의 규율 — 문서 7 §7.5). 수치는 메커니즘 확인용이다 — 상수·문턱의 근거가
아니다(CLAUDE.md §3 [정정] 50).

    python tests/fixtures/make_b89.py
      → tests/fixtures/raw/NL01.xlsx   ① 시트 「사양」 — 짧은 1레벨 절 셋 + 2레벨 제목을 품은
                                         긴 1레벨 절(셀 안 줄바꿈 → 행 수는 적어도 글 줄은
                                         80 초과 · 한 2레벨 부분이 3,000자 초과) + 「해당 없음」 반복 행
      → tests/fixtures/raw/SKIP01.xlsx ② 「사양」(산문) + 「메모」(한 칸 「상동」 — 관문 G38을
                                         FAIL시키는 시트 · 역할 `skip`이면 관문이 보지 않아야 한다)
      → tests/fixtures/raw/MIX01.xlsx  ③ 표 모양 셋(가격·일정·도면목록) + 산문 하나(사양) —
                                         통합문서 전체로 판정하면 table로 기운다
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

RAW = Path(__file__).resolve().parent / "raw"
NL = RAW / "NL01.xlsx"
SKIP = RAW / "SKIP01.xlsx"
MIX = RAW / "MIX01.xlsx"

_S = "설비는 규격 범위 안에서 연속 운전하며 이상이 생기면 경보를 내고 안전하게 정지한다"
NONE = "해당 없음"


def _cell_lines(tag, n):
    """셀 하나에 든 여러 줄 — 실물 사양서의 「한 칸에 항목 여럿」 모양."""
    return "\n".join(f"{tag} 항목 {i} — {_S}." for i in range(1, n + 1))


def _short(title, tag):
    return [title] + [f"{tag} 본문 {i} — {_S}." for i in range(1, 9)]


#: ① 시트 「사양」 — 로마 번호가 레벨 1 · 점 번호가 레벨 2(첫 등장 순서)
NL_ROWS = (_short("Ⅰ. 개요", "개요") + _short("Ⅱ. 적용 범위", "적용") + _short("Ⅲ. 인수 조건", "인수")
           + ["Ⅳ. 상세 사양",
              _cell_lines("공통", 3)]                               # 첫 2레벨 제목 앞 — 줄바꿈 셀
           + ["1. 기계"] + [_cell_lines(f"기계 {k}", 7) for k in range(1, 9)]   # 이 부분 혼자 3,000자 초과
           + ["2. 전장"] + [_cell_lines(f"전장 {k}", 4) for k in range(1, 6)]
           + ["3. 비고"] + [NONE, _cell_lines("비고", 2), NONE, NONE])
NL_HEADS = ("Ⅰ.", "Ⅱ.", "Ⅲ.", "Ⅳ.", "1. ", "2. ", "3. ")

#: ② 시트 「사양」(산문 · 짧다) + 「메모」(한 칸 — 상동 기호 그대로)
SKIP_SPEC = _short("1. 개요", "개요") + _short("2. 성능", "성능") + _short("3. 인수", "인수")

#: ③ 표 모양 셋 — 머리 행 + 짧은 값 행(숫자 많음)
PRICE = [["품목", "수량", "단가", "금액", "비고"]] + [
    [f"P{i:02d}", i, 1000 * i, 1000 * i * i, "-"] for i in range(1, 41)]
PLAN = [["단계", "시작", "종료", "담당", "진척"]] + [
    [f"S{i}", f"2026-0{1 + i % 9}-01", f"2026-0{1 + i % 9}-20", f"팀 {i % 3}", i * 2]
    for i in range(1, 41)]
DRAW = [["도면번호", "명칭", "개정", "매수", "비고"]] + [
    [f"D{i:03d}", f"조립도 {i}", f"R{i % 4}", i % 5 + 1, "-"] for i in range(1, 41)]
MIX_SPEC = _short("1. 개요", "개요") + _short("2. 요구 사양", "사양")    # 산문 시트 혼자면 prose


def _prose(wb, name, rows, heads=()):
    """본문 행은 셀 들여쓰기 1 · 줄바꿈 칸은 `wrap_text`(make_b87과 같은 모양)."""
    ws = wb.create_sheet(title=name)
    for r, text in enumerate(rows, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = Font(size=10)
        if not str(text).startswith(heads):
            cell.alignment = Alignment(indent=1, wrap_text="\n" in str(text))
    ws.column_dimensions["A"].width = 80


def _table(wb, name, rows):
    ws = wb.create_sheet(title=name)
    for r, vals in enumerate(rows, start=1):
        for c, v in enumerate(vals, start=1):
            if v != "":
                ws.cell(row=r, column=c, value=v).font = Font(size=10, bold=(r == 1))


def build():
    RAW.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    _prose(wb, "사양", NL_ROWS, NL_HEADS)
    wb.save(NL)

    wb = Workbook()
    wb.remove(wb.active)
    _prose(wb, "사양", SKIP_SPEC, ("1. ", "2. ", "3. "))
    wb.create_sheet(title="메모").cell(row=1, column=1, value="상동")
    wb.save(SKIP)

    wb = Workbook()
    wb.remove(wb.active)
    _table(wb, "가격", PRICE)
    _table(wb, "일정", PLAN)
    _table(wb, "도면목록", DRAW)
    _prose(wb, "사양", MIX_SPEC, ("1. ", "2. "))
    wb.save(MIX)
    return NL, SKIP, MIX


if __name__ == "__main__":
    print("만들었다 —", *build())
