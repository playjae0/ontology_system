# -*- coding: utf-8 -*-
"""G6 — **PPT 캡처가 실제로 돈다** — 캡처 가능 판정 한 자리 · 실호출 PPT · PDF는 파싱 전에 그림 입력과 캡처 길을 확인 (B128).

**대역 없는 실경로**: 렌더 · 판정 함수(`parser.render`)를 가짜로 바꾸지 않는다 — python-pptx · PyMuPDF · LibreOffice가 실물이고
LLM 답(덱 구조 · 추출 · 채팅 왕복 · 그림 입력 거부)만 가짜 게이트웨이다. 창작 표본(`make_ppt.build_report` — 수치는 근거가 아니다
[정정] 50). soffice가 없는 기계는 「soffice 없는 길」을 하위 실행의 PATH에서 빼서 잰다. 대표 사례 하나씩:
  ⓐ 옆 PDF — soffice를 PATH에서 뺀 하위 실행: PPTX + 더 최신 PDF → 추출 호출마다 캡처 1장 · PDF 문서도 같다
  ⓑ 멈춤 — soffice 없음 · 옆 PDF 없음 · 실호출 → 사전 점검에서 멈춤(파싱 0 · 빠진 것 문구) · `--no-capture`면 진행 + 기록
  ⓒ 그림 입력 — 게이트웨이가 그림을 거부하면 박힌 그림이 없는 PPT도 사전 점검에서 멈춘다
  ⓓ soffice — 있으면 PPTX만으로 호출마다 캡처 1장 · 숨김 슬라이드가 있어도 쪽이 맞는다(없는 기계면 그 사실을 적는다)
  ⓔ 환경 점검 — `doctor --env`에 PyMuPDF · soffice · 「PPT 캡처」 줄 · `import pymupdf`(폐기 경고 0)
네 벌(ⓕ)은 마감 한 줄의 네 벌 단계가 잰다(mock은 사전 점검 · 캡처를 타지 않는다).
"""
from __future__ import annotations

import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done   # noqa: F401

import fake_gateway as FG                                 # noqa: E402
from fake_gateway import CALLS, Live                      # noqa: E402
from g6_common import SCHEMA                              # noqa: E402
from fixtures.make_ppt import build_report               # noqa: E402
import run as RUN                                        # noqa: E402
from cli import ingest as IG                             # noqa: E402
from core.build import deck as DK, extract as EX, extract_unit as EU   # noqa: E402
from core.llm import check as CK, gateway, narrow        # noqa: E402
from parser import pipeline, render, tagger as TG        # noqa: E402
from parser.adapters import basic_pdf as PDF, basic_ppt as PX   # noqa: E402

DT = "b128_ppt"
NO_OFFICE = os.pathsep.join(d for d in os.environ.get("PATH", "").split(os.pathsep)
                            if not any((Path(d) / n).exists() for n in ("soffice", "libreoffice")))
REJECT = []                                              # 그림 입력을 거부하는 게이트웨이(ⓒ)


def _quiet():
    return contextlib.redirect_stdout(io.StringIO())


def _extract(body):
    return {"entities": [], "relations": [], "about": []}


def _post(url, payload, key, timeout):
    """가짜 전송 — 채팅 왕복은 답 하나 · 그림이 든 채팅은 `REJECT`면 HTTP 400 · 임베딩은 가짜 게이트웨이 그대로."""
    if not str(url).endswith("/chat/completions"):
        return FG._post(url, payload, key, timeout)
    imgs = any(isinstance(m.get("content"), list) for m in payload.get("messages") or [])
    if imgs and REJECT:
        raise gateway.GatewayError(400, "image input not supported", "fake")
    return {"choices": [{"message": {"content": "ok"}}], "usage": {"prompt_tokens": 1, "completion_tokens": 1}}


@contextlib.contextmanager
def _live():
    with Live(by={"deck": lambda b: {"slides": [], "overview": None}, "extract": _extract,
                  "outline": lambda b: {"headings": []}, "reinforce": lambda b: {"pick": None}}) as lv:
        gateway._post = _post
        gateway.chat = lambda m, json_schema=None, point=None, **kw: ({"text": "그림 요약(가짜)"} if point == "image_summary"
                                                                      else lv.chat(m, json_schema, point, **kw))
        yield lv


def _images_per_call(path, adapter, did):
    """실경로 파싱(캡처 켬) → 추출 — 추출 호출마다 붙은 그림 수 · 캡처 사유."""
    EU._PAGES.clear()
    EX.invalidate(did)
    with _live(), _quiet():
        res = pipeline.parse(adapter, did, str(path), closed_list=TG.closed_list("process"), layer="process",
                             deck=DK.decker(["process"]) if adapter is PX else None, capture=True)
        env = res.envelope
        env["doc_type"] = DT
        for c in env["chunks"]:
            c["doc_type"] = DT
        run_document(dict(env))
        calls = [b.get("_images", 0) for p, b in CALLS if p == "extract"]
    return calls, EU._PAGES.get(did, ({}, None))[1], (env.get("capture") or {}).get("pages")


if len(sys.argv) > 2 and sys.argv[1] == "--child":           # 하위 실행 — soffice 없는 PATH에서(ⓐ)
    _path = Path(sys.argv[2])
    _ad = PDF if _path.suffix == ".pdf" else PX
    print(json.dumps(dict(zip(("calls", "why", "pages"), _images_per_call(_path, _ad, _path.stem))), ensure_ascii=False))
    sys.exit(0)


def _register():
    p = _P.schemas(f"{DT}.json")
    _P.ensure(p)
    p.write_text(json.dumps({**SCHEMA, "doc_type": DT, "layer": "process", "payload_kind": "prose"}, ensure_ascii=False),
                 encoding="utf-8")
    ad = _P.adapters(f"{DT}.py")
    _P.ensure(ad)
    ad.write_text("from parser.adapters import basic_ppt\n"
                  f"ADAPTER = {{**basic_ppt.ADAPTER, \"doc_type\": {DT!r}}}\nextract = basic_ppt.extract\n", encoding="utf-8")
    dts = store.read(store.DOC_TYPES, {})
    dts[DT] = {"doc_type": DT, "status": "registered", "layer": "process", "schema": f"schemas/{DT}.json",
               "adapter": str(ad), "schema_version": 1}
    store.write(store.DOC_TYPES, dts)


def _unregister():
    dts = store.read(store.DOC_TYPES, {})
    dts.pop(DT, None)
    store.write(store.DOC_TYPES, dts)
    _P.schemas(f"{DT}.json").unlink(missing_ok=True)
    _P.adapters(f"{DT}.py").unlink(missing_ok=True)


def _deck(path):
    """보고서 꼴 덱 — 문서마다 표지 제목에 이름을 붙인다(같은 내용 문서는 중복 거부라 시험마다 다른 문서)."""
    from pptx import Presentation
    p = build_report(path)
    prs = Presentation(str(p))
    prs.slides[0].shapes.title.text += f" — {p.stem}"
    prs.save(p)
    return p


def _child(path):
    r = subprocess.run([sys.executable, str(Path(__file__).resolve()), "--child", str(path)], capture_output=True,
                       text=True, cwd=str(ROOT), env={**os.environ, "PATH": NO_OFFICE, "USE_MOCK": "1"},
                       stdin=subprocess.DEVNULL, timeout=300)
    try:
        return json.loads(r.stdout.strip().splitlines()[-1])
    except (ValueError, IndexError):
        return {"calls": None, "why": (r.stderr or r.stdout)[-300:]}


TMP = Path(tempfile.mkdtemp(prefix="b128_"))
with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register()
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")

# ────────────────────────────────────────────────────────────── ⓐ 옆 PDF · PDF 문서 (soffice 없는 PATH)
print("\n■ ⓐ 옆 PDF · PDF 문서 — soffice 없이")
_pa = _deck(TMP / "B128A.pptx")
_pdf = TMP / "B128A.pdf"
if SOFFICE:                                                 # 옆 PDF는 사람이 PowerPoint에서 저장한 것 — 여기서는 LibreOffice가 만든다
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(TMP), str(_pa)],
                   capture_output=True, timeout=180, stdin=subprocess.DEVNULL)
_beside = _pdf.exists()
if _beside:
    os.utime(_pdf, (_pa.stat().st_mtime + 60, _pa.stat().st_mtime + 60))   # 원본보다 새것
    _pd = TMP / "B128P.pdf"
    shutil.copy(_pdf, _pd)
    _ra, _rp = _child(_pa), _child(_pd)
    show("ⓐ soffice를 뺀 PATH — PPTX + 더 최신 PDF → 추출 호출마다 캡처 1장 · PDF 문서도 같다(PyMuPDF만으로)",
         _ra.get("calls") and set(_ra["calls"]) == {1} and _ra.get("why") is None
         and _rp.get("calls") and set(_rp["calls"]) == {1} and _rp.get("why") is None,
         f"PPTX 호출별 그림 {_ra.get('calls')} · PDF {_rp.get('calls')} · 사유 {_ra.get('why')} / {_rp.get('why')}")
else:
    show("ⓐ 옆 PDF — 이 기계는 PDF를 만들 soffice가 없다(허브가 실경로를 돈다)", True, "건너뜀 — soffice 없음")

# ────────────────────────────────────────────────────────────── ⓑ 멈춤 · ⓒ 그림 입력 (사전 점검 — 실호출 갈래)
print("\n■ ⓑ 캡처 길이 없으면 파싱 전에 멈춤 · ⓒ 그림 입력")
_pb = _deck(TMP / "B128B.pptx")
_path0, _narrow0, _chat0 = os.environ.get("PATH", ""), narrow.narrow_choice, CK.chat_roundtrip
narrow.narrow_choice = lambda: ("overlap", "시험")
CK.chat_roundtrip = lambda cfg, add: (add("②", "도달", True), {"choices": [{"message": {"content": "ok"}}],
                                                               "usage": {"prompt_tokens": 1, "completion_tokens": 1}})[1]
try:
    os.environ["PATH"] = NO_OFFICE
    with _live():
        _o1 = io.StringIO()
        with contextlib.redirect_stdout(_o1):
            _r1 = IG.ingest_file(_pb, DT, ask=False)
        _parsed1 = _P.parsed("B128B.json").exists()
        _o2 = io.StringIO()
        with contextlib.redirect_stdout(_o2):
            _r2 = IG.ingest_file(_pb, DT, ask=False, no_capture=True)
        _x2 = [b for p, b in CALLS if p == "extract"]
    _reg = store.read(store.DOC_REGISTRY, {}).get("B128B") or {}
    show("ⓑ soffice 없음 · 옆 PDF 없음 · 실호출 → 사전 점검 「캡처」에서 멈춤(파싱 0 · 빠진 것 · 다음 줄) · --no-capture면 진행 + 기록",
         _r1.get("preflight_fail") and not _parsed1 and "**캡처**에서 멈춘다" in _o1.getvalue()
         and "LibreOffice" in _o1.getvalue() and "--no-capture" in _o1.getvalue()
         and _r2.get("status") == IG.OK and _reg.get("capture_off") and _x2 and not any(b.get("_images") for b in _x2),
         " / ".join(ln.strip() for ln in _o1.getvalue().splitlines() if "사전 점검" in ln or "빠진 것" in ln)[:300])
    os.environ["PATH"] = _path0
    REJECT.append(True)
    from cli.preflight import has_images
    _pc = TMP / "B128C.pptx"
    from pptx import Presentation
    _prs = Presentation()
    _prs.slides.add_slide(_prs.slide_layouts[1]).shapes.title.text = "그림 없는 슬라이드"
    _prs.save(_pc)
    with _live():
        _o4 = io.StringIO()
        with contextlib.redirect_stdout(_o4):
            _ok4 = IG._preflight_once(str(_pc), None, False, DT)
    show("ⓒ 게이트웨이가 그림을 거부하면 박힌 그림이 없는 PPT도 사전 점검 「그림」에서 멈춘다(다음 줄 --no-capture)",
         not has_images(_pc) and _ok4 is False and "**그림**에서 멈춘다" in _o4.getvalue()
         and "--no-capture" in _o4.getvalue(),
         " / ".join(ln.strip() for ln in _o4.getvalue().splitlines() if "사전 점검" in ln or "다음 줄" in ln)[:300])
finally:
    os.environ["PATH"] = _path0
    narrow.narrow_choice, CK.chat_roundtrip = _narrow0, _chat0
    REJECT.clear()

# ────────────────────────────────────────────────────────────── ⓓ soffice (있으면 PPTX만으로 · 숨김 슬라이드)
print("\n■ ⓓ soffice — PPTX만으로 · 숨김 슬라이드")
if render.capture_path()[0]:
    _pdd = _deck(TMP / "B128D.pptx")
    from pptx import Presentation
    _prs = Presentation(str(_pdd))
    _hid = _prs.slides.add_slide(_prs.slide_layouts[1])
    _hid.shapes.title.text = "숨긴 슬라이드"
    _hid._element.set("show", "0")
    _prs.save(_pdd)
    _cd, _whyd, _pg = _images_per_call(_pdd, PX, "B128D")
    show("ⓓ soffice가 있으면 PPTX만으로 추출 호출마다 캡처 1장 · 숨김 슬라이드가 있어도 렌더 쪽 = 보이는 슬라이드",
         _cd and set(_cd) == {1} and _whyd is None and _pg == 5, f"호출별 그림 {_cd} · 보이는 쪽 {_pg} · 사유 {_whyd}")
else:
    show("ⓓ soffice — 이 기계에는 soffice(Impress)가 없다(허브가 실경로를 돈다)", True, f"{render.capture_path()[2]}")

# ────────────────────────────────────────────────────────────── ⓔ 환경 점검
print("\n■ ⓔ 환경 점검 — doctor --env")
_env = subprocess.run([sys.executable, str(ROOT / "doctor.py"), "--env"], capture_output=True, text=True, cwd=str(ROOT),
                      stdin=subprocess.DEVNULL, timeout=300).stdout
with warnings.catch_warnings(record=True) as _w:
    warnings.simplefilter("always")
    _mod = render.pymupdf()
show("ⓔ doctor --env에 PyMuPDF · soffice · 「PPT 캡처」 줄 · `import pymupdf`로 돈다(폐기 경고 0)",
     "pymupdf" in _env and "soffice" in _env and "PPT 캡처 —" in _env and _mod is not None
     and _mod.__name__ == "pymupdf" and not [x for x in _w if issubclass(x.category, DeprecationWarning)],
     " / ".join(ln.strip() for ln in _env.splitlines() if "pymupdf" in ln or "soffice" in ln or "PPT 캡처" in ln)[:300])
_unregister()

done()
