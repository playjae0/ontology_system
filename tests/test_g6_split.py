# -*- coding: utf-8 -*-
"""G6 — **구조는 규칙 나무가 기준 · LLM은 고칠 곳만 이유와 함께 · 시트마다 확인 · 구조까지만** (B126).

창작 표본(시트 셋 · 수치는 근거가 아니다 [정정] 50) · `edit` 답은 가짜 게이트웨이(실호출 갈래 — 입력 조립 · 검사는 실물) ·
`--split-only`는 mock 힌트(`{"edits": …}` 꼴). LLM 품질은 사내에서 잰다. 대표 사례 하나씩:
  ⓐ 고침 없는 `edit` = `rules` · `rules` = 구조 함수 없는 파싱(B118 청크)
  ⓑ 고침(제목 더함 · 부모 바꿈 · 대상) → 사슬에 유닛 · 이유가 md에 · 어긋난 고침만 버린다
  ⓒ 확인 고리 — 힌트 → 둘째 호출 입력에 앞 출력 + 힌트 → 다시 찍기 → `c` → 확정 기록 → 보존분 없이 다시 파싱해도 구조 LLM 0 · 같은 나무
  ⓓ 형태 단서 — 「가.」 아래 「1.」 다시 시작 · 「1)~11)」 이어짐이 LLM 입력에
  ⓔ 규칙 나무가 빈 시트만 llm · `structure_mode=llm`이면 `edit` 호출 0
  ⓕ `--split-only` — 추출 0 · 그래프 쓰기 0 · `export/split_<doc_id>.md`
  B127 ⓓ 확정한 시트에 힌트 → 확정을 비켜 두고 기록된 힌트 + 새 힌트로 다시 → `c`면 새 나무 · 쌓인 힌트 · ⓔ `q`면 기록 그대로
네 벌(ⓖ)은 마감 한 줄의 네 벌 단계가 잰다.
"""
from __future__ import annotations

import contextlib
import io
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done   # noqa: F401

from fake_gateway import CALLS, Live                     # noqa: E402
from g6_common import _register, _run, _unregister, data_hash   # noqa: E402
from openpyxl import Workbook                            # noqa: E402
import run as RUN                                        # noqa: E402
from cli import _screen, split_report, structure_screen as SS   # noqa: E402
from core.build import structure as ST                   # noqa: E402
from parser import pipeline, struct_map                  # noqa: E402
from parser.adapters import basic_prose_xlsx as BX       # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="b126_"))
HINTS = ROOT / "tests" / "fixtures" / "extract_hints"
DID = "B126S"
E, N, M = "설비 사양", "공통사양", "메모"
SHEETS = {E: ["가. 공급부", "공급부는 소재를 넣는다.", "1. 기본 사양", "속도 기준은 50이다.", "2. 지정 사양", "하중 기준은 3이다.",
              "나. 적층부", "1. 기본 사양", "장력 기준은 30이다.", "설비 메모", "메모 본문이다."],
          N: ["2. 공통사양"] + [f"{i}) 항목{i}의 기준은 {i}이다." for i in range(1, 12)],
          M: ["회의에서 정한 것", "속도는 다음 회의에서 다시 본다.", "하중은 그대로 둔다."]}
WHY = {"대상": "줄 머리의 이름이 설비 단위다", "더함": "번호 없는 제목 줄", "부모": "메모는 시트 전체에 걸린다",
       "다시": "지정 사양은 기본 사양과 한 덩어리"}
AGAIN = "지정 사양은 제목이 아니다"
EDITS = [{"row": 1, "what": "대상", "name": "공급부", "category": "Unit", "why": WHY["대상"]},
         {"row": 10, "what": "제목 더함", "parent": 7, "role": "주제", "why": WHY["더함"]},
         {"row": 10, "what": "부모 바꿈", "parent": None, "why": WHY["부모"]},
         {"row": 99, "what": "제목 뺌", "why": "없는 줄"}]                 # 시트에 없는 줄 — 그 고침만 버린다
_E = lambda e: {k: e.get(k) for k in ("row", "what", "parent", "to_row", "role", "name", "category", "why")}


def _xlsx():
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in SHEETS.items():
        ws = wb.create_sheet(title=name)
        for i, v in enumerate(rows, 1):
            ws.cell(row=i, column=1, value=v)
    p = TMP / f"{DID}.xlsx"
    wb.save(p)
    return p


def _parse(p, structure="edit", confirmed=None):
    st = None if structure is None else ST.structurer(["process"])
    res = pipeline.parse(BX, DID, str(p), closed_list=[], sheet_roles={n: "prose" for n in SHEETS},
                         structure=st, structure_confirmed=confirmed)
    return res


def _by(res, frame):
    return [(c["source_locator"], c["text"]) for c in split_report.chunks_of(res.envelope, frame)]


def _chain(res, frame, needle):
    c = next(c for c in split_report.chunks_of(res.envelope, frame) if needle in c["text"])
    return [x.get("name") or x.get("text") for x in (c.get("meta") or {}).get("structure") or []]


def _structure(body):
    """가짜 ⑪ — `edit`이면 시트 E만 고침(힌트가 있으면 둘째 판) · 그 밖은 고침 0 · llm이면 제목 0(판정 없음 아님 — 빈 나무)."""
    if body.get("mode") != "edit":
        return {"headings": []}
    if body["frame"] != E:
        return {"edits": [], "headings": []}
    eds = [_E(e) for e in EDITS[:1]] if not body.get("hints") else [_E(e) for e in EDITS]
    if any(AGAIN in h for h in body.get("hints") or []):      # 확정한 시트를 다시 연 힌트(B127 ③) — 새 답
        eds.append(_E({"row": 5, "what": "제목 뺌", "why": WHY["다시"]}))
    return {"edits": eds, "headings": [{"row": 7, "role": "대상", "name": "적층부", "category": "Unit", "why": "둘째 설비 단위"}]}


with contextlib.redirect_stdout(io.StringIO()):
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_x = _xlsx()
ST.confirmed_path().unlink(missing_ok=True)

# ────────────────────────────────────────────────────────────── ⓐ 바닥 — rules = 구조 함수 없음 · 고침 없는 edit = rules
print("\n■ ⓐ 규칙 나무가 바닥")
ST.MODE = "rules"
_r0 = _parse(_x, structure=None)
struct_map.invalidate(DID)
_r1 = _parse(_x)
ST.MODE = "edit"
struct_map.invalidate(DID)
CALLS.clear()
with Live(by={"structure": lambda b: {"edits": [], "headings": []}}):
    _r2 = _parse(_x)
_modes = {f: x.get("mode") for f, x in (_r2.report["structure"]["프레임"]).items()}
show("ⓐ `rules` = 구조 함수 없는 파싱(B118 청크) · 고침 없는 `edit` = `rules`(규칙 나무 시트) · 구조 LLM 0(rules)",
     all(_by(_r1, f) == _by(_r0, f) for f in (E, N)) and all(_by(_r2, f) == _by(_r1, f) for f in (E, N))
     and _modes.get(E) == "edit" and (_r1.report.get("structure") or {}).get("방식", {}).get(E) == "rules"
     and _r1.report["structure"]["프레임"][E]["headings"],                 # rules도 구조 확인에 규칙 나무가 선다
     f"청크 E {len(_by(_r0, E))}/{len(_by(_r1, E))}/{len(_by(_r2, E))} · 방식 {_modes}")

# ────────────────────────────────────────────────────────────── ⓑ 고침 → 사슬 · 이유 · 버림
print("\n■ ⓑ 고칠 곳만 이유와 함께")
struct_map.invalidate(DID)
CALLS.clear()
with Live(by={"structure": lambda b: {**_structure({**b, "hints": ["다 고친다"]})}}):
    _r3 = _parse(_x)
_f3 = _r3.report["structure"]["프레임"][E]
_md = split_report.write_md(DID, _r3.report, _r3.envelope).read_text(encoding="utf-8")
show("ⓑ 대상 · 제목 더함 · 부모 바꿈 → 「속도」 사슬에 유닛 공급부 · 「메모」는 시트 바로 아래 · 이유가 md에 · 없는 줄 고침만 버림",
     _chain(_r3, E, "속도 기준")[:1] == ["공급부"] and _chain(_r3, E, "메모 본문")[-1:] == ["설비 메모"]
     and "적층부" not in _chain(_r3, E, "메모 본문") and len(_f3["edits"]) == 3 and len(_f3["edits_dropped"]) == 1
     and all(WHY[k] in _md for k in ("대상", "더함", "부모")) and "시트에 없는 줄" in _md,
     f"사슬 {_chain(_r3, E, '속도 기준')} / {_chain(_r3, E, '메모 본문')} · 고침 {len(_f3['edits'])} · 버림 {len(_f3['edits_dropped'])}")

# ────────────────────────────────────────────────────────────── ⓒ 확인 고리 — 힌트 · 다시 · 확정 · 재생
print("\n■ ⓒ 시트마다 확인 — 힌트 → 다시 → c → 확정 → 재생")
struct_map.invalidate(DID)
CALLS.clear()
HINT_TEXT = "설비 메모는 시트 전체에 걸린다"
_answers = iter([HINT_TEXT, "c", "c", "c"])
_ask0 = _screen.ask
_screen.ask = lambda prompt: next(_answers)
_buf = io.StringIO()
with Live(by={"structure": _structure}):
    _r4 = _parse(_x)
    with contextlib.redirect_stdout(_buf):
        _r5, _go = SS.gate(_r4, DID, lambda: _parse(_x, confirmed=ST.confirmed_of(DID)), layers=["process"])
_screen.ask = _ask0
_calls = [b for p, b in CALLS if p == "structure" and b.get("frame") == E]
_cf = ST.confirmed_of(DID)
_scr = _buf.getvalue()
struct_map.invalidate(DID)                               # 보존분 없이 — 확정 기록만으로
CALLS.clear()
with Live(by={"structure": _structure}):
    _r6 = _parse(_x, confirmed=ST.confirmed_of(DID))
_rows = lambda r: [(h["row"], h.get("parent"), h.get("role"), h.get("name")) for h in r.report["structure"]["프레임"][E]["headings"]]
show("ⓒ 힌트 → 둘째 호출 입력에 앞 출력 + 힌트 · 다시 찍은 나무에 고침 · c → 확정 기록(힌트) · 보존분 없이 다시 파싱해도 구조 LLM 0 · 같은 나무",
     _go and len(_calls) == 2 and _calls[1].get("previous") and _calls[1].get("hints") == [HINT_TEXT]
     and _scr.count(f"구조 확인 — 시트 '{E}'") == 2 and WHY["부모"] in _scr
     and _cf.get(E, {}).get("hints") == [HINT_TEXT]
     and not [p for p, _b in CALLS if p == "structure"] and _rows(_r6) == _rows(_r5)
     and _chain(_r6, E, "메모 본문")[-1:] == ["설비 메모"],
     f"호출 {len(_calls)} · 확정 {sorted(_cf)} · 재생 호출 {len([p for p, _b in CALLS if p == 'structure'])}")

# ────────────────────────────────────────────────────────────── ⓓ 형태 단서
_form = _calls[0]["form"]["번호 이어짐"] if _calls else []
show("ⓓ 형태 단서가 LLM 입력에 — 「1.」이 「나.」 아래에서 다시 시작 · 「1)」 11개 이어짐(공통사양)",
     any("'1.' 8행 — '가.'(7행) 아래에서 다시 시작" in x for x in _form)
     and any("'1)' 2~12행 11개 이어짐" in x for x in split_report.SE.run_lines(split_report.SE.number_runs(
         _r4.report["structure"]["프레임"][N]["labels"]))),
     f"{_form}")

# ────────────────────────────────────────────────────────────── B127 ③ 확정한 시트도 힌트로 다시 연다
print("\n■ B127 ③ 확정한 시트에 힌트 — 확정을 비켜 두고 다시 · c면 덮고 q면 그대로")


def _reopen(answers):
    """확정 기록만으로 재생한 파싱(보존분 없음) → 구조 확인 고리 — 돌려주는 것은 `(E 시트 호출 본문들, 확정 기록 바이트)`."""
    struct_map.invalidate(DID)
    CALLS.clear()
    it = iter(answers)
    _screen.ask = lambda prompt: next(it)
    with Live(by={"structure": _structure}), contextlib.redirect_stdout(io.StringIO()):
        SS.gate(_parse(_x, confirmed=ST.confirmed_of(DID)), DID, lambda: _parse(_x, confirmed=ST.confirmed_of(DID)),
                layers=["process"])
    _screen.ask = _ask0
    return [b for p, b in CALLS if p == "structure" and b.get("frame") == E], ST.confirmed_path().read_bytes()


_old = ST.confirmed_of(DID)[E]
_c7, _b7 = _reopen([AGAIN, "c", "c", "c"])
_new = ST.confirmed_of(DID)[E]
show("B127 ⓓ 확정한 시트에 힌트 → 다시 묻는 입력의 힌트 = 기록된 힌트 + 새 힌트 · 앞 출력 = 확정된 나무 → c → 기록이 새 나무 · 쌓인 힌트",
     len(_c7) == 1 and _c7[0].get("hints") == [HINT_TEXT, AGAIN]
     and [h["row"] for h in (_c7[0].get("previous") or {}).get("headings") or []] == [h["row"] for h in _old["headings"]]
     and 5 in [h["row"] for h in _old["headings"]] and 5 not in [h["row"] for h in _new["headings"]]
     and _new["hints"] == [HINT_TEXT, AGAIN] and not ST.SUSPENDED,
     f"호출 {len(_c7)} · 힌트 {_c7[0].get('hints') if _c7 else None} · 기록 제목 {[h['row'] for h in _new['headings']]}")
_c8, _b8 = _reopen(["하나 더 바꿔 본다", "q"])
show("B127 ⓔ 확정한 시트에 힌트를 준 뒤 q → 확정 기록이 바이트까지 그대로",
     len(_c8) == 1 and _b8 == _b7 and not ST.SUSPENDED, f"호출 {len(_c8)} · 같음 {_b8 == _b7}")

# ────────────────────────────────────────────────────────────── ⓔ 빈 규칙 나무 → llm · structure_mode=llm
print("\n■ ⓔ 규칙 나무가 빈 시트만 llm · structure_mode=llm")
_m4 = {f: x.get("mode") for f, x in _r4.report["structure"]["프레임"].items()}
struct_map.invalidate(DID)
ST.confirmed_path().unlink(missing_ok=True)
ST.MODE = "llm"
CALLS.clear()
with Live(by={"structure": _structure}):
    _r7 = _parse(_x)
ST.MODE = "edit"
show("ⓔ 규칙 나무가 빈 시트(메모)만 llm · 나머지 edit · `structure_mode=llm`이면 시트 전부 llm(B124 입력 — 규칙 나무 · edit 0)",
     _m4.get(M) == "llm" and _m4.get(E) == "edit" and _m4.get(N) == "edit"
     and all(x.get("mode") == "llm" for x in _r7.report["structure"]["프레임"].values())
     and CALLS and all(b.get("mode") != "edit" and "rules_tree" not in b for p, b in CALLS if p == "structure"),
     f"edit 기본 {_m4} · llm {len([1 for p, _b in CALLS if p == 'structure'])}호출")

# ────────────────────────────────────────────────────────────── ⓕ --split-only
print("\n■ ⓕ ingest-file --split-only")
_register("process", "b126_process")
(HINTS / f"{DID}@structure.json").write_text(json.dumps({E: {"edits": EDITS[:1], "headings": []}}, ensure_ascii=False),
                                             encoding="utf-8")
struct_map.invalidate(DID)
_md_p = _P.export(f"split_{DID}.md")
_md_p.unlink(missing_ok=True)
_h0 = data_hash()
_out = _run("ingest-file", str(_x), "--doc-type", "b126_process", "--allow-mock", "--coord", "노칭",
            "--sheets", "1-3:prose", "--split-only")
(HINTS / f"{DID}@structure.json").unlink(missing_ok=True)
_unregister("b126_process")                              # 등록 단은 클린이 남긴다 — 시험 몫을 치운다
ST.confirmed_path().unlink(missing_ok=True)
show("ⓕ `--split-only` — 구조까지(추출 0 · 그래프 쓰기 0) · md가 선다(edit 고침 · 이유)",
     data_hash() == _h0 and _md_p.exists() and WHY["대상"] in _md_p.read_text(encoding="utf-8")
     and "구조까지" in _out and "추출" not in "".join(ln for ln in _out.splitlines() if ln.strip().startswith("[단계]")),
     " / ".join(ln.strip() for ln in _out.splitlines() if "구조까지" in ln or "split_" in ln)[:300])

done()
