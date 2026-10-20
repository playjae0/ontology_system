# -*- coding: utf-8 -*-
"""G6 — **구조 v2** — 시트 전문과 형태 힌트를 보고 LLM이 나무 · 청크 구간 · 주인을 함께 · 역할은 층 중립 · 층 힌트는 config (B124).

창작 표본(시트 넷 · 수치는 근거가 아니다 [정정] 50) · 구조 호출은 mock 힌트(`<doc>@structure.json` — 나무 + 청크 구간)로 ·
층 중립(ⓓ)만 가짜 게이트웨이(실호출 갈래 — 지시문 조립 · 검사는 실물). LLM 품질은 사내에서 잰다. 대표 사례 하나씩:
  ⓐ 형제 · 부모 — 「2. 공통사양」 아래 1)~11) · 10)에만 가) 나): 1)~11)이 같은 단의 형제 · 한 구간 · 가) 나)는 10)의 자식 ·
     권장값을 넘으면 형제 경계에서만 나뉜다
  ⓑ 이름 줄 — 「가. ○○」 + 요약 줄 + 빈 줄 + 「1. 기본 사양 …」: 대상 이름 = 「가.」 줄 · 「1.」의 사슬 = 시트 › 가. › 1. · 빈 줄 표시
  ⓒ 시트 = 대상 — 모든 청크의 주인이 시트(사슬에 시트 대상)
  ⓓ 중립 — 층 힌트가 없으면 구조 지시문에 층 어휘 0 · quality 층에서 힌트 없이 구조 호출이 선다(대상 카테고리 = 그 층 목록) ·
     같은 번호 꼴 형제의 단 · 역할이 갈리면 약한 구조 신호
  ⓔ 폴백 · 재생 — 구간 검사가 어긋난 시트만 지금 방식(결함 로그) · 같은 파일이면 구조 LLM 0 · 같은 청크
네 벌은 마감 한 줄의 네 벌 단계가 잰다.
"""
from __future__ import annotations

import contextlib
import io
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done   # noqa: F401

from fake_gateway import Live                            # noqa: E402
from openpyxl import Workbook                            # noqa: E402
import run as RUN                                        # noqa: E402
from core.build import structure as ST                   # noqa: E402
from core.llm import gateway                             # noqa: E402
from parser import pipeline, struct_map, struct_rule as SR, tuning   # noqa: E402
from parser.adapters import basic_prose_xlsx as BX       # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="b124_"))
HINTS = ROOT / "tests" / "fixtures" / "extract_hints"
DID = "B124S"
A, B, C, D = "공통사양", "6.공정별 상세사양", "언와인더 unit", "깨진 구간"
_A = ["2. 공통사양"] + [f"{i}) 항목{i}의 기준은 {i}이다." for i in range(1, 10)] + \
     ["10) 신호타이밍", "가) 입력 신호는 PLC 기준이다.", "나) 출력 신호는 100ms 안에 낸다.", "11) 통신은 이더넷이다."]
SHEETS = {A: _A,
          B: ["가. carrier 이송 물류", "carrier를 공정 사이로 옮기는 물류 전체를 말한다.", None,
              "1. 기본 사양 및 물류 운용 concept", "이송 속도는 1m/s이다.", "나. 검사 물류", "검사 물류는 별도 라인이다."],
          C: ["1) 기본사양", "속도는 100ppm이다.", "2) 지정사양", "장력은 30N이다."],
          D: ["1. 개요", "개요 본문이다.", "2. 범위", "범위 본문이다."]}
_kids = [{"text": t, "depth": 2, "role": "주제"} for t in _A[1:10]]
HINT = {A: {"headings": [{"text": "2. 공통사양", "depth": 1, "role": "주제"}] + _kids
            + [{"text": "10) 신호타이밍", "depth": 2, "role": "주제"},
               {"text": "가) 입력 신호는 PLC 기준이다.", "depth": 3, "role": "주제"},
               {"text": "나) 출력 신호는 100ms 안에 낸다.", "depth": 3, "role": "주제"},
               {"text": "11) 통신은 이더넷이다.", "depth": 2, "role": "주제"}],
            "chunks": [{"from_row": 1, "to_row": 14, "owner_row": 1}]},
        B: {"headings": [{"text": "가. carrier 이송 물류", "depth": 1, "role": "대상", "name": "carrier 이송 물류",
                          "category": "Unit"},
                         {"text": "1. 기본 사양 및 물류 운용 concept", "depth": 2, "role": "묶음"},
                         {"text": "나. 검사 물류", "depth": 1, "role": "대상", "name": "검사 물류", "category": "Unit"}],
            "chunks": [{"from_row": 1, "to_row": 2, "owner_row": 1}, {"from_row": 4, "to_row": 5, "owner_row": 4},
                       {"from_row": 6, "to_row": 7, "owner_row": 6}]},
        C: {"headings": [{"text": C, "depth": 0, "role": "대상", "name": C, "category": "Unit"},
                         {"text": "1) 기본사양", "depth": 1, "role": "묶음"}, {"text": "2) 지정사양", "depth": 1, "role": "묶음"}],
            "chunks": [{"from_row": 1, "to_row": 4, "owner_row": 0}]},
        D: {"headings": [{"text": "1. 개요", "depth": 1, "role": "주제"}, {"text": "2. 범위", "depth": 1, "role": "주제"}],
            "chunks": [{"from_row": 1, "to_row": 2, "owner_row": 1}]}}      # 3~4행이 구간 밖 — 검사가 버린다


def _xlsx():
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in SHEETS.items():
        ws = wb.create_sheet(title=name)
        for i, v in enumerate(rows, 1):
            if v is not None:
                ws.cell(row=i, column=1, value=v)
    p = TMP / f"{DID}.xlsx"
    wb.save(p)
    return p


def _parse(p):
    res = pipeline.parse(BX, DID, str(p), closed_list=[], sheet_roles={n: "prose" for n in SHEETS},
                         structure=ST.structurer(["process"]))
    return res, {n: [c for c in res.envelope["chunks"] if (c.get("meta") or {}).get("frame") == n] for n in SHEETS}


with contextlib.redirect_stdout(io.StringIO()):
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
(HINTS / f"{DID}@structure.json").write_text(json.dumps(HINT, ensure_ascii=False), encoding="utf-8")
struct_map.invalidate(DID)
_x = _xlsx()
_res, _by = _parse(_x)
_fr = _res.report["structure"]["프레임"]

print("\n■ ⓐ 형제 · 부모 · ⓑ 이름 줄 · ⓒ 시트 = 대상")
_hA = {h["row"]: h for h in _fr[A]["headings"]}
_dA = dict(zip([h["row"] for h in _fr[A]["headings"]], ST._depth_list(_fr[A]["headings"])))
_sib = [r for r in range(2, 15) if r not in (12, 13)]
_k1 = struct_map.CHUNK_MAX_CHARS
tuning.apply({"chunk_max_chars": 40})                    # 권장값을 작게 — 그 구간이 넘는다(창작 손잡이 — 근거가 아니다)
struct_map.invalidate(DID)
_res2, _by2 = _parse(_x)
tuning.apply({})
_starts = [int(c["source_locator"].split("!R")[-1].split("-")[0]) for c in _by2[A]]
show("ⓐ 1)~11)이 같은 단의 형제 · 가) 나)는 10)의 자식 · 한 구간이 청크 하나 · 권장값을 넘으면 형제 줄에서만 나뉜다",
     len({_dA[r] for r in _sib}) == 1 and _hA[12]["parent"] == 11 and _hA[13]["parent"] == 11
     and len(_by[A]) == 1 and _by[A][0]["meta"]["chunk_range"] == [1, 14]
     and len(_by2[A]) > 1 and all(s in _sib for s in _starts)
     and "\n".join(c["text"] for c in _by2[A]) == _by[A][0]["text"]        # 나뉘어도 글은 그대로(형제 줄도 본문에)
     and all(c["meta"]["chunk_range"] == [1, 14] and c["meta"].get("range_split") for c in _by2[A]),
     f"형제 단 {sorted({_dA[r] for r in _sib})} · 가) 부모 {_hA[12]['parent']} · 청크 {len(_by[A])} · 나뉜 시작 {_starts}")

_cB = {tuple(c["meta"]["chunk_range"]): c for c in _by[B]}
_ch = [x.get("name") or x["text"] for x in _cB[(4, 5)]["meta"]["structure"]]
_labB = {x["row"]: x for x in _fr[B]["labels"]}
show("ⓑ 대상 이름 = 「가.」 줄(요약 줄은 본문) · 「1.」의 사슬 = 가. › 1. · 빈 줄 자리가 이름표에 선다",
     _ch == ["carrier 이송 물류", "1. 기본 사양 및 물류 운용 concept"] and "carrier를" in _cB[(1, 2)]["text"]
     and "빈 줄" in _labB[4]["marks"] and _cB[(4, 5)]["section"].endswith("1. 기본 사양 및 물류 운용 concept"),
     f"사슬 {_ch} · 경로 {_cB[(4, 5)]['section']} · 4행 표시 {_labB[4]['marks']}")
show("ⓒ 시트 하나가 대상 하나 — 모든 청크의 주인이 시트(사슬에 시트 대상)",
     _by[C] and all((c["meta"].get("structure_frame") or {}).get("name") == C and not c["meta"]["structure"]
                    for c in _by[C] + _by2[C]),
     f"청크 {len(_by[C])}(나뉜 판 {len(_by2[C])}) · 시트 대상 {[(c['meta'].get('structure_frame') or {}).get('name') for c in _by[C]]}")

print("\n■ ⓓ 층 중립 · ⓔ 폴백 · 재생")
_vocab = {k for lay in ("process", "quality") for k in load_config(lay)["categories"]} | {"설비", "유닛", "사양 묶음"}
_bare = gateway.prompt_for(ST.PROMPT, ST._cfgs(["quality"]), "structure")
_with = gateway.prompt_for(ST.PROMPT, ST._cfgs(["process"]), "structure")
_qlines = [(1, "1. 불량 유형"), (2, "1) 버는 날카로운 돌기다."), (3, "2) 찍힘은 눌린 자국이다."), (4, "3) 주름은 접힌 자리다.")]
_qlabs = SR.labels("불량 유형", _qlines, [], {})


def _qstructure(body):
    cats = list(body["categories"])
    heads = [{"row": 1, "parent": None, "role": "주제", "name": None, "category": None}]
    heads += [{"row": r, "parent": 1, "role": "대상", "name": t.split(")")[1].split("는")[0].split("은")[0].strip(),
               "category": cats[0]} for r, t in _qlines[1:3]]           # 3) 하나만 형제와 다르게 둔다(창작)
    return {"headings": heads, "chunks": [{"from_row": 1, "to_row": 4, "owner_row": 1}]}


with Live(by={"structure": _qstructure}):
    _qst = ST.structurer(["quality"])("B124Q", "불량 유형", _qlabs)
_qfr = {"불량 유형": {"labels": _qlabs, "headings": _qst["headings"], "chunks": _qst.get("chunks") or []}}
_qw = ST.weak({"프레임": _qfr})
_qcats = set(load_config("quality")["categories"])
show("ⓓ 층 힌트가 없으면 구조 지시문에 층 어휘 0(힌트가 있으면 「층 어휘 — 층」 절) · quality 층에서 힌트 없이 구조 호출이 선다"
     "(대상 카테고리 = 그 층 목록) · 같은 번호 꼴 형제의 단 · 역할이 갈리면 약한 구조 신호",
     not any(w in _bare for w in _vocab) and "## 층 어휘 — process" in _with
     and {h["category"] for h in _qst["headings"] if h["role"] == "대상"} <= _qcats and _qst.get("chunks")
     and any("같은 번호 꼴 「1)」" in x for x in _qw.get("불량 유형") or []),
     f"층 어휘 {[w for w in _vocab if w in _bare]} · 대상 {[(h['name'], h['category']) for h in _qst['headings'] if h['role'] == '대상']} · "
     f"신호 {_qw.get('불량 유형')}")

_defects = store.path(store.DEFECTS).read_text(encoding="utf-8") if store.path(store.DEFECTS).exists() else ""
_res3, _by3 = _parse(_x)
_st3 = _res3.report["structure"]
show("ⓔ 구간 검사가 어긋난 시트만 지금 방식(결함 로그 · 구간 0) · 다른 시트는 구간 그대로 · 같은 파일이면 구조 LLM 0 · 같은 청크",
     _fr[D].get("chunks_dropped") and not any(c["meta"].get("chunk_range") for c in _by[D]) and len(_by[D]) >= 1
     and "청크 구간을 버렸다" in _defects and all(c["meta"].get("chunk_range") for n in (A, B, C) for c in _by[n])
     and _st3.get("재사용") == len(SHEETS) and not _st3.get("호출")
     and [c["text"] for c in _res3.envelope["chunks"]] == [c["text"] for c in _res.envelope["chunks"]],
     f"버림 {_fr[D].get('chunks_dropped')} · 재사용 {_st3.get('재사용')} · 호출 {_st3.get('호출', 0)}")

(HINTS / f"{DID}@structure.json").unlink(missing_ok=True)
done()
