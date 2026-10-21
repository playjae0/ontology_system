# -*- coding: utf-8 -*-
"""G6 — **계층 사다리** — 코드는 재료(행 전체 · 들여쓰기 깊이 · 꼴 서명 · 서수)와 적용 · 검사 · LLM은 사다리(문서당 1)와 시트 확인(시트당 1) (B132).

창작 표본(시트 둘 · 수치는 근거가 아니다 [정정] 50) — 한 시트에 사내 세 경우를 섞는다: 번호+글(띄어쓰기 섞임) · 단계마다 다른 열 ·
앞 공백. 사다리 · 시트 확인 답은 가짜 게이트웨이(실호출 갈래 — 입력 조립 · 검사 · 적용은 실물). LLM 판단의 품질은 사내
`--split-only`에서 잰다. 완료판정 하나에 대표 어서션 하나:
  ⓐ 재료 — 시작 열 · 들여쓰기 깊이 · 서명 · 서수 · 빈 행 뒤 · 다른 열의 제목이 줄에 · 한 열 시트는 줄이 셀 글 그대로
  ⓑ 서명 — 여섯 예가 정해 둔 꼴 목록 없이 · 두 벌 패턴 0
  ⓒ 사다리 — `9. › 가.(대상 · Unit) › 9) › A.`로 목차 · 단의 역할이 줄에 · 사다리 호출 문서당 1 · 입력에 줄 전문 0
  ⓓ 판단 필요 — 넷이 하나씩 표시 · 시트 확인 입력에 이유와 앞뒤 두 줄 · 답으로 바뀐다 · 코드가 혼자 정한 애매한 줄 0
  ⓔ 시트 확인 — 시트당 1 · 골격 보기(본문 접음 · 줄 전문 0) · 「대개 맞다」 0 · 사다리 고침이 둘째 시트에도
  ⓕ 폴백 — 사다리 호출 실패 · mock 힌트 없음이면 코드 사다리(`rules`와 같은 목차)
  ⓖ 화면 — 사다리 한 줄 · 들여쓰기 목차 · 판단 필요와 답 · 경로 (md 같게)
  ⓗ 방식 — `structure_mode` 값 둘 · 옛 값은 문면으로 멈춘다 · 읽는 곳 0이 된 경로 0
  B134 — 시트 확인 답 꼴이 어긋나면 그 칸만 버리고 사다리 목차(결함 로그 · 화면) · 사다리 답이 문자열이면 코드 사다리 · 구조 함수 안의
     TypeError는 파싱 실패 · 지도 훅을 받지 않는 옛 꼴 어댑터는 그대로 파싱
(B124 「구조 v2」 조각 — LLM 청크 구간 · 구간 검사 · 폴백은 경로를 지워(B132 ⑧) 이 조각으로 바꿨다 · ⓘ는 test_site · ⓙ는 네 벌 단계)
"""
from __future__ import annotations

import contextlib
import io
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done   # noqa: F401

from fake_gateway import CALLS, Live                     # noqa: E402
from g6_common import _run                               # noqa: E402
from openpyxl import Workbook                            # noqa: E402
import run as RUN                                        # noqa: E402
from cli import _screen, split_report, structure_screen as SS   # noqa: E402
from core.build import structure as ST, structure_edit as SE   # noqa: E402
from core.llm import gateway                             # noqa: E402
from core.state import knobs as KN                       # noqa: E402
from parser import pipeline, sig as SG, struct_map, struct_rule as SR, tuning   # noqa: E402
from parser.reader import read                           # noqa: E402
from parser.adapters import basic_prose_xlsx as BX       # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="b132_"))
DID, S1, S2 = "B132L", "공정별 상세사양", "검사 사양"
LONG = "다.패키징 unit — 파우치 포밍 · 스택 삽입 · 사이드 실링 · 전해액 주액"
#: `(행, 열, 글)` — 단계마다 다른 열(A 번호 · B 소제목 · C 본문) · 번호 뒤 띄어쓰기 섞임 · 앞 공백 · 판단 필요 넷(ⓑ 8행 · ⓐ 10행 ·
#: ⓒ 17행 · ⓓ 20행)
ROWS = {S1: [(1, "A", "6.공정별 상세사양"), (2, "A", "가."), (2, "B", "노칭 unit"), (3, "B", "1)기본사양"),
             (4, "C", "속도는 60ppm이다."), (5, "B", "2)지정사양"), (6, "B", "   A.입고전 검수"), (7, "C", "외관을 본다."),
             (8, "A", "B.입고후 검수"), (9, "C", "치수를 잰다."), (10, "A", "■ 참고"), (11, "C", "참고 본문이다."),
             (12, "A", "나.스태킹 unit"), (13, "B", "1) 기본사양"), (14, "C", "장력은 30N이다."), (15, "B", "2) 지정사양"),
             (16, "C", "폭은 20mm다."), (17, "B", "4) 검수방법"), (18, "C", "검수 본문이다."), (20, "C", "비고"),
             (21, "C", "비고 본문이다."), (22, "A", LONG), (23, "C", "패키징 본문이다.")],
        S2: [(1, "A", "7.검사 사양"), (2, "A", "가.검사 unit"), (3, "B", "1)주기"), (4, "C", "주기는 매일이다."),
             (5, "A", "■ 참고"), (6, "C", "둘째 참고 본문이다.")]}
FOLDED = "속도는 60ppm이다."                                 # 어느 판단 필요의 앞뒤도 아닌 본문 줄 — 구조 입력에 없어야 한다
_N = lambda why: {"rows": [], "why": why}
LADDER = {"levels": [{"sigs": ["9."], "role": "주제", "category": None, **_N("문서 장")},
                     {"sigs": ["가."], "role": "대상", "category": "Unit", **_N("유닛 단")},
                     {"sigs": ["9)"], "role": None, "category": None, **_N("소제목")},
                     {"sigs": ["A."], "role": None, "category": None, **_N("세목")}], "unnumbered": []}
RULE = {"shape": "9)", "role": "묶음", "category": None, "parent_shape": None, "said": "'9)' 줄은 사양 묶음"}
ANSWER = {8: {"heading": True, "parent": 5, "why": "들여쓰기만 얕다 — A.의 형제"}, 10: {"heading": False, "why": "참고는 본문"},
          17: {"heading": True, "parent": 12, "why": "3)이 빠진 발췌"}, 20: {"heading": True, "parent": None, "role": "주제",
                                                                       "why": "빈 행 뒤 비고는 주제"},
          5: {"heading": False, "why": "참고는 본문"}}


def _xlsx():
    wb = Workbook()
    wb.remove(wb.active)
    for name, cells in ROWS.items():
        ws = wb.create_sheet(title=name)
        for r, c, v in cells:
            ws[f"{c}{r}"] = v
    p = TMP / f"{DID}.xlsx"
    wb.save(p)
    return p


def _llm(body):
    """가짜 ⑪ — 사다리는 표본의 답 · 시트 확인은 판단 필요마다 답(빠짐없이) + 첫 시트에서 사다리 고침 한 줄."""
    if body.get("task") == "ladder":
        return LADDER
    ans = [{"row": n["row"], "parent": None, "role": None, "name": None, "category": None, **ANSWER[n["row"]]}
           for n in body.get("needs") or []]
    return {"answers": ans, "edits": [], "rules": [RULE] if body["frame"] == S1 else []}


def _parse(p, confirmed=None):
    return pipeline.parse(BX, DID, str(p), closed_list=[], sheet_roles={S1: "prose", S2: "prose"},
                          structure=ST.structurer(["process"]), structure_confirmed=confirmed)


def _tree(res, frame):
    return [(h["row"], h.get("parent")) for h in res.report["structure"]["프레임"][frame]["headings"]]


with contextlib.redirect_stdout(io.StringIO()):
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
ST.confirmed_path().unlink(missing_ok=True)
_x = _xlsx()

# ────────────────────────────────────────────────────────────── ⓐ 재료 · ⓑ 서명
print("\n■ ⓐ 재료 · ⓑ 꼴 서명")
_sh = next(s for s in read(str(_x))["sheets"] if s["name"] == S1)
_l, _rw, _m = BX._rows_of(_sh)
_lab = {x["row"]: x for x in SR.labels(S1, _l, _rw, BX._marks(_m), True, _m)}
_toc = next(s for s in read(str(ROOT / "tests" / "fixtures" / "raw" / "TOC01.xlsx"))["sheets"])
_col = BX._content_column(_toc)
_one = [(int(a[len(_col):]), str(v).strip()) for a, v in (_toc.get("cells") or {}).items()
        if a.startswith(_col) and a[len(_col):].isdigit() and str(v).strip()]
show("ⓐ 재료 — 시작 열 · 들여쓰기 깊이(열 · 앞 공백) · 서명 · 서수 · 빈 행 뒤 · 다른 열의 제목이 줄에(「가. | 노칭 unit」) · "
     "한 열 시트는 줄이 셀 글 그대로",
     [(_lab[r]["col"], _lab[r]["depth"]) for r in (1, 3, 6, 4)] == [("A", 0), ("B", 1), ("B", 2), ("C", 3)]
     and (_lab[3]["sig"], _lab[3]["ord"], _lab[8]["sig"], _lab[8]["ord"]) == ("9)", 1, "A.", 2)
     and _lab[20]["gap"] and not _lab[21]["gap"] and _lab[2]["text"] == "가. | 노칭 unit"
     and sorted(BX._rows_of(_toc)[0]) == sorted(_one),
     f"{[(r, _lab[r]['col'], _lab[r]['depth'], _lab[r]['sig']) for r in (1, 2, 3, 6, 8, 20)]}")
_six = {"1.기본사양": ("9.", 1), "(가) 범위": ("(가)", 1), "4)검수방법": ("9)", 4), "1-1 공급부": ("9-9", 1),
        "■ 개요": ("■", None), "b. 지정": ("a.", 2), "1. 기본사양": ("9.", 1)}
show("ⓑ 서명 — 여섯 예(띄어쓰기 · 길이와 무관)가 서명 · 서수로 선다 · 번호 패턴 두 벌(`HEADING_FAMILIES` · `NUMBER_SHAPES`) 0",
     all(SG.signature(t) == v for t, v in _six.items()) and SG.signature("1.기본사양 " + "가" * 80) == ("9.", 1)
     and not any(hasattr(struct_map, n) for n in ("HEADING_FAMILIES", "NUMBER_SHAPES", "heading_family")),
     str({t: SG.signature(t) for t in _six}))

# ────────────────────────────────────────────────────────────── ⓒ 사다리 · ⓓ 판단 필요 · ⓔ 시트 확인
print("\n■ ⓒ 계층 사다리 · ⓓ 판단 필요 · ⓔ 시트 확인 (가짜 게이트웨이)")
struct_map.invalidate(DID)
CALLS.clear()
with Live(by={"structure": _llm}):
    _r = _parse(_x)
_st = _r.report["structure"]
_f1, _f2 = _st["프레임"][S1], _st["프레임"][S2]
_lad = [b for p, b in CALLS if p == "structure" and b.get("task") == "ladder"]
_chk = {b["frame"]: b for p, b in CALLS if p == "structure" and b.get("task") == "check"}
_h1 = {h["row"]: h for h in _f1["headings"]}
_want = [(1, None), (2, 1), (3, 2), (5, 2), (6, 5), (8, 5), (12, 1), (13, 12), (15, 12), (17, 12), (20, None), (22, 1)]
_lin = json.dumps(_lad[0], ensure_ascii=False) if _lad else ""
show("ⓒ 사다리 답(`9. › 가.(대상 · Unit) › 9) › A.`)으로 목차가 맞고 단의 역할이 줄에(대상 이름 = 번호를 뗀 글) · 사다리 호출 문서당 1 · "
     "입력에 줄 전문 0(본문 · 긴 제목의 전문 없음)",
     _tree(_r, S1) == _want and (_h1[2]["role"], _h1[2]["name"], _h1[2]["category"]) == ("대상", "노칭 unit", "Unit")
     # 기대 변경(B137 ② — 1행은 시트 제목 · 역할은 시트 확인이 정한다 — 사다리 단의 역할 「주제」를 받지 않는다)
     and _h1[12]["name"] == "스태킹 unit" and _h1[1]["by"] == "시트 제목" and len(_lad) == 1
     and FOLDED not in _lin and LONG not in _lin and "stats" in _lad[0],
     f"목차 {_tree(_r, S1)} · 사다리 호출 {len(_lad)} · 2행 {(_h1[2].get('role'), _h1[2].get('name'))}")
_need = {n["row"]: n["kind"] for n in _f1["needs"]}
_t1 = {e.get("row"): e for e in _chk.get(S1, {}).get("toc") or []}
_toc0 = {h["row"] for h in _f1["toc0"]}
show("ⓓ 판단 필요 넷이 하나씩(ⓑ 8 · ⓐ 10 · ⓒ 17 · ⓓ 20) · 시트 확인 입력에 이유와 앞뒤 두 줄 · 답으로 바뀐다(20행 제목 · 10행 본문) · "
     "코드가 혼자 정한 애매한 줄 0(ⓐ · ⓓ 줄은 코드 목차에 없다 · 답 빠짐 0)",
     _need == {8: "ⓑ", 10: "ⓐ", 17: "ⓒ", 20: "ⓓ"}
     and all(_t1.get(r, {}).get("kind") == "판단 필요" and _t1[r].get("why") for r in _need)
     and _t1.get(9, {}).get("kind") == "앞뒤" and _t1.get(11, {}).get("kind") == "앞뒤" and _t1.get(19) is None
     and 20 in _h1 and _h1[20]["role"] == "주제" and 10 not in _h1 and not ({10, 20} & _toc0)
     and not _f1["missing"] and not _f2["missing"],
     f"판단 필요 {_need} · 입력 {[(r, _t1.get(r, {}).get('kind')) for r in (8, 9, 10, 11, 17, 20)]} · 빠짐 {_f1['missing']}")
_cin = json.dumps(_chk.get(S1) or {}, ensure_ascii=False)
_prompt = (ROOT / "prompts" / "2.4_ladder.md").read_text(encoding="utf-8")
show("ⓔ 시트 확인 — 시트당 1 · 입력은 골격 보기(본문 접음 · 줄 전문 0 · 긴 제목은 앞 40자 + 「…총 n자」) · 시트 확인 지시문에 「대개 맞다」 0 · "
     "사다리 고침(9) → 묶음)이 둘째 시트에도",
     sorted(_chk) == sorted([S1, S2]) and sum(1 for p, b in CALLS if b.get("task") == "check") == 2
     and any(e.get("kind") == "본문" for e in _chk[S1]["toc"]) and FOLDED not in _cin and LONG not in _cin
     and f"…총 {len(LONG)}자" in _cin and "대개 맞다" not in _prompt
     and [h.get("role") for h in _f2["headings"] if h["row"] == 3] == ["묶음"] and _h1[3]["role"] == "묶음",
     f"확인 호출 {sorted(_chk)} · 둘째 시트 3행 역할 {[h.get('role') for h in _f2['headings'] if h['row'] == 3]}")

# ────────────────────────────────────────────────────────────── ⓕ 폴백
print("\n■ ⓕ 폴백 — 코드 사다리")
ST.DOC.pop(DID, None)
ST.MODE = "rules"
struct_map.invalidate(DID)
_rules = _parse(_x)
ST.MODE = "ladder"
struct_map.invalidate(DID)


def _boom(_b):
    raise gateway.GatewayError(500, "가짜 실패", "fake")


with Live(by={"structure": _boom}):
    _fail = _parse(_x)
struct_map.invalidate(DID)
_bare = _parse(_x)                                       # mock · 힌트 없음
show("ⓕ 사다리 호출 실패 · mock 힌트 없음이면 코드 사다리로 목차가 선다(`rules`와 같은 목차) · 구조가 안 서는 일 0",
     all(_tree(x, f) == _tree(_rules, f) and _tree(x, f) for x in (_fail, _bare) for f in (S1, S2))
     and _fail.report["structure"]["사다리"] == "코드(호출 실패)" and _bare.report["structure"]["사다리"] == "코드(힌트 없음)",
     f"사다리 {_fail.report['structure'].get('사다리')} / {_bare.report['structure'].get('사다리')} · 목차 같음 "
     f"{[_tree(x, S1) == _tree(_rules, S1) for x in (_fail, _bare)]}")

# ────────────────────────────────────────────────────────────── 답 꼴 · 코드의 오류 (B134)
print("\n■ B134 — 어긋난 답 꼴은 그 답만 버린다 · 코드의 TypeError는 파싱 실패 · 훅 없는 옛 꼴 어댑터")
from core.state import store as _STO                     # noqa: E402 — 결함 로그 자리(그 모듈이 안다)
_dfp = _STO.path(_STO.DEFECTS)
_df0 = _dfp.read_text(encoding="utf-8") if _dfp.exists() else ""


def _int_answers(body):
    """첫 시트의 시트 확인 답 `answers`가 정수 — 나머지는 정상 답."""
    return {"answers": 5, "edits": [], "rules": []} if body.get("task") == "check" and body["frame"] == S1 else _llm(body)


ST.DOC.pop(DID, None)
struct_map.invalidate(DID)
with Live(by={"structure": _int_answers}):
    _ri = _parse(_x)
_fi1, _fi2 = (_ri.report.get("structure") or {}).get("프레임", {}).get(S1) or {}, \
    (_ri.report.get("structure") or {}).get("프레임", {}).get(S2) or {}
_new = (_dfp.read_text(encoding="utf-8") if _dfp.exists() else "")[len(_df0):]
_blk = split_report.sheet_block(S1, _fi1, []) if _fi1 else []
show("B134 ⓐ 시트 확인 답 answers가 정수 → 그 시트는 사다리 목차로 서고 구조 표가 있다 · 결함 로그 한 줄(문서 · 시트 · 칸) · 화면에 답 버림 · "
     "다른 시트는 정상 답대로",
     _ri.ok and _fi1 and [(h["row"], h.get("parent")) for h in _fi1["headings"]] == [(h["row"], h.get("parent")) for h in _fi1["toc0"]]
     and _fi1["needs"] and _fi1["shape_dropped"] == ["answers(int)"]
     and sum(1 for ln in _new.splitlines() if "답 꼴 어긋남" in ln and f"'{S1}'" in ln and "answers(int)" in ln) == 1
     and any("시트 확인 답 버림" in ln and "answers(int)" in ln for ln in _blk)
     and not _fi2.get("shape_dropped") and not _fi2.get("missing"),
     f"ok {_ri.ok} · 버림 {_fi1.get('shape_dropped')} · 둘째 시트 빠짐 {_fi2.get('missing')}")
ST.DOC.pop(DID, None)
struct_map.invalidate(DID)
with Live(by={"structure": lambda b: json.dumps(LADDER) if b.get("task") == "ladder" else _llm(b)}):
    _rs = _parse(_x)
show("B134 ⓑ 사다리 답이 문자열 → 코드 사다리로 선다(「코드(답 버림)」) · 파싱 성공",
     _rs.ok and (_rs.report.get("structure") or {}).get("사다리") == "코드(답 버림)",
     f"ok {_rs.ok} · 사다리 {(_rs.report.get('structure') or {}).get('사다리')} · {_rs.failures[:1]}")
ST.DOC.pop(DID, None)
struct_map.invalidate(DID)


def _type_error(_b):
    raise TypeError("B134 시험 — 구조 함수 안의 TypeError")


with Live(by={"structure": _type_error}):
    _rt = _parse(_x)
show("B134 ⓒ 구조 함수 안에서 TypeError → 파싱 실패이고 사유에 TypeError가 보인다(구조 없는 성공 0)",
     not _rt.ok and any("TypeError" in f["reason"] for f in _rt.failures) and not _rt.envelope,
     f"ok {_rt.ok} · {[f['reason'][:80] for f in _rt.failures]}")
ST.DOC.pop(DID, None)
struct_map.invalidate(DID)


class _OldAdapter:
    """지도 훅을 받지 않는 옛 꼴 어댑터(시험 사본 — `extract(raw)`)."""
    ADAPTER = BX.ADAPTER

    @staticmethod
    def extract(raw):
        return BX.extract(raw)


_ro = pipeline.parse(_OldAdapter, DID, str(_x), closed_list=[], sheet_roles={S1: "prose", S2: "prose"},
                     structure=ST.structurer(["process"]))
show("B134 ⓓ 지도 훅을 받지 않는 옛 꼴 어댑터가 지금처럼 파싱된다(훅 없이 · 청크가 선다)",
     _ro.ok and len(_ro.envelope["chunks"]) > 0 and "structure" not in _ro.report,
     f"ok {_ro.ok} · 청크 {len((_ro.envelope or {}).get('chunks') or [])} · {_ro.failures[:1]}")
ST.DOC.pop(DID, None)
struct_map.invalidate(DID)

# ────────────────────────────────────────────────────────────── ⓖ 화면
print("\n■ ⓖ 구조 확인 화면")
ST.DOC.pop(DID, None)
struct_map.invalidate(DID)
tuning.apply({"chunk_rows": [1, 2]})                     # 작은 표본이 깊은 단까지 잘리게(창작 손잡이 — 근거가 아니다)
_ask0 = _screen.ask
_it = iter(["c", "c"])
_screen.ask = lambda prompt: next(_it)
_buf = io.StringIO()
with Live(by={"structure": _llm}), contextlib.redirect_stdout(_buf):
    _rg, _go = SS.gate(_parse(_x), DID, lambda: _parse(_x), layers=["process"])
_screen.ask = _ask0
tuning.apply({})
_scr = _buf.getvalue()
_ind = {r: next((len(ln) - len(ln.lstrip()) for ln in _scr.splitlines() if ln.lstrip().startswith(f"{r}행 ")), -1) for r in (5, 6)}
_md = split_report.write_md(DID, _rg.report, _rg.envelope).read_text(encoding="utf-8")
# 기대 변경(B137 ⑥ — 사다리는 첫 서수 표기 · 판단 필요는 따로 목록이 아니라 목차 줄에 표지 + 답)
LINE = "계층 사다리: 시트 › 1.(주제) › 가.(대상 · Unit) › 1)(묶음) › A."
_l20 = [ln for ln in (_scr + "\n" + _md).splitlines() if ln.lstrip(" -").startswith("20행 ")]
show("ⓖ 화면 — 사다리 한 줄 · 들여쓰기 목차(A.가 2)보다 한 단 깊다) · 판단 필요와 답(목차 줄) · 청크마다 전체 경로 · md에 같은 줄",
     _go and LINE in _scr and LINE in _md and _ind[6] == _ind[5] + 2 > 2
     and len(_l20) == 2 and all(split_report.NEED_MARK in ln and "맨 위(LLM)" in ln for ln in _l20)
     and f"경로 {S1} › 6.공정별 상세사양 › 가. | 노칭 unit › 2)지정사양 › A.입고전 검수" in _scr,
     " / ".join(ln.strip()[:90] for ln in _scr.splitlines() if "계층 사다리" in ln or ln.lstrip().startswith("20행"))[:300])
ST.confirmed_path().unlink(missing_ok=True)

# ────────────────────────────────────────────────────────────── ⓗ 방식
print("\n■ ⓗ 방식 둘 · 옛 값")
_kn = _P.home() / "knobs.json"
_kn.write_text(json.dumps({"structure_mode": "edit"}), encoding="utf-8")
_out = _run("show", "knobs")
_kn.unlink()
show("ⓗ `structure_mode` 값 둘(rules · ladder) · 옛 값 edit는 명령이 문면으로 멈추고 바꿀 값을 말한다 · 읽는 곳 0이 된 경로 0"
     "(LLM 청크 구간 · 나눠 부르기 · 규칙 나무 입력)",
     KN.KNOBS["structure_mode"][2] == "choice:rules|ladder" and ST.MODES == ("rules", "ladder")
     and '"edit"는 폐기한 값' in _out and '"structure_mode": "ladder"' in _out
     and not any(hasattr(SR, n) for n in ("range_chunks", "check_chunks", "CUT_ROLES"))
     and not any(hasattr(SE, n) for n in ("number_runs", "run_lines", "rules_tree", "ask_edit"))
     and not any(hasattr(ST, n) for n in ("check", "table", "range_lines", "ask_edit")),
     " / ".join(ln.strip() for ln in _out.splitlines() if "폐기" in ln or "다음 줄" in ln)[:200])

# ────────────────────────────────────────────────────────────── ⓚ 개요 · 덱의 판은 구조와 따로
print("\n■ ⓚ 개요 · 덱 — 판이 구조와 따로 · 옛 보존분 재사용")
from core.build import deck as DK                         # noqa: E402
from core.build.extract import prompt_version             # noqa: E402
from parser import deck as PD                             # noqa: E402
_cf = ST._cfgs(["process"])
_k1 = DK.key(_cf)
ST.MODE = "rules"
_k2 = DK.key(_cf)
ST.MODE = "ladder"
_slides = [{"slide": 1, "layout": "제목", "boxes": [{"kind": "text", "text": "노칭 점검"}]}]
_asked = []
_fake = lambda doc_id, slides, doc_coord=None: _asked.append(doc_id) or {"slides": {}, "overview": ""}
_fake.key = lambda: _k1
_old = {"labels": PD.labels_hash(_slides, None), "key": {**_k1, "mode": "edit"}, "answer": {"slides": {}, "overview": "옛 개요"}}
_got = PD.hook(_fake, "B132K", kept=_old, made={}, tally={})(_slides)
# 기대 변경(B135 ④ — 덱은 제 지시문 `2.4_deck.md` · 판은 개요의 판을 이어받는다 · 개요 지시문에 고침 · 덱 절 0)
from core.build import outline as _OLK                    # noqa: E402
_oh = [{"id": 1, "text": "노칭", "parent": None, "depth": 1, "path": "노칭"}]
_onow = _OLK._now_key(_oh, _cf["process"])
_op = TMP / "b134_outline.json"                           # B134 판 개요 체크포인트 — 문면 칸 없음
_op.write_text(json.dumps({**_onow, "headings": [], "of": {}}, ensure_ascii=False), encoding="utf-8")
_otxt = gateway.prompt("outline")
show("ⓚ 개요 · 덱은 구조 지시문과 판이 따로 · 덱 열쇠는 구조 방식과 무관 · 옛 덱 보존분(열쇠에 방식 칸)이 그대로 재사용(LLM 0) / B135 ⓓ 덱은 "
     "제 지시문(판은 개요 판 그대로) · B134 판 개요 체크포인트도 재사용 · 개요 지시문에 고침 · 덱 절 0",
     prompt_version("outline") != prompt_version(ST.PROMPT) and _k1 == _k2 and "mode" not in _k1 and not _asked and _got is not None
     and DK.PROMPT == "deck" and gateway.prompt_path("deck") != gateway.prompt_path("outline")
     and prompt_version("deck") == prompt_version("outline") and _OLK._reuse("B135K", None, _onow, _op)[0] is not None
     and "고침 입력" not in _otxt and "덱 입력" not in _otxt and '"mode": "edit"' not in _otxt,
     f"개요 판 {prompt_version('outline')} · 덱 판 {prompt_version('deck')} · 구조 판 {prompt_version(ST.PROMPT)} · 덱 다시 물음 {len(_asked)}")

# ────────────────────────────────────────────────────────────── B137 — 초안 · 첫 행 · 판단 필요 · 규칙 · 화면 · 토큰
print("\n■ B137 — 첫 서수 · 이어짐 초안 · 첫 행 시트 제목 · 서식 보조 · 규칙 한 줄 · 구조 화면 · 단계마다 토큰")
from openpyxl.styles import Font                          # noqa: E402
from parser import ladder as PLD                          # noqa: E402
D7, SA, SB = "B137D", "공정별 상세사양", "검수 사양"
#: 사내 꼴 표본(창작) — `(행, 열, 글, 서식)` · A: 첫 행 = 시트 이름(번호 붙음) · 가. · 1. · 2. · 나. · 1. / B: 판단 필요 넷(ⓑ 3 · ⓔ 5 · ⓕ 10 ·
#: ⓔ 12 — 첫 서수인데 얕은 줄 · 같은 꼴 형제와 서식이 다른 줄 · 붙인 다단 번호 · 굵게 · 색만 다른 짧은 줄) + 접히는 형제 넷(6~9)
B7 = {SA: [(1, "A", "6.공정별 상세사양", 14), (2, "A", "가. 이송 물류", 12), (3, "B", "1. 기본사양 및 물류 컨셉", None),
           (4, "C", "속도는 60ppm이다.", None), (5, "B", "2. 지정사양", None), (6, "C", "폭은 20mm다.", None),
           (7, "A", "나. 적재", 12), (8, "B", "1. 기본사양", None), (9, "C", "장력은 30N이다.", None)],
      SB: [(1, "A", "7.검수 사양", 14), (2, "B", "가. 입고 검수", None), (3, "A", "1. 외관", None), (4, "C", "외관 본문이다.", None),
           (5, "B", "2. 치수", 14), (6, "C", "1) 길이", None), (7, "C", "2) 폭", None), (8, "C", "3) 두께", None), (9, "C", "4) 무게", None),
           (10, "C", "3.2출하검수", None), (11, "C", "출하 본문이다.", None), (12, "C", "특기사항", "red"), (13, "C", "특기 본문이다.", None)]}
WANT_A = [(1, None), (2, 1), (3, 2), (5, 2), (7, 1), (8, 7)]
LAD7 = {"levels": [{"sigs": ["9."], "role": "주제", "category": None, **_N("사내 답 — 장")},
                   {"sigs": ["가.", "9)"], "role": None, "category": None, **_N("절 · 항")}], "unnumbered": []}
RULE7 = {"shape": "9.", "role": "묶음", "category": None, "parent_shape": "가.", "scope": "첫 서수 목록",
         "said": "가. 아래 '1.'로 시작하는 목록은 가.의 묶음"}
ANS7 = {3: {"heading": True, "parent": 2}, 5: {"heading": True, "parent": 2}, 10: {"heading": True, "parent": 5},
        12: {"heading": False}, 8: {"heading": True, "parent": 7}}


def _x7():
    wb = Workbook()
    wb.remove(wb.active)
    for name, cells in B7.items():
        ws = wb.create_sheet(title=name)
        for r, c, v, f in cells:
            ws[f"{c}{r}"] = v
            if f == "red":
                ws[f"{c}{r}"].font = Font(bold=True, color="FFFF0000")
            elif f:
                ws[f"{c}{r}"].font = Font(bold=True, sz=f)
    p = TMP / f"{D7}.xlsx"
    wb.save(p)
    return p


def _llm7(body):
    """가짜 ⑪ — 사다리는 사내 답(「9. › 가.」) · 시트 확인은 A에 규칙 한 줄 · B에 판단 필요 답 + 시트에 없는 줄 고침 하나(버려진다) ·
    부를 때마다 사용량(사다리 100/10 · 시트 확인 200/20 — 창작 수)."""
    lad = body.get("task") == "ladder"
    gateway._account("structure", {"usage": {"prompt_tokens": 100 if lad else 200, "completion_tokens": 10 if lad else 20,
                                             "total_tokens": 110 if lad else 220}})
    if lad:
        return LAD7
    ans = [{"row": n["row"], "parent": None, "role": None, "name": None, "category": None, "why": "답",
            **ANS7.get(n["row"], {"heading": False})} for n in body.get("needs") or []]
    if body["frame"] == SA:
        return {"answers": ans, "edits": [], "rules": [RULE7]}
    return {"answers": ans, "rules": [], "edits": [{"row": 99, "what": "제목 뺌", "parent": None, "to_row": None, "role": None,
                                                    "name": None, "category": None, "why": "없는 줄"}]}


def _p7(confirmed=None):
    return pipeline.parse(BX, D7, str(_x7()), closed_list=[], sheet_roles={SA: "prose", SB: "prose"},
                          structure=ST.structurer(["process"]), structure_confirmed=confirmed)


struct_map.invalidate(D7)
_c7 = _p7().report["structure"]["프레임"]                  # mock · 힌트 없음 — 코드 사다리 초안(LLM 0)
_lab = lambda name, first: [{"row": 0, "text": name}, {"row": 1, "text": first}]
_ttl = [("공정별 상세사양", "6.공정별 상세사양"), ("6.공정별 상세사양", "공정별 상세사양"), ("6.공정별 상세사양", "6. 공정별 상세사양"),
        ("공정별상세사양", "6 . 공정별 상세 사양")]
ST.DOC.pop(D7, None)
struct_map.invalidate(D7)
_ask0 = _screen.ask
_it = iter(["c", "c"])
_screen.ask = lambda prompt: next(_it)
_buf = io.StringIO()
_u0 = dict(gateway.usage_by(D7).get(gateway.point_label("structure")) or {})
with Live(by={"structure": _llm7}), contextlib.redirect_stdout(_buf):
    gateway.set_doc(D7)
    _g7, _go7 = SS.gate(_p7(), D7, lambda: _p7(confirmed=ST.confirmed_of(D7)), layers=["process"])
    gateway.set_doc(None)
_screen.ask = _ask0
_scr7 = _buf.getvalue()
_s7 = _g7.report["structure"]
_fa, _fb = _s7["프레임"][SA], _s7["프레임"][SB]
_tr = lambda fr, k="headings": [(h["row"], h.get("parent")) for h in fr[k]]
show("B137 ⓐ 사내 꼴 표본 — 코드 초안(LLM 0)이 맞다(6. 맨 위 · 가. · 나. 그 아래 · 1. · 2.가 가. · 나. 아래) · 사다리 답이 「9. › 가.」여도 "
     "초안이 같다(꼴 하나가 두 깊이에)",
     _tr(_c7[SA]) == WANT_A and _tr(_fa, "toc0") == WANT_A and _fa["ladder_source"] == "LLM",
     f"코드 {_tr(_c7[SA])} · 사다리 답 초안 {_tr(_fa, 'toc0')}")
show("B137 ⓑ 첫 행 = 시트 제목 — 번호를 붙인 이름 · 뗀 이름 · 공백이 다른 이름이 모두 시트 제목 · 다른 글이면 지금처럼(시트 제목 없음)",
     all(PLD.title_row(_lab(n, f)) == 1 for n, f in _ttl) and PLD.title_row(_lab("공정별 상세사양", "가. 이송 물류")) is None
     and [h["by"] for h in _c7[SA]["headings"] if h["row"] == 1] == ["시트 제목"],
     str([PLD.title_row(_lab(n, f)) for n, f in _ttl]))
_nb = {n["row"]: n["kind"] for n in _fb["needs"]}
_hb = {h["row"]: h for h in _fb["headings"]}
_lb = {x["row"]: x for x in _fb["labels"]}
show("B137 ⓒ 판단 필요 — 첫 서수인데 얕은 줄(ⓑ) · 같은 꼴 형제와 서식이 다른 줄(ⓔ) · 붙인 다단 번호(ⓕ) · 굵게 · 글자색만 다른 짧은 줄(ⓔ)이 "
     "하나씩 · 답으로 바뀐다(10행 제목 · 12행 본문) · 코드가 혼자 정한 줄 0 · 리더가 글자색을 읽는다(실경로 openpyxl)",
     _nb == {3: "ⓑ", 5: "ⓔ", 10: "ⓕ", 12: "ⓔ"} and _hb.get(10, {}).get("parent") == 5 and 12 not in _hb
     and not ({10, 12} & {h["row"] for h in _fb["toc0"]}) and not _fb["missing"] and "글자색" in _lb[12]["marks"],
     f"판단 필요 {_nb} · 10행 부모 {_hb.get(10, {}).get('parent')} · 12행 표시 {_lb[12]['marks']}")
_ra = {h["row"]: h for h in _fa["headings"]}
show("B137 ⓓ 규칙 한 줄 — 「가. 아래 첫 서수로 시작하는 '1.' 목록」이 맞는 줄 전부(셋)에 같게 · 6.(시트 제목)은 제자리 · 줄 고침 0",
     _tr(_fa) == WANT_A and all(_ra[r]["role"] == "묶음" and _ra[r]["by"] == "규칙" for r in (3, 5, 8))
     and _ra[1]["role"] is None and _fa["line_edits"] == 0 and [x["적용"] for x in _fa["rules"]] == [3],
     f"목차 {_tr(_fa)} · 역할 {[(r, _ra[r].get('role'), _ra[r].get('by')) for r in (1, 3, 5, 8)]} · 규칙 적용 {[x['적용'] for x in _fa['rules']]}")
_md7 = split_report.write_md(D7, _g7.report, _g7.envelope).read_text(encoding="utf-8")
_l10 = next((ln for ln in _scr7.splitlines() if ln.lstrip().startswith("10행 ")), "")
show("B137 ⓔ 화면 — 사다리 첫 서수 표기 · 접은 형제 줄 · 목차 줄의 판단 필요 표지와 LLM 답 · 시트마다 LLM 한 줄 · ✎ · md 끝의 버린 고침",
     _go7 and "시트 › 1.(주제) › 가. · 1)" in _scr7 and "외 3줄" in _scr7 and split_report.NEED_MARK in _l10 and "(LLM)" in _l10
     and "✎" in _l10 and sum(1 for ln in _scr7.splitlines() if ln.strip().startswith("LLM —")) == 2
     and "고침 상세" in _md7 and "버림 — 99행" in _md7 and "버림 — 99행" not in _scr7,
     f"10행 줄 {_l10.strip()[:90]}")
_ub = gateway.usage_by(D7).get(gateway.point_label("structure")) or {}
show("B137 ⓕ 토큰 — 사다리 · 시트 확인 사용량이 갈려 멈춤마다 한 줄 · 초안과 다른 줄 수 · 합은 지점 ⑪ 누계와 같다",
     _s7["사다리_사용"]["prompt"] == 100 and _fa["usage"]["prompt"] == _fb["usage"]["prompt"] == 200
     and _ub.get("prompt", 0) - _u0.get("prompt", 0) == 500 and _fa["changed"] == [3, 5, 8] and _fb["changed"] == [10]
     and "사다리 — LLM · 호출 1 · 입력 100" in _scr7 and "초안과 다른 줄 1/제목" in _scr7 and "구조 합계 — 호출 3 · 입력 500" in _scr7,
     f"사다리 {_s7['사다리_사용']} · 시트 {_fa['usage']} · 다른 줄 {_fa['changed']} / {_fb['changed']}")
ST.confirmed_path().unlink(missing_ok=True)

# ────────────────────────────────────────────────────────────── B138 — Word 목차 재료 (판단은 B137 함수 그대로)
print("\n■ B138 — Word: 자동 번호 되살리기 · 목록 수준 · 스타일 표시 · 목차 필드 거름 · 문서 제목 · 개요 먼저")
import zipfile                                            # noqa: E402
from parser.adapters import basic_docx as BD              # noqa: E402
_WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _wp(text="", style=None, num=None, runs=None):
    """문단 하나 — `num`은 `(numId, ilvl)` · `runs`는 날 런 XML(필드 문자 · 명령)."""
    ppr = (f'<w:pStyle w:val="{style}"/>' if style else "") + (
        f'<w:numPr><w:ilvl w:val="{num[1]}"/><w:numId w:val="{num[0]}"/></w:numPr>' if num else "")
    return f"<w:p>{f'<w:pPr>{ppr}</w:pPr>' if ppr else ''}{runs or ''}{f'<w:r><w:t>{text}</w:t></w:r>' if text else ''}</w:p>"


def _lvl(il, fmt, txt):
    return f'<w:lvl w:ilvl="{il}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{txt}"/></w:lvl>'


def _wdoc(path, body, styles, numbering, title=None):
    """표준 라이브러리로 docx(창작 표본 — python-docx 0) · 번호 정의 · 스타일 · 문서 속성 제목."""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
                   'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/></Types>')
        z.writestr("word/document.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{_WNS}"><w:body>{body}'
                   '</w:body></w:document>')
        z.writestr("word/styles.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:styles xmlns:w="{_WNS}">{styles}</w:styles>')
        z.writestr("word/numbering.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:numbering xmlns:w="{_WNS}">{numbering}'
                   '</w:numbering>')
        if title:
            z.writestr("docProps/core.xml", '<?xml version="1.0" encoding="UTF-8"?><cp:coreProperties xmlns:cp="http://'
                       'schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/'
                       f'1.1/"><dc:title>{title}</dc:title></cp:coreProperties>')
    return path


_STY = ('<w:style w:type="paragraph" w:default="1" w:styleId="a"><w:name w:val="바탕글"/></w:style>'
        '<w:style w:type="paragraph" w:styleId="L2"><w:name w:val="목록 2"/><w:basedOn w:val="a"/>'
        '<w:pPr><w:numPr><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr></w:pPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="H1"><w:name w:val="제목 1"/><w:pPr><w:outlineLvl w:val="0"/>'
        '<w:numPr><w:numId w:val="3"/></w:numPr></w:pPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="H2"><w:name w:val="제목 2"/><w:pPr><w:outlineLvl w:val="1"/>'
        '<w:numPr><w:ilvl w:val="1"/><w:numId w:val="3"/></w:numPr></w:pPr></w:style>')
_NUM = ('<w:abstractNum w:abstractNumId="1">' + _lvl(0, "decimal", "%1.") + _lvl(1, "ganada", "%2.") + _lvl(2, "decimal", "%3)")
        + '</w:abstractNum><w:abstractNum w:abstractNumId="2">' + _lvl(0, "koreanCounting", "%1.") + '</w:abstractNum>'
        '<w:abstractNum w:abstractNumId="3">' + _lvl(0, "decimal", "%1.") + _lvl(1, "decimal", "%1.%2.") + '</w:abstractNum>'
        '<w:abstractNum w:abstractNumId="4">' + _lvl(0, "decimal", "%1)") + '</w:abstractNum>'
        + "".join(f'<w:num w:numId="{n}"><w:abstractNumId w:val="{n}"/></w:num>' for n in (1, 2, 3, 4)))
_FLD = ('<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>')
# A — 개요 없음 · 자동 번호(문단 직접 1 · 스타일 상속 2 — 목록 2) · 문서 앞 목차 필드 둘 · 첫 문단 = 파일 이름(공백 다름) · 모르는 꼴 하나
_xa = _wdoc(TMP / "이송사양서.docx", "".join([
    _wp("이송 사양서"), _wp("1. 이송 물류 …… 2", runs=_FLD), _wp("가. 기본사양 …… 2"),
    _wp(runs='<w:r><w:fldChar w:fldCharType="end"/></w:r>'),
    _wp("이송 물류", num=(1, 0)), _wp("기본사양", style="L2"), _wp("속도는 60ppm이다."), _wp("검수", num=(1, 2)),
    _wp("검수는 매일 한다."), _wp("지정사양", style="L2"), _wp("폭은 20mm다."), _wp("적재", num=(1, 0)), _wp("기본사양", style="L2"),
    _wp("장력은 30N이다."), _wp("특수 항목", num=(2, 0)), _wp("특수 본문이다.")]), _STY, _NUM)
# B — 개요 수준(제목 1 · 2 — 스타일에 건 번호 「1.」 · 「1.1.」) · 개요 없는 손 번호 「3. 결론」 · 본문 번호 목록 「1)」 / C — 손 목차 셋(문서 앞)
_xc = _wdoc(TMP / "B138C.docx", "".join([_wp("목차"), _wp("1. 개요 …… 1"), _wp("2. 사양 …… 2"), _wp("3. 결론 …… 3")]
                                        + [_wp(f"본문 {i}째 줄이다.") for i in range(1, 7)]), _STY, _NUM)
_xb = _wdoc(TMP / "B138B.docx", "".join([
    _wp("개요", style="H1"), _wp("개요 본문이다."), _wp("범위", style="H2"), _wp("범위 본문이다."), _wp("사양", style="H1"),
    _wp("속도", num=(4, 0)), _wp("속도는 30이다."), _wp("폭", num=(4, 0)), _wp("3. 결론"), _wp("결론 본문이다.")]), _STY, _NUM)
_ra = read(str(_xa))
_pa = {p["index"]: p for p in _ra["paragraphs"]}
_txt = [p["text"] for p in _ra["paragraphs"]]
show("B138 ⓐ 자동 번호 — 「1. › 가. › 1)」(문단 직접 · 스타일 상속)이 보이는 번호 글로 · 아래 단은 위 단이 바뀌면 다시(「2. 적재」 뒤 「가.」) · "
     "모르는 꼴은 「자동 번호 · 수준 1」 + 판단 필요",
     [t for t in _txt if t.endswith(("물류", "기본사양", "검수", "지정사양", "적재"))]
     == ["1. 이송 물류", "가. 기본사양", "1) 검수", "나. 지정사양", "2. 적재", "가. 기본사양"]
     and any(p.get("num_unknown") and p["text"] == "특수 항목" for p in _pa.values())
     and "자동 번호 · 수준 1" in (BD._marks(_ra).get(next(i for i, p in _pa.items() if p["text"] == "특수 항목")) or []),
     str(_txt[:9]))
struct_map.invalidate("B138A")
_wa = pipeline.parse(BD, "B138A", str(_xa), closed_list=[], structure=ST.structurer(["process"]))
_fa8 = _wa.report["structure"]["프레임"][BD.FRAME]
_ix = {p["text"]: p["index"] for p in _ra["paragraphs"]}
_par8 = {h["row"]: h.get("parent") for h in _fa8["headings"]}
_want8 = {_ix["이송 사양서"]: None, _ix["1. 이송 물류"]: _ix["이송 사양서"], _ix["2. 적재"]: _ix["이송 사양서"],
          _ix["1) 검수"]: [p["index"] for p in _ra["paragraphs"] if p["text"] == "가. 기본사양"][0]}
struct_map.invalidate("B138B")
_wb8 = pipeline.parse(BD, "B138B", str(_xb), closed_list=[], structure=ST.structurer(["process"]))
_rb = read(str(_xb))
_ib = {p["text"]: p["index"] for p in _rb["paragraphs"]}
_fb8 = _wb8.report["structure"]["프레임"][BD.FRAME]
_hb8 = {h["row"]: h.get("parent") for h in _fb8["headings"]}
_nb8 = {n["row"]: n for n in _fb8["needs"]}
show("B138 ⓑ 목차 — 개요 없는 자동 번호 문서가 B137 함수로 맞는 초안(LLM 0) · 개요 문서는 개요대로 서고(1. 개요 › 1.1. 범위 · 2. 사양) "
     "어긋난 번호 줄(손 번호 「3. 결론」)만 판단 필요",
     all(_par8.get(r, "없음") == p for r, p in _want8.items()) and _hb8.get(_ib["1.1. 범위"]) == _ib["1. 개요"]
     and _hb8.get(_ib["1. 개요"]) is None and _hb8.get(_ib["2. 사양"]) is None
     and set(_nb8) == {_ib["3. 결론"]} and "개요 수준" in _nb8[_ib["3. 결론"]]["why"],
     f"A {[(_par8.get(r), p) for r, p in _want8.items()]} · B 판단 필요 {[(r, n['kind'], n['why'][:40]) for r, n in sorted(_nb8.items())]}")
_ch8 = " ".join(c["text"] for c in _wa.envelope["chunks"])
_md8 = split_report.write_md("B138A", _wa.report, _wa.envelope).read_text(encoding="utf-8")
struct_map.invalidate("B138C")
_fc8 = pipeline.parse(BD, "B138C", str(_xc), closed_list=[], structure=ST.structurer(["process"])).report["structure"]["프레임"][BD.FRAME]
_lc = {x["row"] for x in _fc8["labels"] if "……" in str(x.get("text"))}
show("B138 ⓒ 목차 필드 — 문서 앞 목차 필드의 줄이 목차 · 청크에 0 · 구조 확인(md)에 「목차 필드 — 2줄 거름」 · 필드가 아닌 손 목차 줄은 "
     "판단 필요(「손 목차?」)",
     _ra.get("toc_field_lines") == 2 and "…… 2" not in _ch8 and not any("……" in str(x.get("text")) for x in _fa8["labels"])
     and "목차 필드 — 2줄 거름" in _md8 and len(_lc) == 3 and {n["row"] for n in _fc8["needs"] if n["kind"] == "ⓖ"} == _lc,
     f"거름 {_ra.get('toc_field_lines')} · 청크에 남은 목차 줄 {'…… 2' in _ch8} · 손 목차 {[(n['row'], n['kind']) for n in _fc8['needs']]}")
_one = lambda t, title=None, name="이송사양서.docx": {"path": name, "title": title, "paragraphs": [  # noqa: E731
    {"index": 1, "kind": "para", "text": t}]}
show("B138 ⓓ 문서 제목 — 파일 이름과 같은 첫 문단(공백 다름 · 번호 붙임)과 문서 속성 제목과 같은 첫 문단이 맨 위 · 다른 글이면 없음",
     _par8.get(_ix["이송 사양서"], "없음") is None and [h["by"] for h in _fa8["headings"] if h["row"] == _ix["이송 사양서"]] == ["시트 제목"]
     and BD._title(_one("1. 이송사양서")) == 1 and BD._title(_one("이송 설비 사양", "이송 설비 사양", "x.docx")) == 1
     and BD._title(_one("개요")) is None, f"문서 제목 줄 {[h['row'] for h in _fa8['headings'] if h.get('by') == '시트 제목']}")

done()
