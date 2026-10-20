# -*- coding: utf-8 -*-
"""G6 ㉔ 연결 보강 — 다른 공정 이름은 노드가 아니다 · 빈칸 표시 정규화 · 재연결은 더 붙는 것이 없을 때까지 (B113).

창작 표본(mock · 시험 주입)으로 메커니즘을 잰다 — 수치는 근거가 아니다([정정] 50). 형태마다 대표 한 벌이 층 둘에
걸친다 — 표는 process · 산문은 quality(집이 process인 개체 · 엣지는 문서 층 — B100). 층 · 형태 곱은 층 중립 자리라 대표
한 벌이다(B121 ②ⓒ — 층 호환은 층 어휘 정적 점검 · 네 벌이 잰다). 골격 이름 · 빈칸 표시는 데이터다(코드 분기 0).

잠그는 성질:
  ⓐ 다른 공정 이름: 좌표 계보 밖 골격 이름(산문 개체 · 표 개체 열) ⇒ 노드 0 · 그 공정 노드로 관련 링크(`about`) ·
     대장 판정 `mention` · 맞춤표 「다른 공정 언급」 칸 · 자기 공정과 그 하위 · 조상 이름은 지금 그대로
  ⓑ 빈칸 표시: 개체 칸(표 · 산문)의 `-` · `N/A` · `해당 없음` ⇒ 개체 0 · 대장 판정 `blank` · 맞춤표 칸 · 좌표 칸은 비워
     문서 좌표를 물려받는다 · 값 칸의 「없음」은 그래프에 그대로 · 대장에 표지 · 공통 config 목록을 바꾸면 따른다
  ⓒ 재연결: ①좌표 미해소 보류 행이 문서 좌표 기록 뒤 손 명령(`ops relink --doc --coord --apply`)으로 서고 붙는다(추출
     LLM 0 · 체크포인트 그대로) ②뒤 청크가 만든 유닛에 앞 청크의 개체가 다음 바퀴에 붙는다 ③바퀴가 「붙음 0」에서 멈춘다 ·
     같은 입력이면 같은 그래프(클린 2회 — 첫 회는 ②의 실행)
"""
from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, _rw, done, overlay   # noqa: F401
from g6_common import _register, _unregister, _run, DT   # noqa: E402

from openpyxl import Workbook, load_workbook             # noqa: E402
from cli import doc_coord as DC, ingest as IG             # noqa: E402
from core.build import extract as EX, ledger as LG, retry as RT   # noqa: E402
from core.state.world import World                        # noqa: E402
from parser import normalizer as NM, pipeline, struct_map, tagger as TG   # noqa: E402
from parser.adapters import basic_prose_xlsx as BX        # noqa: E402
from router import discover                               # noqa: E402

HINTS = ROOT / "tests" / "fixtures" / "extract_hints"
RAWS = ROOT / "tests" / "fixtures" / "raw"
TMP = Path(__import__("tempfile").mkdtemp(prefix="b113_"))
_made = []
_KEEP = {p: (p.read_bytes() if p.exists() else None) for p in [DC._file()]}


def _book(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for r, v in enumerate(rows, 1):
            ws.cell(row=r, column=1, value=v)
    _P.ensure(path)
    wb.save(path)
    return path


def _hint(doc_id, by_loc):
    _made.append(doc_id)
    (HINTS / f"{doc_id}.json").write_text(json.dumps(by_loc, ensure_ascii=False), encoding="utf-8")


def _setup():
    """덧칠(Component · 겸 렌즈) + `nest_categories` + 소속 엣지의 카테고리쌍 — 상태 루트에만(B102 시험과 같은 바닥)."""
    overlay(component=True)
    _rw(_P.common(), lambda c: c.setdefault("canonical_scope", {}).__setitem__("nest_categories", ["Component"]))
    pats = [{"src": s_, "rel": "part_of", "dst": d_, "symmetric": False, "정의문": f"{s_}가 {d_}에 속한다"}
            for s_, d_ in (("Component", "Unit"), ("Component", "Process"), ("Unit", "Process"))]
    pairs = {f"{p_['src']},{p_['dst']}": "part_of" for p_ in pats}

    def _q(c):
        c["categories"]["Component"] = "(품질층의 렌즈) 불량이 일어나는 구성 기계"
        c["relation_patterns"] += pats
        c.setdefault("category_pair_map", {}).update(pairs)

    def _p(c):
        have = {(p_["src"], p_["rel"], p_["dst"]) for p_ in c["relation_patterns"]}
        c["relation_patterns"] += [p_ for p_ in pats if (p_["src"], p_["rel"], p_["dst"]) not in have]
        for k, v in pairs.items():
            c.setdefault("category_pair_map", {}).setdefault(k, v)
    _rw(_P.layers("quality", "config.json"), _q)
    _rw(_P.layers("process", "config.json"), _p)
    for lay in ("process", "quality"):
        bootstrap(lay, echo=False)


def _prose(doc_id, lay, sheets, ents_of, doc_coord="탭용접", path=None):
    """창작 엑셀 → 봉투(문서 좌표) · 청크마다 힌트(`ents_of(청크) -> [개체]`)."""
    path = path or _book(TMP / f"{doc_id}.xlsx", sheets)
    struct_map.invalidate(doc_id)
    env = pipeline.parse(BX, doc_id, str(path), closed_list=TG.closed_list("process"), doc_coord=doc_coord,
                         sheet_roles={n: "prose" for n in sheets}).envelope
    env["doc_type"] = "ppt_" + lay
    for c in env["chunks"]:
        c["doc_type"] = env["doc_type"]
    _hint(doc_id, {c["source_locator"]: {"entities": ents_of(c), "relations": [], "attach": []}
                   for c in env["chunks"] if ents_of(c)})
    return env


def _rows(doc_id, role="entity"):
    return [r for r in (LG.read(doc_id) or {}).get("rows") or [] if r.get("role") == role]


def _end(notes):
    return next(n for n in notes if n.get("단계") == "끝")["맞춤"]


def _cell(mt, label, cell):
    ln = next(x for x in mt["줄"] if x["줄"].startswith(label))
    return dict((k, n) for k, n, _d in ln["칸"]).get(cell), ln["맞음"]


def _live(name):
    return [n for _l, n in World().nodes() if n["canonical"] == name and ops.is_live(n)]


def _about(doc_id):
    ch = store.read(store.CHUNKS, {"chunks": {}})
    ids = {cid for cid, c in ch["chunks"].items() if c.get("doc_id") == doc_id}
    return [(a["chunk_id"], World().get(a["node_id"])["canonical"]) for a in ch.get("about") or []
            if a["chunk_id"] in ids]


# ────────────────────────────────────────────────────────────── ⓐ
print("\n■ B113 ① 다른 공정의 골격 이름은 노드가 아니다 — 관련 링크로 (산문 × quality · 표 × process)")
TEXT = {"사양": ["1. 연동", "앞 공정 스태킹 설비와 연동한다. 탭용접의 cutting 단계는 빠르다.", "2. 구동",
                 "모터는 서보 모터를 쓴다."]}


def _ents(c):
    t = c.get("text") or ""
    out = []
    if "스태킹" in t:
        out += [{"surface": "스태킹", "category": "Unit"}, {"surface": "cutting", "category": "Unit"},
                {"surface": "탭용접", "category": "Process"}]
    if "모터" in t:
        out.append({"surface": "모터", "category": "Component"})
    return out


_oa = []
for _lay in ("quality",):                       # 산문은 quality 한 벌 — 표(아래)가 process(B121 ②ⓒ)
    _setup()
    _did = f"B113A{_lay[0].upper()}"
    _env = _prose(_did, _lay, TEXT, _ents)
    _nt = []
    run_document(_env, notice=_nt.append)
    _rw_ = {r["surface"]: r for r in _rows(_did)}
    _n_other = _cell(_end(_nt), "값(개체 행)", "다른 공정 언급")
    _ab = _about(_did)
    _oa.append((_lay, _rw_["스태킹"]["verdict"] == "mention" and _rw_["스태킹"].get("node_id") is None
                and not _live("탭용접::스태킹") and ("스태킹" in {c for _ch, c in _ab})
                and _rw_["cutting"]["verdict"] != "mention" and _rw_["탭용접"]["verdict"] == "match"
                and _n_other == (1, True),
                f"{_lay}: 스태킹 {_rw_['스태킹']['verdict']} · cutting {_rw_['cutting']['verdict']} · 탭용접 "
                f"{_rw_['탭용접']['verdict']} · 노드 탭용접::스태킹 {len(_live('탭용접::스태킹'))} · 관련 링크 {_ab} · "
                f"맞춤 칸 {_n_other}"))
show("ⓐ 산문 — 계보 밖 공정 이름(스태킹)은 노드 0 · 그 공정으로 관련 링크 · 대장 mention · 맞춤 「다른 공정 언급」 "
     "· 자기 공정(탭용접)과 그 하위(cutting)는 지금 그대로", all(o for _l, o, _d in _oa),
     " ‖ ".join(d for _l, _o, d in _oa))

_setup()
_cp = {**TABLE, "doc_id": "B113ACP",
       "records": [dict(CPREC, source_locator="AQ-R1", 설비="스태킹", 관리항목="노칭 정밀도",
                        대응계획="앞 공정 설비와 연동해 점검한다"),          # 행의 청크(내용 칸) — 링크의 출발
                   dict(CPREC, source_locator="AQ-R2", 설비="노칭 프레스", 관리항목="노칭 정밀도")]}
_nt = []
run_document(_cp, notice=_nt.append)
_tr = {(r["locator"], r["surface"]): r for r in _rows("B113ACP")}
show("ⓐ 표 — 개체 열 값이 계보 밖 공정 이름이면 노드 0 · 그 행의 청크(내용 칸)에서 그 공정으로 관련 링크 · 대장 mention · 맞춤 칸 · "
     "같은 행의 다른 값은 그대로",
     _tr[("AQ-R1", "스태킹")]["verdict"] == "mention" and not _live("노칭::스태킹")
     and "스태킹" in {c for _ch, c in _about("B113ACP")}
     and _tr[("AQ-R2", "노칭 프레스")]["verdict"] in ("new", "match")
     and _cell(_end(_nt), "값(개체 행)", "다른 공정 언급") == (1, True),
     f"{_tr[('AQ-R1', '스태킹')]['verdict']} · {_tr[('AQ-R2', '노칭 프레스')]['verdict']} · 관련 링크 "
     f"{_about('B113ACP')} · {_cell(_end(_nt), '값(개체 행)', '다른 공정 언급')}")

# ────────────────────────────────────────────────────────────── ⓑ
print("\n■ B113 ② 빈칸 표시 정규화 (개체 칸 · 좌표 칸 · 값 칸 · 목록은 공통 config)")
_ob = []
for _lay in ("quality",):                       # 산문은 quality 한 벌 — 표(아래 파일)가 process(B121 ②ⓒ)
    _setup()
    _did = f"B113B{_lay[0].upper()}"
    _env = _prose(_did, _lay, {"사양": ["모터는 서보 모터를 쓴다."]},
                  lambda c: [{"surface": "모터", "category": "Component"}, {"surface": "-", "category": "Unit"},
                             {"surface": "해당 없음", "category": "Component",
                              "belongs_to": {"name": "N/A", "category": "Unit", "from": "본문"}}])
    _nt = []
    run_document(_env, notice=_nt.append)
    _bl = [r for r in _rows(_did) if r["verdict"] == "blank"]
    _ob.append((_lay, len(_bl) == 2 and not _live("탭용접::-") and _cell(_end(_nt), "값(개체 행)", "빈칸 표시") == (2, True),
                f"{_lay}: blank {[r['surface'] for r in _bl]} · 맞춤 {_cell(_end(_nt), '값(개체 행)', '빈칸 표시')}"))
show("ⓑ 산문 — 개체 표기가 빈칸 표시면 개체 0 · 대장 blank · 맞춤 「빈칸 표시」", all(o for _l, o, _d in _ob),
     " ‖ ".join(d for _l, _o, d in _ob))

# 표 — 실제 파일(파서 정규화)로: 개체 칸 「-」 · 좌표 칸 「N/A」 · 값 칸 「없음」
init.init(fresh_=True, all_=True)
for _lay in ("process", "quality"):
    bootstrap(_lay, echo=False)
_wb = load_workbook(RAWS / "CP01.xlsx")
_ws = _wb.worksheets[0]
_hr, _hdr = next((r[0].row, {str(c.value).strip(): c.column for c in r if c.value})
                 for r in _ws.iter_rows(min_row=1, max_row=8) if {"설비", "관리항목"} <= {str(c.value or "").strip() for c in r})


def _first(name):
    """병합 칸이 아닌 그 공정 이름의 첫 행(창작 CP01 — 칸 자리는 머리행 이름으로 찾는다)."""
    return next(r for r in range(_hr + 1, _ws.max_row + 1) if _ws.cell(row=r, column=_hdr["공정명"]).value == name
                and _ws.cell(row=r + 1, column=_hdr["설비"]).value != "〃")


_ws.cell(row=_first("탭용접"), column=_hdr["설비"], value="-")          # 개체 칸
_ws.cell(row=_first("사이드 실링"), column=_hdr["공정명"], value="N/A")   # 좌표 칸
_ws.cell(row=_first("스태킹"), column=_hdr["규격"], value="없음")         # 값 칸
_XB = TMP / "B113_빈칸.xlsx"
_wb.save(_XB)
_did_b = IG.doc_id_of(_XB)
_made.append(_did_b)
_ob1 = _run("ingest-file", str(_XB), "--doc-type", "cp", "--allow-mock", "--coord", "노칭")
_rows_b = (LG.read(_did_b) or {}).get("rows") or []
_env_b = json.loads(_P.parsed(f"{_did_b}.json").read_text(encoding="utf-8"))
_blank_ent = [r for r in _rows_b if r.get("verdict") == "blank" and r.get("role") == "entity"]
_coord_row = next((r for r in _env_b["records"] if "process_ref" in (r.get("meta") or {}).get("null_marked", [])), {})
_val_rows = [r for r in _rows_b if r.get("blank") and r.get("role") == "attribute"]
show("ⓑ 표(파일 · 파서 정규화) — 개체 칸 「-」 개체 0 · 대장 blank / 좌표 칸 「N/A」는 비워 문서 좌표(노칭)를 물려받는다 / "
     "값 칸 「없음」은 그대로 · 대장에 표지",
     len(_blank_ent) == 1 and _blank_ent[0]["surface"] == "-"
     and _coord_row.get("process_ref") == "노칭" and (_coord_row.get("meta") or {}).get("coord_from_doc")
     and _val_rows and _val_rows[0]["blank"] == "없음" and "[성공]" in _ob1,
     f"blank 개체 {[r['surface'] for r in _blank_ent]} · 좌표 칸 {_coord_row.get('process_ref')} "
     f"({(_coord_row.get('meta') or {}).get('null_marked')}) · 값 칸 {[(r.get('field'), r['blank']) for r in _val_rows]}")
_mo = re.search(r"조각 좌표 [\d,]+ = .*?\(([^)]*빈칸 표시로 비움[^)]*)\)", _ob1)
show("ⓑ 맞춤표 — 조각 좌표 줄 곁에 「빈칸 표시로 비움 n」 · 값 줄 「빈칸 표시 n」", _mo and re.search(r"빈칸 표시 1", _ob1),
     (_mo.group(1) if _mo else "") + " / " + " ".join(re.findall(r"빈칸 표시 \d+", _ob1)[:2]))
_rw(_P.common(), lambda c: c.__setitem__("null_markers", ["#"]))
show("ⓑ 목록은 공통 config `null_markers` — 바꾸면 따른다(「-」는 더는 빈칸이 아니다 · 「#」은 빈칸이다)",
     NM.null_marker("-") is None and NM.null_marker(" # ") == "#", f"{NM.null_markers()}")
_rw(_P.common(), lambda c: c.pop("null_markers", None))
show("ⓑ 목록이 없으면 기본값 — `-` · `—` · `–` · `N/A` · `n/a` · `해당 없음` · `해당없음` · `없음` (공백 무시)",
     NM.null_markers() == NM.NULL_MARKERS and NM.null_marker("해당없음") and NM.null_marker("해 당 없 음")
     and NM.null_marker("노칭") is None, str(NM.null_markers()))

# ────────────────────────────────────────────────────────────── ⓒ
print("\n■ B113 ③ 재연결 — 더 붙는 것이 없을 때까지 (좌표 미해소 보류 · 뒤 청크의 유닛 · 바퀴 · 같은 결과)")
_setup()
_register("process", DT)
_XC = _book(TMP / "B113_좌표없음_사양.xlsx", {"사양": ["장력은 매 시간 기록한다.", "커버를 단다."]})
_did_c = IG.doc_id_of(_XC)
_env_c = pipeline.parse(BX, _did_c, str(_XC), closed_list=TG.closed_list("process"), doc_coord=None,
                        sheet_roles={"사양": "prose"}).envelope
struct_map.invalidate(_did_c)
_hint(_did_c, {c["source_locator"]: {"entities": [{"surface": "장력", "category": "Property"}], "relations": [],
                                      "attach": []} for c in _env_c["chunks"] if "장력" in c.get("text", "")})
_oc = _run("ingest-file", str(_XC), "--doc-type", DT, "--allow-mock", "--sheets", "1:prose")
_h0 = [r for r in _rows(_did_c) if r["surface"] == "장력"]
_ck0 = EX.checkpoint_path(_did_c).stat().st_mtime_ns if EX.checkpoint_path(_did_c).exists() else None
_plan = _run("ops", "relink", "all", "--doc", _did_c, "--coord", "노칭", "--actor", "시험")
_apply = _run("ops", "relink", "all", "--doc", _did_c, "--coord", "노칭", "--actor", "시험", "--apply")
_h1 = [r for r in _rows(_did_c) if r["surface"] == "장력"]
_prop = _live("노칭::장력")
_es = [e for lay in discover() for e in open_graph(lay).edges if _prop and _prop[0]["id"] in (e["src"], e["dst"])]
show("ⓒ① 좌표 미해소 보류 행 — 문서 좌표 기록(ops relink --doc --coord · 계획 쓰기 0) 뒤 --apply로 서고 붙는다 · "
     "추출 LLM 0(체크포인트 그대로) · 대장 행 by retry",
     _h0 and _h0[0]["verdict"] == "pending" and "계획만 — 쓰기 0" in _plan
     and _prop and _es and _h1 and _h1[0]["verdict"] in ("new", "match") and _h1[0].get("by") == "retry"
     and (_h1[0].get("attached") or [{}])[0].get("path") == "폴백"
     and (EX.checkpoint_path(_did_c).stat().st_mtime_ns if EX.checkpoint_path(_did_c).exists() else None) == _ck0
     and (DC.recorded(_did_c) or {}).get("coord") == "노칭" and "재연결 — 바퀴" in _apply,
     f"전 {[r['verdict'] for r in _h0]} → 뒤 {[(r['verdict'], r.get('by')) for r in _h1]} · 노드 {len(_prop)} · 엣지 "
     f"{len(_es)} · " + " / ".join(ln.strip() for ln in _apply.splitlines() if "바퀴" in ln))
_again = _run("ops", "relink", "all", "--actor", "시험")
show("ⓒ① 되살린 노드는 사전에 남는다(재시도 끝 저장이 덮지 않는다) · 다시 치면 볼 것 0(지문이 실행 사이에 같다)",
     _prop and _prop[0]["id"] in Dictionary.open().lookup("노칭::장력") and "이번에 볼 것 0 " in _again,
     f"사전 {Dictionary.open().lookup('노칭::장력')} · " + " ".join(_again.split()[:12]))


def _units(c):
    t = c.get("text") or ""
    out = []
    if "모터" in t:
        out.append({"surface": "모터", "category": "Component",
                    "belongs_to": {"name": "이송 유닛", "category": None, "from": "본문"}})
    if "이송 유닛은" in t:
        out.append({"surface": "이송 유닛", "category": "Unit"})
    return out


def _snap():
    out = {}
    for lay in discover():
        g = open_graph(lay)
        out[lay] = (sorted(n["canonical"] for n in g.nodes.values()),
                    sorted((g.get(e["src"])["canonical"], e["rel"], g.get(e["dst"])["canonical"]) for e in g.edges
                           if g.get(e["src"]) and g.get(e["dst"])))
    return out


UNITS = {"앞": ["모터는 이송 유닛의 구동을 맡는다."], "뒤": ["이송 유닛은 전극을 옮긴다."]}
_oc2, _s = [], []
for _lay in ("quality",):                       # 층 하나 — 바퀴 · 붙음은 층 중립(B121 ②ⓒ) · 유닛의 집은 process
    _setup()
    _did = f"B113C{_lay[0].upper()}"
    _env = _prose(_did, _lay, UNITS, _units, doc_coord="노칭")
    run_document(_env)
    _q0 = [x for x in store.read(store.QUEUE, []) if x.get("doc_id") == _did and x["kind"] == "orphan_attach"]
    from core.build.entry import finalize      # noqa: E402
    finalize()
    _rounds = [r["붙음"] for r in RT.LAST.get("rounds") or []]
    _mot = next((n for _l, n in World().nodes() if n["canonical"].endswith("::모터") and ops.is_live(n)), None)
    _unit = next((n for _l, n in World().nodes() if n["canonical"].endswith("이송 유닛") and n["category"] == "Unit"), None)
    _pe = [e for lay in discover() for e in open_graph(lay).edges
           if _mot and _unit and {e["src"], e["dst"]} == {_mot["id"], _unit["id"]} and e["rel"] == "part_of"]
    _q1 = [x for x in store.read(store.QUEUE, []) if x.get("doc_id") == _did and x["kind"] == "orphan_attach"]
    _oc2.append((_lay, _q0 and _pe and not _q1 and _rounds and _rounds[-1] == 0 and sum(_rounds) >= 1,
                 f"{_lay}: 앞 청크 소속 미해소 큐 {len(_q0)} → 바퀴 붙음 {_rounds} · part_of 엣지 {len(_pe)} · 남은 큐 {len(_q1)}"))
    _s.append(_snap())                          # ⓒ③의 첫 회 — 같은 입력의 클린 실행을 다시 돌리지 않는다(B121 ②ⓓ)
show("ⓒ② 뒤 청크가 만든 유닛에 앞 청크의 개체가 다음 바퀴에 붙는다 · ③ 바퀴는 「붙음 0」에서 멈춘다",
     all(o for _l, o, _d in _oc2), " ‖ ".join(d for _l, _o, d in _oc2))

_setup()                                        # ⓒ③의 둘째 회 — 클린 · 같은 입력(문서 · 층 · 힌트)
run_document(_prose(_did, _lay, UNITS, _units, doc_coord="노칭"))
finalize()
_s.append(_snap())
show("ⓒ③ 같은 입력이면 같은 결과 — 클린 2회 그래프(노드 · 엣지 이름) 같다", _s[0] == _s[1],
     f"노드 {sum(len(v[0]) for v in _s[0].values())} · 엣지 {sum(len(v[1]) for v in _s[0].values())}")

for _d in set(_made):
    (HINTS / f"{_d}.json").unlink(missing_ok=True)
for _p, _b in _KEEP.items():
    if _b is None:
        _p.unlink(missing_ok=True)
    else:
        _p.write_bytes(_b)
_unregister(DT)
shutil.rmtree(TMP, ignore_errors=True)
done()
