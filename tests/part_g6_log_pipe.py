# -*- coding: utf-8 -*-
"""G6 ㉖ 명령 로그 보존 · 화면이 닫혀도 트레이스백 없이 (B109 ① ②).

창작 표본(mock)으로 메커니즘을 잰다 — 수치는 근거가 아니다([정정] 50). 대표 둘(B121 ②ⓒ — 형태 · 층을 한 번씩):
표 CP01(cp · process) · 산문 B109SQ(prose 엑셀 사본 · quality · 힌트 개체 Failure). 화면 닫힘은 둘로 잰다 — 진짜 파이프(자식 프로세스의 화면을 첫 쓰기 전에 닫는다 ·
`| head -0`과 같다)와 줄 위치(공통 진입의 화면 스트림이 k번째 줄에서 `BrokenPipeError` — 파이프가 그 줄에서 닫힌 것과 같다).

잠그는 성질:
  ⓐ 로그 보존: 인입(표 · 산문) → 기본 `init --fresh` ⇒ `work/logs/`의 이전 로그가 바이트 그대로 · 화면 「남긴 것」에 명령 로그 /
     `rebuild --yes` ⇒ 재구축 전 인입 로그가 바이트 그대로 · 그 실행의 로그도(계획 · 끝) / `init --fresh --all` ⇒ 지금의 클린
     (③④⑤ 잔재 0 — 명령 로그도 지운다 · 그 실행의 로그만 다시 열린다)
  ⓑ 화면 닫힘: 읽기 명령(진입 갈래 셋 — `run.py` · `python -m cli.*` · `doctor.py`)은 트레이스백 0 · 종료 코드 141 · 로그에
     멈춤 줄 / 쓰기 명령(`ingest-file` — 표·산문 × 층 둘)은 트레이스백 0 · 141 · 멈춤 줄 · ③은 셋 중 하나(시작 전 · 문서 끝 —
     마무리 전 · 마무리까지) · 결함 장부에 화면 닫힘 0 / 줄 위치 표본 전부가 그 셋 중 하나 · 「되돌림」 줄이면 시작 전 · 「저장 뒤
     멈춤」 줄이면 끝난 문서 / 재인입 도중이면 첫 인입 뒤 판 그대로(회수한 그래프도 되돌린다) / 화면에 쓰는 스레드도 트레이스백 0
"""
from __future__ import annotations

import contextlib
import io
import shutil
import subprocess
import tempfile
import threading
import time
from pathlib import Path

sys_path = __import__("sys").path
sys_path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done   # noqa: F401
from g6_common import DT, RFQ, _register, _run, _unregister   # noqa: E402 — 산문 엑셀 표본 · 등록 · CLI

import run as RUN                                                  # noqa: E402
from cli import _entry, ingest as IG                               # noqa: E402
from core.graph import GraphStore                                  # noqa: E402 — 그래프 파일 이름은 그 모듈이 안다
from core.state import log as LG                                   # noqa: E402
from core.state.world import World                                 # noqa: E402

FX = ROOT / "tests" / "fixtures"
from core.state import fixtures as _FXR                   # noqa: E402 — 픽스처 뿌리(회귀는 스위트마다 사본 — B133)
HINTS = _FXR.EXTRACT_HINTS
QDT = "b109pq"
SH = "2-4:prose 5:ref *:skip"
TMP = Path(tempfile.mkdtemp(prefix="b109_"))
shutil.copy(RFQ, TMP / "B109SP.xlsx")
shutil.copy(RFQ, TMP / "B109SQ.xlsx")
_HINTED = {"B109SP": {"surface": "B109 산문 속성", "category": "Property"},
           "B109SQ": {"surface": "B109 산문 고장", "category": "Failure"}}
for _d, _e in _HINTED.items():
    (HINTS / f"{_d}.json").write_text(json.dumps({"사양_기계!R8-R17": {"entities": [_e], "relations": []}},
                                                ensure_ascii=False), encoding="utf-8")
#: 대표 둘 — (원본, doc_type, 추가 인자) · 되돌림 경계는 형태 · 층과 무관한 한 자리(`run_document`)라 곱하지 않는다(B121 ②ⓒ —
#: 표 · process와 산문 · quality가 두 형태 · 두 층을 한 번씩 · 층 호환은 층 어휘 정적 점검 + 네 벌이 잰다)
CASES = {"표 · process": (FX / "raw" / "CP01.xlsx", None, []),
         "산문 · quality": (TMP / "B109SQ.xlsx", QDT, ["--sheets", SH, "--coord", "노칭"])}
TS = ("created", "at", "first_seen", "last_seen", "last_tried", "last_fp", "attempts", "obsoleted_at", "confirmed_at",
      "parsed_at", "first_ingested_at", "ingested_at", "updated_at", "last_hit")
ENV = {**os.environ, "USE_MOCK": "1"}


def _argv(case):
    src, dt, extra = CASES[case]
    return [str(src), "--allow-mock"] + (["--doc-type", dt] if dt else []) + extra


def _base():
    """클린 → bootstrap → 산문 doc_type 둘 등록(층 둘)."""
    with contextlib.redirect_stdout(io.StringIO()):
        init.init(fresh_=True, all_=True)
        RUN.cmd_bootstrap([])
    _register("process", DT)
    _register("quality", QDT)


from core.build import ledger as _LDG                                  # noqa: E402 — 중간 저장 꼬리(B165)


def _state():
    """③ + 판정 대장 — 그래프(층마다) · 사전 · 큐 · 청크 · 문서 대장 · 대장 행 · data 파일 목록(id는 이름으로 · 시각 뺌)."""
    nm = {n["id"]: f"{lay}:{n['canonical']}" for lay, n in World().nodes()}

    def canon(x):
        if isinstance(x, dict):
            return {k: canon(v) for k, v in x.items() if k not in TS}
        if isinstance(x, list):
            return sorted((canon(v) for v in x), key=lambda v: json.dumps(v, ensure_ascii=False, sort_keys=True))
        return nm.get(x, x) if isinstance(x, str) else x

    led = _P.work("ingest_log")
    return json.dumps(canon({
        "graph": {lay: {"nodes": [dict(n, id=None) for n in g.nodes.values()], "edges": g.edges}
                  for lay, g in World().graphs.items()},
        "dict": store.read(store.DICTIONARY, {}), "queue": store.read(store.QUEUE, []),
        "chunks": store.read(store.CHUNKS, {}), "docs": store.read(store.DOC_REGISTRY, {}),
        # 기대 변경(B165 — 멈춘 실행의 중간 저장은 다시 칠 때의 재생 재료다(④ · 체크포인트와 같은 결) · 흔적 비교 밖)
        "ledger": {f.name: json.loads(f.read_text(encoding="utf-8")) for f in sorted(led.glob("*.json"))
                   if not f.name.endswith(_LDG.PARTIAL)} if led.exists() else {},
        "files": sorted(_P.rel_to_home(x) for x in _P.data().rglob("*") if x.is_file() and not x.name.endswith(".lock"))}),
        ensure_ascii=False, sort_keys=True)


class _Cut(io.StringIO):
    """k번째 줄(또는 `mark`가 든 줄)에서 닫히는 화면 — 파이프가 그 줄에서 닫힌 것과 같다(`write`가 `BrokenPipeError`)."""

    def __init__(self, k, mark=None):
        super().__init__()
        self.k, self.mark, self.lines = k, mark, 0

    def write(self, s):
        self.lines += s.count("\n")
        if (self.k is not None and self.lines >= self.k) or (self.mark and self.mark in s):
            raise BrokenPipeError(32, "Broken pipe")
        return super().write(s)


def _last_run_log(command):
    """그 명령의 오늘 로그에서 마지막 실행 한 덩이."""
    f = LG.log_dir() / f"{command}_{time.strftime('%Y%m%d')}.log"
    txt = f.read_text(encoding="utf-8") if f.exists() else ""
    return txt[txt.rfind("===== 실행 시작"):]


def _inproc(case, k):
    """공통 진입으로 인입 — 화면이 k번째 줄에서 닫힌다(None이면 안 닫힌다). `(종료 코드, 화면 줄 수, 그 실행의 로그)`."""
    real, cut = sys.stdout, _Cut(k)
    sys.stdout = cut
    try:
        rc = _entry.run("ingest-file", IG.main, _argv(case))
    finally:
        sys.stdout = real
    return rc, cut.lines, _last_run_log("ingest-file")


def _closed(argv, script="run.py", module=None):
    """자식 프로세스의 화면을 **첫 쓰기 전에** 닫는다(`| head -0`) — `(종료 코드, stderr)`."""
    cmd = [sys.executable, "-m", module, *argv] if module else [sys.executable, str(ROOT / script), *argv]
    p = subprocess.Popen(cmd, cwd=str(ROOT), env=ENV, stdin=subprocess.DEVNULL,
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    p.stdout.close()
    err = p.stderr.read().decode("utf-8", "replace")
    p.stderr.close()
    return p.wait(), err


def _no_tb(err):
    return "Traceback" not in err and "Exception ignored" not in err and "Exception in thread" not in err


def _stop_line(lg):
    return "화면이 닫혀 멈췄다 — " in lg or "화면이 닫혔다(명령은 끝까지 돌았다" in lg


def _logs():
    return {f.name: f.read_bytes() for f in sorted(LG.log_dir().glob("*.log"))} if LG.log_dir().exists() else {}


# ────────────────────────────────────────────────────────────── ⓐ 로그 보존
print("\n■ B109 ① 명령 로그 보존 — 기본 fresh · 재구축이 남기고 클린만 지운다 (표 CP01 · 산문 B109SQ)")
_base()
_o1 = _run("ingest-file", *_argv("표 · process"))
_o2 = _run("ingest-file", *_argv("산문 · quality"))
L0 = _logs()
_ob = _run("rebuild", "--allow-mock", "--yes")
L1 = _logs()
_rb = _last_run_log("rebuild")
_plan = next((ln.strip() for ln in _ob.splitlines() if ln.strip().startswith("남길 것 —")), "")
show("ⓐ 인입(표 · 산문) → rebuild --yes ⇒ 재구축 전 인입 로그가 바이트 그대로 · 그 실행의 로그도(계획 · 끝) · 계획 「남길 것」에 "
     "명령 로그", L0 and all(L1.get(n) == b for n, b in L0.items()) and "■ 재구축 계획" in _rb and "■ 재구축 끝" in _rb
     and "명령 로그" in _plan and "성공" in _o1 and "성공" in _o2,
     f"그 전 로그 {sorted(L0)} → 그대로 {sum(L1.get(n) == b for n, b in L0.items())}/{len(L0)} · 재구축 로그 "
     f"계획 {'■ 재구축 계획' in _rb} · 끝 {'■ 재구축 끝' in _rb}")
_oi = _run("init", "--fresh")          # 시험 대상: 기본 fresh — 클린 호출이 아니다(B109 ① 로그 보존)
L2 = _logs()
_kept = next((ln.strip() for ln in _oi.splitlines() if "남긴 것 —" in ln), "")
show("ⓐ 기본 init --fresh ⇒ 그 전 명령 로그(인입 · 재구축)가 바이트 그대로 · 화면 「남긴 것」에 명령 로그",
     all(L2.get(n) == b for n, b in L1.items()) and "명령 로그" in _kept,
     f"그 전 로그 {sorted(L1)} → 그대로 {sum(L2.get(n) == b for n, b in L1.items())}/{len(L1)} · 「{_kept}」")
_oc = _run("init", "--fresh", "--all")
_left = sorted(_P.rel_to_home(x) for t in ("data", "work", "export") for x in getattr(_P, t)().rglob("*") if x.is_file()
               and not x.name.endswith(".lock"))
_empty = {_P.rel_to_home(store.path(n)) for n in init.EMPTY}
_res = [f for f in _left if f not in _empty and not f.endswith("/" + GraphStore.FILENAME)
        and not (f.startswith("work/logs/init_") and f.endswith(".log"))]
show("ⓐ init --fresh --all ⇒ 지금의 클린 — 잔재 0(명령 로그도 지운다 · 그 실행의 로그만 다시 열린다) · 그래프는 빈 그래프",
     not _res and all(not g.nodes and not g.edges for g in World().graphs.values())
     and [f for f in _left if f.startswith("work/")] == [f"work/logs/init_{time.strftime('%Y%m%d')}.log"],
     f"남은 파일 {len(_left)}(빈 상태 · 빈 그래프 · 그 실행 로그) · 잔재 {_res[:3]}")

# ────────────────────────────────────────────────────────────── ⓑ 화면 닫힘 — 읽기
print("\n■ B109 ② 화면이 닫혀도 — 읽기 명령(진입 갈래 셋)")
_reads = [("run.py show aliases", "show", _closed(["show", "aliases"])),
          ("python -m cli.platform queue", "platform", _closed(["queue"], module="cli.platform")),
          ("doctor.py --quick", "doctor", _closed(["--quick"], script="doctor.py"))]
_rd = [(lbl, rc, _no_tb(err), _stop_line(_last_run_log(cmd))) for lbl, cmd, (rc, err) in _reads]
show("ⓑ 읽기 명령 파이프 닫힘 — 트레이스백 0 · 종료 코드 141 · 로그에 멈춤 줄(run.py · python -m cli.* · doctor.py)",
     all(rc == LG.SCREEN_CLOSED and tb and sl for _l, rc, tb, sl in _rd),
     " · ".join(f"{lbl} rc {rc}{'' if tb else ' 트레이스백!'}{' 멈춤 줄' if sl else ' 멈춤 줄 없음!'}" for lbl, rc, tb, sl in _rd))


def _thread_main(_argv):
    """화면에 쓰는 스레드(30초 줄 자리) — 화면이 닫힌 뒤 쓴다 · 본 흐름은 다음 줄에서 멈춘다."""
    print("첫 줄")
    t = threading.Thread(target=lambda: print("스레드 줄"))
    t.start()
    t.join()
    print("본 흐름 다음 줄")
    return 0


_err = io.StringIO()
with contextlib.redirect_stderr(_err):
    _real, sys.stdout = sys.stdout, _Cut(2)            # 첫 줄은 나가고 스레드의 줄에서 닫힌다
    try:
        _trc = _entry.run("b109-thread", _thread_main, [])
    finally:
        sys.stdout = _real
show("ⓑ 화면에 쓰는 스레드도 닫힌 화면에서 트레이스백 0 — 본 흐름은 멈춤 줄 · 141",
     _trc == LG.SCREEN_CLOSED and _no_tb(_err.getvalue()) and _stop_line(_last_run_log("b109-thread")),
     f"rc {_trc} · stderr {len(_err.getvalue().splitlines())}줄")
(LG.log_dir() / f"b109-thread_{time.strftime('%Y%m%d')}.log").unlink(missing_ok=True)

# ────────────────────────────────────────────────────────────── ⓑ 화면 닫힘 — 쓰기(표·산문 × 층 둘)
print("\n■ B109 ② 화면이 닫혀도 — 쓰기 명령 ingest-file (표 · 산문 × 층 둘)")
REF, TOTAL, WHERE = {}, {}, {}
for case in CASES:
    _base()
    s0 = _state()
    src, dt, extra = CASES[case]
    with contextlib.redirect_stdout(io.StringIO()):
        _r = IG.ingest_file(str(src), dt, ask=False, finalize_after=False,
                            sheets=SH if extra else None, coord="노칭" if extra else None)
    sdoc = _state()
    _base()
    _rc, TOTAL[case], _lg = _inproc(case, None)
    REF[case] = {"시작 전": s0, "문서 끝": sdoc, "마무리까지": _state()}
    assert _r.get("status") == "성공" and _rc == 0, (case, _r, _rc)


def _kind(case, s):
    """셋 중 어디인가 — 문서 끝 = 마무리 전(마무리가 바꿀 것이 없으면 마무리까지와 같다)."""
    return next((k for k, v in REF[case].items() if v == s), "어긋남")


_def = store.path(store.DEFECTS)
_wr = []
for case in CASES:
    _base()
    rc, err = _closed(["ingest-file", *_argv(case)])
    _wr.append((case, rc, _no_tb(err), _stop_line(_last_run_log("ingest-file")), _kind(case, _state()),
                not (_def.exists() and any(w in _def.read_text(encoding="utf-8") for w in ("ScreenClosed", "Broken pipe")))))
show("ⓑ 쓰기 명령 파이프 닫힘(표 · 산문 × 층 둘) — 트레이스백 0 · 141 · 멈춤 줄 · ③은 셋 중 하나 · 결함 장부에 화면 닫힘 0",
     all(rc == LG.SCREEN_CLOSED and tb and sl and k != "어긋남" and dfc for _c, rc, tb, sl, k, dfc in _wr),
     " · ".join(f"{c} rc {rc} {k}{'' if tb else ' 트레이스백!'}{'' if sl else ' 멈춤 줄 없음!'}{'' if dfc else ' 결함!'}"
                for c, rc, tb, sl, k, dfc in _wr))

_cuts = {}
for case in CASES:
    n = TOTAL[case]
    picks = sorted(set(range(1, n, max(1, n // 3))) | set(range(max(1, n - 9), n + 1)))   # B131 ④ — 고른 셋 + 끝 열(저장 경계)
    for k in picks:
        _base()
        rc, _l, lg = _inproc(case, k)
        _cuts.setdefault(case, []).append((k, rc, _kind(case, _state()), "되돌림 — " in lg, "저장 뒤 멈춤 — " in lg,
                                          _stop_line(lg)))
_bad = [(c, k, kd) for c, rows in _cuts.items() for k, rc, kd, back, after, sl in rows
        if rc != LG.SCREEN_CLOSED or kd == "어긋남" or not sl or (back and kd != "시작 전")
        or (after and kd == "시작 전")]
_both = all(any(b and kd == "시작 전" for _k, _rc, kd, b, _a, _s in rows)
            and any(a and kd != "시작 전" for _k, _rc, kd, _b, a, _s in rows) for rows in _cuts.values())
show("ⓑ 줄 위치 표본(표 · 산문 × 층 둘) — 전부 셋 중 하나 · 「되돌림」 줄이면 시작 전 · 「저장 뒤 멈춤」 줄이면 끝난 문서 · "
     "종류마다 둘 다 나온다", not _bad and _both,
     f"어긋남 {len(_bad)} {_bad[:3]}")
for case, rows in _cuts.items():
    agg = {}
    for k, _rc, kd, b, a, _s in rows:
        agg.setdefault(kd + (" · 되돌림 줄" if b else " · 저장 뒤 줄" if a else ""), []).append(k)
    print(f"      {case} — 화면 {TOTAL[case]}줄 · 끊은 자리 {len(rows)}: "
          + " · ".join(f"{kd} {len(v)}({v[0]}~{v[-1]})" for kd, v in agg.items()))

def _revise(path, kind):
    """개정본 — 같은 자리 · 같은 doc_id · 내용이 다르다(해시가 바뀌어 재인입이 회수를 탄다)."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["관리계획서"] if kind == "표" else wb["요구사항"]
    c = next(c for row in ws.iter_rows(min_row=4) for c in row if isinstance(c.value, str) and c.value.strip())
    c.value = c.value + " (B109 개정)"
    wb.save(path)


_re = []
#: 구축 안(되돌림 경계 안 · 저장 전)에서만 찍히는 줄 — 표는 판정 예고 · 산문은 추출 끝
_INSIDE = {"표": "판정 예고 — entity", "산문": "추출 끝 —"}
for case, did in (("표 · process", "B109RP"), ("산문 · quality", "B109RQ")):
    src, dt, extra = CASES[case]
    form = case.split(" ")[0]
    at, rev = TMP / f"{did}.xlsx", TMP / f"{did}_개정.xlsx"
    shutil.copy(src, at)
    shutil.copy(src, rev)
    _revise(rev, form)
    if dt:
        (HINTS / f"{did}.json").write_text((HINTS / f"{src.stem}.json").read_text(encoding="utf-8"), encoding="utf-8")
        _HINTED[did] = None
    _base()
    with contextlib.redirect_stdout(io.StringIO()):
        r1 = IG.ingest_file(str(at), dt, ask=False, sheets=SH if extra else None, coord="노칭" if extra else None)
    s1 = _state()
    shutil.copy(rev, at)                                   # 같은 자리 · 같은 doc_id · 내용이 다르다
    real, cut = sys.stdout, _Cut(None, _INSIDE[form])
    sys.stdout = cut
    try:
        rc = _entry.run("ingest-file", IG.main, [str(at), "--allow-mock"] + (["--doc-type", dt] if dt else []) + extra)
    finally:
        sys.stdout = real
    _re.append((case, rc, r1.get("status"), "되돌림 — " in _last_run_log("ingest-file"), _state() == s1))
show("ⓑ 재인입(개정본) 도중에 닫히면 첫 인입 뒤 판 그대로 — 회수(withdraw)가 저장한 층 그래프도 되돌린다(표 · 산문 · 층 둘)",
     all(rc == LG.SCREEN_CLOSED and st == "성공" and b and same for _c, rc, st, b, same in _re),
     " · ".join(f"{c} rc {rc} · 첫 인입 {st} · 되돌림 줄 {b} · 첫 인입 뒤와 같음 {same}" for c, rc, st, b, same in _re))

for _d in _HINTED:
    (HINTS / f"{_d}.json").unlink(missing_ok=True)
for _dt in (DT, QDT):
    _unregister(_dt)
shutil.rmtree(TMP, ignore_errors=True)
with contextlib.redirect_stdout(io.StringIO()):
    init.init(fresh_=True, all_=True)
done()
