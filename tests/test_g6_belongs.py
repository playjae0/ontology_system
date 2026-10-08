# -*- coding: utf-8 -*-
"""G6 ⑳ 소속 — 맥락 줄 · 문서 좌표 · 소속 계약 · 이름에 소속 · 좌표 없는 스코프 개체 0 · 재사용 조건 · 붙은 곳
화면 · 보존물·동명 (B102).

창작 표본(mock · 시험 주입)으로 메커니즘을 잰다 — 수치는 근거가 아니다([정정] 50). 산문 표본은 **엑셀**(시트가
유닛 단위 · 제목 있는 시트 「이송 유닛」 · 제목 없는 시트 「가압 유닛」)과 **Word**(제목 경로)를 실제 파일로 만들어
파서로 읽는다. 추출은 mock 힌트(문서 7 §7.5 대체 표 「추출」 행)다. 시험마다 형태(표·산문) × 층(process ·
quality)에서 같은 성질을 잰다(CLAUDE.md §4 범용성).

잠그는 성질:
  ⓐ 맥락 줄: 엑셀은 제목이 있어도 시트명이 실린다 · 문서 · 공정 칸 · 포맷 넷(엑셀·Word·PDF·PPT)이 한 함수 ·
     실호출 입력이 그 함수의 출력이다
  ⓑ 문서 좌표: 다른 공정의 이름인 시트명 · 공유 별칭 시트 — 문서 좌표면 서브트리 밖 좌표 0 · 없으면 공유 별칭 좌표 0 ·
     파일명 제안(하나일 때만) · 기록 재사용 · 비대화형 한 줄 · 표도 같은 함수
  ⓒ 소속: 「모터」가 시트·경로·본문의 유닛 이름을 소속으로 갖는다(from) · 근거 없으면 null(지어낸 이름은 버림) ·
     경로·시트 이름은 개체 후보가 아니다 · 옛 체크포인트 `attach` 호환
  ⓓ 이름에 소속: 두 유닛의 「모터」가 다른 노드 · 소속 없는 「모터」는 좌표 접두 · `Property`는 공정 스코프(CP와 같은
     키로 만난다) · `nest_categories`가 비면 지금과 같다
  ⓔ 좌표 없는 스코프 개체 0(표·산문 × 층 둘) · 재료는 큐가 든다 · 좌표를 주고 다시 넣으면 붙는다 · 엣지 없는 노드 0
  ⓕ 재사용 조건: 지시문 판·층 config 판이 바뀌면 다시 뽑는다(이유 줄)
  ⓖ 대장 `target`·`attached`·`from` · 「표기 → 붙은 노드」 · 「└ 관계 → 상대」 · `--trace` · 좌표 줄 · 끝 요약 ·
     목록 · 재시도 분포 · 표 진행 줄 분모 = 실제 값 수
  ⓗ 해시 없는 보존물 재사용 0 · 동명 다른 경로 멈춤 / `--revise`면 개정
"""
from __future__ import annotations

import contextlib
import io
import re
import zipfile
from pathlib import Path

sys_path = __import__("sys").path
sys_path.insert(0, str(Path(__file__).resolve().parent))
sys_path.insert(0, str(Path(__file__).resolve().parent / "fixtures"))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, _rw, done, overlay   # noqa: F401

import make_b88 as MB                                   # noqa: E402 — Word 조립 재료(창작)
from openpyxl import Workbook                           # noqa: E402
from cli import doc_coord as DC                         # noqa: E402
from cli import ingest as IG                            # noqa: E402
from core.build import extract as EX, extract_ctx as XC, ledger as LG   # noqa: E402
from core.llm import gateway                            # noqa: E402
from core.state.world import World                      # noqa: E402
from parser import pipeline, struct_map, tagger as TG   # noqa: E402
from parser.adapters import basic_docx as DX, basic_prose_xlsx as BX   # noqa: E402

HINTS = ROOT / "tests" / "fixtures" / "extract_hints"
TMP = Path(__import__("tempfile").mkdtemp(prefix="b102_"))   # 클린(init --fresh)이 지우지 않는 자리
XLSX_SHEETS = {"이송 유닛": ["1. 구동부", "모터는 서보 모터를 쓴다.", "2. 안전", "커버를 단다."],
               "가압 유닛": ["모터는 유압 모터를 쓴다.", "가압력은 매 시간 기록한다."]}


def _xlsx(path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in XLSX_SHEETS.items():
        ws = wb.create_sheet(title=name)
        for r, v in enumerate(rows, 1):
            ws.cell(row=r, column=1, value=v)
    _P.ensure(path)
    wb.save(path)
    return path


def _docx(path):
    body = "".join([MB._p("이송 유닛", "1"), MB._p("모터는 서보 모터를 쓴다."),
                    MB._p("가압 유닛", "1"), MB._p("모터는 유압 모터를 쓴다. 가압력은 매 시간 기록한다.")])
    _P.ensure(path)
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", MB._CT)
        z.writestr("_rels/.rels", MB._ROOT_RELS)
        z.writestr("word/document.xml", MB._doc(body))
        z.writestr("word/styles.xml", MB._STYLES)
        z.writestr("word/_rels/document.xml.rels", MB._rels([("rIdSty", "styles", "styles.xml")]))
    return path


FILES = {"엑셀": _xlsx(TMP / "B102_노칭_사양.xlsx"), "Word": _docx(TMP / "B102_노칭_설명.docx")}
ROLES = {"이송 유닛": "prose", "가압 유닛": "prose"}


def _parse(form, doc_id, doc_coord=None, layer=None):
    ad = BX if form == "엑셀" else DX
    struct_map.invalidate(doc_id)
    kw = {"sheet_roles": ROLES} if form == "엑셀" else {}
    r = pipeline.parse(ad, doc_id, str(FILES[form]), closed_list=TG.closed_list(layer or "process"),
                       doc_coord=doc_coord, **kw)
    return r.envelope


def _setup(nest=("Component",)):
    """덧칠(Component · 겸 렌즈) + 공통 `nest_categories` + 소속 엣지에 필요한 카테고리쌍 — 상태 루트에만."""
    overlay(component=True)

    def _c(c):
        c.setdefault("canonical_scope", {})["nest_categories"] = list(nest)
    _rw(_P.common(), _c)

    pats = [{"src": s_, "rel": "part_of", "dst": d_, "symmetric": False, "정의문": f"{s_}가 {d_}에 속한다"}
            for s_, d_ in (("Component", "Unit"), ("Component", "Process"), ("Unit", "Process"))]
    pairs = {f"{p_['src']},{p_['dst']}": "part_of" for p_ in pats}

    def _q(c):                                    # 품질 층도 Component·Unit을 렌즈로 쓴다(층 둘)
        c["categories"]["Component"] = "(품질층의 렌즈) 불량이 일어나는 구성 기계"
        c["relation_patterns"] += pats
        c.setdefault("category_pair_map", {}).update(pairs)

    def _p(c):                                    # 공정 층 — 이미 있는 쌍은 그대로(덧칠은 더하기만)
        have = {(p_["src"], p_["rel"], p_["dst"]) for p_ in c["relation_patterns"]}
        c["relation_patterns"] += [p_ for p_ in pats if (p_["src"], p_["rel"], p_["dst"]) not in have]
        for k, v in pairs.items():
            c.setdefault("category_pair_map", {}).setdefault(k, v)
    _rw(_P.layers("quality", "config.json"), _q)
    _rw(_P.layers("process", "config.json"), _p)
    for lay in ("process", "quality"):
        bootstrap(lay, echo=False)


def _hint(doc_id, by_loc):
    (HINTS / f"{doc_id}.json").write_text(json.dumps(by_loc, ensure_ascii=False), encoding="utf-8")


def _unit_of(c):
    """그 청크가 어느 유닛의 것인가 — 표본 정답(시트 또는 제목 경로에서)."""
    m = c.get("meta") or {}
    return m.get("sheet") or (m.get("section_path") or c.get("section") or "").split(" > ")[0]


def _from_of(c):
    return "시트" if (c.get("meta") or {}).get("sheet") else "경로"


def _prose_env(form, lay, doc_id, belongs="unit", doc_coord="노칭"):
    env = _parse(form, doc_id, doc_coord=doc_coord)
    dt = "ppt_" + lay
    env["doc_type"] = dt
    for c in env["chunks"]:
        c["doc_type"] = dt
    hints = {}
    for c in env["chunks"]:
        if "모터" not in c.get("text", ""):
            continue
        u = _unit_of(c)
        b = ({"name": u, "category": "Unit", "from": _from_of(c)} if belongs == "unit" else
             {"name": "없는 유닛 B102", "category": "Unit", "from": "본문"} if belongs == "invented" else None)
        ents = [{"surface": "모터", "category": "Component", "belongs_to": b}]
        if "가압력" in c.get("text", "") and "Property" in load_config(lay)["categories"]:
            ents.append({"surface": "가압력", "category": "Property",
                         "belongs_to": {"name": u, "category": "Unit", "from": _from_of(c)}})
        hints[c["source_locator"]] = {"entities": ents, "relations": [], "attach": []}
    if belongs == "legacy":                        # 옛 체크포인트 — attach만 있고 belongs_to 없음
        hints = {k: {"entities": [{"surface": "모터", "category": "Component"}], "relations": [],
                     "attach": [{"surface": "모터", "attach_to": {"name": "노칭", "category": "Process"}}]}
                 for k in hints}
    _hint(doc_id, hints)
    return env


def _live(name_part, cat=None):
    return [n for _l, n in World().nodes() if name_part in n["canonical"] and ops.is_live(n)
            and (cat is None or n["category"] == cat)]


def _all_edges():
    out = []
    for lay in ("process", "quality"):
        out += [(e["rel"], e["src"], e["dst"]) for e in open_graph(lay).edges]
    return out


def _canon(nid):
    n = World().get(nid)
    return (n or {}).get("canonical")


PROSE_MATRIX = [(f, l) for f in ("엑셀", "Word") for l in ("process", "quality")]
_made = []

# ────────────────────────────────────────────────────────────── ⓐ
print("\n■ B102 ① 맥락 줄 — 한 함수 (엑셀 · Word · PDF · PPT)")
_setup()
_ok, _det = True, []
for form in ("엑셀", "Word"):
    env = _parse(form, f"B102CTX{form[0]}", doc_coord="노칭")
    doc = XC.doc_info(env)
    for c in env["chunks"]:
        line = XC.context_line(c, doc)
        need = [f"문서: {FILES[form].name}", "공정: 노칭"]
        if form == "엑셀":
            need.append(f"시트: {c['meta']['sheet']}")
        else:
            need.append(f"경로: {c['meta']['section_path']}")
        _ok &= all(x in line for x in need)
        _det.append(line)
_titled = [l for l in _det if "시트: 이송 유닛" in l]
_ok &= bool(_titled) and all("경로: " in l for l in _titled)          # 제목이 있어도 시트명
# PDF · PPT 조각 — 같은 함수(포맷 차이는 조각이 든 값이 말한다)
_pp = XC.context_line({"process_ref": None, "section": "슬라이드 3", "meta": {"section_path": "슬라이드 3"}},
                      {"name": "a.pptx", "coord": "스태킹"})
_pd = XC.context_line({"process_ref": "적층", "meta": {"section_path": "2. 사양 > 2.1 적층"}}, {"name": "b.pdf"})
_ok &= _pp == "[문서: a.pptx · 공정: 스태킹 · 경로: 슬라이드 3]" and _pd == "[문서: b.pdf · 공정: 적층 · 경로: 2. 사양 › 2.1 적층]"
# 실호출 입력 = 그 함수의 출력(가짜 게이트웨이 — 보낸 것을 잡는다)
_sent = []
_keep = (gateway.use_mock, gateway.chat, gateway.prompt)
gateway.use_mock = lambda: False
gateway.chat = lambda msgs, **k: (_sent.append(json.loads(msgs[1]["content"])), {"entities": [], "relations": []})[1]
gateway.prompt = lambda name: "지시문"
try:
    env = _parse("엑셀", "B102CTXL", doc_coord="노칭")
    EX._DOC.clear()
    EX._DOC.update(XC.doc_info(env))
    c0 = env["chunks"][0]
    EX._candidates_for("x", c0, load_config("process"), {})
finally:
    gateway.use_mock, gateway.chat, gateway.prompt = _keep
_ok &= bool(_sent) and _sent[0]["chunk"] == XC.with_context(c0, XC.doc_info(env)) \
    and _sent[0]["chunk"].startswith(XC.context_line(c0, XC.doc_info(env)))
show("ⓐ 맥락 줄 — 엑셀은 제목이 있어도 시트명 · 문서 · 공정 칸 · 엑셀·Word·PDF·PPT 한 함수 · 실호출 입력이 그 출력",
     _ok, " / ".join(_det[:3]) + f" / {_pp} / {_pd}")

# ────────────────────────────────────────────────────────────── ⓑ
print("\n■ B102 ② 문서 좌표 — 서브트리 · 공유 별칭 · 파일명 제안 · 기록 · 표도 같은 함수")
nodes = TG.closed_list("process")
pcs = [{"source_locator": "S1", "section": "적층"},            # 다른 공정(스태킹)의 이름인 시트명
       {"source_locator": "S2", "section": "비전 검사"},       # 공유 별칭(노칭·스태킹 둘)
       {"source_locator": "S3", "section": "노칭 타발"}]
_with = {p["source_locator"]: p for p in TG.coord_from_section(pcs, nodes=nodes, doc_coord="노칭")}
_none = {p["source_locator"]: p for p in TG.coord_from_section(pcs, nodes=nodes)}
_sub_idx = TG.scoped_index(nodes, "노칭")[0]            # 노칭 서브트리 안에서 하나로 풀리는 표기
_ok_b1 = (all(p["process_ref"] in _sub_idx for p in _with.values())
          and _with["S1"]["meta"].get("coord_ignored") == ["적층"] and _with["S1"]["meta"].get("coord_from_doc")
          # 기대 변경(B104 ①) — 범위 안에서 하나인 공유 별칭은 그 노드 canonical로 쓰고 원 표기는 meta에
          and _with["S2"]["process_ref"] == _sub_idx["비전 검사"]["canonical"]
          and _with["S2"]["meta"].get("coord_tag_from") == "비전 검사" and _with["S3"]["process_ref"] == "노칭 타발")
_ok_b2 = (_none["S2"].get("process_ref") is None and _none["S2"]["meta"].get("coord_shared_skip") == ["비전 검사"]
          and _none["S1"]["process_ref"] == "적층")
show("ⓑ 문서 좌표면 서브트리 밖 좌표 0(밖 이름은 무시 표시 · 문서 좌표 물려받음) · 없으면 공유 별칭 좌표 0",
     _ok_b1 and _ok_b2, f"좌표 있음 {[(k, p['process_ref']) for k, p in _with.items()]} · "
     f"없음 {[(k, p.get('process_ref')) for k, p in _none.items()]}")
_s1 = TG.suggest_doc_coord("B102_노칭_사양.xlsx", nodes)
_s2 = TG.suggest_doc_coord("노칭_스태킹_비교.xlsx", nodes)
_s3 = TG.suggest_doc_coord("비전 검사_기준.xlsx", nodes)
show("ⓑ 파일명 제안 — 조각 하나가 노드 하나에 맞을 때만(둘이면 없음 · 공유 별칭 없음)",
     _s1[0] == "노칭" and _s2[0] is None and _s3[0] is None, f"{_s1} · {_s2} · {_s3}")
_rec = DC._file()
_rec0 = _rec.read_bytes() if _rec.exists() else None
_b = io.StringIO()
with contextlib.redirect_stdout(_b):
    _d1 = DC.decide("x/B102DC.xlsx", "B102DC", "노칭", ask=False)
    _d2 = DC.decide("x/B102DC.xlsx", "B102DC", None, ask=False)
    _d3 = DC.decide("x/B102_노칭.xlsx", "B102DC2", None, ask=False)
    _d4 = DC.decide("x/B102무명.xlsx", "B102DC3", None, ask=False)
_lines = [l for l in _b.getvalue().splitlines() if l.strip()]
show("ⓑ 기록 재사용(사람 → 기록) · 파일명 · 없음 — 비대화형은 묻지 않고 한 줄씩",
     (_d1, _d2, _d3, _d4) == (("노칭", "사람"), ("노칭", "기록"), ("노칭", "파일명"), (None, "없음"))
     and len(_lines) == 4 and all(l.strip().startswith("문서 좌표 — ") for l in _lines), " / ".join(_lines))
if _rec0 is None:
    _rec.unlink(missing_ok=True)
else:
    _rec.write_bytes(_rec0)
_recs = [{"source_locator": "R1", "process_ref": None}, {"source_locator": "R2", "process_ref": "노칭 타발"}]
_inh = TG.inherit_doc_coord(_recs, "노칭")
show("ⓑ 표도 같은 함수 — 빈 행만 문서 좌표를 물려받는다(있는 값은 덮지 않는다)",
     [r["process_ref"] for r in _inh] == ["노칭", "노칭 타발"] and _inh[0]["meta"].get("coord_from_doc")
     and not (_inh[1].get("meta") or {}).get("coord_from_doc"), str([r["process_ref"] for r in _inh]))

# ────────────────────────────────────────────────────────────── ⓒ ⓓ ⓔ(산문)
print("\n■ B102 ③④⑤ 소속 · 이름에 소속 · 좌표 없는 스코프 개체 0 (엑셀·Word × 층 둘)")
_oc, _od, _oe, _ol = [], [], [], []
for form, lay in PROSE_MATRIX:
    _setup()
    did = f"B102{form[0]}{lay[0].upper()}"
    _made.append(did)
    env = _prose_env(form, lay, did)
    run_document(env)
    rows = [r for r in (LG.read(did) or {}).get("rows") or [] if r.get("role") == "entity"]
    motors = {r["target"] for r in rows if r["surface"] == "모터"}
    froms = {r.get("from") for r in rows if r["surface"] == "모터"}
    want = {f"노칭::{u}::모터" for u in ROLES}
    # 모터 → 유닛(소속) · 유닛 → 골격(좌표 폴백)
    e = _all_edges()
    units = {u: next((n["id"] for n in _live(f"노칭::{u}", "Unit") if n["canonical"] == f"노칭::{u}"), None)
             for u in ROLES}
    skel = next(n["id"] for _l, n in World().nodes() if n["canonical"] == "노칭")
    mot_ids = {n["canonical"]: n["id"] for n in _live("::모터", "Component")}
    edge_ok = all(any(r == "part_of" and s == mot_ids.get(f"노칭::{u}::모터") and d == units[u] for r, s, d in e)
                  and any(units[u] in (s, d) and skel in (s, d) for r, s, d in e) for u in ROLES)
    res = (LG.read(did) or {}).get("result") or {}
    _oc.append((form, lay, motors == want and froms == {_from_of(env["chunks"][0])} and edge_ok
                and not any(r["surface"] in ROLES for r in rows),
                f"모터 → {sorted(motors)} · from {sorted(map(str, froms))} · 엣지 {edge_ok}"))
    _od.append((form, lay, len(motors) == 2
                and {r["target"] for r in rows if r["surface"] == "가압력"} == ({"노칭::가압력"} if lay == "process" else set()),
                f"가압력 → {sorted({r['target'] for r in rows if r['surface'] == '가압력'})}"))
    _ol.append((form, lay, res.get("엣지 없는 노드") == 0 and (res.get("부착") or {}).get(f"소속 {_from_of(env['chunks'][0])}", 0) >= 2,
                f"엣지 없는 노드 {res.get('엣지 없는 노드')} · 부착 {res.get('부착')}"))
    # 근거 없음 → null(지어낸 이름은 버린다) · 옛 attach
    run_document(_prose_env(form, lay, did + "N", belongs="invented"), allow_duplicate=True)
    _made.append(did + "N")
    rn = [r for r in (LG.read(did + "N") or {}).get("rows") or [] if r.get("surface") == "모터"]
    dfx = store.path(store.DEFECTS).read_text(encoding="utf-8") if store.path(store.DEFECTS).exists() else ""
    _null_ok = (rn and all(r.get("from") is None and r["target"] == "노칭::모터" for r in rn)
                and "소속 출처 없음" in dfx and not _live("없는 유닛 B102")
                and ((LG.read(did + "N") or {}).get("result") or {}).get("소속 없음 비율"))
    run_document(_prose_env(form, lay, did + "L", belongs="legacy"), allow_duplicate=True)
    _made.append(did + "L")
    rl = [r for r in (LG.read(did + "L") or {}).get("rows") or []
          if r.get("surface") == "모터" and r.get("role") == "entity"]
    _leg_ok = rl and all(r.get("from") == "옛 attach" and r["target"] == "노칭::모터" for r in rl)
    _oc[-1] = (*_oc[-1][:2], _oc[-1][2] and _null_ok and _leg_ok,
               _oc[-1][3] + f" · 지어낸 소속 → null {bool(_null_ok)} · 옛 attach {bool(_leg_ok)}")
# 경로·시트 이름은 개체 후보가 아니다 — mock 갈래는 본문만 본다(맥락 줄은 실호출의 소속 근거)
_c0 = next(c for c in _parse("엑셀", "B102MK")["chunks"] if (c.get("meta") or {}).get("sheet") == "가압 유닛")
_mk = EX._mock_candidates("x", _c0["text"], load_config("process"), {"가압 유닛": "Unit", "모터": "Component"})
show("ⓒ 소속 — 「모터」가 시트·경로의 유닛 이름을 소속으로(from) · 소속 엣지 · 유닛은 좌표에 · 근거 없으면 null"
     "(지어낸 이름은 버리고 결함 로그) · 옛 attach 호환 (엑셀·Word × 층 둘)",
     all(o for _f, _l, o, _d in _oc), " ‖ ".join(f"{f}×{l} {d}" for f, l, _o, d in _oc))
show("ⓒ 경로·시트 이름은 개체 후보가 아니다 — 본문에 없는 시트명은 후보 0 · 지시문 규칙 판 e-1.4",
     [x["surface"] for x in _mk["entities"]] == ["모터"] and EX.prompt_version() == "e-1.4",
     f"후보 {[x['surface'] for x in _mk['entities']]} · 지시문 {EX.prompt_version()}")

# 이름에 소속 — CP 표본과 같은 키 · nest가 비면 지금과 같다
_setup()
run_document(_prose_env("엑셀", "process", "B102KEY"))
_cp = {**TABLE, "doc_id": "B102CPK", "records": [dict(CPREC, source_locator="K1", 관리항목="가압력")]}
run_document(_cp)
_pk = _live("노칭::가압력", "Property")
_setup(nest=())
run_document(_prose_env("엑셀", "process", "B102FLAT"))
_flat = {r["target"] for r in (LG.read("B102FLAT") or {}).get("rows") or [] if r.get("surface") == "모터"}
_made += ["B102KEY", "B102FLAT"]
show("ⓓ 이름에 소속 — 두 유닛의 「모터」는 다른 노드 · Property는 공정 스코프(CP 행과 한 노드) · nest가 비면 지금과 같다(스코프 밖 이름 하나)",
     all(o for _f, _l, o, _d in _od) and len(_pk) == 1 and _flat == {"모터"},
     " ‖ ".join(f"{f}×{l} {d}" for f, l, _o, d in _od) + f" ‖ CP·산문 가압력 노드 {len(_pk)} · nest 빈 목록 → {sorted(_flat)}")

# 좌표 없는 스코프 개체 0 — 표·산문 × 층 둘
for form, lay in [("표", "process"), ("표", "quality"), ("엑셀", "process"), ("Word", "quality")]:
    _setup()
    did = f"B102NC{form[0]}{lay[0].upper()}"
    _made.append(did)
    if form == "표":
        dt = f"b102t{lay[0]}"
        p = _P.schemas(f"{dt}.json")
        _P.ensure(p)
        p.write_text(json.dumps({"doc_type": dt, "schema_version": 1, "layer": lay,
                                 "use_blocks": ["common_core", "process_coord"],
                                 "fields": {"개체": {"role": "entity", "category": "Unit"}}, "edges": []},
                                ensure_ascii=False), encoding="utf-8")
        dts = store.read(store.DOC_TYPES, {})
        dts[dt] = {"doc_type": dt, "status": "registered", "layer": lay, "schema": f"schemas/{dt}.json",
                   "adapter": "-", "schema_version": 1}
        store.write(store.DOC_TYPES, dts)
        base = {k: v for k, v in CPREC.items() if k not in ("설비", "관리항목")}
        mk = lambda ref: {**TABLE, "doc_id": did, "doc_type": dt, "records": [  # noqa: E731
            dict(base, source_locator="N1", doc_type=dt, process_ref=ref, process_group=ref and "조립",
                 개체="B102 이송 유닛")]}
        env0, env1 = mk(None), mk("노칭")
    else:
        env0 = _prose_env(form, lay, did, belongs=None, doc_coord=None)
        env1 = _prose_env(form, lay, did, belongs=None, doc_coord="노칭")
    run_document(env0)
    q = [x for x in store.read(store.QUEUE, []) if x["doc_id"] == did and x["kind"] == "missing_field"]
    mat = [d for x in q for d in (x["payload"].get("dropped_entities") or x["payload"].get("items") or [])]
    no_node = not _live("모터", "Component") and not _live("B102 이송 유닛")
    run_document(env1)
    res = (LG.read(did) or {}).get("result") or {}
    got = _live("노칭::모터", "Component") or _live("노칭::B102 이송 유닛")
    unscoped = [n for _l, n in World().nodes() if n["category"] in ("Component", "Unit", "Property")
                and "::" not in n["canonical"] and ops.is_live(n)]
    _oe.append((form, lay, no_node and bool(q) and bool(mat) and bool(got) and not unscoped
                and res.get("엣지 없는 노드") == 0,
                f"좌표 없음 → 노드 0 {no_node} · missing_field {len(q)} · 재료 {len(mat)} → 좌표 주고 다시: "
                f"{[n['canonical'] for n in got][:2]} · 엣지 없는 노드 {res.get('엣지 없는 노드')}"))
    if form == "표":
        dts = store.read(store.DOC_TYPES, {})
        dts.pop(dt, None)
        store.write(store.DOC_TYPES, dts)
        _P.schemas(f"{dt}.json").unlink(missing_ok=True)
show("ⓔ 좌표 없는 스코프 개체 0 — 노드를 만들지 않고 재료는 missing_field가 든다 · 좌표를 주고 다시 넣으면 붙는다 · "
     "스코프 카테고리에 좌표 접두 없는 노드 0 · 엣지 없는 노드 0 (표·산문 × 층 둘)",
     all(o for _f, _l, o, _d in _oe), " ‖ ".join(f"{f}×{l} {d}" for f, l, _o, d in _oe))
show("ⓔ 폴백 한 함수 — 소속 대상으로 앉은 유닛도 좌표에 닿는다 · 엣지 없는 노드 0 (엑셀·Word × 층 둘)",
     all(o for _f, _l, o, _d in _ol), " ‖ ".join(f"{f}×{l} {d}" for f, l, _o, d in _ol))

# ────────────────────────────────────────────────────────────── ⓕ
print("\n■ B102 ⑥ 추출 재사용 조건 — 지시문 판 · 층 config 판")
_setup()
_of = []
for lay in ("process", "quality"):
    did = f"B102RU{lay[0].upper()}"
    _made.append(did)
    env = _prose_env("엑셀", lay, did)
    cfg = load_config(lay)
    _pv = EX.prompt_version
    EX.prompt_version = lambda name="extract": "e-1.3"
    try:
        run_document(env)
    finally:
        EX.prompt_version = _pv
    r1 = EX.reuse_check(env, cfg=cfg)
    notes = []
    run_document(env, notice=notes.append)
    r2 = EX.reuse_check(env, cfg=cfg)
    cfg2 = dict(cfg, config_version=f"{cfg.get('config_version')}-b102")
    r3 = EX.reuse_check(env, cfg=cfg2)
    why = [n.get("다시") for n in notes if n.get("단계") == "추출예고"]
    _of.append((lay, r1 == (False, "지시문 e-1.3 → e-1.4") and r2[0] and not r3[0] and r3[1].startswith("층 config ")
                and why and why[0] == "지시문 e-1.3 → e-1.4", f"{r1} · 다시 뽑은 뒤 {r2} · config 판 바뀜 {r3} · 예고 {why[:1]}"))
show("ⓕ 지시문 판이 바뀌면 다시 뽑는다(예고 줄에 이유) · 다시 뽑은 뒤 재사용 · 층 config 판이 바뀌어도 다시 (층 둘 · "
     "표는 추출이 없다)", all(o for _l, o, _d in _of), " ‖ ".join(f"{l} {d}" for l, _o, d in _of))

# ────────────────────────────────────────────────────────────── ⓖ
print("\n■ B102 ⑦ 붙은 곳을 말한다 — 대장 · 화면 (인입 명령 그대로)")
from g6_common import _register, _unregister, _run, DT   # noqa: E402
_setup()
_register("process", DT)
xl = FILES["엑셀"]
did = IG.doc_id_of(xl)
_made.append(did)
_hint(did, {c["source_locator"]: {"entities": [{"surface": "모터", "category": "Component",
                                               "belongs_to": {"name": _unit_of(c), "category": "Unit", "from": "시트"}}],
                                  "relations": [], "attach": []}
            for c in _parse("엑셀", did, doc_coord="노칭")["chunks"] if "모터" in c.get("text", "")})
out = _run("ingest-file", str(xl), "--doc-type", DT, "--allow-mock", "--coord", "노칭",
           "--sheets", "1-2:prose")
rows = [r for r in (LG.read(did) or {}).get("rows") or [] if r.get("role") == "entity"]
att = [a for r in rows for a in r.get("attached") or []]
_ok_g1 = (rows and all(r.get("target") and r.get("from") == "시트" for r in rows)
          and any(a.get("path", "").startswith("소속") and a.get("rel") == "part_of" for a in att))
show("ⓖ 대장 — target(붙은 노드) · attached(관계 · 상대 · 경로) · from", _ok_g1,
     json.dumps([{k: r.get(k) for k in ("surface", "target", "from", "attached")} for r in rows][:2], ensure_ascii=False))
_need = ["문서 좌표 — 노칭 (출처: 사람)", "   좌표 — 조각 ", "부착 — 소속(", "     └ 모터 · part_of"]
show("ⓖ 인입 화면 — 문서 좌표 줄 · 좌표 줄 · 「└ 관계 → 상대」 · 끝 요약 부착(소속 없음 비율은 없는 것이 있을 때 — ⓒ)",
     all(x in out for x in _need), " / ".join(l.strip() for l in out.splitlines()
                                              if any(x.strip() in l for x in _need))[:600])
_tr = _run("show", "report", did, "--trace")
_ex = _run("show", "extract", did, "--full")
show("ⓖ show report --trace(표기 → 붙은 노드 · 관계 → 상대) · show extract --full(맥락 줄 · 소속)",
     "노칭::이송 유닛::모터" in _tr and "part_of" in _tr and "[문서: " in _ex and "소속" in _ex and "이송 유닛" in _ex,
     " / ".join([l.strip() for l in _tr.splitlines() if "모터" in l][:2] + [l.strip() for l in _ex.splitlines() if "[문서" in l][:1]))
_unregister(DT)
# 표 진행 줄 분모 = 실제 값 수(빈 칸은 세지 않는다 — 행 × 2 고정이 아니다)
_cpenv = load("CP01")
_cpenv["records"] = [dict(r, 관리항목=None) if i % 2 else r for i, r in enumerate(_cpenv["records"][:6])]
_v = IG._value_count(_cpenv, {"doc_type": "cp"})
_ent = sum(1 for r in _cpenv["records"] for f in ("설비", "관리항목") if r.get(f))
show("ⓖ 표 진행 줄 분모 = 판정 계획의 값 수(빈 칸 제외 — 행 × 2 고정 아님)", 0 < _v <= _ent < len(_cpenv["records"]) * 2,
     f"분모 {_v} · 값 있는 칸 {_ent} · 행×2 {len(_cpenv['records']) * 2}")

# ────────────────────────────────────────────────────────────── ⓗ
print("\n■ B102 ⑧ 보존물 · 동명")
_kp = struct_map.keep_path("B102KEEP")
_P.ensure(_kp)
_kp.write_text(json.dumps({"doc_id": "B102KEEP", "rules": {"x": 1}}, ensure_ascii=False), encoding="utf-8")
_h1 = struct_map.load_kept("B102KEEP", "abc")
_kp.write_text(json.dumps({"doc_id": "B102KEEP", "source_hash": "abc", "rules": {"x": 1}}), encoding="utf-8")
_h2 = struct_map.load_kept("B102KEEP", "abc")
_h3 = struct_map.load_kept("B102KEEP", "zzz")
_kp.unlink(missing_ok=True)
show("ⓗ 원본 해시 없는 보존물은 재사용 0 · 같은 해시면 재사용 · 다르면 0", _h1 is None and _h2 and _h3 is None,
     f"해시 없음 {_h1} · 같음 {bool(_h2)} · 다름 {_h3}")
_setup()
_register("process", DT)
a_dir, b_dir = TMP / "a", TMP / "b"
fa, fb = _xlsx(a_dir / "B102SAME.xlsx"), _xlsx(b_dir / "B102SAME.xlsx")
_made.append("B102SAME")
o1 = _run("ingest-file", str(fa), "--doc-type", DT, "--allow-mock", "--sheets", "1-2:prose")
g0 = sorted(n["id"] for _l, n in World().nodes())
o2 = _run("ingest-file", str(fb), "--doc-type", DT, "--allow-mock", "--sheets", "1-2:prose")
g1 = sorted(n["id"] for _l, n in World().nodes())
src2 = (store.read(store.DOC_REGISTRY, {}).get("B102SAME") or {}).get("source_path")
o3 = _run("ingest-file", str(fb), "--doc-type", DT, "--allow-mock", "--sheets", "1-2:prose", "--revise")
src3 = (store.read(store.DOC_REGISTRY, {}).get("B102SAME") or {}).get("source_path")
_unregister(DT)
show("ⓗ 같은 doc_id · 다른 원본 경로 — 상태 거부 + 다음 줄(이름 변경 · --revise) · 그래프·등록 그대로 / --revise면 개정",
     "[상태] 같은 doc_id(B102SAME)" in o2 and "--revise" in o2 and Path(str(src2)).parent.name == "a"
     and g1 == g0 and "개정 — " in o3 and Path(str(src3)).parent.name == "b",
     " / ".join([l.strip() for l in o2.splitlines() if "[상태]" in l or "--revise" in l][:3])
     + f" · 등록 경로 {src2} → {src3}")

# ── 정리 — 표본·힌트·등록은 남기지 않는다
for d in set(_made):
    (HINTS / f"{d}.json").unlink(missing_ok=True)
import shutil                                       # noqa: E402
shutil.rmtree(TMP, ignore_errors=True)
done()
