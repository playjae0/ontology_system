# -*- coding: utf-8 -*-
"""G6 ㉑ 공유 별칭의 범위 안 해소 · 질의 하이브리드 — 임베딩 후보 + LLM 선별 링킹 · 문서 검색 채널 · 짧은 답과
링크 결과 · 채점 축 (B104).

창작 표본(mock 그래프 + 시험이 만든 엑셀 · 시험 주입)으로 메커니즘을 잰다 — 수치는 근거가 아니다([정정] 50).
실호출 갈래는 **가짜 게이트웨이**(`tests/fake_gateway.py`)로 잰다: 임베딩은 동의어 표를 아는 결정적 벡터(전송
`_post` 대역) · 선별·답변은 결정적 규칙(`chat` 대역). 형태(표·산문) × 층(process · quality)에서 같은 성질을 잰다(CLAUDE.md §4).

잠그는 성질:
  ⓐ 공유 별칭: 문서 좌표 범위 안에서 하나인 공유 별칭 시트·행의 조각 좌표 = 그 노드 canonical(원 표기 meta) ·
     구축 보류 0 · 그 시트의 개체가 그 노드 아래 · 범위 안에서도 여럿이면 지금처럼(산문은 건너뜀 · 표는
     보류) · 보류 문면 「표기 모호(골격 노드 여럿)」 + 다음 줄 `--coord`
  ⓑ 링킹: 사전에 없는 동의 표현이 임베딩 후보 + 선별로 노드에 닿는다 · 후보 밖 id 0 · LLM 후보 수 ≤ k(노드 전부
     아님) · 보충은 사전 노드 유지 + 추가 · 폴백은 사전이 잡으면 선별 호출 0 · 미스율은 사전 단 기준
  ⓒ 문서 채널: 링킹 0인 질문이 문서 검색 청크로 답(「근거 없음」 아님) · 두 채널이 다 비면 「근거 없음」 ·
     캐시 재사용(두 번째 질의 벡터 재계산 0) · 청크 추가 시 그 몫만
  ⓓ 답: 5줄 이내 표시 · used_facts·used_chunks 표시 · CLI가 ⓐ~ⓕ를 나란히 · `--json`과 CLI가 같은 수
  ⓔ 채점: 링킹 사전만 vs 보충 · 문서 검색 doc@k가 BM25 대조군 옆에
  ⓕ 쓰기 0(③진실·②등록 해시 불변 — 캐시는 `work/vectors/`)
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import re
import tempfile
from pathlib import Path

sys_path = __import__("sys").path
sys_path.insert(0, str(Path(__file__).resolve().parent))
sys_path.insert(0, str(Path(__file__).resolve().parent / "fixtures"))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, done, overlay   # noqa: F401

from openpyxl import Workbook                           # noqa: E402
from cli import golden as GD, query as CQ               # noqa: E402
from cli.platform import orphan_next_lines              # noqa: E402
from core.build import ledger as LG                     # noqa: E402
from core.build.result import hold_reason               # noqa: E402
from core.query import vectors as V                     # noqa: E402
from core.state import knobs as KB                      # noqa: E402
from parser import pipeline, struct_map, tagger as TG   # noqa: E402
from parser.adapters import basic_prose_xlsx as BX      # noqa: E402
from fake_gateway import CALLS, Live                    # noqa: E402 — 실호출 갈래의 대역(전송·채팅·설정만)

HINTS = ROOT / "tests" / "fixtures" / "extract_hints"
TMP = Path(tempfile.mkdtemp(prefix="b104_"))
_made = []


def _state_hash():
    h = hashlib.sha256()
    for d in (_P.data(), _P.registry()):
        for p in sorted(d.rglob("*")) if d.exists() else []:
            if p.is_file() and not p.name.endswith(".lock"):
                h.update(p.relative_to(d).as_posix().encode("utf-8"))
                h.update(p.read_bytes())
    return h.hexdigest()


def _knobs(d):
    if d:
        _P.knobs().write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    else:
        _P.knobs().unlink(missing_ok=True)
    KB.apply()


# ────────────────────────────────────────────────────────────── ⓐ
print("\n■ B104 ① 공유 별칭의 범위 안 해소 (엑셀 산문 · 표 × 층 둘)")


def _xlsx(path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in {"비전 검사": ["카메라는 B104 비전 카메라를 쓴다."],
                       "노칭 타발": ["타발 금형은 매 교대 점검한다."]}.items():
        ws = wb.create_sheet(title=name)
        for r, v in enumerate(rows, 1):
            ws.cell(row=r, column=1, value=v)
    wb.save(path)
    return path


XL = _xlsx(TMP / "B104_노칭_사양.xlsx")
ROLES = {"비전 검사": "prose", "노칭 타발": "prose"}


def _prose_env(doc_id, lay, doc_coord):
    struct_map.invalidate(doc_id)
    env = pipeline.parse(BX, doc_id, str(XL), closed_list=TG.closed_list("process"),
                         sheet_roles=ROLES, doc_coord=doc_coord).envelope
    dt = "ppt_" + lay
    env["doc_type"] = dt
    for c in env["chunks"]:
        c["doc_type"] = dt
    hints = {c["source_locator"]: {"entities": [{"surface": "B104 비전 카메라", "category": "Unit"}],
                                   "relations": [], "attach": []}
             for c in env["chunks"] if "카메라" in c.get("text", "")}
    (HINTS / f"{doc_id}.json").write_text(json.dumps(hints, ensure_ascii=False), encoding="utf-8")
    _made.append(doc_id)
    return env


def _table_env(doc_id, lay, ref, doc_coord):
    dt = f"b104t{lay[0]}"
    p = _P.schemas(f"{dt}.json")
    _P.ensure(p)
    p.write_text(json.dumps({"doc_type": dt, "schema_version": 1, "layer": lay,
                             "use_blocks": ["common_core", "process_coord"],
                             "fields": {"개체": {"role": "entity", "category": "Unit"}}, "edges": []},
                            ensure_ascii=False), encoding="utf-8")
    dts = store.read(store.DOC_TYPES, {})
    dts[dt] = {"doc_type": dt, "status": "registered", "layer": lay, "schema": f"schemas/{dt}.json",
               "adapter": "-", "schema_version": 1}
    store.write(store.DOC_TYPES, dts)
    base = {k: v for k, v in CPREC.items() if k not in ("설비", "관리항목")}
    recs = [dict(base, source_locator="T1", doc_type=dt, process_ref=ref, process_group=None,
                 개체="B104 비전 카메라")]
    # 파서와 같은 함수(좌표 태깅 — 문서 좌표 범위 안 해소)를 지난 행이 구축으로 간다
    recs = TG.tag(recs, nodes=TG.closed_list("process"), doc_coord=doc_coord)
    return {**TABLE, "doc_id": doc_id, "doc_type": dt, "records": recs}


_oa, _ob = [], []
for form, lay in (("산문", "process"), ("산문", "quality"), ("표", "process"), ("표", "quality")):
    overlay()                                         # 품질층이 Unit을 렌즈로 쓴다(층 둘)
    did = f"B104S{form[0] == '표' and 'T' or 'P'}{lay[0].upper()}"
    env = _prose_env(did, lay, "노칭") if form == "산문" else _table_env(did, lay, "비전 검사", "노칭")
    pcs = env.get("chunks") or env.get("records")
    one = [p for p in pcs if (p.get("meta") or {}).get("coord_tag_from") == "비전 검사"]
    run_document(env)
    rows = (LG.read(did) or {}).get("rows") or []
    holds = [x for x in store.read(store.QUEUE, []) if x["doc_id"] == did and x["kind"] == "orphan_anchor"]
    ents = [r for r in rows if r.get("role") == "entity" and r.get("surface") == "B104 비전 카메라"]
    _oa.append((form, lay, one and all(p["process_ref"] == "노칭::비전검사" for p in one) and not holds
                and ents and all((r.get("target") or "").startswith("노칭::비전검사::") for r in ents)
                and not any(r.get("hold") for r in rows),
                f"좌표 {sorted({p['process_ref'] for p in one})}(원 표기 비전 검사) · 보류 {len(holds)} · "
                f"개체 → {sorted({r.get('target') for r in ents})}"))
    # 범위 안에서도 여럿 — 문서 좌표 「조립」(노칭·스태킹이 다 그 아래)
    did2 = did + "M"
    env2 = _prose_env(did2, lay, "조립") if form == "산문" else _table_env(did2, lay, "비전 검사", "조립")
    pcs2 = env2.get("chunks") or env2.get("records")
    run_document(env2, allow_duplicate=True)
    rows2 = (LG.read(did2) or {}).get("rows") or []
    q2 = [x for x in store.read(store.QUEUE, []) if x["doc_id"] == did2 and x["kind"] == "orphan_anchor"]
    if form == "산문":
        sk = [p for p in pcs2 if "비전 검사" in ((p.get("meta") or {}).get("coord_shared_skip") or [])]
        ok2 = sk and all(p["process_ref"] == "조립" for p in sk) and not q2
        det = f"건너뜀 {len(sk)} · 문서 좌표 물려받음 {sorted({p['process_ref'] for p in sk})}"
    else:
        hr = [hold_reason(r) for r in rows2 if r.get("hold")]
        nxt = orphan_next_lines(q2[0]) if q2 else ""
        ok2 = (pcs2[0]["process_ref"] == "비전 검사" and q2 and "표기 모호" in q2[0]["reason"]
               and hr and all(h == "표기 모호(골격 노드 여럿)" for h in hr) and "--coord" in nxt)
        det = f"보류 {q2[0]['reason'] if q2 else '-'} · 대장 {sorted(set(hr))} · 다음 줄 --coord {'--coord' in nxt}"
    _ob.append((form, lay, ok2, det))
    if form == "표":
        dts = store.read(store.DOC_TYPES, {})
        dts.pop(f"b104t{lay[0]}", None)
        store.write(store.DOC_TYPES, dts)
        _P.schemas(f"b104t{lay[0]}.json").unlink(missing_ok=True)
show("ⓐ 범위 안에서 하나인 공유 별칭 — 조각 좌표가 그 노드 canonical(원 표기 meta) · 보류 0 · 개체가 그 노드 아래 "
     "(엑셀 산문 · 표 × 층 둘)", all(o for _f, _l, o, _d in _oa), " ‖ ".join(f"{f}×{l} {d}" for f, l, _o, d in _oa))
show("ⓐ 범위 안에서도 여럿이면 지금처럼 — 산문은 건너뛰고 문서 좌표 · 표는 보류이되 문면 「표기 모호(골격 노드 여럿)」 "
     "+ 다음 줄 --coord (층 둘)", all(o for _f, _l, o, _d in _ob), " ‖ ".join(f"{f}×{l} {d}" for f, l, _o, d in _ob))

# ────────────────────────────────────────────────────────────── ⓑ
print("\n■ B104 ② 링킹 — 임베딩 후보 + LLM 선별 (가짜 게이트웨이)")
with contextlib.redirect_stdout(io.StringIO()):
    fresh()                                           # mock 그래프(CP01·PFMEA01·PPT01~03·QPPT01)
_h0 = _state_hash()
_knobs({"query_link_top_k": 5})
_nall = sum(1 for lay in ("process", "quality") for n in open_graph(lay).nodes.values() if ops.is_live(n))
_lm = store.path(store.LINK_MISS)
_lm0 = _lm.read_text(encoding="utf-8").count("\n") if _lm.exists() else 0
with Live():
    r1 = CQ.answer("틈새는 어떻게 관리해?")
    c1 = [b for p, b in CALLS if p == "link"]
    r2 = CQ.answer("노칭 틈새는 어떻게 관리해?")
    _knobs({"query_link_top_k": 5, "query_link_mode": "폴백"})
    del CALLS[:]
    r3 = CQ.answer("노칭 틈새는 어떻게 관리해?")
    c3 = [p for p, _b in CALLS if p == "link"]
_knobs({})
_lm1 = _lm.read_text(encoding="utf-8").count("\n") if _lm.exists() else 0
L1 = r1["trace"]["linking"]
show("ⓑ 사전에 없는 동의 표현(「틈새」)이 임베딩 후보 + 선별로 노드에 닿는다 — method embed+llm · 점수 · 이유 · 어디의 무엇",
     [x["canonical"] for x in L1] == ["노칭::금형 클리어런스"] and L1[0]["method"] == "embed+llm"
     and L1[0].get("why") and L1[0].get("score") is not None and L1[0].get("where"),
     json.dumps(L1, ensure_ascii=False)[:260])
show("ⓑ 후보 밖 id 0 · LLM이 받은 후보 수 ≤ k(5) < 노드 전부 — 노드 전부를 보내지 않는다",
     not any("NOPE" in x["node_id"] for x in L1) and c1 and len(c1[0]["candidates"]) == 5 < _nall
     and len(r1["trace"]["link_stage"]["candidates"]) == 5,
     f"후보 {len(c1[0]['candidates']) if c1 else '-'} / 노드 {_nall} · 고른 {r1['trace']['link_stage'].get('picked')}")
m2 = {x["canonical"]: x["method"] for x in r2["trace"]["linking"]}
m3 = {x["canonical"]: x["method"] for x in r3["trace"]["linking"]}
show("ⓑ 보충(기본) — 사전 노드는 그대로 + 선별이 더한다 · 폴백 — 사전이 잡으면 선별 호출 0(지금 규칙)",
     m2.get("노칭") == "dict" and m2.get("노칭::금형 클리어런스") == "embed+llm"
     and m3 == {k: v for k, v in m2.items() if v == "dict"} and not c3,
     f"보충 {m2} · 폴백 {m3} · 폴백 선별 호출 {len(c3)}")
show("ⓑ 링킹 미스율은 사전 단 기준 — 선별이 찾았어도 사전이 못 찾은 질문은 link_miss에 적재(사전이 잡은 질문은 0)",
     _lm1 - _lm0 == 1 and r1["trace"]["miss"] == ["틈새는 어떻게 관리해?"] and not r2["trace"]["miss"],
     f"link_miss +{_lm1 - _lm0} · trace.miss {r1['trace']['miss']} / {r2['trace']['miss']}")

# ────────────────────────────────────────────────────────────── ⓒ
print("\n■ B104 ③ 문서 검색 채널 · 벡터 캐시")
_r = CQ.answer("육안 점검 교체 주기는?")
_g = CQ.answer("리튬이온 배터리 동작 원리 설명해줘")
show("ⓒ 링킹 0인 질문이 문서 검색 청크로 답한다(「근거 없음」 아님 · path chunk) · 두 채널이 다 비면 「근거 없음」",
     not _r["trace"]["linking"] and _r["path"] == "chunk" and _r["doc_search"] and not _r.get("note")
     and _g["path"] == "general_knowledge" and not _g["doc_search"] and "근거를 찾지 못했다" in (_g["note"] or ""),
     f"{_r['path']} · 문서 검색 {len(_r['doc_search'])} {[(c['doc_id'], c['by']) for c in _r['doc_search'][:3]]} · "
     f"리튬이온 → {_g['path']}")
with Live():
    _rl = CQ.answer("틈새는 어떻게 관리해?")
_emb = [c for c in _rl["doc_search"] if "embed" in c["by"]]
show("ⓒ 실호출 갈래 — 임베딩이 동의 표현으로 청크를 찾는다(문턱 이상 · BM25와 합쳐 중복 제거)",
     _emb and any("클리어런스" in c["text"] for c in _emb)
     and len({c["chunk_id"] for c in _rl["doc_search"]}) == len(_rl["doc_search"]),
     f"임베딩 {len(_emb)} · BM25 {sum(1 for c in _rl['doc_search'] if 'bm25' in c['by'])} · "
     f"{[(c['doc_id'], c['embed']) for c in _emb[:3]]}")
CQ.answer("육안 점검 교체 주기는?")              # 판이 바뀌었다(가짜 게이트웨이 → mock) — 통째로 다시
_s0 = {k: dict(v) for k, v in V.STATS.items()}
CQ.answer("육안 점검 교체 주기는?")
_s1 = {k: dict(v) for k, v in V.STATS.items()}
_extra = {**PROSE, "doc_id": "B104ADD", "source_path": "B104ADD.pptx", "doc_type": "ppt_quality",
          "chunks": [{**C1, "source_locator": "B104ADD-C001", "text": "육안 점검 기준은 사진 견본과 대조한다."},
                     {**C1, "source_locator": "B104ADD-C002", "text": "교체 주기는 타발 횟수로 정한다."}]}
(HINTS / "B104ADD.json").write_text("{}", encoding="utf-8")
_made.append("B104ADD")
run_document(_extra)
CQ.answer("육안 점검 교체 주기는?")
_s2 = {k: dict(v) for k, v in V.STATS.items()}
show("ⓒ 캐시 재사용 — 두 번째 질의는 벡터 재계산 0 · 청크를 더하면 그 몫만 다시(노드 그대로면 0) · "
     "판(갈래·모델)이 바뀌면 통째로 다시",
     _s1["chunks"]["made"] == 0 and _s1["nodes"]["made"] == 0 and _s2["chunks"]["made"] == 2
     and _s2["chunks"]["reused"] == _s1["chunks"]["reused"] and _s2["nodes"]["made"] == 0
     and _s0["chunks"]["reused"] == 0 == _s0["nodes"]["reused"] and V.cost_line().startswith("벡터 — "),
     f"판 바뀐 뒤 {_s0['chunks']['made']}/{_s0['nodes']['made']} 다시 · 두 번째 {_s1['chunks']} · "
     f"청크 2 추가 뒤 {_s2['chunks']} · {V.cost_line()}")

# ────────────────────────────────────────────────────────────── ⓓ
print("\n■ B104 ④ 답 — 짧게 · 쓴 근거 · 링크 결과와 나란히")
with Live():
    _out = io.StringIO()
    with contextlib.redirect_stdout(_out):
        CQ.main(["노칭 틈새는 어떻게 관리해?"])
    _txt = _out.getvalue()
    _jo = io.StringIO()
    with contextlib.redirect_stdout(_jo), contextlib.redirect_stderr(io.StringIO()):
        CQ.main(["노칭 틈새는 어떻게 관리해?", "--json"])
_j = json.loads(_jo.getvalue())
_jt = _j["trace"]
_need = ["ⓐ 링킹", "[링킹]", "[링킹 추가]", "ⓑ 확장", "[그래프 사실]", "[문서 근거]", "ⓔ 문서 검색",
         "[문서 검색]", "ⓕ 답 — LLM(live)"]
_cnt = lambda t, k: sum(1 for l in t.splitlines() if l.strip().startswith(k))   # noqa: E731
show("ⓓ CLI가 링크 결과(ⓐ~ⓔ)와 LLM 답(ⓕ)을 나란히 낸다",
     all(k in _txt for k in _need), " / ".join(l.strip()[:50] for l in _txt.splitlines()
                                                if any(l.strip().startswith(k) for k in ("ⓐ", "ⓑ", "ⓔ", "ⓕ"))))
_ds = [c for c in _j["doc_search"] if not c.get("in_graph")]
_tick = lambda k: _cnt(_txt, k + " ✓")                                          # noqa: E731
show("ⓓ `--json`과 CLI가 같은 수 — 사실(trace 전부 · ✓ = 답이 쓴 묶음 facts) · 노드 근거 · 문서 검색 · 링킹 추가 · "
     "쓴 청크 ✓ = used",
     _cnt(_txt, "[그래프 사실]") == len(_jt["facts"]) and _tick("[그래프 사실]") == len(_j["facts"])
     and _cnt(_txt, "[문서 근거]") == len(_j["chunks"]) and _cnt(_txt, "[문서 검색]") == len(_ds)
     and _cnt(_txt, "[링킹 추가]") == sum(1 for x in _jt["linking"] if x["method"] == "embed+llm")
     and _tick("[문서 근거]") + _tick("[관련 원문]") + _tick("[문서 검색]") == len(_jt["answer"]["used_chunks"]),
     f"사실 {_cnt(_txt, '[그래프 사실]')}={len(_jt['facts'])}(✓ {_tick('[그래프 사실]')}={len(_j['facts'])}) · "
     f"노드 근거 {_cnt(_txt, '[문서 근거]')}={len(_j['chunks'])} · 문서 검색 {_cnt(_txt, '[문서 검색]')}={len(_ds)} · "
     f"쓴 청크 ✓ {_tick('[문서 근거]') + _tick('[관련 원문]') + _tick('[문서 검색]')}")
_a = _jt["answer"]
_used_c = [c for c in _jt["collection"] if c.get("used")] + [c for c in _jt["doc_search"] if c.get("used")]
show("ⓓ 답 — 5줄 이내(줄 수 표시) · used_facts·used_chunks가 사실·청크에 표시된다 · 지시문 a-1.2 · l-1.1",
     _a["mode"] == "live" and _a["lines"] == 3 <= CQ.ANSWER_LINES and _a["used_chunks"] == [0]
     and len(_used_c) == 1 and sum(1 for f in _jt["facts"] if f["used"]) >= 1
     and "5줄" in (ROOT / "prompts/4.4_answer.md").read_text(encoding="utf-8")
     and re.search(r"^version: a-1\.2", (ROOT / "prompts/4.4_answer.md").read_text(encoding="utf-8"), re.M)
     and re.search(r"^version: l-1\.1", (ROOT / "prompts/4.1_link.md").read_text(encoding="utf-8"), re.M),
     f"{_a['lines']}줄 · 쓴 청크 {_a['used_chunks']} → 표시 {[c['chunk_id'][-10:] for c in _used_c]}")

# ────────────────────────────────────────────────────────────── ⓔ
print("\n■ B104 ⑤ 채점 — 링킹 사전만 vs 보충 · 문서 검색 doc@k (BM25 대조군 옆)")
_qs = [{"id": "B104-1", "type": "Q1", "q": "틈새는 어떻게 관리해?", "expected_path": "chunk",
        "expected_linked": ["process:노칭::금형 클리어런스"], "expected_docs": ["CP01"]},
       {"id": "B104-2", "type": "Q1", "q": "육안 점검 교체 주기는?", "expected_path": "chunk",
        "expected_docs": ["PPT03"]}]
with Live():
    _rows = GD.score_set(_qs)
_agg = GD.aggregate(_rows)
_ren = GD.render(_agg, _rows, src="B104", is_mock=True, skipped=[], k=8)
show("ⓔ 링킹 두 칸(사전만 0 → 보충 1) · 문서 검색 doc@k가 근거@k · BM25@k 옆에 계산된다",
     _agg["linking_dict_recall"] == 0.0 and _agg["linking_recall"] == 1.0 and _agg["doc_n"] == 2
     and _agg["doc_at_k"] is not None and _agg["bm25_at_k"] is not None
     and "link사전" in _ren and "link보충" in _ren and "doc@k" in _ren and "bm25@k" in _ren,
     f"링킹 사전 {_agg['linking_dict_recall']} → 보충 {_agg['linking_recall']} · doc@k {_agg['doc_at_k']} · "
     f"evid@k {_agg['evidence_at_k']} · bm25@k {_agg['bm25_at_k']}")

# ────────────────────────────────────────────────────────────── ⓕ
print("\n■ B104 ⑥ 쓰기 0")
show("ⓕ 질의는 ③진실·②등록에 쓰지 않는다(해시 불변 — 캐시는 work/vectors) — 시험이 더한 문서(B104ADD)는 제외",
     _P.vectors("chunks.json").exists() and _P.vectors("nodes.json").exists()
     and not (_P.data() / "vectors").exists(), str(sorted(p.name for p in _P.vectors().iterdir())))
_hq = _state_hash()
with Live():
    CQ.answer("노칭 틈새는 어떻게 관리해?")
CQ.answer("육안 점검 교체 주기는?")
show("ⓕ 질의 전후 ③진실·②등록 해시 같다 (mock · 가짜 게이트웨이 둘 다)", _state_hash() == _hq)

for d in set(_made):
    (HINTS / f"{d}.json").unlink(missing_ok=True)
import shutil                                       # noqa: E402
shutil.rmtree(TMP, ignore_errors=True)
done()
