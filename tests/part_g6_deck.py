# -*- coding: utf-8 -*-
"""G6 — **PPT: 슬라이드가 추출 단위 · 캡처를 추출에 · 덱 구조 1회 · 청크 공정은 목록** (B125).

대표 사례 하나씩(가벼운 회차 · 창작 표본 `make_ppt.build_report` — 수치는 근거가 아니다 [정정] 50 · 덱 구조 · 추출은 가짜
게이트웨이 — 지시문 · 조립 · 검사 · 구축은 실물):
  ⓐ 상자 — 글 상자 제목 + 작은 상자 여섯 슬라이드: 본문 근거 청크 1 · 추출 호출 1 · 상자 분할 0
  ⓑ 같은 슬라이드 — 본문 + 표 + 노트: 추출 입력 하나에 셋 다 · 표 위치 `S4-TB1` 그대로 · 그림 요약에만 있는 값은 버림
  ⓒ 캡처 — 그림 없는 슬라이드에도 추출 입력에 캡처(`capture=page`) · 렌더가 없으면 `none` + 화면 수 · 인입은 선다
  ⓓ 덱 구조 — 글 상자 제목 → 그 글이 제목 · 「(계속)」 → 앞 슬라이드의 경로 · 공정 · 덱 개요가 모든 추출 맥락 줄에 · 골격 밖 공정 버림 셈
  ⓔ 다공정 — 노칭 · 스태킹 슬라이드: 개체가 각자 제 공정 아래 · 언급만 한 공정은 노드 0 + 관련 링크 · `--coord` 없이 `ingest-dir` 성공
  ⓕ 고침 · 재생 — 공정 고침 → 기록 → 다시 파싱이 그대로(덱 구조 LLM 0)
  B151 ⓐ~ⓕ 덱 판단은 층 · 골격 · 문서 좌표를 모른다(두 상태의 입력이 같다 · 층 config를 바꿔도 재사용) · 다루는 것 · 공정 표기 검사 ·
       공정 걸기는 표와 같은 길(정확 · ⑨ 표기마다 한 번 · 목록 원소마다 · 못 맞으면 그대로 → 큐) · 공정 없는 덱 · 옛 언급 고침은 읽지 않는다
네 벌(A · B 같음 · C 바뀐 몫)은 마감 한 줄의 네 벌 단계가 잰다.
"""
from __future__ import annotations

import contextlib
import io
import json
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from g65_common import *          # noqa: F401,F403 — 바닥은 하나다
from g65_common import _P, _iso, done   # noqa: F401

from fake_gateway import CALLS, Live                      # noqa: E402
from g6_common import SCHEMA                              # noqa: E402
from fixtures.make_ppt import build_report               # noqa: E402
import run as RUN                                        # noqa: E402
from cli import structure_screen as STS                  # noqa: E402
from core.build import deck as DK, extract as EX, extract_unit as EU, ledger as LG, structure as ST, values as VL  # noqa: E402
from parser import pipeline, render, struct_map, tagger as TG   # noqa: E402
from parser.adapters import basic_ppt as PX              # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="b125_"))
DT, DID = "b125_ppt", "B125D"
PATH = build_report(TMP / f"{DID}.pptx")


def _quiet():
    return contextlib.redirect_stdout(io.StringIO())


def _register():
    p = _P.schemas(f"{DT}.json")
    _P.ensure(p)
    p.write_text(json.dumps({**SCHEMA, "doc_type": DT, "layer": "process", "payload_kind": "prose"}, ensure_ascii=False),
                 encoding="utf-8")
    ad = _P.adapters(f"{DT}.py")
    _P.ensure(ad)
    ad.write_text("from parser.adapters import basic_ppt\n"
                  f"ADAPTER = {{**basic_ppt.ADAPTER, \"doc_type\": {DT!r}}}\nextract = basic_ppt.extract\n", encoding="utf-8")
    dts = store.read(store.DOC_TYPES, {})
    dts[DT] = {"doc_type": DT, "status": "registered", "layer": "process", "schema": f"schemas/{DT}.json",
               "adapter": str(ad), "schema_version": 1}
    store.write(store.DOC_TYPES, dts)


# ── 가짜 LLM의 답(창작) — 덱 구조 · 추출 · 그림 요약
# 기대 변경(B151 ① — 덱 판단의 칸: 언급 공정 0 · 다루는 것 `subject` · 공정은 문서 표기)
DECK = {"slides": [
    {"slide": 1, "title": "노칭 · 스태킹 점검 보고", "section": None, "continues": None, "subject": None, "processes": ["없는공정"],
     "no_entities": True},
    {"slide": 2, "title": "노칭 금형 점검 결과", "section": None, "continues": None, "subject": "노칭 금형", "processes": ["노칭"],
     "no_entities": False},
    {"slide": 3, "title": "노칭 금형 점검 결과 (계속)", "section": None, "continues": 2, "subject": None, "processes": [],
     "no_entities": False},
    {"slide": 4, "title": "커터 사양", "section": None, "continues": None, "subject": "커터 사양", "processes": ["노칭"],
     "no_entities": False},
    {"slide": 5, "title": "노칭 · 스태킹 연동 점검", "section": None, "continues": None, "subject": "노칭 · 스태킹 연동 점검",
     "processes": ["노칭", "스태킹"], "no_entities": False}],
    "overview": "노칭 금형과 스태킹 연동을 점검한 보고서다."}


def _ent(surface, cat, parent=None, value=None):
    return {"surface": surface, "category": cat, "parent": parent, "belongs_to": None, "value": value}


def _extract(body):
    t = body["chunk"]
    ents = []
    if "마모 0.05mm" in t:
        ents = [_ent("금형", "Unit")]
    elif "다이 균열" in t:
        ents = [_ent("다이", "Unit")]
    elif "아래 표는" in t:
        ents = [_ent("장력", "Property", value="30N"), _ent("압력", "Property", value="5bar"),
                _ent("온도", "Property", value="80도")]
    elif "매거진" in t:
        # 기대 변경(B151 ② — 다른 공정 언급은 덱 판단이 아니라 그래프 단계의 산문 함수가 잡는다: 본문의 「탭용접」이 개체로 뽑히면 가드)
        ents = [_ent("커터", "Unit", parent="노칭"), _ent("매거진", "Unit", parent="스태킹"), _ent("탭용접", "Unit")]
    return {"entities": ents, "relations": [], "about": []}


def _live():
    return Live(by={"deck": lambda b: DECK, "extract": _extract, "outline": lambda b: {"headings": []},
                    "reinforce": lambda b: {"pick": None}})


def _parse(**kw):
    with _live():
        res = pipeline.parse(PX, DID, str(PATH), closed_list=TG.closed_list("process"), layer="process",
                             deck=DK.decker(), capture=True, structure_fixes=ST.fixes_of(DID),
                             summarize=lambda ref, **k: "도면 — 온도 80도", **kw)
        calls = [p for p, _b in CALLS]
    env = res.envelope
    env["doc_type"] = DT
    for c in env["chunks"]:
        c["doc_type"] = DT
    return res, env, calls


def _run(env):
    notes = []
    with _live(), _quiet():
        run_document(dict(env), notice=notes.append)
        calls = list(CALLS)
    return notes, calls


with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register()
_keep = (render.capture_path, render.render_pages)
render.capture_path = lambda path=None, probe=False: (True, "시험", None)      # 기대 변경(B128 ① — 판정 한 자리의 이름 · 답 꼴)
render.render_pages = lambda path: {n: f"PNG{n}".encode() for n in range(1, 6)}   # 보이는 슬라이드 5장의 가짜 렌더
struct_map.invalidate(DID)
RES, ENV, PCALLS = _parse()
NOTES, CALLS1 = _run(ENV)
_by = {c["source_locator"]: c for c in ENV["chunks"]}
_ext = [b for p, b in CALLS1 if p == "extract"]

# ────────────────────────────────────────────────────────────── ⓐ 상자
print("\n■ ⓐ 상자 — 슬라이드 = 추출 1회")
_s2 = [b for b in _ext if "마모 0.05mm" in b["chunk"]]
show("ⓐ 글 상자 제목 + 작은 상자 여섯 → 본문 근거 청크 1 · 추출 호출 1 · 상자 분할 0",
     "슬라이드 2" in _by and not any(k.startswith("슬라이드 2#") for k in _by) and len(_s2) == 1
     and all(x in _s2[0]["chunk"] for x in ("→ 버 발생", "점검 완료")),
     f"슬라이드 2 청크 {[k for k in _by if k.startswith('슬라이드 2')]} · 호출 {len(_s2)}")

# ────────────────────────────────────────────────────────────── ⓑ 같은 슬라이드
print("\n■ ⓑ 같은 슬라이드 — 본문 · 표 · 노트 · 그림 요약이 한 입력")
_s4 = next((b for b in _ext if "아래 표는" in b["chunk"]), {"chunk": ""})
_rows = {r.get("surface"): r for r in (LG.read(DID) or {}).get("rows") or [] if r.get("role") == "entity"}
_v = {k: (_rows.get(k) or {}).get("value", {}).get("verdict") for k in ("장력", "압력", "온도")}
show("ⓑ 추출 입력 하나에 본문 · 표 · 노트 · 표 위치 S4-TB1 그대로 · 표 · 노트의 값은 앉고 그림 요약에만 있는 값은 버린다",
     all(x in _s4["chunk"] for x in ("아래 표는", "[표 S4-TB1]", "30N", "[노트]", "5bar")) and "S4-TB1" in _by
     and _v == {"장력": VL.LANDED, "압력": VL.LANDED, "온도": VL.DROPPED},
     f"값 {_v} · 입력에 표 {'[표 S4-TB1]' in _s4['chunk']} · 노트 {'[노트]' in _s4['chunk']}")

# ────────────────────────────────────────────────────────────── ⓒ 캡처
print("\n■ ⓒ 캡처 — 그림 유무와 무관하게 추출에")
_ck = json.loads(EX.checkpoint_path(DID).read_text(encoding="utf-8"))
_cap = {c["chunk_id"]: c.get("capture") for c in _ck["candidates"] if c.get("capture")}
_pre = next((n for n in NOTES if n.get("단계") == "추출예고"), {})
EX.invalidate(DID)                                   # 렌더가 없는 환경 — 같은 문서를 다시 뽑는다
EU._PAGES.clear()
render.capture_path = lambda path=None, probe=False: (False, "없음", "soffice 없음")
_n2, _c2 = _run(ENV)
_pre2 = next((n for n in _n2 if n.get("단계") == "추출예고"), {})
_ck2 = json.loads(EX.checkpoint_path(DID).read_text(encoding="utf-8"))
render.capture_path, render.render_pages = _keep
show("ⓒ 그림 없는 슬라이드에도 캡처가 추출 입력에(capture=page) · 렌더가 없으면 none + 화면 수(사유) · 인입은 선다",
     _s2 and _s2[0].get("_images") == 1 and set(_cap.values()) == {"page"}
     and (_pre.get("캡처") or {}).get("있음") == 4
     and (_pre2.get("캡처") or {}).get("없음") == 4 and (_pre2["캡처"].get("사유") == "soffice 없음")
     and {c.get("capture") for c in _ck2["candidates"] if c.get("capture")} == {"none"}
     and not any(b.get("_images") for p, b in _c2 if p == "extract"),
     f"캡처 {_pre.get('캡처')} → 렌더 없음 {_pre2.get('캡처')}")

# ────────────────────────────────────────────────────────────── ⓓ 덱 구조
print("\n■ ⓓ 덱 구조 — 제목 · 이어짐 · 개요 · 버림")
_dk = RES.report.get("deck") or {}
show("ⓓ 글 상자 제목이 제목 · (계속)은 앞 슬라이드의 경로 · 공정 · 덱 개요가 모든 추출 맥락 줄에 · 골격 밖 공정 버림 셈",
     _by["슬라이드 2"]["section"] == "노칭 금형 점검 결과"
     and _by["슬라이드 3"]["meta"]["section_path"] == _by["슬라이드 2"]["meta"]["section_path"]
     and _by["슬라이드 3"]["process_ref"] == "노칭"
     and _ext and all("개요: 노칭 금형과 스태킹 연동을 점검한 보고서다." in b["chunk"] for b in _ext)
     and (_dk.get("버림") or {}).get("공정") == 1 and _dk["수"]["이어짐"] == 1
     and (_dk["수"]["공정둘이상"] == 1) and PCALLS.count("deck") == 1
     and any(n.get("개체없음") == 1 for n in NOTES if n.get("단계") == "추출예고"),
     f"제목 {_by['슬라이드 2']['section']!r} · 경로 {_by['슬라이드 3']['meta']['section_path']!r} · 수 {_dk.get('수')} · "
     f"버림 {_dk.get('버림')}")

# ────────────────────────────────────────────────────────────── B143 ⓑ 덱 블록
print("\n■ B143 ⓑ 덱 블록 — 덱 판단이 parsed JSON 한 곳에")
from parser import deck as PDK                           # noqa: E402
_blk = (ENV.get("deck") or {}).get("slides") or []
_vals = {i: list(x.values()) for i, x in enumerate(_blk, 1)}
_bad = PDK.mismatch({**ENV, "chunks": [dict(c, process_ref="스태킹") if c["source_locator"] == "슬라이드 2" else c
                                       for c in ENV["chunks"]]})
# 기대 변경(B151 ① — 언급 칸 0 · 다루는 것 칸)
show("B143 ⓑ 봉투 하나에 슬라이드마다 이어짐 · 제목 초안과 판단 · 공정 목록 · 다루는 것 · 개체 없음 · 버림 · 근거 위치 · 청크 칸과 "
     "어긋나지 않는다(어긋나게 바꾸면 결함으로 잡힌다) · 개요도 블록에(옛 doc_overview 0)",
     len(_blk) == 5 and not PDK.mismatch(ENV) and _bad and "doc_overview" not in ENV
     and 2 in _vals[3] and ["노칭", "스태킹"] in _vals[5] and "노칭 · 스태킹 연동 점검" in _vals[5] and ["공정"] in _vals[1]
     and not any("mentions" in x for x in _blk)
     and "노칭 금형 점검 결과" in _vals[2] and ["슬라이드 4", "S4-TB1", "S4-P1"] in _vals[4]
     and (ENV.get("deck") or {}).get("overview") == DECK["overview"],
     f"슬라이드 {len(_blk)} · 어긋남 {len(PDK.mismatch(ENV))} → 바꾸면 {len(_bad)} · 3번 {_blk[2] if len(_blk) > 2 else None}")

# ────────────────────────────────────────────────────────────── ⓔ 다공정
print("\n■ ⓔ 다공정 — 개체는 제 공정 아래 · 언급은 관련 링크")
_g = open_graph("process")
_names = {n.get("canonical") for n in _g.nodes.values()}
_tab = next(nid for nid, n in _g.nodes.items() if n.get("canonical") == "탭용접")
_ch = store.read(store.CHUNKS, {})
_s5 = next(cid for cid, c in _ch["chunks"].items() if c.get("doc_id") == DID and c.get("source_locator") == "슬라이드 5")
# `--coord` 없이 폴더째 넣기 — mock(덱 구조 힌트) 하위 실행
_dir = TMP / "dir"
_dir.mkdir()
(_dir / f"{DID}.pptx").write_bytes(PATH.read_bytes())
from core.state import fixtures as _FXR                   # noqa: E402 — 픽스처 뿌리(회귀는 스위트마다 사본 — B133)
_hint = _FXR.EXTRACT_HINTS / f"{DID}@deck.json"
_hint.write_text(json.dumps(DECK, ensure_ascii=False), encoding="utf-8")
with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register()


def _sub(*a):
    return subprocess.run([sys.executable, "run.py", *a], cwd=str(ROOT), capture_output=True, text=True,
                          stdin=subprocess.DEVNULL, env={**_iso.isolated_env(), "USE_MOCK": "1"}, timeout=600)


# B143 ⓐ — 공정 목록(슬라이드 5 노칭 · 스태킹)이 든 덱으로 등록 실경로 — 좌표 쌍 표가 목록을 받는다
_rg = _sub("register", "generate", "b143_reg", "process", str(_dir / f"{DID}.pptx"), "--use-basic", "--allow-mock")
_rs = _sub("register", "status", "b143_reg", "--allow-mock")
show("B143 ⓐ 공정 목록이 둘 이상인 슬라이드가 든 덱 — register generate --use-basic → status가 끝까지 돈다(좌표 쌍 표 · 관문 줄)",
     _rg.returncode == 0 and "좌표 쌍 —" in _rg.stdout and "[결함]" not in _rg.stdout + _rg.stderr
     and _rs.returncode == 0 and "기계 관문(하네스): PASS" in _rs.stdout,
     f"generate rc {_rg.returncode} · status rc {_rs.returncode} · "
     + next((ln.strip() for ln in _rg.stdout.splitlines() if "좌표 쌍 —" in ln), (_rg.stdout + _rg.stderr)[-120:]))
# B143 ⓓ · ⓔ — PPT --split-only: 추출 0 · 그래프 쓰기 0 · parsed JSON과 md · md는 봉투만으로 찍은 글
from cli import envelope_md as EMD                       # noqa: E402
_so = _sub("ingest-file", str(_dir / f"{DID}.pptx"), "--doc-type", DT, "--allow-mock", "--split-only")
_md = _P.export(f"split_{DID}.md")
_pj = json.loads(_P.parsed(f"{DID}.json").read_text(encoding="utf-8")) if _P.parsed(f"{DID}.json").exists() else {}
_mdt = _md.read_text(encoding="utf-8") if _md.exists() else ""
import re as _re                                         # noqa: E402
show("B143 ⓓ PPT --split-only — 추출 호출 0 · 그래프 쓰기 0 · parsed JSON과 md가 남는다 · md 슬라이드 표 행 수 = 덱 블록 슬라이드 수",
     _so.returncode == 0 and "구조까지" in _so.stdout and not EX.checkpoint_path(DID).exists()
     and not open_graph("process").nodes.keys() - {n for n, x in open_graph("process").nodes.items() if x.get("status") == "seed"}
     and _pj.get("deck") and len([ln for ln in _mdt.splitlines() if _re.match(r"^\| \d+ \|", ln)]) == len(_pj["deck"]["slides"]),
     f"rc {_so.returncode} · md {_md.exists()} · 표 행 {len([ln for ln in _mdt.splitlines() if _re.match(r'^[|] [0-9]+ [|]', ln)])} · "
     f"덱 블록 {len((_pj.get('deck') or {}).get('slides') or [])}")
show("B143 ⓔ md ⊆ JSON — md 글자 = parsed JSON(봉투) 하나로 찍은 글(md를 쓰는 함수는 봉투 밖을 읽지 않는다)",
     bool(_mdt) and _mdt == "\n".join(EMD.lines(_pj)) + "\n", f"md {len(_mdt)}자")
_p = _sub("ingest-dir", str(_dir), "--doc-type", DT, "--allow-mock")
_hint.unlink(missing_ok=True)
show("ⓔ 노칭 · 스태킹 슬라이드의 개체가 각자 제 공정 아래 · 언급만 한 탭용접은 노드 0 + 관련 링크 · --coord 없이 ingest-dir 성공",
     {"노칭::커터", "스태킹::매거진"} <= _names and not any(x.endswith("::탭용접") for x in _names if x)
     and {"chunk_id": _s5, "node_id": _tab} in (_ch.get("about") or [])
     and _p.returncode == 0 and "덱 구조 — 슬라이드 5" in _p.stdout,
     f"노드 {sorted(x for x in _names if x and ('커터' in x or '매거진' in x))} · ingest-dir rc {_p.returncode} · "
     + (_p.stdout + _p.stderr).strip().splitlines()[-1][:120])

# ────────────────────────────────────────────────────────────── B146 ⓔ 중간 파일 지도
print("\n■ B146 ⓔ 중간 파일 지도 — 표 · 엑셀 산문 · PPT를 하나씩 넣은 상태 루트 ↔ 상태 폴더 가이드 표")
from g6_common import RFQ as _RFQ6, _register as _reg6, DT as _DT6   # noqa: E402 — 산문 엑셀 표본(B83)
with _quiet():
    init.init(fresh_=True, all_=True)                     # 클린 — 앞 절의 산출이 섞이지 않게
    RUN.cmd_bootstrap([])
_register()
_reg6()
_hint.write_text(json.dumps(DECK, ensure_ascii=False), encoding="utf-8")
_ins6 = [_sub("ingest-file", str(_dir / f"{DID}.pptx"), "--doc-type", DT, "--allow-mock").returncode,
         _sub("ingest-file", str(ROOT / "tests/fixtures/raw/CP01.xlsx"), "--doc-type", "cp", "--allow-mock").returncode,
         _sub("ingest-file", str(_RFQ6), "--doc-type", _DT6, "--allow-mock", "--sheets", "2-4:prose 5:ref *:skip").returncode]
_hint.unlink(missing_ok=True)
_files = sorted(str(f.relative_to(_P.home())) for f in _P.home().rglob("*")
                if f.is_file() and not any(x.startswith(".") or x == "__pycache__" for x in f.relative_to(_P.home()).parts))
_gd = (ROOT / "docs/가이드/상태_폴더_가이드.md").read_text(encoding="utf-8")
_rows6 = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in _gd.split("### 2-1.")[1].split("\n## ")[0].splitlines()
          if ln.startswith("| `")]


def _rx146(tok, doc=None):
    """지도의 자리 → 정규식 — `<doc_id>`는 그 문서 · 다른 `<…>`는 한 칸 · `/`로 끝나면 그 아래 전부."""
    parts = _re.split(r"(<[^>]+>)", tok)
    return _re.compile("".join((_re.escape(doc) if doc and x == "<doc_id>" else "[^/]+") if x.startswith("<") else _re.escape(x)
                               for x in parts) + (".*" if tok.endswith("/") else ""))


_toks6 = [(r, _re.findall(r"`([^`]+)`", r[0])) for r in _rows6]
_loose6 = [f for f in _files if not any(_rx146(t).fullmatch(f) for _r, ts in _toks6 for t in ts)]
_miss6 = [(t, d) for r, ts in _toks6 for col, d in ((6, "CP01"), (7, "RFQ01"), (8, DID)) if r[col] == "●"
          for t in ts if not any(_rx146(t, d).fullmatch(f) for f in _files)]
show("B146 ⓔ 중간 파일 지도 — 표 · 엑셀 산문 · PPT를 하나씩 넣은 상태 루트의 파일이 전부 가이드 표의 자리에 들고 · 표가 「생긴다(●)」고 한 "
     "자리는 그 포맷 문서의 파일이 있다",
     _ins6 == [0, 0, 0] and _files and not _loose6 and not _miss6,
     f"인입 rc {_ins6} · 파일 {len(_files)} · 표 밖 {_loose6[:4]} · 없음 {_miss6[:4]}")

# B146 ⓐ~ⓓ — 구조 블록(엑셀 · Word) · 덱 블록과 같은 읽기 함수 · md는 봉투 하나에서 · 문서 해시 · 확정 시트
from core.state import ids as _IDS6                         # noqa: E402
from parser.adapters import basic_docx as _BD6               # noqa: E402
from parser.struct_rule import labels_hash as _LH6           # noqa: E402
_rq = json.loads(_P.parsed("RFQ01.json").read_text(encoding="utf-8"))
_mdq = _P.export("split_RFQ01.md").read_text(encoding="utf-8")
_P.work("struct_maps", "RFQ01.json").unlink(missing_ok=True)       # 보고 · 캐시 없이 — 봉투(JSON)만
struct_map.invalidate("B146W")
_ew = pipeline.parse(_BD6, "B146W", str(ROOT / "tests/fixtures/raw/DOC01.docx"), closed_list=[],
                     structure=ST.structurer()).envelope
_ewj = json.loads(json.dumps(_ew, ensure_ascii=False))
_uq = (_rq.get("structure") or {}).get("units") or []
show("B146 ⓐ 구조 블록 — 엑셀 산문(시트 둘 이상) · Word 봉투에 단위마다 목차 나무 · 청크 경계 · 판단 출처 · md는 봉투(JSON) 하나로 찍은 글과 같다",
     len(_uq) >= 2 and all(u.get("headings") and u.get("chunks") and u.get("by") and u.get("source") for u in _uq)
     and EMD.text("RFQ01", _rq) == _mdq and (_ewj.get("structure") or {}).get("units")
     and EMD.text("B146W", _ewj) == EMD.text("B146W", _ew) and "## 시트" in EMD.text("B146W", _ewj),
     f"단위 {[(u['unit'], u['by'], len(u['chunks'])) for u in _uq]} · Word 단위 {len((_ewj.get('structure') or {}).get('units') or [])}")
_pj6 = json.loads(_P.parsed(f"{DID}.json").read_text(encoding="utf-8"))
show("B146 ⓑ PPT md는 B143 꼴 그대로 · 덱 블록과 구조 블록이 같은 읽기 함수(units)로 읽힌다(단위 → 그 근거 청크)",
     EMD.text(DID, _pj6) == "\n".join(EMD.lines(_pj6)) + "\n" and EMD.units(_pj6) and EMD.units(_rq)
     and all(cs for _u, cs in EMD.units(_pj6) + EMD.units(_rq)),
     f"PPT 단위 {len(EMD.units(_pj6))} · 엑셀 단위 {len(EMD.units(_rq))}")
_u0 = _uq[0]
ST.confirm("RFQ01", _u0["unit"], {**_u0, "labels": _LH6(_u0.get("labels") or [])}, ST.doc_state("RFQ01")["hints"])
_cf0 = ST.confirmed_path().read_bytes()
_x7 = _sub("ingest-file", str(_RFQ6), "--doc-type", _DT6, "--allow-mock", "--sheets", "2-4:prose 5:ref *:skip")
_rq2 = json.loads(_P.parsed("RFQ01.json").read_text(encoding="utf-8"))
show("B146 ⓒ ⓓ 같은 문서를 다시 넣으면 추출 체크포인트 재사용(문서 해시 = 구조 블록 밖) · 사람이 확정한 시트는 구조 블록에 출처 「확정」 · "
     "확정 기록 파일은 그대로",
     "체크포인트 재사용" in _x7.stdout and _IDS6.doc_hash(_rq2) == _IDS6.doc_hash({k: v for k, v in _rq2.items() if k != "structure"})
     and _rq2["structure"]["units"][0]["by"] == "확정" and ST.confirmed_path().read_bytes() == _cf0,
     f"rc {_x7.returncode} · 첫 단위 {_rq2['structure']['units'][0]['by']} · "
     + next((ln.strip() for ln in _x7.stdout.splitlines() if "체크포인트" in ln), (_x7.stdout + _x7.stderr)[-160:]))
ST.confirmed_path().unlink(missing_ok=True)

# ────────────────────────────────────────────────────────────── ⓕ 고침 · 재생
print("\n■ ⓕ 고침 · 재생")
with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register()
struct_map.invalidate(DID)
_r1, _e1, _c1 = _parse()                               # 덱 구조 1회(보존)
_got = STS.parse_deck_fix("4 공정 스태킹", (_r1.report.get("deck") or {}).get("표") or [], STS.deck_text_of(_e1))
DK.record_fix(DID, *_got)
_r2, _e2, _c2b = _parse()                              # 다시 파싱 — 보존분 위에 고침(LLM 0)
_s4b = next(c for c in _e2["chunks"] if c["source_locator"] == "슬라이드 4")
show("ⓕ 구조 표의 공정 고침 → 기록(@덱) → 다시 파싱이 그대로(덱 구조 LLM 0 · 재구축도 같은 기록)",
     _got == (4, {"processes": ["스태킹"]}) and "deck" not in _c2b and _s4b["process_ref"] == "스태킹"
     and (_r2.report["deck"]["수"]["고침"] == 1) and "@덱" in json.loads(ST.fixes_path().read_text(encoding="utf-8"))[DID],
     f"고침 {_got} · 다시 파싱 호출 {_c2b} · 슬라이드 4 공정 {_s4b['process_ref']}")

# ────────────────────────────────────────────────────────────── B143 ⓒ 추출 입력은 JSON + 캡처만
print("\n■ B143 ⓒ 추출 입력 — 덱 블록 고침이 맥락 줄로 · 원본은 캡처에서만")
from core.build import extract_ctx as XC                 # noqa: E402
_d4 = next((x for x in (_e2.get("deck") or {}).get("slides") or [] if 4 in x.values()), {})
_line4 = XC.context_line(_s4b, XC.doc_info(_e2))
_opened, _rendered = [], []
import builtins as _bi                                    # noqa: E402
_open0 = _bi.open
_bi.open = lambda f, *a, **k: (_opened.append(str(f)), _open0(f, *a, **k))[1]
render.capture_path = lambda path=None, probe=False: (True, "시험", None)
render.render_pages = lambda path: (_rendered.append(str(path)), {n: b"PNG" for n in range(1, 6)})[1]
EU._PAGES.clear()
EX.invalidate(DID)
try:
    _run(_e2)
finally:
    _bi.open = _open0
    render.capture_path, render.render_pages = _keep
show("B143 ⓒ 덱 블록의 공정 고침(사람)이 추출 입력 맥락 줄을 따른다 · 원본 파일은 캡처 렌더에서만 열린다(추출은 봉투 + 캡처)",
     ["스태킹"] in _d4.values() and any(isinstance(v, list) and "processes" in v for v in _d4.values())
     and "공정: 스태킹" in _line4
     and "개요: " + DECK["overview"] in _line4 and _rendered == [str(PATH)] and str(PATH) not in _opened,
     f"맥락 줄 {_line4[:80]} · 렌더 {len(_rendered)} · 원본 열기 {sum(1 for x in _opened if x == str(PATH))}")

# ────────────────────────────────────────────────────────────── B151 덱 중립화
print("\n■ B151 — 덱 판단은 층 · 골격 · 문서 좌표를 모른다 · 공정은 문서 표기 · 공정 걸기는 표와 같은 길")
from pptx import Presentation                            # noqa: E402
from pptx.util import Inches                             # noqa: E402
from cli import parse as CP                              # noqa: E402
from core.build import extract_ctx as XC151             # noqa: E402
from g65_common import _rw                               # noqa: E402
SEEN151 = []


class _Live151(Live):
    """덱 호출의 입력 실물(지시문 + 사용자 글 — 메시지 전부)을 그대로 모은다."""

    def chat(self, messages, json_schema=None, point=None, **kw):
        if point == "deck":
            SEEN151.append([dict(m) for m in messages])
        return super().chat(messages, json_schema=json_schema, point=point, **kw)


def _deck151(path, slides):
    """창작 덱 — 슬라이드마다 글 상자들(첫 상자가 제목 자리)."""
    prs = Presentation()
    for boxes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for i, text in enumerate(boxes):
            s.shapes.add_textbox(Inches(0.5), Inches(0.3 + i * 1.2), Inches(8), Inches(0.5)).text_frame.text = text
    _P.ensure(path)
    prs.save(path)
    return path


D151 = [["B151 점검 보고"], ["노칭 커터 점검", "커터 마모 0.03mm"], ["스태킹 매거진 점검", "매거진 정렬 0.1mm"],
        ["스태킹 매거진 점검 (계속)", "정렬 핀 교체"], ["스태킹 공정 정렬 기준", "정렬 기준 0.2mm"],
        ["스태킹 공정 정렬 기준 (계속)", "기준 갱신 3월"], ["노칭 · 스태킹 공정 연동", "커터와 매거진을 함께 본다"], ["일정", "3월 착수"]]
A151 = {"slides": [
    {"slide": 1, "title": "B151 점검 보고", "section": None, "continues": None, "subject": None, "processes": [],
     "no_entities": True},
    {"slide": 2, "title": "노칭 커터 점검", "section": None, "continues": None, "subject": "노칭 커터 · 없는 말",
     "processes": ["노칭", "없는공정"], "no_entities": False},
    {"slide": 3, "title": "스태킹 매거진 점검", "section": None, "continues": None, "subject": "스태킹 매거진",
     "processes": ["스태킹"], "no_entities": False},
    {"slide": 4, "title": "스태킹 매거진 점검 (계속)", "section": None, "continues": 3, "subject": None, "processes": [],
     "no_entities": False},
    {"slide": 5, "title": "스태킹 공정 정렬 기준", "section": None, "continues": None, "subject": "정렬 기준",
     "processes": ["스태킹 공정"], "no_entities": False},
    {"slide": 6, "title": "스태킹 공정 정렬 기준 (계속)", "section": None, "continues": 5, "subject": None, "processes": [],
     "no_entities": False},
    {"slide": 7, "title": "노칭 · 스태킹 공정 연동", "section": None, "continues": None, "subject": "커터 · 매거진",
     "processes": ["노칭", "스태킹 공정"], "no_entities": False},
    {"slide": 8, "title": "일정", "section": None, "continues": None, "subject": "일정", "processes": [], "no_entities": False}],
    "overview": "노칭 · 스태킹 점검 보고서다."}
P151 = _deck151(TMP / "B151D.pptx", D151)


def _p151(doc_id, path, answer, pick=None, coord="노칭", keep=False):
    if not keep:                                         # 보존분(덱 판단은 구조 지도 보존분에 산다)을 지우고 처음부터
        struct_map.invalidate(doc_id)
    del SEEN151[:]
    with _Live151(by={"deck": lambda b: answer, "extract": lambda b: {"entities": [], "relations": [], "about": []},
                      "outline": lambda b: {"headings": []}, "reinforce": lambda b: {"pick": None}}):
        res = pipeline.parse(PX, doc_id, str(path), closed_list=TG.closed_list("process"), layer="process", deck=DK.decker(),
                             structure_fixes=ST.fixes_of(doc_id), doc_coord=coord, pick_coord=pick)
        calls = [p_ for p_, _b in CALLS]
    env = res.envelope
    env["doc_type"] = DT
    for c in env["chunks"]:
        c["doc_type"] = DT
    return res, env, calls, [list(m) for m in SEEN151]


def _judged(env):
    return [{k: s.get(k) for k in ("slide", "title", "section", "continues", "subject", "processes", "no_entities")}
            for s in (env.get("deck") or {}).get("slides") or []]


def _sl(env, n):
    return next(c for c in env["chunks"] if c["source_locator"] == f"슬라이드 {n}")


with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register()
_ra, _ea, _ca, _ma = _p151("B151D", P151, A151)
_rw(_P.layers("process", "config.json"), lambda c: c.setdefault("prompts", {}).__setitem__("structure", "층 힌트 B151 — 매거진은 설비다"))
_rc151, _eb, _cb, _mb = _p151("B151D", P151, A151, keep=True)   # 같은 상태 · 층 config만 바뀜 — 보존분 재사용
_sk = _P.layers("process", "skeleton.json")
_rw(_sk, lambda s: s["ALIASES"].setdefault("스태킹", []).append("스태킹 B151"))   # 골격 판이 바뀐다
with _quiet():
    RUN.cmd_bootstrap([])                                # 별칭을 골격에(fresh 없이) — 덱 판단 보존분은 아래가 지운다
_r2b, _e2b, _c2b2, _m2b = _p151("B151D", P151, A151, coord="조립")   # 층 힌트 · 골격 판 · 문서 좌표가 다른 상태에서 다시 묻는다
_sys = [m[0]["content"] for m in _ma + _m2b]
_usr = [json.loads(m[1]["content"]) for m in _ma + _m2b]
from tools import layer_words as _LW151                  # noqa: E402
_lw = [v for v in _LW151.scan(ROOT) if "2.4_deck" in str(v)] if hasattr(_LW151, "scan") else []
show("B151 ⓐ 층 config(층 힌트) · 골격 판 · 문서 좌표를 바꾼 두 상태에서 같은 덱 — 덱 블록의 판단 칸이 같다 · 덱 호출 입력(실물)에 층 힌트 · 골격 "
     "이름 목록 · 문서 좌표 0(두 상태의 입력이 글자까지 같다) · 덱 지시문 층 어휘 0",
     _ma and _m2b and _judged(_ea) == _judged(_e2b) and _ma == _m2b
     and all(set(u) == {"deck", "previous", "slides"} for u in _usr)
     and not any(x in s for s in _sys for x in ("층 힌트 B151", "## 층 어휘", "탭용접", "패키징"))
     and not _lw,
     f"호출 {len(_ma)} · {len(_m2b)} · 같은 입력 {_ma == _m2b} · 판단 같음 {_judged(_ea) == _judged(_e2b)} · 입력 열쇠 "
     f"{sorted(_usr[0]) if _usr else None}")
show("B151 ⓑ 층 config를 바꾼 뒤 같은 덱 재파싱 → 덱 판단 보존분 재사용(LLM 0)",
     "deck" not in _cb and (_rc151.report["deck"]["수"]["재사용"] == 1) and _judged(_eb) == _judged(_ea),
     f"호출 {_cb} · 재사용 {_rc151.report['deck']['수'].get('재사용')}")
_s2 = next(s for s in _ea["deck"]["slides"] if s["slide"] == 2)
show("B151 ⓒ 다루는 것 · 공정 검사 — 슬라이드 글에 없는 명사구 · 덱 글에 없는 공정 표기는 버려지고 dropped에 · 있는 것은 덱 블록 · "
     "청크 meta · 맥락 줄(「다루는 것」 칸)에",
     _s2["subject"] == "노칭 커터" and _s2["processes"] == ["노칭"] and {"다루는 것", "공정"} <= set(_s2["dropped"])
     and _sl(_ea, 2)["meta"].get("subject") == "노칭 커터" and "다루는 것: 노칭 커터" in XC151.context_line(_sl(_ea, 2))
     and _sl(_ea, 4)["meta"].get("subject") == "스태킹 매거진",
     f"슬라이드 2 {_s2} · 맥락 줄 {XC151.context_line(_sl(_ea, 2))[:120]}")
#: 표 · 차트 머리도 덱 글이다 — 덱 입력이 준 머리에만 있는 공정 표기 · 다루는 것은 남는다(제목 검사만 머리를 뺀다)
_hd = DK.check("B151H", [{"slide": 1, "layout": "", "boxes": [{"kind": "title", "text": "검사 결과"},
                                                              {"kind": "table", "text": "스태킹 공정 | 정렬 | 판정"}]}],
               {"slides": [{"slide": 1, "title": "검사 결과", "subject": "정렬", "processes": ["스태킹"]}]})["slides"][1]
show("B151 ⓒ 표 머리의 공정 표기 · 다루는 것 — 덱 입력이 준 표 머리에만 있어도 검사를 지난다",
     _hd.get("processes") == ["스태킹"] and _hd.get("subject") == "정렬" and not _hd.get("dropped"), _hd)
# ⓓ 공정 걸기 — mock(⑨ 없음) 그래프 단계 · ⑨ 주입(시험 pick)
with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register()
_rd, _ed, _cd, _md = _p151("B151D", P151, A151)
_n151, _cr151 = _run(_ed)
_q151 = [q for q in store.read(store.QUEUE, []) if q.get("doc_id") == "B151D"]
_mis = {(q.get("payload") or {}).get("provenance", "").split("#")[-1] for q in _q151 if q.get("kind") == "coord_mismatch"}
_orp = [q for q in _q151 if q.get("kind") == "orphan_anchor"]
_ASK = []
_rp, _ep, _cp, _mp = _p151("B151P", P151, A151, pick=lambda ref, cands: (_ASK.append(ref), "스태킹" if ref == "스태킹 공정" else None)[1])
_s7 = _sl(_ep, 7)
show("B151 ⓓ 공정 걸기(문서 좌표 「노칭」) — 슬라이드 4는 3을 이어 「스태킹」 · 태깅에서 골격에 맞음 · 문서 좌표 밖이라 coord_mismatch에 "
     "그 슬라이드들 · 골격 · 별칭에 없는 「스태킹 공정」은 ⑨ 없는 mock에서 그대로 → orphan_anchor 큐 · ⑨ 주입이면 같은 표기를 한 번 묻고 "
     "채택은 canonical + 원 표기 meta(목록 조각은 원소마다) · 공정 빈 슬라이드는 문서 좌표를 물려받는다",
     _sl(_ed, 3)["process_ref"] == "스태킹" and _sl(_ed, 4)["process_ref"] == "스태킹"
     and {"슬라이드 3", "슬라이드 4"} <= _mis and _sl(_ed, 5)["process_ref"] == "스태킹 공정"
     and any("스태킹 공정" in json.dumps(q, ensure_ascii=False) for q in _orp)
     and _ASK == ["스태킹 공정"] and _sl(_ep, 5)["process_ref"] == "스태킹" and _sl(_ep, 6)["process_ref"] == "스태킹"
     and (_sl(_ep, 5)["meta"].get("coord_tag_source"), _sl(_ep, 5)["meta"].get("coord_tag_from")) == ("live", "스태킹 공정")
     and _s7["process_ref"] == ["노칭", "스태킹"]
     and [(x["from"], x.get("source")) for x in _s7["meta"].get("coord_tags") or []] == [("노칭", None), ("스태킹 공정", "live")]
     and not PDK.mismatch(_ep) and _sl(_ed, 8)["process_ref"] == "노칭" and _sl(_ed, 8)["meta"].get("coord_from_doc"),
     f"3·4 {_sl(_ed, 3)['process_ref']} · {_sl(_ed, 4)['process_ref']} · 어긋남 {sorted(_mis)} · 고아 {len(_orp)} · 물음 {_ASK} · "
     f"7 {_s7['process_ref']} {_s7['meta'].get('coord_tags')} · 8 {_sl(_ed, 8)['process_ref']}")
# B157 리뷰 고침 — 덱 판단이 「공정 없음」(빈 목록 — 어댑터는 좌표를 싣지 않는다)이라 한 슬라이드는 제목이 줄기 밖 공정 이름이어도 지금 그대로
_p157 = _deck151(TMP / "B157D.pptx", [["B157 보고"], ["2. 스태킹", "적층 속도 10ppm"]])
_a157 = {"slides": [{"slide": 1, "title": "B157 보고", "section": None, "continues": None, "subject": None, "processes": [],
                     "no_entities": True},
                    {"slide": 2, "title": "2. 스태킹", "section": None, "continues": None, "subject": None, "processes": [],
                     "no_entities": False}], "overview": "보고"}
_s157 = _sl(_p151("B157D", _p157, _a157)[1], 2)
show("B157 리뷰 고침 ⓐ 덱 판단 「공정 없음」 슬라이드(제목 「2. 스태킹」 · 문서 좌표 「노칭」) — 다른 공정 구간을 세우지 않고 문서 좌표를 물려받는다",
     _s157["process_ref"] == "노칭" and _s157["meta"].get("coord_from_doc") and not _s157["meta"].get("coord_other"),
     f"{_s157['process_ref']} · {_s157.get('meta', {}).get('coord_other')}")
# B154 ① 좌표 쌍 다음 줄 — 상위의 출처로(ⓓ의 덱 = 문서 좌표 상위 · 행 상위가 어긋난 표 조각 = 행)
from cli import coord_queue as CQU154, doc_coord as DC154   # noqa: E402
_ps154, _qs154, _ts154 = io.StringIO(), io.StringIO(), io.StringIO()
with contextlib.redirect_stdout(_ps154):
    DC154.pair_screen(_ed)
with contextlib.redirect_stdout(_qs154):
    CQU154.show()
with contextlib.redirect_stdout(_ts154):
    DC154.pair_screen({"doc_id": "B154T", "doc_type": "cp", "source_path": "B154T.xlsx",
                       "records": [{"source_locator": "CP!R4", "process_group": "스태킹", "process_ref": "노칭"}]})
_nx = lambda b: [ln.strip() for ln in b.getvalue().splitlines() if ln.strip().startswith("▶ 다음 줄(")]   # noqa: E731
show("B154 ⓐ 좌표 쌍 다음 줄 — 덱(문서 좌표 상위)은 문서 좌표 줄 · 열 매핑 0 · 파싱 화면과 큐 화면이 같은 줄 · 행 상위가 어긋난 표 "
     "조각은 열 매핑 줄",
     _nx(_ps154) and _nx(_ps154) == _nx(_qs154) and all("문서 좌표 상위" in x and "열 매핑" not in x for x in _nx(_ps154))
     and "--coord <상위 공정>" in _nx(_ps154)[0] and _nx(_ts154) and all("열 매핑을 바로잡는다" in x for x in _nx(_ts154)),
     f"파싱 {_nx(_ps154)[:1]} · 큐 {len(_nx(_qs154))}줄 · 표 {[x[:40] for x in _nx(_ts154)]}")
from cli import report_cause as RC154                      # noqa: E402
_h154 = RC154.head_count(RC154.material("B151D"))[:2]
store.append_defect("B151D: 관계 후보 끝점 미해소 — '브래킷' -part_of-> '노칭' @ B151D:b154")   # 관계 끝점 표본(결함 로그의 꼴 그대로)
_l154 = [ln.strip() for ln in RC154.report("B151D")[0]]
show("B154 ⓑ 행 밖 — 값은 다 붙었고(못 붙은 행 0) 골격에 없는 공정 표기 「스태킹 공정」 · 끝점 없는 관계가 원인 묶음 아래 「행 밖」에 · "
     "머리 수 그대로(행 밖은 합에 없다)",
     _h154[0] == 0 and f"원인 묶음 — 못 붙은 행 0 / 전체 {_h154[1]}" in _l154
     and any(x.startswith("공정 표기 — 골격에 없음 '스태킹 공정' · 조각 ") for x in _l154)
     and "관계 — 끝점 없음 '브래킷' · 못 선 관계 1 · B151D:b154" in _l154
     and RC154.head_count(RC154.material("B151D"))[:2] == _h154,
     f"머리 {_h154} · {[x for x in _l154 if '—' in x][:4]}")
# B155 ⓒ 추출까지만 — 덱(B151 덱 표본) · 추출 md(맥락 줄 · 원문 · 개체 표 · 판정 예고 · 한눈에)
def _x155(body):
    t = body["chunk"]
    if "커터 마모" in t:                                  # 사전 일치 · 판정에 오름(글에 없음) · 값 글자 버림
        return {"entities": [_ent("노칭", "Process"), _ent("커터 홀더", "Unit"), _ent("마모량", "Property", value="0.09mm")],
                "relations": [], "about": []}
    if "정렬 기준 0.2mm" in t:                            # 좌표 「스태킹」 아래 — 판정에 오름(판정 절에 새 노드)
        return {"entities": [_ent("정렬 지그", "Unit")], "relations": [], "about": []}
    return {"entities": [], "relations": [], "about": []}   # 일정 슬라이드 — 부른 단위인데 개체 0


_r155, _e155, _c155, _m155 = _p151("B155D", P151, A151)
_P.ensure(_P.parsed("B155D.json"))
_P.parsed("B155D.json").write_text(json.dumps(_e155, ensure_ascii=False), encoding="utf-8")   # 파싱이 남기는 봉투 자리
_g155 = json.dumps(store.read(store.DOC_REGISTRY, {}), sort_keys=True)
with Live(by={"extract": _x155, "outline": lambda b: {"headings": []}}), _quiet():
    _rx155 = run_document(dict(_e155), extract_only=True)
from cli import extract_md as XMD155                    # noqa: E402
_md155 = XMD155.text("B155D") or ""
_need155 = ("다루는 것: 노칭 커터", "커터 마모 0.03mm", "| 커터 홀더 | Unit |", "글에 없음", "| 노칭 | Process |", "사전 일치",
            "판정에 오름", "검사 버림(값 글자)", "⚠ 부른 단위인데 개체 0", "## 판정 — 층 process")
show("B155 ⓒ (B167 고침) 덱 판정까지 — md에 슬라이드마다 맥락 줄(다루는 것) · 원문 · 개체 표(글에서 찾은 자리 · 「글에 없음」 · "
     "판정 예고 칸(사전 일치 · 판정에 오름 · 검사 버림)) · 판정 절 · 한눈에 ⚠ · 문서 대장 0",
     getattr(_rx155[0], "extract_stop", False) and all(x in _md155 for x in _need155)
     and json.dumps(store.read(store.DOC_REGISTRY, {}), sort_keys=True) == _g155,
     f"{_rx155[0].reason} · 빠진 것 {[x for x in _need155 if x not in _md155]}")
#: 리뷰 고침(PR #111) — ①개체 부모가 주 좌표를 이긴다(좌표 없는 청크라도 부모가 서면 보류 예상이 아니다 — 구축과 같은 손) ②봉투가
#: 체크포인트와 다른 판이면 그 사실을 말하고 저장된 청크로 찍는다
from core.build.build import Builder as _BD155                                    # noqa: E402
_m155 = XMD155.material("B155D")
_m155.update(doc=XC151.doc_info(_m155["env"]), closed={"process": XMD155._closed("process")},
             builders={"process": _BD155(open_graph("process"), load_config("process"), None, "(시험)", "process")})
_pl155 = XMD155._plan(_m155, "process", {"source_locator": "x", "text": "노칭 신규 커터 점검", "process_ref": None},
                      _ent("신규 커터", "Unit", parent="노칭"))
_pe155 = json.loads(_P.parsed("B155D.json").read_text(encoding="utf-8"))
_pe155["chunks"][1]["text"] += " (고친 문서)"
_P.parsed("B155D.json").write_text(json.dumps(_pe155, ensure_ascii=False), encoding="utf-8")   # --split-only로 다시 파싱한 꼴
_st155 = XMD155.material("B155D")
show("B155 리뷰 고침 — 판정 예고는 개체 부모를 먼저 본다(좌표 없는 청크 + 부모 「노칭」 → 판정에 오름 · 보류 예상 0) · 봉투가 체크포인트와 "
     "다른 판이면 md 머리가 그 사실을 말한다(저장된 청크로 찍는다)",
     _pl155 == ("판정에 오름", []) and _st155["stale"] and not XMD155.material("B155X")["stale"]
     and "다른 판이다" in (XMD155.text("B155D") or ""),
     f"예고 {_pl155} · 다른 판 {_st155['stale']} · 청크 {len(_st155['chunks'])}")
# ⓔ 공정이 없는 덱
D151E = [["투자 계획"], ["예산 3억", "3월 집행"], ["일정", "4월 착수"]]
A151E = {"slides": [{"slide": 1, "title": "투자 계획", "section": None, "continues": None, "subject": None, "processes": [],
                     "no_entities": True},
                    {"slide": 2, "title": "예산 3억", "section": None, "continues": None, "subject": "예산", "processes": [],
                     "no_entities": False},
                    {"slide": 3, "title": "일정", "section": None, "continues": None, "subject": "일정", "processes": [],
                     "no_entities": False}], "overview": "투자 계획서다."}
_re, _ee, _ce, _me = _p151("B151E", _deck151(TMP / "B151E.pptx", D151E), A151E, coord=None)
_ne, _cre = _run(_ee)
_qe = [q for q in store.read(store.QUEUE, []) if q.get("doc_id") == "B151E"]
show("B151 ⓔ 공정이 없는 덱 — 모든 슬라이드 공정 빈 목록 · 다루는 것은 채워진다 · 파싱 · 구축 오류 0 · 큐 0",
     _re.ok and all(s["processes"] == [] for s in _ee["deck"]["slides"])
     and [s["subject"] for s in _ee["deck"]["slides"]] == [None, "예산", "일정"]
     and not any(c.get("process_ref") for c in _ee["chunks"]) and not _qe and not PDK.mismatch(_ee),
     f"공정 {[s['processes'] for s in _ee['deck']['slides']]} · 다루는 것 {[s['subject'] for s in _ee['deck']['slides']]} · 큐 {len(_qe)}")
# ⓕ 옛 @덱 고침(공정 · 언급 칸 포함)
DK.record_fix("B151D", 3, {"title": "스태킹 매거진 점검", "processes": ["노칭"], "mentions": ["탭용접"]})
_rf, _ef, _cf, _mf = _p151("B151D", P151, A151, keep=True)
_buf151 = io.StringIO()
with contextlib.redirect_stdout(_buf151):
    CP.print_split_notes(_rf.report)
_s3f = next(s for s in _ef["deck"]["slides"] if s["slide"] == 3)
show("B151 ⓕ 옛 @덱 고침(공정 · 언급 칸 포함) — 제목 · 공정은 고침대로 · 언급 칸은 읽지 않고 화면에 한 줄",
     _s3f["title"] == "스태킹 매거진 점검" and _s3f["processes"] == ["노칭"] and "mentions" not in _s3f
     and "읽지 않은 옛 언급 고침 1" in _buf151.getvalue() and "deck" not in _cf,
     f"슬라이드 3 {_s3f.get('title')} · {_s3f.get('processes')} · 줄 "
     f"{[ln.strip() for ln in _buf151.getvalue().splitlines() if '옛 언급' in ln]}")
ST.fixes_path().unlink(missing_ok=True)

# ────────────────────────────────────────────────────────────── B156 쪽 판독
print("\n■ B156 — 쪽 판독(메타 · 층 무관) · 상자 위치 · 크기 · 판독 표(다시 · 지움) · 추출 · 그래프 불변")
from core.build import reading as RD156                  # noqa: E402
from g6_common import _run as _pty156                    # noqa: E402 — pty 도우미(관문은 tty에서만 산다)
with _quiet():
    init.init(fresh_=True, all_=True)
    RUN.cmd_bootstrap([])
_register()
#: mock 판독 힌트(창작) — 쪽 4는 검사에 걸리는 칸 넷(없는 상자 · 목록 밖 관계 · 근거 없는 요약 · 보임/추론 밖) · `again`은 사람 힌트의 답
R156 = {"pages": {
    "2": {"figures": [], "relations": [], "summary": [{"text": "노칭 금형 마모와 클리어런스를 점검했다.", "boxes": ["S2-B1", "S2-B2"]}]},
    "3": {"figures": [], "relations": [], "summary": [{"text": "앞 장에 이어 다이 균열을 봤다.", "boxes": ["S3-B2"]}]},
    "4": {"figures": [{"box": "S4-P1", "what": "커터 장력 기준을 보이는 단면 도면이다", "certainty": "높음", "grounds": "그림 속 글자",
                       "kind": "도면", "labels": ["장력", "T-1"],
                       "seen": [{"text": "그림 속 글자 T-1", "basis": "보임"}, {"text": "커터 단면으로 보인다", "basis": "추론"},
                                {"text": "느낌", "basis": "느낌"}],
                       "meaning": "표의 장력 기준을 그림으로 보인다", "unreadable": False},
                      {"box": "S9-P9", "kind": "사진", "labels": [], "seen": [], "meaning": None, "unreadable": False}],
          "relations": [{"from": "S4-B1", "to": "S4-TB1", "relation": "설명"}, {"from": "S4-TB1", "to": "S4-P1", "relation": "무관"}],
          "summary": [{"text": "커터 사양 표와 도면이다.", "boxes": ["S4-TB1", "S4-P1"]}, {"text": "근거 없는 문장", "boxes": ["없는상자"]}],
          "again": {"figures": [], "relations": [], "summary": [{"text": "도면의 치수 T-1을 다시 읽었다.", "boxes": ["S4-P1"]}]}},
    "5": {"figures": [], "relations": [], "summary": [{"text": "노칭과 스태킹 연동을 함께 본다.", "boxes": ["S5-B2"]}]}}}
_rh156 = _FXR.EXTRACT_HINTS / f"{DID}@reading.json"
_hint.write_text(json.dumps(DECK, ensure_ascii=False), encoding="utf-8")
_rh156.write_text(json.dumps(R156, ensure_ascii=False), encoding="utf-8")
_SEEN156, _NOTE156 = [], []


def _p156(reading=True):
    """mock 파싱(덱 · 판독 힌트) — 판독 함수에 들어간 입력(실물)을 모은다."""
    fn = RD156.reader(notice=_NOTE156.append) if reading else None
    if fn is not None:
        inner = fn

        def fn(doc_id, inp, **kw):
            _SEEN156.append(json.loads(json.dumps(inp, ensure_ascii=False)))
            return inner(doc_id, inp, **kw)
        for k in ("key", "live", "pages", "capture_ref", "notice", "drift"):
            setattr(fn, k, getattr(inner, k))
    with _quiet():
        res = pipeline.parse(PX, DID, str(PATH), closed_list=TG.closed_list("process"), layer="process", deck=DK.decker(),
                             structure_fixes=ST.fixes_of(DID), reading=fn)
    return res, res.envelope


struct_map.invalidate(DID)
_ra6, _ea6 = _p156()
_in6 = list(_SEEN156)
_rn6, _en6 = _p156(reading=False)                       # 판독 없는 파싱 — B156 전과 같은 길
_rd6 = {e["page"]: e for e in _ea6["reading"]["pages"]}
_p4 = _rd6[4]
_bx4 = next(b for pg in _ea6["boxes"]["pages"] if pg["page"] == 4 for b in pg["boxes"] if b["id"] == "S4-P1")
_md6 = EMD.text(DID, _ea6)
show("B156 ⓐ PPT 덱 — 봉투 reading(쪽마다 그림 · 그림 속 글자 · 보임/추론 · 배치 관계 · 쪽 요약 + 근거 id) · 개체 없음 장 건너뜀 · 검사(없는 "
     "상자 · 목록 밖 관계 · 근거 없는 요약 문장 · 보임/추론 밖 → dropped) · split md에 판독 이름표 줄과 [그림 요약]이 나란히(B169 고침) · "
     "상자 위치 · 크기",
     _rd6[1].get("skipped") == "개체 없음" and _p4["figures"][0]["box"] == "S4-P1" and _p4["figures"][0]["labels_in_text"] == ["장력"]
     and [s["basis"] for s in _p4["figures"][0]["seen"]] == ["보임", "추론"] and len(_p4["figures"]) == 1
     and _p4["relations"] == [{"from": "S4-B1", "to": "S4-TB1", "relation": "설명"}]
     and [s["boxes"] for s in _p4["summary"]] == [["S4-TB1", "S4-P1"]]
     and _p4["dropped"] == {"보이는 것": 1, "그림": 1, "관계": 1, "요약": 1}
     and all(isinstance(_bx4[k], int) for k in ("left", "top", "width", "height")) and 0 < _bx4["rel"][0] < 1
     and "- **판독**: 출처 mock" in _md6 and "[그림 요약 S4-P1]" in _md6 and "- **그림 S4-P1**:" in _md6 and "  - 종류 도면" in _md6,
     f"쪽 {[(e['page'], e.get('skipped') or len(e.get('summary') or [])) for e in _ea6['reading']['pages']]} · 버림 {_p4['dropped']} · "
     f"상자 {_bx4}")
show("B156 ⓕ 추출 · 그래프 불변 — 판독이 있든 없든 같은 덱의 청크(그림 요약 글 포함) · 추출 입력의 맥락 줄이 같다 · 판독 블록은 청크 밖",
     json.dumps(_en6["chunks"], sort_keys=True, ensure_ascii=False) == json.dumps(_ea6["chunks"], sort_keys=True, ensure_ascii=False)
     and [XC.context_line(c, XC.doc_info(_en6)) for c in _en6["chunks"]]
     == [XC.context_line(c, XC.doc_info(_ea6)) for c in _ea6["chunks"]] and "reading" not in _en6 and "boxes" in _en6,
     f"청크 {len(_ea6['chunks'])} · 판독 없는 봉투 키 {sorted(k for k in _en6 if k in ('reading', 'boxes'))}")
_rb6, _eb6 = _p156()                                    # 같은 덱 재파싱 — 보존분 재사용
_kb6 = _rb6.report["reading"]["수"]
_kn6 = _P.registry(RD156.KNOWLEDGE)
_P.ensure(_kn6)
_kn6.write_text("# 읽기 지식\n노칭 커터는 은색 날이다.\n", encoding="utf-8")
del _NOTE156[:]
_rk6, _ek6 = _p156()                                    # 읽기 지식을 고치면 다시 부른다(예고에 수)
_kn6.unlink()
_rw(_P.layers("process", "config.json"), lambda c: c.setdefault("prompts", {}).__setitem__("structure", "층 힌트 B156 — 매거진은 설비다"))
struct_map.invalidate(DID)
del _SEEN156[:]
_rl6, _el6 = _p156()                                    # 층 config가 바뀐 상태에서 처음부터 — 판독 입력이 같다
show("B156 ⓑ 같은 덱 재파싱 → 판독 재사용(부름 0) · 읽기 지식을 고치면 다시 부름(예고에 「읽기 지식이 바뀌어 다시」 수) · 층 config를 바꿔도 "
     "판독 입력(실물)이 같다 · 이어진 장은 앞 장 쪽 요약을 받는다",
     (_kb6["부름"], _kb6["재사용"]) == (0, 4) and _NOTE156 and _NOTE156[0]["지식바뀜"] == 4
     and _rk6.report["reading"]["수"]["부름"] == 4 and _ek6["reading"]["knowledge"]
     and _SEEN156 == _in6 and next(x for x in _in6 if x["page"] == 3)["previous"] == "노칭 금형 마모와 클리어런스를 점검했다.",
     f"재파싱 {(_kb6['부름'], _kb6['재사용'])} · 예고 {_NOTE156[:1]} · 입력 같음 {_SEEN156 == _in6}")
RD156.record_fix(DID, 4, {"action": "다시", "hint": "도면의 치수를 읽어라"})
_rx6, _ex6 = _p156()
RD156.record_fix(DID, 5, {"action": "지움"})
_ry6, _ey6 = _p156()
_kx6, _ky6 = _rx6.report["reading"]["수"], _ry6.report["reading"]["수"]
_x4 = next(e for e in _ex6["reading"]["pages"] if e["page"] == 4)
_y5 = next(e for e in _ey6["reading"]["pages"] if e["page"] == 5)
ST.fixes_path().unlink(missing_ok=True)
_o6 = _pty156("ingest-file", str(PATH), "--doc-type", DT, "--allow-mock", "--split-only", "--step",
              answers="c\nc\n4 다시 도면의 치수를 읽어라\n5 지움\nc\n" + "c\n" * 3)   # [1/8] · 덱 표 · 판독 표(다시 · 지움 · 계속)
_fx6 = (json.loads(ST.fixes_path().read_text(encoding="utf-8")) if ST.fixes_path().exists() else {}).get(DID, {}).get("@판독", {})
_pj156 = json.loads(_P.parsed(f"{DID}.json").read_text(encoding="utf-8")) if _P.parsed(f"{DID}.json").exists() else {}
_pr156 = {e["page"]: e for e in (_pj156.get("reading") or {}).get("pages") or []}
show("B156 ⓒ 판독 고침 — 「4 다시 <힌트>」는 그 장만 다시(다른 장 재사용 · 출처 hint) · 「5 지움」은 그 장 판독 0(지움 표시) · --step(pty) "
     "판독 표의 동사가 @판독 프레임에 기록 · 그 파싱 산출이 재생",
     (_kx6["부름"], _kx6["재사용"]) == (1, 3) and _x4["source"] == "hint" and _x4["summary"][0]["text"].startswith("도면의 치수")
     and _y5.get("deleted") and not _y5.get("summary") and (_ky6["부름"], _ky6["지움"], _ky6["고침"]) == (0, 1, 2)
     and (_fx6.get("4") or {}).get("action") == "다시" and (_fx6.get("4") or {}).get("hint") == "도면의 치수를 읽어라"
     and (_fx6.get("5") or {}).get("action") == "지움" and "구조까지" in _o6
     and (_pr156.get(4) or {}).get("source") == "hint" and (_pr156.get(5) or {}).get("deleted"),
     f"다시 {(_kx6['부름'], _kx6['재사용'])} · 지움 {(_ky6['부름'], _ky6['지움'])} · 기록 {sorted(_fx6)} · pty 끝 "
     f"{[ln.strip()[:50] for ln in _o6.splitlines() if '고침 —' in ln or '구조까지' in ln][:3]}")
ST.fixes_path().unlink(missing_ok=True)
_hint.unlink(missing_ok=True)
_rh156.unlink(missing_ok=True)

# ────────────────────────────────────────────────────────────── B168 판독 입력
print("\n■ B168 — 판독 입력: 그림 원본은 필요한 것만 · 표시 상자(판독 입력에만) · 고침 동사 「자세히 · 빼기」")
from PIL import Image as _Im168                          # noqa: E402 — 시험이 덱을 짓는다(창작 · 새 픽스처 파일 0)
from pptx.dml.color import RGBColor as _RGB168           # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE     # noqa: E402
from cli import reading_screen as RSC168                  # noqa: E402
from parser import reader as RDR168, reader_marks as RM168, reading as RDG168   # noqa: E402


def _png168(c):
    b = io.BytesIO()
    _Im168.new("RGB", (24, 24), c).save(b, "PNG")
    return b.getvalue()


def _deck168(name, builds):
    prs = Presentation()
    for fn in builds:
        fn(prs.slides.add_slide(prs.slide_layouts[6]))
    path = TMP / name
    _P.ensure(path)
    prs.save(path)
    return path


def _pic168(s, data, x, y):
    return s.shapes.add_picture(io.BytesIO(data), Inches(x), Inches(y), Inches(0.6), Inches(0.6))


def _read168(path, fixes=None):
    """판독 훅을 가짜 판독(입력 · 보낸 그림 id를 모은다)으로 — `(쪽 → (입력, 원본 id), 셈, 판독 블록)`."""
    seen, tally = {}, {}

    def ask(doc_id, inp, *, capture=None, images=(), hint=None, text=""):
        seen[inp["page"]] = (inp, [r for r, _b, _m in images])
        return {"figures": [], "relations": [], "summary": [{"text": "시험", "boxes": [inp["boxes"][0]["id"]]}],
                "dropped": {}, "calls": 0, "source": "mock"}
    blk = RDG168.hook(ask, "B168", str(path), RDR168.read(str(path)), fixes=fixes, tally=tally)
    return seen, (tally.get("수") or {}).get("그림") or {}, blk


_red168, _logo168, _blue168 = _png168((200, 30, 30)), _png168((30, 30, 200)), _png168((30, 200, 30))


def _ten(s):
    s.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(6), Inches(0.5)).text_frame.text = "같은 그림 열 장"
    for i in range(10):
        _pic168(s, _red168, 0.5 + (i % 5) * 1.2, 1.0 + (i // 5) * 1.2)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.8), Inches(0.9), Inches(0.8), Inches(0.8))   # 셋째 그림에 빨간 원
    o.fill.solid()
    o.fill.fore_color.rgb = _RGB168(0xFF, 0, 0)


_A168 = _deck168("B168A.pptx", [_ten])
_sa, _ga, _ = _read168(_A168)
_ia, _oa = _sa[1]
_marks168 = [b for b in _ia["boxes"] if b["kind"] == RM168.KIND]
_seen168 = [b["id"] for b in _ia["boxes"] if str(b.get("image") or "").startswith("캡처로 봄")]
show("B168 ⓐ 같은 그림 10장 중 1장에 빨간 원(글 없는 도형) → 상자 목록에 표시 1 · 원본은 대표 1 + 그 1장(표시가 겹친 예외) · 나머지 8 「캡처로 봄」",
     len(_marks168) == 1 and _oa == ["S1-P1", "S1-P3"] and len(_seen168) == 8 and (_ga["원본"], _ga["같은 꼴"]) == (2, 8),
     f"표시 {[(m['id'], m.get('shape'), m.get('color')) for m in _marks168]} · 원본 {_oa} · 셈 {_ga}")


def _logo(text):
    def f(s):
        s.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(6), Inches(0.5)).text_frame.text = text
        _pic168(s, _logo168, 8, 6)
    return f


_sb, _gb, _ = _read168(_deck168("B168B.pptx", [_logo("쪽 하나"), _logo("쪽 둘"), _logo("쪽 셋")]))
show("B168 ⓑ 모든 쪽에 같은 로고 → 장식 · 원본 입력 0",
     all(not v[1] for v in _sb.values()) and len(_sb) == 3 and (_gb["장식"], _gb["원본"]) == (3, 0), f"셈 {_gb}")


def _arrow(s):
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(3), Inches(0.5))
    tb.text_frame.text = "이 그림을 보라"
    _pic168(s, _blue168, 1, 3)
    p2 = _pic168(s, _blue168, 5, 3)                       # 같은 꼴 — 연결선이 가리켜 원본(예외)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1), Inches(5), Inches(3))
    c.begin_connect(tb, 2)
    c.end_connect(p2, 0)


_sc, _gc, _ = _read168(_deck168("B168C.pptx", [_arrow]))
_cx = [b for b in _sc[1][0]["boxes"] if b["kind"] == RM168.KIND]
show("B168 ⓒ 연결선 하나(글 상자 → 그림) → 표시 상자에 from → to · 그 그림은 원본(같은 꼴이어도 — 예외)",
     [(b.get("shape"), b.get("from"), b.get("to")) for b in _cx] == [("연결선", "S1-B1", "S1-P2")]
     and _sc[1][1] == ["S1-P1", "S1-P2"], f"표시 {_cx} · 원본 {_sc[1][1]}")

_keep168 = (RM168.collector, RM168.finish)


def _env168(path, did):
    with _quiet():
        res = pipeline.parse(PX, did, str(path), closed_list=TG.closed_list("process"), layer="process")
    return json.dumps({k: v for k, v in (res.envelope or {}).items() if k != "parsed_at"}, ensure_ascii=False, sort_keys=True)


_d168 = [(p, _env168(p, f"B168D{i}")) for i, p in enumerate((_FXR.RAW / "PPT_basic.pptx", _FXR.RAW / "PPT_shapes.pptx", _A168,
                                                               TMP / "B168C.pptx"))]
RM168.collector, RM168.finish = (lambda: None), (lambda m: {"marks": [], "groups": []})     # 표시를 모으지 않는 리더(B168 전 모양)
try:
    _d168b = [_env168(p, f"B168D{i}") for i, (p, _e) in enumerate(_d168)]
finally:
    RM168.collector, RM168.finish = _keep168
show("B168 ⓓ 판독 없는 덱(PPT_basic · PPT_shapes · 표시 · 연결선이 있는 창작 덱 둘)의 봉투(분할 · 청크 · 상자 목록 · 덱 칸 — 추출 입력의 "
     "재료)가 표시를 모으지 않는 리더와 바이트 같다 · 봉투에 표시 0",
     [e for _p, e in _d168] == _d168b and not any(RM168.KIND in e for _p, e in _d168),
     f"같음 {[a == b for (_p, a), b in zip(_d168, _d168b)]}")

_benv = {"boxes": RDG168.boxes(RDR168.read(str(_A168))), "reading": _read168(_A168)[2]}
for _ans in ("1 자세히 S1-P5", "1 빼기 S1-P1"):
    _pg, _fx = RSC168.parse_fix(_ans, _benv)
    RD156.record_fix("B168", _pg, _fx)
_fxe = ST.fixes_of("B168").get(RDG168.FRAME) or {}
_se, _ge, _be = _read168(_A168, fixes=_fxe)
_im = _be["pages"][0].get("images") or {}
show("B168 ⓔ 「1 자세히 S1-P5」 → 그 그림 원본 입력 · 「1 빼기 S1-P1」 → 원본 0(사람이 뺌) · 기록은 @판독 그 쪽 칸의 그림 목록",
     "S1-P5" in _se[1][1] and "S1-P1" not in _se[1][1] and _im.get("S1-P1", "").endswith("사람이 뺌")
     and (_fxe.get(1) or {}).get("자세히") == ["S1-P5"] and (_fxe.get(1) or {}).get("빼기") == ["S1-P1"],
     f"원본 {_se[1][1]} · 기록 {_fxe.get(1)} · 셈 {_ge}")
ST.fixes_path().unlink(missing_ok=True)


def _grp(s):
    g = s.shapes.add_group_shape()
    for i, c in enumerate((_red168, _blue168)):
        g.shapes.add_picture(io.BytesIO(c), Inches(1 + i), Inches(1), Inches(0.6), Inches(0.6))


_sg, _gg, _bg = _read168(_deck168("B168G.pptx", [_grp]))
show("B168 리뷰 고침 글 없이 그림 조각 그룹만 있는 쪽 → 빈 쪽으로 건너뛰지 않고 읽는다(빈 쪽은 봉투 상자로 · 입력은 그룹 상자)",
     1 in _sg and any(b["kind"] == RM168.GROUP for b in _sg[1][0]["boxes"]),
     f"부른 쪽 {sorted(_sg)} · 건너뜀 {[p.get('skipped') for p in (_bg or {}).get('pages') or []]}")

# ────────────────────────────────────────────────────────────── B169 판독 지시문 r-1.1 · md 이름표
print("\n■ B169 — 그림마다 첫 줄 「무엇인가 · 확신 · 근거」 · 일반 지식 절 · md 이름표 · 추출 확인 md에 판독")
_f169 = {"kind": "사진", "labels": [], "seen": [], "meaning": None, "unreadable": False}
_k169 = RD156.check("B169", {"page": 1, "boxes": [{"id": "S1-P1", "kind": "picture"}, {"id": "S1-P2", "kind": "picture"}]},
                    {"figures": [dict(_f169, box="S1-P1", what="파우치 셀의 측면 사진이다", certainty="높음", grounds="보이는 것"),
                                 dict(_f169, box="S1-P2", what="같은 셀의 윗면이다", certainty="아주 높음", grounds=None)],
                     "relations": [], "summary": []})
show("B169 ⓐ 판독 산출의 그림마다 첫 줄 · 확신(닫힌 셋) · 근거 · 셋 밖 확신은 그 칸만 버림(첫 줄은 남는다)",
     [(f["what"], f["certainty"], f["grounds"]) for f in _k169["figures"]]
     == [("파우치 셀의 측면 사진이다", "높음", "보이는 것"), ("같은 셀의 윗면이다", None, None)] and _k169["dropped"] == {"확신": 1},
     f"그림 {[(f['box'], f['certainty']) for f in _k169['figures']]} · 버림 {_k169['dropped']}")
_pr169 = (ROOT / "prompts" / "2.4_reading.md").read_text(encoding="utf-8")
_lw169 = subprocess.run([sys.executable, str(ROOT / "tools" / "layer_words.py")], capture_output=True, text=True, cwd=str(ROOT),
                        stdin=subprocess.DEVNULL)
_ah169 = subprocess.run([sys.executable, str(ROOT / "tests" / "asset_hashes.py")], capture_output=True, text=True, cwd=str(ROOT),
                        stdin=subprocess.DEVNULL)
show("B169 ⓑ 지시문 r-1.1 · 첫 줄 · 확신 · 근거 · 제조 문서를 읽는 일반 지식 절 · 층 어휘 정적 점검 위반 0 · 자산 해시 일치",
     "\nversion: r-1.1\n" in _pr169 and "## 제조 문서를 읽는 일반 지식" in _pr169 and "`certainty`" in _pr169
     and "위반 0건" in _lw169.stdout and _ah169.returncode == 0, f"층 어휘 {_lw169.stdout.strip()[-40:]} · 해시 rc {_ah169.returncode}")
from core.build.ingest import prose_chunk_ids as _pci169    # noqa: E402
_ch169 = dict((cid, c) for c, cid in _pci169(_ea6))
_c4 = next(cid for cid, c in _ch169.items() if (c.get("meta") or {}).get("slide") == 4)
_x169 = XMD155.text(DID, {"doc_id": DID, "env": _ea6, "chunks": _ch169, "stale": False, "rows": [], "partial": False,
                          "checkpoints": [("process", None, {"candidates": [{"chunk_id": _c4, "entities": []}]})]}) or ""
show("B169 ⓒ split md · 추출 확인 md의 이름표 줄 — **맥락 줄** · **쪽 요약** · **배치** · **그림 <id>**: 첫 줄(확신 · 근거) · 추출 확인 md에 "
     "그 단위의 판독",
     all(x in _md6 for x in ("- **그림 S4-P1**: 커터 장력 기준을 보이는 단면 도면이다 (확신 높음 · 근거 그림 속 글자)",
                             "- **배치**: S4-B1 → S4-TB1 설명", "- **쪽 요약**: 커터 사양 표와 도면이다."))
     and all(x in _x169 for x in ("- **맥락 줄**:", "- **판독**: 출처 mock", "- **그림 S4-P1**:", "- **쪽 요약**:")),
     " | ".join(ln for ln in _x169.splitlines() if "**" in ln)[:200])
_lab169 = ("**판독**", "**그림 ", "**배치**", "**쪽 요약**")
_n169 = EMD.text(DID, _en6) or ""
show("B169 ⓓ 판독이 없는 문서 — split md · 추출 확인 md에 판독 줄 0(판독 칸이 비면 줄을 내지 않는다) · 추출 확인 md의 맥락 줄은 이름표",
     _n169 and not any(x in _n169 for x in _lab169) and "- **맥락 줄**:" in _md155 and not any(x in _md155 for x in _lab169),
     f"split {len(_n169)}자 · 추출 확인 판독 줄 {[x for x in _lab169 if x in _md155]}")

# ────────────────────────────────────────────────────────────── B170 판독을 추출에
print("\n■ B170 — 판독이 있는 청크에만 [판독] 묶음 · 그림 칸 · 이름 출처에 그림 속 글자 · 판독 없는 문서는 바이트 그대로")
import hashlib as _hl170                                  # noqa: E402
from core.llm import gateway                              # noqa: E402
from core.build import belongs as BL170, extract_ctx as XC170, extract_reading as ER170   # noqa: E402
from core.build.entry import _vocab as _vocab170          # noqa: E402
_cfg170 = load_config("process")
#: B170 전의 추출 지시문 · 답 꼴(`8dec743`의 `prompts/3.3_extract.md` · `EXTRACT_SCHEMA`) — 판독이 없는 문서가 받는 것과 같아야 한다
_PRE170 = {"prompt": "f1ac41461f918903", "schema": "6f3e6047d44e1cea"}


def _sent170(env):
    """실호출 갈래(가짜 게이트웨이)로 추출 — 호출마다 `(지시문, 답 꼴, 사용자 글)`."""
    sent = []
    EX.invalidate(env["doc_id"])
    with Live(by={"extract": lambda body: {"entities": [], "relations": [], "about": []}}), _quiet():
        inner = gateway.chat

        def chat(messages, json_schema=None, point=None, **kw):
            if point == "extract":
                sent.append((messages[0]["content"], json_schema, messages[1]["content"]))
            return inner(messages, json_schema=json_schema, point=point, **kw)
        gateway.chat = chat
        EX.extract(env, _cfg170, {loc: cid for c, cid in _pci169(env) for loc in [c["source_locator"]]}, _vocab170(_cfg170))
    return sent


_wr170 = _sent170(_ea6)
_sent170(_ea6)                                            # 다시 — 체크포인트만 지운다(호출 보존 그대로)
_run170 = dict(EX.RUN.get((_ea6["doc_id"], "")) or {})
_r4 = next((s for s in _wr170 if "[판독 그림 S4-P1]" in s[2]), None)
show("B170 ⓐ 판독이 있는 쪽의 추출 입력에 [판독] 묶음(쪽 요약 · 배치) · 그 쪽 그림 요약 자리에 판독 그림 칸(첫 줄 · 그림 속 글자) · 그 호출만 "
     "판독 절 · 소속 출처에 판독",
     _r4 is not None and "[판독]" in _r4[2] and "커터 장력 기준을 보이는 단면 도면이다" in _r4[2] and "「T-1」" in _r4[2]
     and "[그림 요약 S4-P1]" not in _r4[2] and "## 판독이 있는 청크" in _r4[0]
     and "판독" in _r4[1]["properties"]["entities"]["items"]["properties"]["belongs_to"]["properties"]["from"]["enum"],
     f"호출 {len(_wr170)} · 판독 붙은 호출 {sum('[판독]' in s[2] for s in _wr170)}")
_wo170 = _sent170(_en6)
_sys0 = gateway.prompt_for("extract", {"process": _cfg170}, "extract")
show("B170 ⓑ 판독이 없는 문서 — 실제 추출 지시문 · 답 꼴 · 사용자 글이 B170 전과 바이트 같다(지시문 파일 · 답 꼴 해시가 그 판 · 판독 0) · "
     "체크포인트에 판독 칸 0",
     _wo170 and all(s[0] == _sys0 and s[1] == EX.EXTRACT_SCHEMA and "판독" not in s[2] for s in _wo170)
     and _hl170.sha256((ROOT / "prompts" / "3.3_extract.md").read_bytes()).hexdigest()[:16] == _PRE170["prompt"]
     and _hl170.sha256(json.dumps(EX.EXTRACT_SCHEMA, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:16] == _PRE170["schema"]
     and "reading" not in json.loads(EX.checkpoint_path(_en6["doc_id"]).read_text(encoding="utf-8")),
     f"호출 {len(_wo170)}")
_ke, _kn = EX._now_key(_ea6, _cfg170), EX._now_key(_en6, _cfg170)
show("B170 ⓒ 판독이 있는 문서만 재사용 열쇠가 바뀐다(판독 해시) · 판독을 지우면 다시 같은 열쇠 · 판독 없는 체크포인트로는 판독 있는 봉투를 "
     "재사용하지 않는다",
     _ke["reading"] and _kn["reading"] is None and EX._now_key({**_ea6, "reading": None}, _cfg170) == _kn
     and EX.reuse_check(_ea6, cfg=_cfg170)[1].startswith("판독") and EX.reuse_check(_en6, cfg=_cfg170)[0],
     f"판독 몫 {_ke['reading']} · 재사용 {EX.reuse_check(_ea6, cfg=_cfg170)}")
_doc170, _s4 = XC170.doc_info(_ea6), _ch169[_c4]
show("B170 ⓓ 그림 속 글자에서 나온 이름은 출처 「판독」 · 외형 추론만 있는 이름(보이는 것의 추론 · 첫 줄)은 출처가 아니다 · 판독 없는 봉투는 판독 칸 0",
     BL170.source_of("T-1", _s4, _doc170) == "판독" and BL170.source_of("커터 단면", _s4, _doc170) is None
     and BL170.source_of("단면 도면", _s4, _doc170) is None and "reading" not in XC170.doc_info(_en6),
     f"T-1 {BL170.source_of('T-1', _s4, _doc170)} · 커터 단면 {BL170.source_of('커터 단면', _s4, _doc170)}")
show("B170 리뷰 고침 판독이 있는 덱을 다시 추출하면(호출 보존 그대로) 판독을 실은 단위도 다시 쓸 단위로 센다 — 부름 0(다시 쓸 단위 · 씨앗의 "
     "열쇠 = 보낸 입력 한 손)",
     _run170.get("부름") == 0 and (_run170.get("재사용") or 0) > 0, f"다시 돈 셈 {_run170}")

dts = store.read(store.DOC_TYPES, {})
dts.pop(DT, None)
store.write(store.DOC_TYPES, dts)
done()
